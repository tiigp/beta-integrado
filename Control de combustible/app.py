from __future__ import annotations

import csv
import calendar
import base64
import json
import os
import re
import secrets
import shutil
import smtplib
import sqlite3
import tempfile
import unicodedata
import zipfile
from functools import wraps
from io import BytesIO, StringIO
from pathlib import Path
from datetime import date, datetime, timedelta
from types import SimpleNamespace
from urllib.parse import quote_plus, urlencode
from email.message import EmailMessage
from xml.sax.saxutils import escape

from flask import Flask, Response, abort, flash, jsonify, redirect, render_template, request, send_file, session, url_for
from flask_login import LoginManager, UserMixin, current_user, login_required, login_user, logout_user
from flask_sqlalchemy import SQLAlchemy
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas
from reportlab.platypus import Flowable, Image, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlalchemy import false, func, or_
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from docx import Document as WordDocument
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image as PILImage
import qrcode

BASE_DIR = Path(__file__).resolve().parent
app = Flask(__name__, root_path=str(BASE_DIR), template_folder=str(BASE_DIR / "templates"), static_folder=str(BASE_DIR / "static"))
secret_key_path = BASE_DIR / "instance" / "app_secret_key"
secret_key_path.parent.mkdir(parents=True, exist_ok=True)
secret_key = os.environ.get("APP_SECRET_KEY")
if not secret_key:
    if secret_key_path.exists():
        secret_key = secret_key_path.read_text(encoding="utf-8").strip()
    else:
        secret_key = os.urandom(32).hex()
        secret_key_path.write_text(secret_key, encoding="utf-8")
app.config["SECRET_KEY"] = secret_key
IMAGENES_DIR = Path(r"C:\Users\PC\OneDrive\Desktop\BETA INTEGRADO\Imagenes")
DB_PATH = BASE_DIR / "instance" / "combustible.db"
DB_PATH.parent.mkdir(parents=True, exist_ok=True)
QUALITY_CERTIFICATES_LOCAL_DIR = Path(os.environ.get("QUALITY_CERTIFICATES_LOCAL_DIR", r"C:\CertificadosCalidad"))
QUALITY_CERTIFICATES_PROJECT_DIR = BASE_DIR / "Reports" / "Backups" / "Certificados"
QUALITY_CERTIFICATES_REMOTE_DIR = Path(os.environ.get("QUALITY_CERTIFICATES_REMOTE_DIR", r"\\100.90.6.40\CertificadosCalidad"))
QUALITY_LABELS_DIR = BASE_DIR / "Reports" / "Backups" / "Etiquetas"
QUALITY_LABELS_IMAGES_DIR = IMAGENES_DIR
QUALITY_LABEL_DELETE_CODE = "4192633"
CERTIFICATE_JSON_RESTORE_PATH = Path(r"E:\b\certificados_calidad_backup.json")
CERTIFICATE_ASSETS_RESTORE_DIR = Path(r"E:\b\c")
QUALITY_ROD_DIAMETERS = ("8", "10", "12", "16", "20", "25")

database_url = os.environ.get("DATABASE_URL", f"sqlite:///{DB_PATH.as_posix()}")
if database_url.startswith("mysql://"):
    database_url = "mysql+pymysql://" + database_url.removeprefix("mysql://")
app.config["SQLALCHEMY_DATABASE_URI"] = database_url
app.config["SQLALCHEMY_TRACK_MODIFICATIONS"] = False
app.config["MAX_CONTENT_LENGTH"] = 25 * 1024 * 1024
app.jinja_env.cache = {}
STATIC_LOGO_PATH = BASE_DIR / "static" / "logo.png"
LOGO_SOURCE_PATH = BASE_DIR / "Imagen" / "logoigp.png"
RRHH_DOCUMENTS_DIR = BASE_DIR / "instance" / "documentos_rrhh"
RRHH_DOCUMENTS_DIR.mkdir(parents=True, exist_ok=True)
RRHH_PHOTOS_DIR = BASE_DIR / "instance" / "fotos_colaboradores"
RRHH_PHOTOS_DIR.mkdir(parents=True, exist_ok=True)
PURCHASE_QUOTES_DIR = BASE_DIR / "instance" / "presupuestos_compras"
PURCHASE_QUOTES_DIR.mkdir(parents=True, exist_ok=True)

db = SQLAlchemy(app)
login_manager = LoginManager()
login_manager.login_view = "login"
login_manager.init_app(app)


class User(db.Model, UserMixin):
    id = db.Column(db.Integer, primary_key=True)
    username = db.Column(db.String(80), unique=True, nullable=False)
    password_hash = db.Column(db.String(255), nullable=False)
    full_name = db.Column(db.String(120), nullable=False)
    role = db.Column(db.String(20), nullable=False, default="despacho")
    correo = db.Column(db.String(160), nullable=True)
    telefono = db.Column(db.String(40), nullable=True)

    @property
    def is_root(self):
        return self.role.strip().lower() == "root"

    @property
    def is_admin(self):
        return self.role.strip().lower() == "admin" or self.is_root

    @property
    def is_operador_horno(self):
        return self.role.strip().lower() == "operador_horno"

    @property
    def is_operador_despunte(self):
        return self.role.strip().lower() == "operador_despunte"

    @property
    def is_dispatcher(self):
        return self.role.strip().lower() == "despacho"

    @property
    def is_deposito(self):
        return self.role.strip().lower() == "deposito"

    @property
    def is_rrhh(self):
        normalized_role = self.role.strip().lower().replace("_", " ").replace("-", " ")
        return normalized_role in {"rrhh", "recursos humanos"}

    @property
    def is_supervisor(self):
        return self.role.strip().lower() == "supervisor"

    @property
    def is_compras(self):
        return self.role.strip().lower() == "compras"

    @property
    def is_gerencia(self):
        return self.role.strip().lower() == "gerencia"

    @property
    def is_calidad(self):
        return self.role.strip().lower() == "calidad"

    @property
    def is_control_calidad(self):
        return self.role.strip().lower() == "control_calidad"

    @property
    def can_access_rrhh(self):
        return self.is_admin or self.is_rrhh or self.is_supervisor or self.is_root

    @property
    def can_access_quality(self):
        return self.is_admin or self.is_calidad or self.is_control_calidad


class EtiquetaCalidad(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    tipo_producto = db.Column(db.String(40), nullable=False)
    calidad = db.Column(db.String(80), nullable=False)
    medida = db.Column(db.String(30), nullable=False, default="")
    longitud = db.Column(db.String(30), nullable=False, default="")
    peso = db.Column(db.String(30), nullable=False, default="")
    colada = db.Column(db.String(60), nullable=False, default="")
    lote = db.Column(db.String(60), nullable=False, default="")
    cantidad = db.Column(db.Integer, nullable=True, default=0)
    fecha = db.Column(db.Date, nullable=True)
    carbono = db.Column(db.String(30), nullable=False, default="")
    silicio = db.Column(db.String(30), nullable=False, default="")
    manganeso = db.Column(db.String(30), nullable=False, default="")
    hornero = db.Column(db.String(120), nullable=False, default="")
    supervisor = db.Column(db.String(120), nullable=False, default="")
    operador_ccm = db.Column(db.String(120), nullable=False, default="")
    aprobado = db.Column(db.Boolean, nullable=False, default=False)
    tipo_b = db.Column(db.Boolean, nullable=False, default=False)
    no_conforme = db.Column(db.Boolean, nullable=False, default=False)
    observaciones = db.Column(db.Text, nullable=False, default="")
    created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    created_by = db.relationship("User", foreign_keys=[created_by_id])


class ResponsableCalidad(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False)
    tipo = db.Column(db.String(30), nullable=False)
    activo = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)


class CertificadoCalidad(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    numero_colada = db.Column(db.String(60), nullable=False, unique=True)
    numero_certificado = db.Column(db.String(30), nullable=False, unique=True)
    tipo_producto = db.Column(db.String(40), nullable=False, default="")
    carbono = db.Column(db.String(30), nullable=False, default="")
    silicio = db.Column(db.String(30), nullable=False, default="")
    manganeso = db.Column(db.String(30), nullable=False, default="")
    medidas = db.Column(db.String(180), nullable=False, default="")
    longitudes = db.Column(db.String(180), nullable=False, default="")
    cantidad_palanquillas = db.Column(db.Integer, nullable=False, default=0)
    longitudes_palanquillas = db.Column(db.String(180), nullable=False, default="")
    pesos_palanquillas = db.Column(db.String(180), nullable=False, default="")
    datos_varillas = db.Column(db.String(300), nullable=False, default="")
    productos_certificados = db.Column(db.String(500), nullable=False, default="")
    alargamiento = db.Column(db.String(30), nullable=False, default="")
    limite_fluencia = db.Column(db.String(30), nullable=False, default="")
    limite_resistencia = db.Column(db.String(30), nullable=False, default="")
    doblado_cumple = db.Column(db.Boolean, nullable=False, default=False)
    fecha_atado = db.Column(db.Date, nullable=True)
    fecha_certificacion = db.Column(db.Date, nullable=True)
    grafico_path = db.Column(db.String(500), nullable=False, default="")
    firma_path = db.Column(db.String(500), nullable=False, default="")
    fecha_creacion = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    created_by = db.relationship("User", foreign_keys=[created_by_id])


class DespachoAcero(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    numero = db.Column(db.String(40), unique=True, nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    destino = db.Column(db.String(180), nullable=False)
    cliente = db.Column(db.String(180), nullable=False, default="")
    cliente_id = db.Column(db.Integer, db.ForeignKey("cliente_despacho.id"), nullable=True)
    chofer = db.Column(db.String(120), nullable=False)
    cedula = db.Column(db.String(40), nullable=False, default="")
    chapa = db.Column(db.String(30), nullable=False, default="")
    responsable = db.Column(db.String(120), nullable=False)
    observacion = db.Column(db.Text, default="")
    created_by_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    created_by = db.relationship("User", foreign_keys=[created_by_id])
    etiquetas = db.relationship("DetalleDespachoAcero", backref="despacho", lazy=True, cascade="all, delete-orphan")


class DetalleDespachoAcero(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    despacho_id = db.Column(db.Integer, db.ForeignKey("despacho_acero.id"), nullable=False)
    etiqueta_id = db.Column(db.Integer, db.ForeignKey("etiqueta_calidad.id"), nullable=False, unique=True)
    etiqueta = db.relationship("EtiquetaCalidad", foreign_keys=[etiqueta_id])


class ClienteDespacho(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(180), unique=True, nullable=False)
    ultimo_destino = db.Column(db.String(180), default="", nullable=False)
    ultimo_chofer = db.Column(db.String(120), default="", nullable=False)
    ultima_cedula = db.Column(db.String(40), default="", nullable=False)
    ultima_chapa = db.Column(db.String(30), default="", nullable=False)
    actualizado = db.Column(db.DateTime, default=datetime.utcnow, onupdate=datetime.utcnow, nullable=False)
    despachos = db.relationship("DespachoAcero", backref="cliente_registrado", lazy=True)


class Equipo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False)
    tipo = db.Column(db.String(80), nullable=False)
    responsable = db.Column(db.String(120), nullable=False)
    estado = db.Column(db.String(50), default="Activo")
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Compra(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    litros = db.Column(db.Float, nullable=False)
    proveedor = db.Column(db.String(120), nullable=False)
    precio_litro = db.Column(db.Float, default=0.0, nullable=False)
    costo_total = db.Column(db.Float, default=0.0, nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    notas = db.Column(db.Text, default="")


class Despacho(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    equipo_id = db.Column(db.Integer, db.ForeignKey("equipo.id"), nullable=False)
    litros = db.Column(db.Float, nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    responsable = db.Column(db.String(120), nullable=False)
    nombre_chofer = db.Column(db.String(120), default="")
    hora = db.Column(db.Float, default=0.0)
    horometro = db.Column(db.Float, default=0.0)
    kilometraje = db.Column(db.Float, default=0.0)
    tipo_carga = db.Column(db.String(30), default="Media")
    observacion = db.Column(db.Text, default="")
    equipo = db.relationship("Equipo", backref="despachos")


class Configuracion(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    clave = db.Column(db.String(80), nullable=False, unique=True)
    valor = db.Column(db.String(255), nullable=False)


class ProductoBascula(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(40), unique=True, nullable=False)
    nombre = db.Column(db.String(160), unique=True, nullable=False)
    unidad = db.Column(db.String(20), nullable=False, default="kg")
    activo = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)


class LaminacionProduccion(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    numero_colada = db.Column(db.String(60), nullable=False)
    cantidad_palanquillas = db.Column(db.Integer, nullable=False, default=0)
    carbono = db.Column(db.String(30), nullable=False, default="")
    manganeso = db.Column(db.String(30), nullable=False, default="")
    silicio = db.Column(db.String(30), nullable=False, default="")
    temperatura = db.Column(db.Float, nullable=False, default=0.0)
    velocidad = db.Column(db.Float, nullable=False, default=0.0)
    estado = db.Column(db.String(20), nullable=False, default="abierta")
    inicio = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    fin = db.Column(db.DateTime, nullable=True)
    creado_por_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    creado_por = db.relationship("User", foreign_keys=[creado_por_id])
    despuntes = db.relationship("LaminacionDespunte", backref="produccion", lazy=True, cascade="all, delete-orphan")
    paradas = db.relationship("LaminacionParada", backref="produccion", lazy=True, cascade="all, delete-orphan")

    @property
    def palanquillas_pasaron(self):
        return sum(item.pasaron for item in self.despuntes)

    @property
    def palanquillas_chatarra(self):
        return sum(item.chatarra for item in self.despuntes)


class LaminacionDespunte(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    produccion_id = db.Column(db.Integer, db.ForeignKey("laminacion_produccion.id"), nullable=False)
    pasaron = db.Column(db.Integer, nullable=False, default=0)
    chatarra = db.Column(db.Integer, nullable=False, default=0)
    fecha = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    observacion = db.Column(db.String(255), nullable=False, default="")
    creado_por_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    creado_por = db.relationship("User", foreign_keys=[creado_por_id])


class LaminacionParada(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    produccion_id = db.Column(db.Integer, db.ForeignKey("laminacion_produccion.id"), nullable=True)
    accion = db.Column(db.String(20), nullable=False, default="parada")
    sector = db.Column(db.String(100), nullable=False)
    problema = db.Column(db.String(160), nullable=False)
    inicio = db.Column(db.DateTime, nullable=False, default=datetime.utcnow)
    fin = db.Column(db.DateTime, nullable=True)
    chatarra = db.Column(db.Integer, nullable=False, default=0)
    observacion = db.Column(db.String(255), nullable=False, default="")
    creado_por_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    creado_por = db.relationship("User", foreign_keys=[creado_por_id])


class OfflineOperation(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    operation_key = db.Column(db.String(80), unique=True, nullable=False)
    user_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    created_at = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)


class CategoriaInsumo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(80), nullable=False, unique=True)
    prefijo = db.Column(db.String(10), nullable=True)
    activo = db.Column(db.Boolean, nullable=False, default=True)


class Insumo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(40), unique=True, nullable=True)
    nombre = db.Column(db.String(120), nullable=False, unique=True)
    categoria = db.Column(db.String(80), nullable=False, default="General")
    unidad = db.Column(db.String(30), nullable=False, default="unidad")
    stock_minimo = db.Column(db.Float, nullable=False, default=0.0)
    vencimiento = db.Column(db.Date, nullable=True)
    activo = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Colaborador(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(40), unique=True, nullable=True)
    legajo = db.Column(db.String(40), unique=True, nullable=True)
    nombre = db.Column(db.String(120), nullable=False)
    documento = db.Column(db.String(40), default="")
    telefono = db.Column(db.String(40), default="")
    area = db.Column(db.String(80), default="")
    estado = db.Column(db.String(30), nullable=False, default="Contratado")
    fecha_ingreso = db.Column(db.Date, nullable=True)
    puesto = db.Column(db.String(120), default="")
    turno_linea = db.Column(db.String(80), default="")
    calle = db.Column(db.String(120), default="")
    barrio = db.Column(db.String(100), default="")
    ciudad = db.Column(db.String(100), default="")
    pais = db.Column(db.String(80), default="Paraguay")
    nacionalidad = db.Column(db.String(80), default="Paraguaya")
    tipo_documento = db.Column(db.String(30), default="CI")
    fecha_nacimiento = db.Column(db.Date, nullable=True)
    estado_civil = db.Column(db.String(30), default="")
    sexo = db.Column(db.String(20), default="")
    cargo = db.Column(db.String(120), default="")
    seccion = db.Column(db.String(120), default="")
    correo = db.Column(db.String(160), default="")
    supervisor_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    vencimiento_cedula = db.Column(db.Date, nullable=True)
    hijos_cantidad = db.Column(db.Integer, nullable=False, default=0)
    hijos = db.Column(db.Boolean, nullable=False, default=False)
    aporte_ips = db.Column(db.String(20), default="SI")
    bonificacion_familiar = db.Column(db.Boolean, nullable=False, default=False)
    foto_archivo = db.Column(db.String(255), nullable=True)
    activo = db.Column(db.Boolean, nullable=False, default=True)
    created_at = db.Column(db.DateTime, default=datetime.utcnow)


class Vacacion(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    colaborador_id = db.Column(db.Integer, db.ForeignKey("colaborador.id"), nullable=False)
    fecha_inicio = db.Column(db.Date, nullable=False)
    fecha_fin = db.Column(db.Date, nullable=False)
    fecha_retorno = db.Column(db.Date, nullable=False)
    dias = db.Column(db.Integer, nullable=False, default=12)
    estado = db.Column(db.String(20), nullable=False, default="Programada")
    observacion = db.Column(db.Text, default="")
    creado_por_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    fecha_creacion = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    colaborador = db.relationship("Colaborador", backref="vacaciones")
    creado_por = db.relationship("User", foreign_keys=[creado_por_id])


class EventoRRHH(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    titulo = db.Column(db.String(160), nullable=False)
    tipo = db.Column(db.String(30), nullable=False, default="Tarea")
    inicio = db.Column(db.DateTime, nullable=False)
    fin = db.Column(db.DateTime, nullable=True)
    lugar = db.Column(db.String(160), default="")
    responsable = db.Column(db.String(120), default="")
    responsable_usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    notificar = db.Column(db.Boolean, nullable=False, default=False)
    colaborador_id = db.Column(db.Integer, db.ForeignKey("colaborador.id"), nullable=True)
    estado = db.Column(db.String(25), nullable=False, default="Pendiente")
    notas = db.Column(db.Text, default="")
    creado_por_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    creado_en = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    colaborador = db.relationship("Colaborador")
    responsable_usuario = db.relationship("User", foreign_keys=[responsable_usuario_id])
    creado_por = db.relationship("User", foreign_keys=[creado_por_id])


class AsignacionEvaluacion(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    colaborador_id = db.Column(db.Integer, db.ForeignKey("colaborador.id"), nullable=False)
    supervisor_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    formulario = db.Column(db.String(20), nullable=False, default="periodica")
    fecha_limite = db.Column(db.Date, nullable=True)
    estado = db.Column(db.String(20), nullable=False, default="Pendiente")
    evaluacion_id = db.Column(db.Integer, db.ForeignKey("evaluacion_personal.id"), nullable=True)
    creado_por_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    creado_en = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    colaborador = db.relationship("Colaborador")
    supervisor = db.relationship("User", foreign_keys=[supervisor_id])
    creado_por = db.relationship("User", foreign_keys=[creado_por_id])


class RegistroDisciplinario(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(40), unique=True, nullable=True)
    colaborador_id = db.Column(db.Integer, db.ForeignKey("colaborador.id"), nullable=False)
    usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    tipo = db.Column(db.String(30), nullable=False, default="Amonestación")
    motivo = db.Column(db.Text, nullable=False)
    dias_suspendidos = db.Column(db.Integer, nullable=False, default=0)
    observacion = db.Column(db.Text, default="")
    colaborador = db.relationship("Colaborador", backref="registros_disciplinarios")
    usuario = db.relationship("User")


class EvaluacionPersonal(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    colaborador_id = db.Column(db.Integer, db.ForeignKey("colaborador.id"), nullable=False)
    usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    formulario = db.Column(db.String(20), nullable=False, default="periodica")
    codigo = db.Column(db.String(20), nullable=False, default="RRHH-FOR-007")
    puntaje_total = db.Column(db.Integer, nullable=False, default=0)
    puntajes = db.Column(db.Text, nullable=False, default="{}")
    fecha_evaluacion = db.Column(db.Date, nullable=True)
    periodo_evaluado = db.Column(db.String(120), default="")
    tipo_evaluacion = db.Column(db.String(30), default="A demanda / especial")
    jefe_directo = db.Column(db.String(120), default="")
    fortalezas = db.Column(db.Text, default="")
    aspectos_mejorar = db.Column(db.Text, default="")
    plan_accion = db.Column(db.Text, default="")
    fecha_revision = db.Column(db.Date, nullable=True)
    recomendacion = db.Column(db.Text, default="")
    promedio = db.Column(db.Float, nullable=False, default=0.0)
    resultado = db.Column(db.String(100), default="")
    estado = db.Column(db.String(30), nullable=False, default="Borrador")
    observaciones = db.Column(db.Text, default="")
    colaborador = db.relationship("Colaborador", backref="evaluaciones")
    usuario = db.relationship("User")


class CarpetaRRHH(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(120), nullable=False, unique=True)
    fecha_creacion = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)


class DocumentoRRHH(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    nombre = db.Column(db.String(160), nullable=False)
    nombre_archivo = db.Column(db.String(255), nullable=False, unique=True)
    tipo = db.Column(db.String(100), nullable=False)
    categoria = db.Column(db.String(50), nullable=False, default="Otros")
    carpeta_id = db.Column(db.Integer, db.ForeignKey("carpeta_rrhh.id"), nullable=True)
    tamano = db.Column(db.Integer, nullable=False, default=0)
    fecha_creacion = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    fecha_modificacion = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    usuario_creacion_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    usuario_modificacion_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    usuario_creacion = db.relationship("User", foreign_keys=[usuario_creacion_id])
    usuario_modificacion = db.relationship("User", foreign_keys=[usuario_modificacion_id])
    carpeta = db.relationship("CarpetaRRHH", backref="documentos")


class Notificacion(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    titulo = db.Column(db.String(160), nullable=False)
    mensaje = db.Column(db.String(255), nullable=False)
    url = db.Column(db.String(255), nullable=False, default="/rrhh")
    leida = db.Column(db.Boolean, nullable=False, default=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    usuario = db.relationship("User")


EVALUATION_FORMS = {
    "prueba": {
        "codigo": "RRHH-FOR-007",
        "nombre": "Evaluación periódica",
        "estado": "Prueba",
        "groups": [
            ("1. Seguridad Industrial y Salud Ocupacional (SISO)", ["1.1 Uso Riguroso de EPP", "1.2 Cumplimiento de Protocolos", "1.3 Orden y Limpieza (5S)"]),
            ("2. Competencia Técnica y Calidad Operativa", ["2.1 Adherencia a Procedimientos", "2.2 Cuidado de Equipos e Insumos", "2.3 Calidad del Trabajo", "2.4 Ritmo y Productividad"]),
            ("3. Disciplina, Asistencia y Actitud", ["3.1 Puntualidad y Asistencia", "3.2 Disposición al Aprendizaje", "3.3 Trabajo en Equipo y Comunicación", "3.4 Respeto a las Normas"]),
        ],
    },
    "periodica": {
        "codigo": "RRHH-FOR-008",
        "nombre": "Evaluación anual",
        "estado": "Contratado",
        "groups": [
            ("1. Seguridad Industrial y Salud Ocupacional (SISO)", ["1.1 Cumplimiento Riguroso de Normas y EPP", "1.2 Cultura Preventiva", "1.3 Orden y Limpieza (5S)"]),
            ("2. Competencia Técnica y Calidad Operativa", ["2.1 Dominio del Puesto", "2.2 Control de Calidad", "2.3 Polivalencia y Flexibilidad", "2.4 Cuidado de Activos e Insumos"]),
            ("3. Disciplina, Asistencia y Actitud", ["3.1 Cumplimiento de Metas", "3.2 Solución de Problemas", "3.3 Autonomía", "4.1 Asistencia y Puntualidad", "4.2 Trabajo en Equipo y Relevo de Turno", "4.3 Disposición al Cambio y Mejora", "4.4 Compromiso con la Mejora Continua"]),
        ],
    },
    "administrativa": {
        "codigo": "RRHH-FOR-010",
        "nombre": "Evaluación de personal administrativo",
        "estado": "Todos",
        "administrativa": True,
        "groups": [
            ("1. Organización y Gestión Administrativa", ["1.1 Organización y planificación", "1.2 Cumplimiento de procedimientos", "1.3 Gestión documental y registros"]),
            ("2. Competencia y Calidad del Trabajo", ["2.1 Dominio de sus funciones", "2.2 Calidad y precisión", "2.3 Cumplimiento de objetivos", "2.4 Manejo de herramientas y sistemas"]),
            ("3. Servicio, Comunicación y Actitud", ["3.1 Atención y comunicación", "3.2 Trabajo en equipo", "3.3 Iniciativa y solución de problemas", "3.4 Confidencialidad y responsabilidad"]),
        ],
    },
    "documento_rrhh": ["carpeta_id"],
}
def es_personal_administrativo(colaborador):
    """Identifica personal administrativo a partir de los datos del padrón."""
    datos = " ".join((getattr(colaborador, campo, "") or "") for campo in ("area", "puesto", "cargo", "seccion"))
    normalizado = unicodedata.normalize("NFKD", datos).encode("ascii", "ignore").decode("ascii").lower()
    return bool(re.search(r"\b(adm|admin|administrativ|administracion|contabil|finanz|tesorer|secretari|recepcion|recursos humanos|rrhh|compras|gerencia|comercial|ventas|sistemas|informatica)\w*", normalizado))


def formulario_evaluacion_colaborador(colaborador):
    """Resuelve el formulario aplicable según antigüedad y área del colaborador."""
    estado = calcular_estado_colaborador(colaborador).lower()
    if estado == "prueba":
        return "prueba"
    return "administrativa" if es_personal_administrativo(colaborador) else "periodica"


class EntradaInsumo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    insumo_id = db.Column(db.Integer, db.ForeignKey("insumo.id"), nullable=False)
    cantidad = db.Column(db.Float, nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    observacion = db.Column(db.Text, default="")
    usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    insumo = db.relationship("Insumo", backref="entradas")
    usuario = db.relationship("User")


class SalidaInsumo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    insumo_id = db.Column(db.Integer, db.ForeignKey("insumo.id"), nullable=False)
    colaborador_id = db.Column(db.Integer, db.ForeignKey("colaborador.id"), nullable=False)
    cantidad = db.Column(db.Float, nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    observacion = db.Column(db.Text, default="")
    usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    retiro_id = db.Column(db.Integer, db.ForeignKey("retiro_insumo.id"), nullable=True)
    insumo = db.relationship("Insumo", backref="salidas")
    colaborador = db.relationship("Colaborador", backref="salidas")
    usuario = db.relationship("User")


class RetiroInsumo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(40), unique=True, nullable=False)
    colaborador_id = db.Column(db.Integer, db.ForeignKey("colaborador.id"), nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    observacion = db.Column(db.Text, default="")
    usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=False)
    colaborador = db.relationship("Colaborador", backref="retiros")
    usuario = db.relationship("User")
    salidas = db.relationship("SalidaInsumo", backref="retiro", order_by="SalidaInsumo.id", lazy=True)


class PedidoInsumo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(40), unique=True, nullable=True)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    estado = db.Column(db.String(20), nullable=False, default="Pendiente")
    estado_aprobacion = db.Column(db.String(25), nullable=False, default="Pendiente")
    evaluado_por_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    fecha_evaluacion = db.Column(db.DateTime, nullable=True)
    observacion_evaluacion = db.Column(db.Text, default="")
    presupuesto_elegido_id = db.Column(db.Integer, db.ForeignKey("presupuesto_pedido.id"), nullable=True)
    solicitante_usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    presupuestos_enviados = db.Column(db.Boolean, nullable=False, default=False)
    rango_monto = db.Column(db.String(20), nullable=True)
    solicitado_por = db.Column(db.String(120), nullable=False)
    observacion = db.Column(db.Text, default="")
    evaluado_por = db.relationship("User", foreign_keys=[evaluado_por_id])
    solicitante_usuario = db.relationship("User", foreign_keys=[solicitante_usuario_id])
    presupuesto_elegido = db.relationship("PresupuestoPedido", foreign_keys=[presupuesto_elegido_id], post_update=True)
    detalles = db.relationship("DetallePedidoInsumo", backref="pedido", cascade="all, delete-orphan")

    @property
    def total_estimado(self):
        return sum(detalle.cantidad * detalle.precio_unitario for detalle in self.detalles)

    @property
    def total_aprobado(self):
        return self.presupuesto_elegido.monto if self.presupuesto_elegido else self.total_estimado


class DetallePedidoInsumo(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    pedido_id = db.Column(db.Integer, db.ForeignKey("pedido_insumo.id"), nullable=False)
    insumo_id = db.Column(db.Integer, db.ForeignKey("insumo.id"), nullable=True)
    descripcion = db.Column(db.String(180), nullable=False, default="")
    cantidad = db.Column(db.Float, nullable=False)
    precio_unitario = db.Column(db.Float, nullable=False, default=0.0)
    insumo = db.relationship("Insumo")


class OrdenCompra(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    codigo = db.Column(db.String(40), unique=True, nullable=True)
    pedido_id = db.Column(db.Integer, db.ForeignKey("pedido_insumo.id"), nullable=False, unique=True)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    estado = db.Column(db.String(20), nullable=False, default="Abierta")
    fecha_cierre = db.Column(db.DateTime, nullable=True)
    observacion_cierre = db.Column(db.Text, default="")
    pedido = db.relationship("PedidoInsumo", backref=db.backref("orden_compra", uselist=False))


class PresupuestoPedido(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    pedido_id = db.Column(db.Integer, db.ForeignKey("pedido_insumo.id"), nullable=False)
    proveedor = db.Column(db.String(160), nullable=False)
    monto = db.Column(db.Float, nullable=False)
    archivo = db.Column(db.String(255), nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    pedido = db.relationship("PedidoInsumo", foreign_keys=[pedido_id], backref="presupuestos")


class HistorialAprobacionPedido(db.Model):
    id = db.Column(db.Integer, primary_key=True)
    pedido_id = db.Column(db.Integer, db.ForeignKey("pedido_insumo.id"), nullable=False)
    decision = db.Column(db.String(25), nullable=False)
    fecha = db.Column(db.DateTime, default=datetime.utcnow, nullable=False)
    usuario_id = db.Column(db.Integer, db.ForeignKey("user.id"), nullable=True)
    observacion = db.Column(db.Text, default="")
    presupuesto_id = db.Column(db.Integer, db.ForeignKey("presupuesto_pedido.id"), nullable=True)
    pedido = db.relationship("PedidoInsumo", backref="historial_aprobaciones")
    usuario = db.relationship("User", foreign_keys=[usuario_id])
    presupuesto = db.relationship("PresupuestoPedido")


@login_manager.user_loader
def load_user(user_id: str):
    return User.query.get(int(user_id))


def admin_required(view):
    @wraps(view)
    @login_required
    def protected_view(*args, **kwargs):
        if not (current_user.is_admin or current_user.is_root):
            abort(403)
        return view(*args, **kwargs)

    return protected_view


def calidad_required(view):
    @wraps(view)
    @login_required
    def protected_view(*args, **kwargs):
        if not current_user.can_access_quality:
            abort(403)
        return view(*args, **kwargs)

    return protected_view


def root_required(view):
    @wraps(view)
    @login_required
    def protected_view(*args, **kwargs):
        if not current_user.is_root:
            abort(403)
        return view(*args, **kwargs)

    return protected_view


def role_required(*roles):
    def decorator(view):
        @wraps(view)
        @login_required
        def protected_view(*args, **kwargs):
            normalized_role = current_user.role.strip().lower()
            if normalized_role not in {role.strip().lower() for role in roles}:
                abort(403)
            return view(*args, **kwargs)

        return protected_view

    return decorator


def offline_operation_key():
    return request.headers.get("X-Offline-Operation", "").strip()


def offline_operation_was_processed():
    key = offline_operation_key()
    return bool(key and OfflineOperation.query.filter_by(operation_key=key).first())


def remember_offline_operation():
    key = offline_operation_key()
    if not key:
        return
    db.session.add(OfflineOperation(operation_key=key, user_id=current_user.id))


def inventory_required(view):
    return role_required("admin", "deposito", "compras")(view)


def rrhh_required(view):
    @wraps(view)
    @login_required
    def protected_view(*args, **kwargs):
        if not (current_user.is_admin or current_user.is_rrhh):
            abort(403)
        return view(*args, **kwargs)

    return protected_view


def evaluacion_required(view):
    return role_required("admin", "rrhh", "supervisor", "root")(view)


def gerencia_required(view):
    return role_required("admin", "gerencia")(view)


def calidad_required(view):
    @wraps(view)
    @login_required
    def protected_view(*args, **kwargs):
        if not current_user.can_access_quality:
            abort(403)
        return view(*args, **kwargs)

    return protected_view


def set_config_value(clave: str, valor: str):
    config = Configuracion.query.filter_by(clave=clave).first()
    if config is None:
        config = Configuracion(clave=clave, valor=valor)
        db.session.add(config)
    else:
        config.valor = valor
    db.session.commit()


def get_config_value(clave: str, default: str = "0") -> str:
    config = Configuracion.query.filter_by(clave=clave).first()
    return config.valor if config else default


def get_quality_certificate_dirs() -> list[Path]:
    configured_dirs: list[Path] = []
    for raw_path in (
        os.environ.get("QUALITY_CERTIFICATES_LOCAL_DIR"),
        str(QUALITY_CERTIFICATES_LOCAL_DIR),
        str(QUALITY_CERTIFICATES_PROJECT_DIR),
        str(QUALITY_CERTIFICATES_REMOTE_DIR),
    ):
        if not raw_path:
            continue
        candidate = Path(raw_path)
        if candidate not in configured_dirs:
            configured_dirs.append(candidate)
    return configured_dirs


def ensure_quality_certificate_dirs() -> tuple[list[Path], str | None]:
    warnings: list[str] = []
    dirs: list[Path] = []
    for candidate in get_quality_certificate_dirs():
        is_unc_network_share = str(candidate).startswith("\\\\")
        if is_unc_network_share and not candidate.exists():
            continue
        try:
            candidate.mkdir(parents=True, exist_ok=True)
            dirs.append(candidate)
        except OSError as exc:
            if is_unc_network_share and not candidate.exists():
                continue
            warnings.append(f"No se pudo preparar la carpeta {candidate}: {exc}")
    if not dirs:
        return [], "No hay carpetas de respaldo de certificados disponibles."
    if warnings:
        return dirs, " ".join(warnings)
    return dirs, None


def list_quality_certificate_files(directories: list[Path] | list[str] | str | None = None) -> list[dict]:
    if directories is None:
        normalized_dirs = get_quality_certificate_dirs()
    elif isinstance(directories, (str, os.PathLike)):
        normalized_dirs = [Path(directories)]
    else:
        normalized_dirs = [Path(directory) for directory in directories]

    files: list[dict] = []
    seen: set[tuple[str, str]] = set()
    for directory in normalized_dirs:
        try:
            if not directory.exists():
                continue
            for item in sorted(directory.iterdir(), key=lambda entry: entry.name.lower()):
                if not item.is_file() or item.suffix.lower() != ".pdf":
                    continue
                signature = (item.name, str(item.resolve()))
                if signature in seen:
                    continue
                seen.add(signature)
                try:
                    file_time = datetime.fromtimestamp(item.stat().st_mtime)
                except OSError:
                    file_time = datetime.utcnow()
                files.append({
                    "name": item.name,
                    "path": str(item),
                    "uri": item.as_uri(),
                    "modified_at": file_time,
                })
        except (OSError, TypeError, ValueError):
            continue
    return sorted(files, key=lambda entry: entry["modified_at"], reverse=True)


@app.context_processor
def inject_dashboard_theme():
    if request.path.startswith("/configuracion/tema"):
        module = request.args.get("module", "combustible").strip().lower()
    elif request.path.startswith("/inventario"):
        module = "inventario"
    elif request.path.startswith("/despachos"):
        module = "despacho"
    else:
        module = "combustible"
    module_theme = get_config_value(f"dashboard_theme_{module}", get_config_value("dashboard_theme", "classic"))
    path = request.path or "/"
    rrhh_notifications = current_user.is_authenticated and (current_user.can_access_rrhh or current_user.is_supervisor or current_user.is_root) and (path.startswith("/rrhh") or path.startswith("/admin") or path.startswith("/supervisor") or current_user.is_rrhh)
    compras_notifications = current_user.is_authenticated and (current_user.is_compras or current_user.is_admin or current_user.is_root) and (path.startswith("/modulo/compras") or path.startswith("/compras") or path.startswith("/gerencia") or current_user.is_compras)
    inventory_notifications = current_user.is_authenticated and (current_user.is_deposito or current_user.is_admin or current_user.is_root) and path.startswith("/inventario")
    notifications_enabled = rrhh_notifications or compras_notifications or inventory_notifications
    unread_notifications = 0
    recent_notifications = []
    if notifications_enabled:
        query = Notificacion.query.filter_by(usuario_id=current_user.id, leida=False)
        if rrhh_notifications:
            if current_user.is_supervisor and not current_user.is_rrhh:
                query = query.filter(Notificacion.url.like("%asignacion_id=%"))
            else:
                query = query.filter(Notificacion.url.startswith("/rrhh") | Notificacion.url.startswith("/admin") | Notificacion.titulo.like("%Evaluar%") | Notificacion.titulo.like("%RRHH%"))
        elif compras_notifications:
            query = query.filter(Notificacion.url.startswith("/modulo/compras") | Notificacion.url.startswith("/compras") | Notificacion.url.startswith("/gerencia") | Notificacion.titulo.like("%Pedido%") | Notificacion.titulo.like("%Solicitud%") | Notificacion.titulo.like("%Compra%"))
        elif inventory_notifications:
            query = query.filter(Notificacion.url.startswith("/inventario") | Notificacion.titulo.like("%Pedido%") | Notificacion.titulo.like("%Inventario%"))
        unread_notifications = query.count()
        recent_notifications = Notificacion.query.filter_by(usuario_id=current_user.id).order_by(Notificacion.fecha.desc()).limit(10).all()
        if rrhh_notifications:
            if current_user.is_supervisor and not current_user.is_rrhh:
                recent_notifications = [n for n in recent_notifications if "asignacion_id=" in (n.url or "")]
            else:
                recent_notifications = [n for n in recent_notifications if n.url.startswith("/rrhh") or n.url.startswith("/admin") or "Evaluar" in n.titulo or "RRHH" in n.titulo]
        elif compras_notifications:
            recent_notifications = [n for n in recent_notifications if n.url.startswith("/modulo/compras") or n.url.startswith("/compras") or n.url.startswith("/gerencia") or "Pedido" in n.titulo or "Solicitud" in n.titulo or "Compra" in n.titulo]
        elif inventory_notifications:
            recent_notifications = [n for n in recent_notifications if n.url.startswith("/inventario") or "Pedido" in n.titulo or "Inventario" in n.titulo]
    return {
        "dashboard_theme": module_theme,
        "theme_module": module,
        "dashboard_palette": get_config_value("dashboard_palette", "blue"),
        "unread_notifications": unread_notifications,
        "recent_notifications": recent_notifications,
        "endpoint_exists": lambda endpoint: endpoint in app.view_functions,
    }


def get_stock_insumo(insumo_id: int) -> float:
    entradas = db.session.query(func.coalesce(func.sum(EntradaInsumo.cantidad), 0)).filter(EntradaInsumo.insumo_id == insumo_id).scalar() or 0
    salidas = db.session.query(func.coalesce(func.sum(SalidaInsumo.cantidad), 0)).filter(SalidaInsumo.insumo_id == insumo_id).scalar() or 0
    return float(entradas - salidas)


app.jinja_env.globals["get_stock_insumo"] = get_stock_insumo


def normalize_prefix(value: str) -> str:
    prefix = re.sub(r"[^A-Z0-9]", "", value.upper())
    return prefix[:10]


def next_insumo_code(categoria: CategoriaInsumo) -> str:
    prefix = normalize_prefix(categoria.prefijo or categoria.nombre) or "INS"
    used_numbers = []
    for codigo in db.session.query(Insumo.codigo).filter(Insumo.codigo.ilike(f"{prefix}%")).all():
        match = re.fullmatch(rf"{re.escape(prefix)}(\d+)", codigo[0] or "")
        if match:
            used_numbers.append(int(match.group(1)))
    return f"{prefix}{max(used_numbers, default=0) + 1:03d}"


def pedido_whatsapp_message(pedido: PedidoInsumo) -> str:
    lines = [
        f"Pedido de insumos {pedido.codigo or f'PED-{pedido.id:06d}'}",
        f"Solicitado por: {pedido.solicitado_por}",
        f"Fecha: {pedido.fecha.strftime('%d/%m/%Y %H:%M')}",
        "",
        "Detalle:",
    ]
    for detalle in pedido.detalles:
        lines.append(f"- {detalle.insumo.codigo} | {detalle.insumo.nombre}: {detalle.cantidad:g} {detalle.insumo.unidad}")
    if pedido.observacion:
        lines.extend(["", f"Observación: {pedido.observacion}"])
    return "\n".join(lines)


def pedido_whatsapp_url(pedido: PedidoInsumo) -> str:
    return f"https://wa.me/595983993332?text={quote_plus(pedido_whatsapp_message(pedido))}"


app.jinja_env.globals["pedido_whatsapp_url"] = pedido_whatsapp_url


WHATSAPP_PEDIDOS_GROUP_URL = "https://chat.whatsapp.com/GvqrVWxlMPm5OfLmNgXTS5?s=cl&p=a&ilr=1"
PEDIDO_ESTADOS = ("Pendiente", "En proceso", "Entregado", "Recibido en empresa", "No gestionado", "Cancelado")
DEPOSITO_PEDIDO_ESTADOS = ("Pendiente", "Entregado", "Cancelado")


def pedido_confirmation_message(pedido: PedidoInsumo) -> str:
    estado_icono = {
        "Entregado": "✅",
        "Recibido": "✅",
        "Recibido en empresa": "✅",
        "Pendiente": "⏳",
        "En proceso": "🚚",
        "Cancelado": "❌",
        "No gestionado": "⚠️",
    }.get(pedido.estado, "📦")
    lines = [
        f"📋 Aviso de pedido {pedido.codigo or f'PED-{pedido.id:06d}'}",
        f"{estado_icono} Estado actual: {pedido.estado}",
        f"👤 Solicitud a nombre de: {pedido.solicitado_por}",
        f"📅 Pedido realizado el: {pedido.fecha.strftime('%d/%m/%Y %H:%M')}",
        "",
        "📦 Producto(s):",
    ]
    for detalle in pedido.detalles:
        lines.append(f"- {detalle.insumo.nombre}: {detalle.cantidad:g} {detalle.insumo.unidad}")
    return "\n".join(lines)


def notificar_estado_pedido(pedido: PedidoInsumo):
    if pedido.solicitante_usuario_id:
        db.session.add(Notificacion(
            usuario_id=pedido.solicitante_usuario_id,
            titulo=f"Actualización del pedido {pedido.codigo}",
            mensaje=f"El estado de su pedido cambió a: {pedido.estado}.",
            url="/supervisor",
        ))


app.jinja_env.globals["pedido_confirmation_message"] = pedido_confirmation_message
app.jinja_env.globals["whatsapp_pedidos_group_url"] = WHATSAPP_PEDIDOS_GROUP_URL


def combustible_whatsapp_url(stock: float, alert_threshold: float) -> str:
    estado = "STOCK BAJO" if stock <= alert_threshold else "stock saludable, pero solicito reposición preventiva"
    message = "\n".join([
        "Solicitud de pedido de combustible",
        f"Estado: {estado}",
        f"Stock actual: {stock:,.0f} L",
        f"Umbral configurado: {alert_threshold:,.0f} L",
        f"Fecha: {datetime.utcnow():%d/%m/%Y %H:%M}",
    ])
    return f"https://wa.me/595983993332?text={quote_plus(message)}"


app.jinja_env.globals["combustible_whatsapp_url"] = combustible_whatsapp_url


def filtrar_reporte_existencias(categoria: str = "", producto: str = "", estado: str = ""):
    resumen = get_inventario_resumen()
    producto = producto.lower()
    return [
        item for item in resumen
        if (not categoria or (item["insumo"].categoria or "General") == categoria)
        and (not producto or producto in f"{item['insumo'].codigo or ''} {item['insumo'].nombre}".lower())
        and (not estado or estado == ("bajo" if item["bajo"] else "normal"))
    ]


def reporte_existencias_query_params():
    return {
        "categoria": request.args.get("categoria", "").strip(),
        "producto": request.args.get("producto", "").strip(),
        "estado": request.args.get("estado", "").strip().lower(),
    }


def reporte_existencias_whatsapp_url(rows, categoria: str = "", producto: str = "", estado: str = "") -> str:
    filtro = categoria or producto or estado or "Todos los productos"
    lines = [
        "Reporte de existencias de inventario",
        f"Filtro: {filtro}",
        f"Fecha: {datetime.utcnow():%d/%m/%Y %H:%M}",
        "",
    ]
    if not rows:
        lines.append("Sin existencias que coincidan con el filtro.")
    else:
        lines.append(f"Productos encontrados: {len(rows)}")
        for item in rows:
            insumo = item["insumo"]
            lines.append(f"- {insumo.codigo or '-'} | {insumo.nombre}: {item['stock']:.2f} {insumo.unidad} ({'Reponer' if item['bajo'] else 'Normal'})")
    return f"https://wa.me/595983993332?text={quote_plus(chr(10).join(lines))}"


app.jinja_env.globals["reporte_existencias_whatsapp_url"] = reporte_existencias_whatsapp_url


def get_inventario_resumen():
    insumos = Insumo.query.filter_by(activo=True).order_by(Insumo.nombre.asc()).all()
    return [{
        "insumo": insumo,
        "stock": get_stock_insumo(insumo.id),
        "bajo": get_stock_insumo(insumo.id) <= insumo.stock_minimo,
    } for insumo in insumos]


def get_deduccion_inventario():
    hoy = datetime.utcnow()
    deducciones = []
    for insumo in Insumo.query.filter_by(activo=True).order_by(Insumo.categoria.asc(), Insumo.nombre.asc()).all():
        stock = get_stock_insumo(insumo.id)
        primera_entrada = EntradaInsumo.query.filter_by(insumo_id=insumo.id).order_by(EntradaInsumo.fecha.asc()).first()
        total_salidas = db.session.query(func.coalesce(func.sum(SalidaInsumo.cantidad), 0)).filter(SalidaInsumo.insumo_id == insumo.id).scalar() or 0
        if not primera_entrada or total_salidas <= 0:
            deducciones.append({"insumo": insumo, "stock": stock, "consumo_diario": None, "dias_minimo": None, "dias_agotamiento": None, "fecha_minimo": None, "fecha_agotamiento": None})
            continue
        dias_observados = max((hoy - primera_entrada.fecha).total_seconds() / 86400, 1)
        consumo_diario = float(total_salidas) / dias_observados
        dias_minimo = max(stock - insumo.stock_minimo, 0) / consumo_diario
        dias_agotamiento = max(stock, 0) / consumo_diario
        deducciones.append({
            "insumo": insumo,
            "stock": stock,
            "consumo_diario": consumo_diario,
            "dias_minimo": dias_minimo,
            "dias_agotamiento": dias_agotamiento,
            "fecha_minimo": hoy + timedelta(days=dias_minimo),
            "fecha_agotamiento": hoy + timedelta(days=dias_agotamiento),
        })
    return deducciones


def get_initial_stock_value() -> float:
    return float(get_config_value("stock_inicial", "0"))


def get_stock_actual() -> float:
    stock_inicial = get_initial_stock_value()
    total_compra = db.session.query(func.coalesce(func.sum(Compra.litros), 0)).scalar() or 0
    total_despacho = db.session.query(func.coalesce(func.sum(Despacho.litros), 0)).scalar() or 0
    return float(stock_inicial + total_compra - total_despacho)


def get_total_consumo() -> float:
    return float(db.session.query(func.coalesce(func.sum(Despacho.litros), 0)).scalar() or 0)


def build_stock_history(days: int = 7):
    labels = []
    values = []
    today = datetime.utcnow().date()
    first_day = today - timedelta(days=days - 1)
    cumulative = get_initial_stock_value()
    cumulative += float(
        db.session.query(func.coalesce(func.sum(Compra.litros), 0))
        .filter(func.date(Compra.fecha) < first_day)
        .scalar() or 0
    )
    cumulative -= float(
        db.session.query(func.coalesce(func.sum(Despacho.litros), 0))
        .filter(func.date(Despacho.fecha) < first_day)
        .scalar() or 0
    )
    for offset in range(days - 1, -1, -1):
        day = today - timedelta(days=offset)
        labels.append(day.strftime("%d/%m"))
        compra_dia = db.session.query(func.coalesce(func.sum(Compra.litros), 0)).filter(func.date(Compra.fecha) == day).scalar() or 0
        despacho_dia = db.session.query(func.coalesce(func.sum(Despacho.litros), 0)).filter(func.date(Despacho.fecha) == day).scalar() or 0
        cumulative += float(compra_dia) - float(despacho_dia)
        values.append(max(float(cumulative), 0.0))
    return labels, values


def build_equipment_consumption():
    rows = db.session.query(
        Equipo.nombre,
        func.coalesce(func.sum(Despacho.litros), 0).label("total")
    ).outerjoin(Despacho, Equipo.id == Despacho.equipo_id).group_by(Equipo.id, Equipo.nombre).order_by(func.sum(Despacho.litros).desc())
    return [{"nombre": r[0], "total": float(r[1] or 0)} for r in rows]


def build_history_entries():
    entries = []
    for compra in Compra.query.order_by(Compra.fecha.desc()).all():
        fecha = compra.fecha or datetime.utcnow()
        entries.append({
            "tipo": "Compra",
            "fecha": fecha,
            "descripcion": f"Compra de {compra.litros} L - {compra.proveedor}",
            "litros": compra.litros,
            "id": compra.id,
            "route": None,
        })
    for despacho in Despacho.query.order_by(Despacho.fecha.desc()).all():
        fecha = despacho.fecha or datetime.utcnow()
        entries.append({
            "tipo": "Despacho",
            "fecha": fecha,
            "descripcion": f"{despacho.equipo.nombre} - {despacho.responsable}",
            "litros": despacho.litros,
            "id": despacho.id,
            "route": url_for("detalle_despacho", despacho_id=despacho.id),
        })
    entries.sort(key=lambda item: item["fecha"], reverse=True)
    return entries


def ensure_database_schema():
    db.create_all()
    inspector = db.inspect(db.engine)
    for table_name, columns in {
        "user": ["role", "correo", "telefono"],
        "insumo": ["codigo", "categoria", "vencimiento"],
        "categoria_insumo": ["prefijo"],
        "colaborador": ["codigo", "legajo", "estado", "fecha_ingreso", "puesto", "turno_linea", "calle", "barrio", "ciudad", "pais", "nacionalidad", "tipo_documento", "fecha_nacimiento", "estado_civil", "sexo", "cargo", "seccion", "correo", "supervisor_id", "vencimiento_cedula", "hijos_cantidad", "hijos", "aporte_ips", "bonificacion_familiar", "telefono", "foto_archivo"],
        "evento_rrhh": ["responsable_usuario_id", "notificar"],
        "pedido_insumo": ["codigo", "estado_aprobacion", "evaluado_por_id", "fecha_evaluacion", "observacion_evaluacion", "presupuesto_elegido_id", "solicitante_usuario_id", "presupuestos_enviados", "rango_monto"],
        "detalle_pedido_insumo": ["precio_unitario", "descripcion"],
        "orden_compra": ["estado", "fecha_cierre", "observacion_cierre"],
        "compra": ["precio_litro", "costo_total"],
        "despacho": ["nombre_chofer", "hora", "horometro", "kilometraje", "tipo_carga"],
        "salida_insumo": ["retiro_id"],
        "documento_rrhh": ["carpeta_id"],
        "evaluacion_personal": ["formulario", "codigo", "puntaje_total", "puntajes", "fecha_evaluacion", "periodo_evaluado", "tipo_evaluacion", "jefe_directo", "fortalezas", "aspectos_mejorar", "plan_accion", "fecha_revision", "recomendacion", "promedio", "resultado"],
        "etiqueta_calidad": ["fecha", "carbono", "silicio", "manganeso", "hornero", "supervisor", "operador_ccm", "aprobado", "tipo_b", "no_conforme"],
        "certificado_calidad": ["numero_colada", "numero_certificado", "tipo_producto", "carbono", "silicio", "manganeso", "medidas", "longitudes", "cantidad_palanquillas", "longitudes_palanquillas", "pesos_palanquillas", "datos_varillas", "productos_certificados", "alargamiento", "limite_fluencia", "limite_resistencia", "doblado_cumple", "fecha_atado", "fecha_certificacion", "grafico_path", "firma_path", "fecha_creacion", "created_by_id"],
        "despacho_acero": ["cliente", "cliente_id", "cedula", "chapa"],
    }.items():
        existing = {col["name"] for col in inspector.get_columns(table_name)}
        for column_name in columns:
            if column_name not in existing:
                if table_name == "user":
                    column_type = "VARCHAR(20) NOT NULL DEFAULT 'despacho'" if column_name == "role" else "VARCHAR(160) NULL" if column_name == "correo" else "VARCHAR(40) NULL"
                elif table_name == "insumo":
                    if column_name == "codigo":
                        column_type = "VARCHAR(40) NULL"
                    else:
                        column_type = "VARCHAR(80) NOT NULL DEFAULT 'General'" if column_name == "categoria" else "DATE NULL"
                elif table_name == "categoria_insumo" and column_name == "prefijo":
                    column_type = "VARCHAR(10) NULL"
                elif table_name == "colaborador" and column_name == "codigo":
                    column_type = "VARCHAR(40) NULL"
                elif table_name == "colaborador" and column_name == "legajo":
                    column_type = "VARCHAR(40) NULL"
                elif table_name == "colaborador" and column_name == "estado":
                    column_type = "VARCHAR(30) NOT NULL DEFAULT 'Contratado'"
                elif table_name == "colaborador" and column_name == "fecha_ingreso":
                    column_type = "DATE NULL"
                elif table_name == "colaborador" and column_name in {"puesto", "turno_linea"}:
                    column_type = "VARCHAR(120) NULL"
                elif table_name == "colaborador" and column_name in {"calle", "barrio", "ciudad", "cargo", "seccion"}:
                    column_type = "VARCHAR(120) NULL"
                elif table_name == "colaborador" and column_name in {"pais", "nacionalidad", "estado_civil", "sexo", "aporte_ips"}:
                    column_type = "VARCHAR(80) NULL"
                elif table_name == "colaborador" and column_name == "tipo_documento":
                    column_type = "VARCHAR(30) NULL"
                elif table_name == "colaborador" and column_name == "correo":
                    column_type = "VARCHAR(160) NULL"
                elif table_name == "colaborador" and column_name == "supervisor_id":
                    column_type = "INTEGER NULL"
                elif table_name == "colaborador" and column_name in {"fecha_nacimiento", "vencimiento_cedula"}:
                    column_type = "DATE NULL"
                elif table_name == "colaborador" and column_name == "hijos_cantidad":
                    column_type = "INTEGER NOT NULL DEFAULT 0"
                elif table_name == "colaborador" and column_name in {"hijos", "bonificacion_familiar"}:
                    column_type = "BOOLEAN NOT NULL DEFAULT 0"
                elif table_name == "evento_rrhh" and column_name == "responsable_usuario_id":
                    column_type = "INTEGER NULL"
                elif table_name == "evento_rrhh" and column_name == "notificar":
                    column_type = "BOOLEAN NOT NULL DEFAULT 0"
                elif table_name == "colaborador" and column_name == "telefono":
                    column_type = "VARCHAR(40) NULL"
                elif table_name == "colaborador" and column_name == "foto_archivo":
                    column_type = "VARCHAR(255) NULL"
                elif table_name == "pedido_insumo":
                    column_type = {"codigo": "VARCHAR(40) NULL", "estado_aprobacion": "VARCHAR(25) NOT NULL DEFAULT 'Pendiente'", "evaluado_por_id": "INTEGER NULL", "fecha_evaluacion": "DATETIME NULL", "observacion_evaluacion": "TEXT NULL", "presupuesto_elegido_id": "INTEGER NULL", "solicitante_usuario_id": "INTEGER NULL", "presupuestos_enviados": "BOOLEAN NOT NULL DEFAULT 0", "rango_monto": "VARCHAR(20) NULL"}[column_name]
                elif table_name == "detalle_pedido_insumo" and column_name == "precio_unitario":
                    column_type = "FLOAT NOT NULL DEFAULT 0"
                elif table_name == "detalle_pedido_insumo" and column_name == "descripcion":
                    column_type = "VARCHAR(180) NOT NULL DEFAULT ''"
                elif table_name == "orden_compra":
                    column_type = {"estado": "VARCHAR(20) NOT NULL DEFAULT 'Abierta'", "fecha_cierre": "DATETIME NULL", "observacion_cierre": "TEXT NULL"}[column_name]
                elif table_name == "salida_insumo" and column_name == "retiro_id":
                    column_type = "INTEGER NULL"
                elif table_name == "documento_rrhh" and column_name == "carpeta_id":
                    column_type = "INTEGER NULL"
                elif table_name == "evaluacion_personal":
                    column_type = {"formulario": "VARCHAR(20) NOT NULL DEFAULT 'periodica'", "codigo": "VARCHAR(20) NOT NULL DEFAULT 'RRHH-FOR-007'", "puntaje_total": "INTEGER NOT NULL DEFAULT 0", "puntajes": "TEXT NOT NULL DEFAULT '{}'", "fecha_evaluacion": "DATE NULL", "periodo_evaluado": "VARCHAR(120) NULL", "tipo_evaluacion": "VARCHAR(30) NULL", "jefe_directo": "VARCHAR(120) NULL", "fortalezas": "TEXT NULL", "aspectos_mejorar": "TEXT NULL", "plan_accion": "TEXT NULL", "fecha_revision": "DATE NULL", "recomendacion": "TEXT NULL", "promedio": "FLOAT NOT NULL DEFAULT 0", "resultado": "VARCHAR(100) NULL"}[column_name]
                elif table_name == "etiqueta_calidad":
                    column_type = {
                        "fecha": "DATE NULL",
                        "carbono": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "silicio": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "manganeso": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "hornero": "VARCHAR(120) NOT NULL DEFAULT ''",
                        "supervisor": "VARCHAR(120) NOT NULL DEFAULT ''",
                        "operador_ccm": "VARCHAR(120) NOT NULL DEFAULT ''",
                        "aprobado": "BOOLEAN NOT NULL DEFAULT 0",
                        "tipo_b": "BOOLEAN NOT NULL DEFAULT 0",
                        "no_conforme": "BOOLEAN NOT NULL DEFAULT 0",
                    }[column_name]
                elif table_name == "certificado_calidad":
                    column_type = {
                        "numero_colada": "VARCHAR(60) NULL",
                        "numero_certificado": "VARCHAR(30) NULL",
                        "tipo_producto": "VARCHAR(40) NOT NULL DEFAULT ''",
                        "carbono": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "silicio": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "manganeso": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "medidas": "VARCHAR(180) NOT NULL DEFAULT ''",
                        "longitudes": "VARCHAR(180) NOT NULL DEFAULT ''",
                        "cantidad_palanquillas": "INTEGER NOT NULL DEFAULT 0",
                        "longitudes_palanquillas": "VARCHAR(180) NOT NULL DEFAULT ''",
                        "pesos_palanquillas": "VARCHAR(180) NOT NULL DEFAULT ''",
                        "datos_varillas": "VARCHAR(300) NOT NULL DEFAULT ''",
                        "productos_certificados": "VARCHAR(500) NOT NULL DEFAULT ''",
                        "alargamiento": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "limite_fluencia": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "limite_resistencia": "VARCHAR(30) NOT NULL DEFAULT ''",
                        "doblado_cumple": "BOOLEAN NOT NULL DEFAULT 0",
                        "fecha_atado": "DATE NULL",
                        "fecha_certificacion": "DATE NULL",
                        "grafico_path": "VARCHAR(500) NOT NULL DEFAULT ''",
                        "firma_path": "VARCHAR(500) NOT NULL DEFAULT ''",
                        "fecha_creacion": "DATETIME NULL",
                        "created_by_id": "INTEGER NULL",
                    }[column_name]
                elif table_name == "despacho_acero":
                    column_type = {"cliente": "VARCHAR(180) NOT NULL DEFAULT ''", "cliente_id": "INTEGER NULL", "cedula": "VARCHAR(40) NOT NULL DEFAULT ''", "chapa": "VARCHAR(30) NOT NULL DEFAULT ''"}[column_name]
                else:
                    column_type = "FLOAT"
                db.session.execute(db.text(f"ALTER TABLE {table_name} ADD COLUMN {column_name} {column_type}"))

    db.session.execute(db.text("UPDATE user SET role = 'admin' WHERE username = 'admin'"))
    if "estado_aprobacion" in {column["name"] for column in db.inspect(db.engine).get_columns("pedido_insumo")} and not Configuracion.query.filter_by(clave="migracion_aprobacion_pedidos_20260828").first():
        db.session.execute(db.text("UPDATE pedido_insumo SET estado_aprobacion = 'Aprobado' WHERE estado_aprobacion = 'Pendiente' AND codigo IS NOT NULL"))
        db.session.add(Configuracion(clave="migracion_aprobacion_pedidos_20260828", valor="1"))
    db.session.commit()


def load_initial_inventory_catalog():
    marker = "catalogo_inicial_20260821"
    if Configuracion.query.filter_by(clave=marker).first():
        return
    catalog = [
        ("ESC001", "Atrapa Escoria", "Minerales", 7070, "kg"),
        ("BUZ002", "Buza Blanca 8mm", "Thundish", 262, "Ud"),
        ("BUZ001", "Buza Negra 9 mm", "Thundish", 114, "Ud"),
        ("BUZ004", "Buza Negra 9,5 mm", "Thundish", 360, "Ud"),
        ("TER001", "Cartucho para Pirometro", "Horno", 5400, "Ud"),
        ("A97001", "Compact Sol A97-15-BR", "Refractario", 13500, "kg"),
        ("B88001", "Compact Sol B88-5-BR", "Refractario", 4500, "kg"),
        ("A95001", "Comprit A95-6-BR", "Refractario", 3500, "kg"),
        ("CRO001", "Cromo", "Minerales", 500, "kg"),
        ("K85001", "Diram Dry K85 MA-4-AR", "Refractario", 12000, "kg"),
        ("K90001", "Diram Dry K90M-5-BR", "Refractario", 3250, "kg"),
        ("B72001", "Displast B72 CRP-5-BR", "Refractario", 875, "kg"),
        ("KOS001", "DivasIl Bom Plast (Kostrosol 3040)", "Fluidos Ligador", 1665, "L"),
        ("GRA001", "Drival Grafito", "Minerales", 14300, "kg"),
        ("682001", "Drive vive 682A", "Refractario", 4000, "kg"),
        ("493001", "DV 493A", "Refractario", 500, "kg"),
        ("MAN001", "FerroManganeso Mn", "Minerales", 14911, "kg"),
        ("MAN002", "FerroManganeso/Silicon75%", "Horno", 30000, "kg"),
        ("SIL002", "Ferrosilicon 75%", "Minerales", 9500, "kg"),
        ("FIB001", "Fibra Ceramica HPS-4-1", "Thundish", 5, "Ud"),
        ("IBA001", "IBAR VALDE", "Refractario", 851, "kg"),
        ("INS003", "Inserto", "Thundish", 300, "—"),
        ("INS002", "Inserto", "Thundish", 300, "—"),
        ("INS001", "Inserto", "Thundish", 300, "Ud"),
        ("MAX002", "Ladrillo Refractario", "Thundish", 230, "—"),
        ("MAX001", "Maxial Tejuelita", "Thundish", 370, "Ud"),
        ("COB002", "Molde de Cobre", "CCM", 13, "Ud"),
        ("COB001", "Molde de Cobre", "CCM", 17, "Ud"),
        ("MYM001", "MYM 82-E-HI", "Refractario", 4000, "—"),
        ("FIB002", "Placa Ceramica", "Thundish", 10, "Ud"),
        ("BUZ003", "Porta Busa", "Thundish", 300, "Ud"),
        ("RAS001", "Rasatund 02", "Refractario", 25075, "kg"),
        ("SIL001", "Silicato de Sodio", "Aglutinante Fundición", 1320, "L"),
        ("SIC001", "Silicato de Sodio (SiCa)", "Minerales", 8900, "kg"),
        ("MICA001", "Sol Term Papelmica", "Horno", 7, "Ud"),
        ("689001", "TCOA T689", "—", 200, "kg"),
    ]
    admin = User.query.filter_by(username="admin").first() or User.query.first()
    categories = {category.nombre: category for category in CategoriaInsumo.query.all()}
    for category_name in sorted({row[2] for row in catalog}):
        if category_name not in categories:
            category = CategoriaInsumo(nombre=category_name, prefijo="MAN" + str(len(categories) + 1))
            db.session.add(category)
            categories[category_name] = category
    db.session.flush()
    for codigo, nombre, categoria, existencia, unidad in catalog:
        insumo = Insumo.query.filter_by(codigo=codigo).first()
        if insumo is None:
            duplicated_names = {"Inserto", "Molde de Cobre"}
            stored_name = f"{nombre} ({codigo})" if nombre in duplicated_names else nombre
            insumo = Insumo(codigo=codigo, nombre=stored_name, categoria=categoria, unidad=unidad, stock_minimo=0)
            db.session.add(insumo)
            db.session.flush()
        if not EntradaInsumo.query.filter_by(insumo_id=insumo.id).first():
            db.session.add(EntradaInsumo(insumo_id=insumo.id, cantidad=existencia, observacion="Carga inicial de catálogo", usuario_id=admin.id))
    db.session.add(Configuracion(clave=marker, valor="1"))


def load_initial_general_outputs():
    marker = "salidas_generales_20260821"
    if Configuracion.query.filter_by(clave=marker).first():
        return
    outputs = {
        "ESC001": 3950, "BUZ002": 0, "BUZ001": 114, "BUZ004": 151, "TER001": 3167,
        "A97001": 5075, "B88001": 4478, "A95001": 3275, "CRO001": 0, "K85001": 8875,
        "K90001": 3000, "B72001": 875, "KOS001": 1337, "GRA001": 2116, "682001": 3625,
        "493001": 300, "MAN001": 14911, "MAN002": 765, "SIL002": 4665, "FIB001": 3,
        "IBA001": 762, "INS003": 0, "INS002": 3, "INS001": 10, "MAX002": 116,
        "MAX001": 325, "COB002": 6, "COB001": 12, "MYM001": 750, "FIB002": 5,
        "BUZ003": 13, "RAS001": 23645, "SIL001": 1189, "SIC001": 1833, "MICA001": 3,
        "689001": 75,
    }
    admin = User.query.filter_by(username="admin").first() or User.query.first()
    general = Colaborador.query.filter_by(codigo="COL-GENERAL").first()
    if general is None:
        general = Colaborador(codigo="COL-GENERAL", legajo="GENERAL", nombre="Carga histórica general", documento="", activo=True)
        db.session.add(general)
        db.session.flush()
    retiro = RetiroInsumo(
        codigo=f"TMP-{secrets.token_hex(8).upper()}",
        colaborador_id=general.id,
        observacion="Salidas históricas cargadas de forma general; pendientes de asignación por colaborador.",
        usuario_id=admin.id,
    )
    db.session.add(retiro)
    db.session.flush()
    retiro.codigo = f"RET-{retiro.fecha:%Y%m%d}-{retiro.id:06d}"
    for codigo, cantidad in outputs.items():
        if cantidad <= 0:
            continue
        insumo = Insumo.query.filter_by(codigo=codigo).first()
        if insumo is not None:
            db.session.add(SalidaInsumo(insumo_id=insumo.id, colaborador_id=general.id, cantidad=cantidad, observacion="Salida histórica general", usuario_id=admin.id, retiro_id=retiro.id))
    db.session.add(Configuracion(clave=marker, valor="1"))

    db.session.execute(db.text("UPDATE compra SET fecha = CURRENT_TIMESTAMP WHERE fecha IS NULL"))
    db.session.execute(db.text("UPDATE despacho SET fecha = CURRENT_TIMESTAMP WHERE fecha IS NULL"))
    db.session.execute(db.text("UPDATE compra SET precio_litro = 0 WHERE precio_litro IS NULL"))
    db.session.execute(db.text("UPDATE compra SET costo_total = 0 WHERE costo_total IS NULL"))
    db.session.execute(db.text("UPDATE compra SET litros = 0 WHERE litros IS NULL"))
    db.session.execute(db.text("UPDATE despacho SET hora = 0 WHERE hora IS NULL"))
    db.session.execute(db.text("UPDATE despacho SET horometro = 0 WHERE horometro IS NULL"))
    db.session.execute(db.text("UPDATE despacho SET kilometraje = 0 WHERE kilometraje IS NULL"))
    db.session.execute(db.text("UPDATE despacho SET litros = 0 WHERE litros IS NULL"))
    for categoria in CategoriaInsumo.query.all():
        if categoria.prefijo is None or (categoria.nombre.strip().lower() == "general" and categoria.prefijo == "GENERAL"):
            categoria.prefijo = "GEN" if categoria.nombre.strip().lower() == "general" else normalize_prefix(categoria.nombre) or "INS"
    for insumo in Insumo.query.filter(Insumo.codigo.is_(None)).order_by(Insumo.id.asc()).all():
        categoria = CategoriaInsumo.query.filter_by(nombre=insumo.categoria).first()
        insumo.codigo = next_insumo_code(categoria) if categoria else f"INS{insumo.id:03d}"
    for colaborador in Colaborador.query.filter(Colaborador.codigo.is_(None)).order_by(Colaborador.id.asc()).all():
        colaborador.codigo = f"COL-{colaborador.id:06d}"
    for colaborador in Colaborador.query.filter(Colaborador.legajo.is_(None)).order_by(Colaborador.id.asc()).all():
        colaborador.legajo = colaborador.codigo or f"LEG-{colaborador.id:06d}"
    for pedido in PedidoInsumo.query.filter(PedidoInsumo.codigo.is_(None)).order_by(PedidoInsumo.id.asc()).all():
        pedido.codigo = f"PED-{pedido.id:06d}"
    db.session.commit()


def get_average_purchase_price() -> float:
    total_litros = db.session.query(func.coalesce(func.sum(Compra.litros), 0)).scalar() or 0
    if total_litros <= 0:
        return 0.0
    total_cost = db.session.query(func.coalesce(func.sum(Compra.costo_total), 0)).scalar() or 0
    return float(total_cost / total_litros)


def get_date_range(period: str = "all", fecha_desde: str | None = None, fecha_hasta: str | None = None):
    if fecha_desde or fecha_hasta:
        start = None
        end = None
        if fecha_desde:
            try:
                start = datetime.strptime(fecha_desde, "%Y-%m-%d")
            except ValueError:
                start = None
        if fecha_hasta:
            try:
                end = datetime.strptime(fecha_hasta, "%Y-%m-%d")
            except ValueError:
                end = None
        if start is not None and end is not None and start > end:
            start, end = end, start
        return start, end

    if period == "dia":
        today = datetime.utcnow().date()
        start = datetime.combine(today, datetime.min.time())
        end = datetime.combine(today, datetime.max.time())
        return start, end
    if period == "semana":
        start = datetime.utcnow() - timedelta(days=6)
        start = start.replace(hour=0, minute=0, second=0, microsecond=0)
        end = datetime.utcnow().replace(hour=23, minute=59, second=59, microsecond=999999)
        return start, end
    if period == "mes":
        start = datetime.utcnow() - timedelta(days=29)
        start = start.replace(hour=0, minute=0, second=0, microsecond=0)
        end = datetime.utcnow().replace(hour=23, minute=59, second=59, microsecond=999999)
        return start, end
    return None, None


def get_machine_report(period: str = "all", equipo_id: int | None = None, fecha_desde: str | None = None, fecha_hasta: str | None = None):
    purchase_price = get_average_purchase_price()
    equipos = Equipo.query.order_by(Equipo.nombre.asc()).all()
    if equipo_id is not None:
        equipos = [equipo for equipo in equipos if equipo.id == equipo_id]
    start_dt, end_dt = get_date_range(period, fecha_desde, fecha_hasta)
    report = []
    for equipo in equipos:
        query = Despacho.query.filter_by(equipo_id=equipo.id)
        if start_dt is not None:
            query = query.filter(Despacho.fecha >= start_dt)
        if end_dt is not None:
            query = query.filter(Despacho.fecha <= end_dt)
        registros = query.order_by(Despacho.fecha.asc()).all()
        total_litros = sum(float(item.litros or 0) for item in registros)
        total_cost = total_litros * purchase_price
        total_hours = 0.0
        if len(registros) > 1:
            for actual, siguiente in zip(registros, registros[1:]):
                if actual.horometro and siguiente.horometro and siguiente.horometro >= actual.horometro:
                    total_hours += max(0.0, float(siguiente.horometro or 0) - float(actual.horometro or 0))
        if registros and registros[-1].horometro and registros[0].horometro:
            total_hours = max(total_hours, max(0.0, float(registros[-1].horometro or 0) - float(registros[0].horometro or 0)))
        consumo_hora = total_litros / total_hours if total_hours > 0 else 0
        consumo_promedio = total_litros / len(registros) if registros else 0
        costo_hora = total_cost / total_hours if total_hours > 0 else 0
        costo_promedio = total_cost / len(registros) if registros else 0
        report.append({
            "equipo": equipo,
            "total_litros": total_litros,
            "total_cost": total_cost,
            "consumo_hora": consumo_hora,
            "consumo_promedio": consumo_promedio,
            "costo_hora": costo_hora,
            "costo_promedio": costo_promedio,
            "kilometraje_total": sum(float(item.kilometraje or 0) for item in registros),
            "registros": len(registros),
            "precio_litro_actual": purchase_price,
        })
    return report


def get_machine_chart_data(period: str = "all", equipo_id: int | None = None, fecha_desde: str | None = None, fecha_hasta: str | None = None):
    rows = get_machine_report(period, equipo_id, fecha_desde, fecha_hasta)
    return {
        "labels": [item["equipo"].nombre for item in rows],
        "litros": [float(item["total_litros"]) for item in rows],
        "costos": [float(item["total_cost"]) for item in rows],
    }


def build_machine_report_csv(period: str = "all", equipo_id: int | None = None, fecha_desde: str | None = None, fecha_hasta: str | None = None) -> bytes:
    rows = get_machine_report(period, equipo_id, fecha_desde, fecha_hasta)
    output = BytesIO()
    writer = csv.writer(output, delimiter=';')
    writer.writerow(["Maquina", "Consumo total (L)", "Consumo por hora (L/h)", "Promedio por despacho (L)", "Costo estimado (Gs)", "Costo por hora (Gs/h)", "Kilometraje (km)"])
    for item in rows:
        writer.writerow([
            item["equipo"].nombre,
            f"{item['total_litros']:.2f}",
            f"{item['consumo_hora']:.2f}",
            f"{item['consumo_promedio']:.2f}",
            f"{item['total_cost']:.0f}",
            f"{item['costo_hora']:.0f}",
            f"{item['kilometraje_total']:.2f}",
        ])
    return output.getvalue()


def build_machine_report_pdf(period: str = "all", equipo_id: int | None = None, fecha_desde: str | None = None, fecha_hasta: str | None = None) -> bytes:
    ensure_logo_exists()
    rows = get_machine_report(period, equipo_id, fecha_desde, fecha_hasta)
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=30,
        leftMargin=30,
        topMargin=30,
        bottomMargin=30,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontSize=16, leading=20, alignment=1, spaceAfter=12)
    normal_style = ParagraphStyle("normal", parent=styles["BodyText"], fontSize=9, leading=11)
    data = [["Máquina", "Consumo total", "Consumo/h", "Promedio", "Costo", "Costo/h", "Km"]]
    for item in rows:
        data.append([
            item["equipo"].nombre,
            f"{item['total_litros']:,.2f} L",
            f"{item['consumo_hora']:,.2f} L/h",
            f"{item['consumo_promedio']:,.2f} L",
            f"{item['total_cost']:,.0f} Gs",
            f"{item['costo_hora']:,.0f} Gs/h",
            f"{item['kilometraje_total']:,.0f} km",
        ])
    table = Table(data, repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8E6FF")),
        ("GRID", (0, 0), (-1, -1), 1, colors.grey),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("PADDING", (4, 4), (-1, -1), 5),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    logo = Image(str(STATIC_LOGO_PATH), width=140, height=60)
    header = Table([[logo, Paragraph("REPORTE DE CONSUMO Y COSTO", title_style)]], colWidths=[180, 350])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    elements = [header, Spacer(1, 12), table]
    if equipo_id is not None:
        equipo = Equipo.query.get(equipo_id)
        if equipo:
            elements = [header, Spacer(1, 12), Paragraph(f"Equipo: {equipo.nombre}", normal_style), Paragraph(f"Tipo: {equipo.tipo}", normal_style), Paragraph(f"Responsable: {equipo.responsable}", normal_style), Spacer(1, 8), table]
    doc.build(elements)
    return buffer.getvalue()


def build_fuel_history_pdf(period: str = "all", fecha_desde: str | None = None, fecha_hasta: str | None = None) -> bytes:
    ensure_logo_exists()
    start_dt, end_dt = get_date_range(period, fecha_desde, fecha_hasta)
    query = Despacho.query.join(Equipo).order_by(Equipo.nombre.asc(), Despacho.fecha.asc(), Despacho.id.asc())
    if start_dt is not None:
        query = query.filter(Despacho.fecha >= start_dt)
    if end_dt is not None:
        query = query.filter(Despacho.fecha <= end_dt)
    despachos = query.all()

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=18, leftMargin=18, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("fuel_history_title", parent=styles["Title"], fontSize=16, leading=19, alignment=1, spaceAfter=4)
    subtitle_style = ParagraphStyle("fuel_history_subtitle", parent=styles["BodyText"], fontSize=8, leading=10, alignment=1, textColor=colors.HexColor("#555555"), spaceAfter=10)
    section_style = ParagraphStyle("fuel_history_section", parent=styles["Heading2"], fontSize=11, leading=13, textColor=colors.HexColor("#123B63"), spaceBefore=8, spaceAfter=5)
    cell_style = ParagraphStyle("fuel_history_cell", parent=styles["BodyText"], fontSize=7, leading=8)
    header_style = ParagraphStyle("fuel_history_header", parent=cell_style, textColor=colors.white, alignment=1)

    period_label = "Total"
    if period == "dia":
        period_label = "Hoy"
    elif period == "semana":
        period_label = "Últimos 7 días"
    elif period == "mes":
        period_label = "Últimos 30 días"
    elif fecha_desde or fecha_hasta:
        period_label = f"{fecha_desde or 'inicio'} al {fecha_hasta or 'hoy'}"

    logo = Image(str(STATIC_LOGO_PATH), width=105, height=45)
    header = Table([[logo, Paragraph("HISTÓRICO DE CARGA DE COMBUSTIBLE", title_style)]], colWidths=[130, 630])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    elements = [header, Paragraph(f"Despachos por máquina | Período: {escape(period_label)} | Generado: {datetime.now().strftime('%d/%m/%Y %H:%M')}", subtitle_style)]

    if not despachos:
        elements.append(Paragraph("No hay despachos registrados para el período seleccionado.", styles["BodyText"]))
    else:
        grouped = {}
        for despacho in despachos:
            grouped.setdefault(despacho.equipo, []).append(despacho)
        for equipo, registros in grouped.items():
            total_litros = sum(float(registro.litros or 0) for registro in registros)
            elements.append(Paragraph(f"{escape(equipo.nombre)} - {escape(equipo.tipo or 'Sin tipo')} | {len(registros)} despacho(s) | Total: {total_litros:,.2f} L", section_style))
            data = [[
                Paragraph("Fecha", header_style), Paragraph("Litros", header_style), Paragraph("Responsable", header_style),
                Paragraph("Chofer", header_style), Paragraph("Hora", header_style), Paragraph("Horómetro", header_style),
                Paragraph("Kilometraje", header_style), Paragraph("Tipo carga", header_style), Paragraph("Observación", header_style),
            ]]
            for registro in registros:
                data.append([
                    Paragraph(registro.fecha.strftime("%d/%m/%Y %H:%M") if registro.fecha else "Sin fecha", cell_style),
                    Paragraph(f"{float(registro.litros or 0):,.2f} L", cell_style),
                    Paragraph(escape(registro.responsable or "-"), cell_style),
                    Paragraph(escape(registro.nombre_chofer or "-"), cell_style),
                    Paragraph(f"{float(registro.hora or 0):,.2f} hs", cell_style),
                    Paragraph(f"{float(registro.horometro or 0):,.2f} hs", cell_style),
                    Paragraph(f"{float(registro.kilometraje or 0):,.2f} km", cell_style),
                    Paragraph(escape(registro.tipo_carga or "Media"), cell_style),
                    Paragraph(escape(registro.observacion or "-"), cell_style),
                ])
            table = Table(data, repeatRows=1, colWidths=[65, 52, 78, 78, 52, 62, 68, 62, 235])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F5A85")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#B8C4CE")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ALIGN", (1, 1), (1, -1), "RIGHT"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F7FA")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 4),
                ("RIGHTPADDING", (0, 0), (-1, -1), 4),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            elements.append(table)
            elements.append(Spacer(1, 8))

    doc.build(elements)
    return buffer.getvalue()


def build_fuel_history_excel(period: str = "all", fecha_desde: str | None = None, fecha_hasta: str | None = None) -> bytes:
    start_dt, end_dt = get_date_range(period, fecha_desde, fecha_hasta)
    query = Despacho.query.join(Equipo).order_by(Equipo.nombre.asc(), Despacho.fecha.asc(), Despacho.id.asc())
    if start_dt is not None:
        query = query.filter(Despacho.fecha >= start_dt)
    if end_dt is not None:
        query = query.filter(Despacho.fecha <= end_dt)
    despachos = query.all()

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen por máquina"
    summary.append(["Histórico de carga de combustible"])
    summary.append(["Período", period])
    summary.append(["Desde", fecha_desde or ""])
    summary.append(["Hasta", fecha_hasta or ""])
    summary.append([])
    summary.append(["Máquina", "Tipo", "Despachos", "Litros totales"])
    for cell in summary[6]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5A85")

    grouped = {}
    for despacho in despachos:
        grouped.setdefault(despacho.equipo, []).append(despacho)
    for equipo, registros in grouped.items():
        summary.append([equipo.nombre, equipo.tipo or "", len(registros), sum(float(item.litros or 0) for item in registros)])
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 20
    summary.column_dimensions["C"].width = 14
    summary.column_dimensions["D"].width = 18
    summary.freeze_panes = "A7"

    detail = workbook.create_sheet("Detalle histórico")
    detail.append(["Máquina", "Tipo", "Fecha", "Litros", "Responsable", "Chofer", "Hora", "Horómetro", "Kilometraje", "Tipo de carga", "Observación"])
    for cell in detail[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F5A85")
    for equipo, registros in grouped.items():
        for registro in registros:
            detail.append([
                equipo.nombre,
                equipo.tipo or "",
                registro.fecha,
                float(registro.litros or 0),
                registro.responsable or "",
                registro.nombre_chofer or "",
                float(registro.hora or 0),
                float(registro.horometro or 0),
                float(registro.kilometraje or 0),
                registro.tipo_carga or "Media",
                registro.observacion or "",
            ])
    widths = [25, 18, 20, 12, 22, 22, 12, 14, 15, 16, 45]
    for index, width in enumerate(widths, start=1):
        detail.column_dimensions[chr(64 + index)].width = width
    for row in detail.iter_rows(min_row=2, min_col=3, max_col=3):
        row[0].number_format = "dd/mm/yyyy hh:mm"
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_compra_pdf(compra: Compra) -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=32, leftMargin=32, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontSize=18, leading=20, alignment=1, spaceAfter=12)
    normal_style = ParagraphStyle("normal", parent=styles["BodyText"], fontSize=10, leading=14)
    logo = Image(str(STATIC_LOGO_PATH), width=140, height=60)
    header = Table([[logo, Paragraph("COMPROBANTE DE COMPRA", title_style)]], colWidths=[180, 350])
    table = Table([
        [Paragraph("Proveedor:", normal_style), Paragraph(compra.proveedor or "No informado", normal_style)],
        [Paragraph("Litros:", normal_style), Paragraph(f"{float(compra.litros or 0):.2f} L", normal_style)],
        [Paragraph("Precio/L:", normal_style), Paragraph(f"{float(compra.precio_litro or 0):,.0f} Gs", normal_style)],
        [Paragraph("Costo total:", normal_style), Paragraph(f"{float(compra.costo_total or 0):,.0f} Gs", normal_style)],
        [Paragraph("Fecha:", normal_style), Paragraph(compra.fecha.strftime("%d/%m/%Y %H:%M") if compra.fecha else "Sin fecha", normal_style)],
        [Paragraph("Notas:", normal_style), Paragraph(compra.notas or "Sin notas", normal_style)],
    ], colWidths=[150, 350])
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 1, colors.grey),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAEAEA")),
        ("PADDING", (6, 6), (-1, -1), 7),
    ]))
    doc.build([header, Spacer(1, 12), table])
    return buffer.getvalue()


def ensure_logo_exists():
    if STATIC_LOGO_PATH.exists():
        return
    if LOGO_SOURCE_PATH.exists():
        STATIC_LOGO_PATH.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(LOGO_SOURCE_PATH, STATIC_LOGO_PATH)
        return


def build_despacho_pdf(despacho: Despacho) -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=32,
        leftMargin=32,
        topMargin=32,
        bottomMargin=32,
    )
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("title", parent=styles["Title"], fontSize=18, leading=22, alignment=1, spaceAfter=12)
    subtitle_style = ParagraphStyle("subtitle", parent=styles["Heading2"], fontSize=11, textColor=colors.black, leading=14, spaceAfter=10)
    normal_style = ParagraphStyle("normal", parent=styles["BodyText"], fontSize=10, leading=15)

    logo = Image(str(STATIC_LOGO_PATH), width=150, height=65)
    header = Table(
        [[logo, Paragraph("COMPROBANTE DE DESPACHO", title_style)]],
        colWidths=[180, 350],
        rowHeights=[80],
    )
    header.setStyle(
        TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 0),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 12),
        ])
    )

    details = [
        [Paragraph("Equipo:", normal_style), Paragraph(despacho.equipo.nombre, normal_style)],
        [Paragraph("Tipo:", normal_style), Paragraph(despacho.equipo.tipo, normal_style)],
        [Paragraph("Responsable:", normal_style), Paragraph(despacho.responsable, normal_style)],
        [Paragraph("Chofer:", normal_style), Paragraph(despacho.nombre_chofer or "No registrado", normal_style)],
        [Paragraph("Hora:", normal_style), Paragraph(f"{despacho.hora or 0:.2f} hs", normal_style)],
        [Paragraph("Horómetro:", normal_style), Paragraph(f"{despacho.horometro or 0:.2f} hs", normal_style)],
        [Paragraph("Kilometraje:", normal_style), Paragraph(f"{despacho.kilometraje or 0:.2f} km", normal_style)],
        [Paragraph("Tipo de carga:", normal_style), Paragraph(despacho.tipo_carga or "Media", normal_style)],
        [Paragraph("Fecha:", normal_style), Paragraph(despacho.fecha.strftime("%d/%m/%Y %H:%M"), normal_style)],
        [Paragraph("Litros entregados:", normal_style), Paragraph(f"{despacho.litros:.2f} L", normal_style)],
        [Paragraph("Observación:", normal_style), Paragraph(despacho.observacion or "Sin observaciones", normal_style)],
    ]
    table = Table(details, colWidths=[170, 320])
    table.setStyle(
        TableStyle([
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EAEAEA")),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("PADDING", (6, 6), (-1, -1), 8),
        ])
    )

    elements = [header, Spacer(1, 15), Paragraph("Detalle del despacho", subtitle_style), table]
    doc.build(elements)
    return buffer.getvalue()


def build_retiro_pdf(retiro: RetiroInsumo) -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=32, leftMargin=32, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("retiro_title", parent=styles["Title"], fontSize=17, leading=21, alignment=1, spaceAfter=12)
    normal_style = ParagraphStyle("retiro_normal", parent=styles["BodyText"], fontSize=9, leading=12)
    logo = Image(str(STATIC_LOGO_PATH), width=140, height=60)
    header = Table([[logo, Paragraph("COMPROBANTE DE RETIRO DE INSUMOS", title_style)]], colWidths=[180, 350])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    details = Table([
        [Paragraph("Código de retiro:", normal_style), Paragraph(retiro.codigo, normal_style)],
        [Paragraph("Colaborador:", normal_style), Paragraph(f"Legajo {retiro.colaborador.legajo} · {retiro.colaborador.nombre}", normal_style)],
        [Paragraph("Fecha:", normal_style), Paragraph(retiro.fecha.strftime("%d/%m/%Y %H:%M"), normal_style)],
        [Paragraph("Observación:", normal_style), Paragraph(retiro.observacion or "Sin observaciones", normal_style)],
    ], colWidths=[150, 350])
    details.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 1, colors.grey), ("PADDING", (0, 0), (-1, -1), 6)]))
    rows = [["Código", "Producto", "Categoría", "Cantidad", "Unidad"]]
    for salida in retiro.salidas:
        rows.append([salida.insumo.codigo, salida.insumo.nombre, salida.insumo.categoria, f"{salida.cantidad:.2f}", salida.insumo.unidad])
    products = Table(rows, repeatRows=1, colWidths=[75, 175, 110, 75, 65])
    products.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8E6FF")),
        ("GRID", (0, 0), (-1, -1), 1, colors.grey),
        ("ALIGN", (3, 1), (-1, -1), "CENTER"),
        ("PADDING", (0, 0), (-1, -1), 6),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
    ]))
    signature = Paragraph("<br/><br/>______________________________<br/>Firma del colaborador", normal_style)
    doc.build([header, Spacer(1, 12), details, Spacer(1, 14), Paragraph("Productos retirados", styles["Heading2"]), products, Spacer(1, 35), signature])
    return buffer.getvalue()


def build_evaluacion_rrhh_pdf(evaluacion: EvaluacionPersonal) -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("rrhh_title", parent=styles["Title"], fontSize=16, leading=20, alignment=1)
    normal = ParagraphStyle("rrhh_normal", parent=styles["BodyText"], fontSize=9, leading=12)
    small = ParagraphStyle("rrhh_small", parent=styles["BodyText"], fontSize=8, leading=10)
    logo = Image(str(STATIC_LOGO_PATH), width=125, height=54)
    colaborador = evaluacion.colaborador
    photo_path = (RRHH_PHOTOS_DIR / colaborador.foto_archivo).resolve() if colaborador.foto_archivo else None
    photo = Image(str(photo_path), width=72, height=86) if photo_path and photo_path.parent == RRHH_PHOTOS_DIR.resolve() and photo_path.is_file() else Paragraph("Sin foto", small)
    header_title = f"{escape(evaluacion.codigo)}<br/>{escape(EVALUATION_FORMS[evaluacion.formulario]['nombre'])}<br/><font size=8>Vigencia 30-07-2026 | Revision 00</font>"
    header = Table([[logo, Paragraph(header_title, title), photo]], colWidths=[145, 315, 70])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    info = Table([
        [Paragraph("Colaborador", normal), Paragraph(escape(colaborador.nombre), normal), Paragraph("Legajo", normal), Paragraph(escape(colaborador.legajo or "-"), normal)],
        [Paragraph("Puesto", normal), Paragraph(escape(colaborador.puesto or "-"), normal), Paragraph("Área", normal), Paragraph(escape(colaborador.area or "-"), normal)],
        [Paragraph("Fecha de ingreso", normal), Paragraph(colaborador.fecha_ingreso.strftime("%d/%m/%Y") if colaborador.fecha_ingreso else "-", normal), Paragraph("Estado", normal), Paragraph(escape(colaborador.estado), normal)],
        [Paragraph("Evaluador", normal), Paragraph(escape(evaluacion.jefe_directo or "-"), normal), Paragraph("Fecha evaluación", normal), Paragraph(evaluacion.fecha_evaluacion.strftime("%d/%m/%Y") if evaluacion.fecha_evaluacion else "-", normal)],
    ], colWidths=[90, 190, 90, 160])
    info.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF2FF")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EAF2FF")), ("PADDING", (0, 0), (-1, -1), 6)]))
    scores = json.loads(evaluacion.puntajes or "{}")
    rows = [["Criterio", "Puntaje"]]
    for group_name, criteria in EVALUATION_FORMS[evaluacion.formulario]["groups"]:
        rows.append([Paragraph(f"<b>{escape(group_name)}</b>", small), ""])
        rows.extend([[Paragraph(escape(criterion), small), str(scores.get(criterion, "-"))] for criterion in criteria])
    criteria_table = Table(rows, colWidths=[450, 80], repeatRows=1)
    criteria_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8E6FF")), ("BACKGROUND", (0, 1), (-1, -1), colors.white), ("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("SPAN", (0, 1), (-1, 1)), ("ALIGN", (1, 1), (1, -1), "CENTER"), ("PADDING", (0, 0), (-1, -1), 5)]))
    summary = Table([[Paragraph("Puntaje total", normal), str(evaluacion.puntaje_total), Paragraph("Promedio", normal), f"{evaluacion.promedio:.2f} / 5"], [Paragraph("Resultado", normal), Paragraph(escape(evaluacion.resultado or "-"), normal), Paragraph("Estado", normal), Paragraph(escape(evaluacion.estado), normal)]], colWidths=[100, 70, 90, 270])
    summary.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF2FF")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EAF2FF")), ("PADDING", (0, 0), (-1, -1), 6)]))
    notes = Table([[Paragraph("Fortalezas", normal), Paragraph(escape(evaluacion.fortalezas or "-"), normal)], [Paragraph("Aspectos a mejorar", normal), Paragraph(escape(evaluacion.aspectos_mejorar or "-"), normal)], [Paragraph("Plan de acción", normal), Paragraph(escape(evaluacion.plan_accion or "-"), normal)], [Paragraph("Recomendación", normal), Paragraph(escape(evaluacion.recomendacion or "-"), normal)], [Paragraph("Observaciones", normal), Paragraph(escape(evaluacion.observaciones or "-"), normal)]], colWidths=[130, 400])
    notes.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF2FF")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 6)]))
    doc.build([header, Spacer(1, 12), info, Spacer(1, 12), criteria_table, Spacer(1, 12), summary, Spacer(1, 12), notes, Spacer(1, 28), Paragraph("Firma del evaluador: ________________________________     Firma del colaborador: ________________________________", normal)])
    return buffer.getvalue()


def build_evaluacion_rrhh_template_pdf(formulario: str) -> bytes:
    form_config = EVALUATION_FORMS[formulario]
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("rrhh_template_title", parent=styles["Title"], fontSize=16, leading=19, alignment=1)
    normal = ParagraphStyle("rrhh_template_normal", parent=styles["BodyText"], fontSize=9, leading=11)
    small = ParagraphStyle("rrhh_template_small", parent=styles["BodyText"], fontSize=8, leading=10)
    logo = Image(str(STATIC_LOGO_PATH), width=125, height=54)
    header_title = f"{escape(form_config['codigo'])}<br/>{escape(form_config['nombre'])}<br/><font size=8>Vigencia 30-07-2026 | Revision 00</font>"
    header = Table([[logo, Paragraph(header_title, title), Paragraph("Foto", small)]], colWidths=[145, 315, 70], rowHeights=[86])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("BOX", (2, 0), (2, 0), 0.5, colors.grey), ("ALIGN", (2, 0), (2, 0), "CENTER")]))
    info = Table([
        [Paragraph("Colaborador", normal), "", Paragraph("Legajo", normal), ""],
        [Paragraph("Puesto", normal), "", Paragraph("Área", normal), ""],
        [Paragraph("Fecha de ingreso", normal), "", Paragraph("Estado", normal), ""],
        [Paragraph("Evaluador", normal), "", Paragraph("Fecha evaluación", normal), ""],
    ], colWidths=[90, 190, 90, 160], rowHeights=[25] * 4)
    info.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF2FF")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EAF2FF")), ("PADDING", (0, 0), (-1, -1), 6)]))
    rows = [["Criterio", "Puntaje (1 a 5)"]]
    for group_name, criteria in form_config["groups"]:
        rows.append([Paragraph(f"<b>{escape(group_name)}</b>", small), ""])
        rows.extend([[Paragraph(escape(criterion), small), ""] for criterion in criteria])
    criteria_table = Table(rows, colWidths=[450, 80], repeatRows=1)
    criteria_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8E6FF")), ("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("SPAN", (0, 1), (-1, 1)), ("ALIGN", (1, 1), (1, -1), "CENTER"), ("PADDING", (0, 0), (-1, -1), 5)]))
    summary = Table([[Paragraph("Puntaje total", normal), "", Paragraph("Promedio", normal), ""], [Paragraph("Resultado", normal), "", Paragraph("Estado", normal), ""]], colWidths=[100, 70, 90, 270], rowHeights=[26, 26])
    summary.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF2FF")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#EAF2FF")), ("PADDING", (0, 0), (-1, -1), 6)]))
    notes = Table([[Paragraph(label, normal), ""] for label in ("Fortalezas", "Aspectos a mejorar", "Plan de acción", "Recomendación", "Observaciones")], colWidths=[130, 400], rowHeights=[32] * 5)
    notes.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#EAF2FF")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 6)]))
    signatures = Table([["Firma del evaluador: __________________________", "Firma del colaborador: __________________________"]], colWidths=[265, 265])
    signatures.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, 0), 22)]))
    doc.build([header, Spacer(1, 12), info, Spacer(1, 12), criteria_table, Spacer(1, 12), summary, Spacer(1, 12), notes, Spacer(1, 18), signatures])
    return buffer.getvalue()


def init_db():
    ensure_database_schema()
    import_legacy_quality_labels()
    normalize_legacy_quality_chemistry()
    ensure_bascula_products()
    if not CategoriaInsumo.query.filter_by(nombre="General").first():
        db.session.add(CategoriaInsumo(nombre="General", prefijo="GEN"))
    if not User.query.filter_by(role="root").first():
        root_user = User(username="root", password_hash=generate_password_hash("root123"), full_name="Superusuario Root", role="root")
        db.session.add(root_user)
    if not User.query.filter_by(role="admin").first():
        user = User(username="admin", password_hash=generate_password_hash("admin123"), full_name="Administrador", role="admin")
        db.session.add(user)
    db.session.flush()
    load_initial_inventory_catalog()
    load_initial_general_outputs()

    if not Configuracion.query.filter_by(clave="stock_inicial").first():
        set_config_value("stock_inicial", "0")
    if not Configuracion.query.filter_by(clave="stock_alert").first():
        set_config_value("stock_alert", "500")
    if not Configuracion.query.filter_by(clave="dashboard_theme").first():
        set_config_value("dashboard_theme", "classic")
    for module in ("combustible", "despacho", "inventario"):
        if not Configuracion.query.filter_by(clave=f"dashboard_theme_{module}").first():
            set_config_value(f"dashboard_theme_{module}", get_config_value("dashboard_theme", "classic"))
    if not Configuracion.query.filter_by(clave="dashboard_palette").first():
        set_config_value("dashboard_palette", "blue")
    if not Configuracion.query.filter_by(clave="umbral_orden_compra_gs").first():
        set_config_value("umbral_orden_compra_gs", "1000000")
    if not Configuracion.query.filter_by(clave="cotizacion_usd").first():
        set_config_value("cotizacion_usd", "7500")

    compra_inicial = Compra.query.filter_by(proveedor="Compra inicial").first()
    if compra_inicial is None:
        if not Compra.query.first() and not Despacho.query.first() and not Equipo.query.first():
            db.session.add(Compra(
                litros=1000,
                proveedor="Compra inicial",
                precio_litro=8690,
                costo_total=8690000,
                fecha=datetime(2026, 8, 13, 8, 0, 0),
                notas="Compra inicial registrada a 8.690 Gs/L.",
            ))
    else:
        compra_inicial.litros = 1000
        compra_inicial.precio_litro = 8690
        compra_inicial.costo_total = 8690000
        compra_inicial.fecha = datetime(2026, 8, 13, 8, 0, 0)
        compra_inicial.notas = "Compra inicial registrada a 8.690 Gs/L."
    db.session.commit()


def ensure_bascula_products():
    productos_iniciales = []
    for tipo, prefijo in (("Varillas Lisas", "VL"), ("Varillas Conformadas", "VC")):
        for medida in ("10", "12", "16", "20", "25"):
            productos_iniciales.append((f"{prefijo}-{medida}", f"{tipo} {medida} mm"))
    for codigo, nombre in productos_iniciales:
        producto = ProductoBascula.query.filter_by(codigo=codigo).first()
        if producto is None:
            db.session.add(ProductoBascula(codigo=codigo, nombre=nombre, unidad="kg"))
        elif not producto.activo:
            producto.activo = True
    db.session.commit()


def import_legacy_quality_labels():
    marker = "importacion_etiquetas_dotnet_20260902_v2"
    if Configuracion.query.filter_by(clave=marker).first():
        return
    source = Path(r"C:\Users\PC\OneDrive\Desktop\BETA INTEGRADO\Control de Calidad 4.0\Release\Database\etiquetas.local.json")
    if not source.exists():
        return
    try:
        records = json.loads(source.read_text(encoding="utf-8-sig"))
    except (OSError, ValueError):
        return
    valid_types = {"AP 500 S": "AP500S", "AP500S": "AP500S", "Varillas Lisas": "VarillasLisas", "Palanquillas": "Palanquillas"}
    imported = 0
    legacy_imported = Configuracion.query.filter_by(clave="importacion_etiquetas_dotnet_20260902").first() is not None
    legacy_counts: dict[str, int] = {}
    for record_index, record in enumerate(records if isinstance(records, list) else [], start=1):
        tipo_producto = valid_types.get((record.get("TipoProducto") or "").strip())
        colada = (record.get("Colada") or "").strip()
        if not tipo_producto or not colada:
            continue
        legacy_number = str(record.get("NumeroEtiqueta") or record.get("Id") or "").strip()
        legacy_counts[legacy_number] = legacy_counts.get(legacy_number, 0) + 1
        source_marker = f"Importado .NET Registro={record_index}"
        if EtiquetaCalidad.query.filter(EtiquetaCalidad.observaciones.like(f"%{source_marker}%")).first():
            continue
        if legacy_imported and legacy_number:
            old_marker = f"Importado .NET NumeroEtiqueta={legacy_number}"
            old_count = EtiquetaCalidad.query.filter(EtiquetaCalidad.observaciones.like(f"%{old_marker}%")).count()
            source_occurrence = legacy_counts[legacy_number]
            if source_occurrence <= old_count:
                continue
        raw_date = str(record.get("FechaHora") or "")
        date_match = re.search(r"/Date\(([-0-9]+)", raw_date)
        created_at = datetime.utcfromtimestamp(int(date_match.group(1)) / 1000) if date_match else datetime.utcnow()
        observaciones = (record.get("Comentario") or "").strip()
        observaciones = f"{observaciones} | {source_marker}".strip(" |")
        db.session.add(EtiquetaCalidad(
            tipo_producto=tipo_producto,
            calidad=str(record.get("Calidad") or "").strip(),
            medida=str(record.get("Medida") or "").strip() if tipo_producto != "Palanquillas" else "",
            longitud=str(record.get("Longitud") or "").strip(),
            peso=str(record.get("Peso") or "").strip(),
            colada=colada,
            lote=str(record.get("Lote") or "").strip(),
            cantidad=int(record.get("Cantidad") or 0),
            fecha=created_at.date(),
            carbono=str(record.get("GradoCarbono") or 0).strip(),
            silicio=str(record.get("Silicio") or 0).strip(),
            manganeso=str(record.get("Manganeso") or 0).strip(),
            hornero=str(record.get("Hornero") or "").strip(),
            supervisor=str(record.get("Supervisor") or "").strip(),
            operador_ccm=str(record.get("OperadorCCM") or "").strip(),
            tipo_b=bool(record.get("TipoB")),
            no_conforme=bool(record.get("NoConforme")),
            observaciones=observaciones,
            created_at=created_at,
        ))
        imported += 1
    db.session.add(Configuracion(clave=marker, valor=str(imported)))
    db.session.commit()


def normalize_legacy_quality_chemistry():
    marker = "normalizacion_quimica_etiquetas_dotnet_20260902"
    if Configuracion.query.filter_by(clave=marker).first():
        return
    imported = EtiquetaCalidad.query.filter(EtiquetaCalidad.observaciones.like("%Importado .NET%"), EtiquetaCalidad.tipo_producto == "Palanquillas").all()
    for etiqueta in imported:
        etiqueta.carbono = etiqueta.carbono or "0"
        etiqueta.silicio = etiqueta.silicio or "0"
        etiqueta.manganeso = etiqueta.manganeso or "0"
    db.session.add(Configuracion(clave=marker, valor=str(len(imported))))
    db.session.commit()


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated and request.method == "GET":
        destination = "quality_dashboard" if (current_user.is_calidad or current_user.is_control_calidad) else "admin_dashboard" if current_user.is_admin else "laminacion_dashboard" if (current_user.is_operador_horno or current_user.is_operador_despunte) else "rrhh" if current_user.is_rrhh else "supervisor_dashboard" if current_user.is_supervisor else "inventario" if current_user.is_deposito else "modulo_compras" if current_user.is_compras else "gerencia_pedidos" if current_user.is_gerencia else "despacho_dashboard"
        return redirect(url_for(destination))
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        if current_user.is_authenticated and current_user.username != username:
            logout_user()
        if username in {"admin", "root"} and not User.query.filter_by(username=username).first():
            init_db()
        user = User.query.filter_by(username=username).first()
        if user and check_password_hash(user.password_hash, password):
            login_user(user)
            destination = "quality_dashboard" if (user.is_calidad or user.is_control_calidad) else "admin_dashboard" if user.is_admin else "laminacion_dashboard" if (user.is_operador_horno or user.is_operador_despunte) else "rrhh" if user.is_rrhh else "supervisor_dashboard" if user.is_supervisor else "inventario" if user.is_deposito else "modulo_compras" if user.is_compras else "gerencia_pedidos" if user.is_gerencia else "despacho_dashboard"
            return redirect(url_for(destination))
        flash("Usuario o contraseña incorrectos.", "error")
    return render_template("login.html")


@app.get("/login-media/<path:filename>")
def login_media(filename: str):
    safe_name = secure_filename(filename)
    target = IMAGENES_DIR / safe_name
    if not target.exists() or not target.is_file():
        abort(404)
    return send_file(target, conditional=True)


@app.get("/sw.js")
def service_worker():
    service_worker_path = BASE_DIR / "static" / "sw.js"
    if not service_worker_path.exists():
        return Response("self.addEventListener('fetch', () => {});", mimetype="application/javascript")
    return send_file(service_worker_path, mimetype="application/javascript")


@app.route("/usuarios", methods=["GET", "POST"])
@admin_required
def usuarios():
    if request.method == "POST":
        username = request.form.get("username", "").strip()
        full_name = request.form.get("full_name", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "despacho").strip().lower()
        if role.replace("_", " ").replace("-", " ") == "recursos humanos":
            role = "rrhh"
        if not username or not full_name or len(password) < 6 or role not in {"admin", "despacho", "deposito", "rrhh", "supervisor", "compras", "gerencia", "calidad", "control_calidad", "root"}:
            flash("Complete los datos y use una contraseña de al menos 6 caracteres.", "error")
        elif User.query.filter_by(username=username).first():
            flash("Ese nombre de usuario ya existe.", "error")
        else:
            user = User(
                username=username,
                full_name=full_name,
                password_hash=generate_password_hash(password),
                role=role,
                correo=request.form.get("correo", "").strip() or None,
                telefono=request.form.get("telefono", "").strip() or None,
            )
            db.session.add(user)
            db.session.flush()
            if role == "supervisor":
                selected_ids = [int(value) for value in request.form.getlist("colaboradores") if value.isdigit()]
                Colaborador.query.filter(Colaborador.id.in_(selected_ids)).update({"supervisor_id": user.id}, synchronize_session=False)
            db.session.commit()
            flash("Usuario creado correctamente.", "success")
            return redirect(url_for("usuarios"))
    role_groups = [
        ("root", "Superusuario"),
        ("admin", "Administrador"),
        ("rrhh", "RRHH"),
        ("compras", "Compras"),
        ("gerencia", "Gerencia"),
        ("supervisor", "Supervisor"),
        ("control_calidad", "Control de Calidad"),
        ("deposito", "Depósito"),
        ("despacho", "Despacho"),
        ("calidad", "Calidad"),
    ]
    users_by_role = {}
    for user in User.query.order_by(User.full_name.asc(), User.username.asc()).all():
        users_by_role.setdefault(user.role.strip().lower(), []).append(user)
    grouped_users = [
        {"role": role, "label": label, "users": users_by_role.pop(role, [])}
        for role, label in role_groups
        if users_by_role.get(role)
    ]
    for role, users in sorted(users_by_role.items()):
        grouped_users.append({"role": role, "label": role.replace("_", " ").title(), "users": users})
    colaboradores = Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all()
    return render_template("usuarios.html", grupos_usuarios=grouped_users, colaboradores=colaboradores)


@app.route("/usuarios/<int:user_id>/editar", methods=["GET", "POST"])
@admin_required
def editar_usuario(user_id: int):
    user = User.query.get_or_404(user_id)
    if request.method == "POST":
        full_name = request.form.get("full_name", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "despacho").strip().lower()
        valid_roles = {"admin", "despacho", "deposito", "rrhh", "supervisor", "compras", "gerencia", "calidad", "control_calidad", "root"}
        if not full_name or role not in valid_roles:
            flash("Complete el nombre y seleccione un rol válido.", "error")
        elif password and len(password) < 6:
            flash("La nueva contraseña debe tener al menos 6 caracteres.", "error")
        elif user.id == current_user.id and role != "admin":
            flash("No puedes quitar el rol de administrador de tu propia cuenta.", "error")
        elif user.is_admin and role != "admin" and sum(candidate.is_admin for candidate in User.query.all()) <= 1:
            flash("Debe existir al menos un administrador.", "error")
        else:
            user.full_name = full_name
            user.role = role
            user.correo = request.form.get("correo", "").strip() or None
            user.telefono = request.form.get("telefono", "").strip() or None
            if role == "supervisor":
                selected_ids = [int(value) for value in request.form.getlist("colaboradores") if value.isdigit()]
                Colaborador.query.filter_by(supervisor_id=user.id).update({"supervisor_id": None}, synchronize_session=False)
                Colaborador.query.filter(Colaborador.id.in_(selected_ids)).update({"supervisor_id": user.id}, synchronize_session=False)
            else:
                Colaborador.query.filter_by(supervisor_id=user.id).update({"supervisor_id": None}, synchronize_session=False)
            if password:
                user.password_hash = generate_password_hash(password)
            db.session.commit()
            flash("Usuario actualizado correctamente.", "success")
            return redirect(url_for("usuarios"))
    colaboradores = Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all()
    return render_template("editar_usuario.html", usuario=user, colaboradores=colaboradores)


@app.post("/usuarios/<int:user_id>/eliminar")
@admin_required
def eliminar_usuario(user_id: int):
    user = User.query.get_or_404(user_id)
    if user.id == current_user.id:
        flash("No puedes eliminar tu propia cuenta.", "error")
    elif user.is_admin and User.query.filter_by(role="admin").count() <= 1:
        flash("Debe existir al menos un administrador.", "error")
    else:
        db.session.delete(user)
        db.session.commit()
        flash("Usuario eliminado correctamente.", "success")
    return redirect(url_for("usuarios"))


def _qr_line(lines: list[str], label: str, value) -> None:
    if value is not None and str(value).strip():
        lines.append(f"{label}: {str(value).replace(chr(10), ' ').strip()}")


def _build_quality_label_qr(etiqueta: EtiquetaCalidad, numero: str) -> str:
    lines: list[str] = []
    palanquilla = None
    if etiqueta.tipo_producto != "Palanquillas":
        palanquilla = EtiquetaCalidad.query.filter_by(tipo_producto="Palanquillas", colada=etiqueta.colada).order_by(EtiquetaCalidad.created_at.desc()).first()
    quimica = palanquilla or etiqueta
    _qr_line(lines, "Tipo de producto", etiqueta.tipo_producto)
    _qr_line(lines, "Calidad", etiqueta.calidad)
    _qr_line(lines, "Medida (mm)", etiqueta.medida)
    _qr_line(lines, "Longitud (m)", etiqueta.longitud)
    _qr_line(lines, "Peso (kg)", etiqueta.peso)
    _qr_line(lines, "Numero de colada", etiqueta.colada)
    _qr_line(lines, "Carbono (%)", quimica.carbono)
    _qr_line(lines, "Silicio (%)", quimica.silicio)
    _qr_line(lines, "Manganeso (%)", quimica.manganeso)
    if etiqueta.tipo_producto == "Palanquillas":
        _qr_line(lines, "Cantidad", etiqueta.cantidad)
        _qr_line(lines, "Lote", etiqueta.lote)
        _qr_line(lines, "Hornero", etiqueta.hornero)
        _qr_line(lines, "Supervisor", etiqueta.supervisor)
        _qr_line(lines, "Operador CCM", etiqueta.operador_ccm)
        if etiqueta.aprobado:
            _qr_line(lines, "Estado", "APROBADO")
    _qr_line(lines, "Comentario", etiqueta.observaciones)
    if etiqueta.no_conforme:
        _qr_line(lines, "Estado", "NO CONFORME")
    elif etiqueta.tipo_b:
        _qr_line(lines, "Clasificacion", "TIPO B")
    else:
        _qr_line(lines, "Clasificacion", "PRIMERA")
    _qr_line(lines, "Numero de etiqueta", numero)
    _qr_line(lines, "Fecha y hora", etiqueta.created_at.strftime("%Y-%m-%d %H:%M"))
    _qr_line(lines, "Usuario", etiqueta.created_by.full_name if etiqueta.created_by else "")
    return "\n".join(lines)


def _build_palanquilla_label_qr(etiqueta: EtiquetaCalidad, numero: str) -> str:
    lines: list[str] = []
    for label, value in (
        ("Carbono %", etiqueta.carbono),
        ("Silicio %", etiqueta.silicio),
        ("Manganeso %", etiqueta.manganeso),
        ("Longitud (mts)", etiqueta.longitud),
        ("Peso (kg)", etiqueta.peso),
        ("Colada", etiqueta.colada),
        ("Lote", etiqueta.lote),
        ("Hornero", etiqueta.hornero),
        ("Supervisor", etiqueta.supervisor),
        ("Operador de CCM", etiqueta.operador_ccm),
        ("Cantidad", etiqueta.cantidad),
    ):
        _qr_line(lines, label, value)
    total_kg, total_ton = _palanquilla_total_weight(etiqueta)
    if total_kg is not None:
        _qr_line(lines, "Peso total colada (kg)", f"{total_kg:.3f}")
        _qr_line(lines, "Peso total colada (t)", f"{total_ton:.3f}")
    _qr_line(lines, "Numero de etiqueta", numero)
    return "\n".join(lines)


def _palanquilla_total_weight(etiqueta: EtiquetaCalidad) -> tuple[float | None, float | None]:
    try:
        cantidad = float(etiqueta.cantidad or 0)
        peso = float(str(etiqueta.peso or "0").replace(",", "."))
        longitud = float(str(etiqueta.longitud or "0").replace(",", "."))
    except (TypeError, ValueError):
        return None, None
    total_kg = cantidad * peso * longitud
    return total_kg, total_kg / 1000


def _draw_quality_label_pdf(etiqueta: EtiquetaCalidad, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    width, height = 100 * mm, 140 * mm
    numero = f"ETQ-{etiqueta.id:07d}"
    pdf = canvas.Canvas(str(output_path), pagesize=(width, height))
    pdf.setTitle(f"Etiqueta {numero}")
    margin = 6 * mm
    navy = colors.HexColor("#0b1f33")
    red = colors.HexColor("#c62828")

    pdf.setFillColor(navy)
    if etiqueta.tipo_producto == "Palanquillas":
        igp_logo = QUALITY_LABELS_IMAGES_DIR / "igp_logo.png.jpg"
        if not igp_logo.exists():
            igp_logo = STATIC_LOGO_PATH
        pdf.setFont("Helvetica-Bold", 9)
        pdf.drawCentredString(width / 2, height - 18 * mm, "PARA OBRAS EXCEPCIONALES")
        pdf.drawCentredString(width / 2, height - 23 * mm, "PRODUCTOS EXCEPCIONALES")
        paraguay_logo = QUALITY_LABELS_IMAGES_DIR / "paraguay_logo.png"
        if paraguay_logo.exists():
            pdf.drawImage(str(paraguay_logo), width / 2 - 6 * mm, height - 45 * mm, 12 * mm, 8 * mm, preserveAspectRatio=True, anchor="c", mask="auto")

        def palanquilla_field(label: str, value, y: float) -> None:
            pdf.setFillColor(colors.black)
            pdf.setFont("Helvetica", 8.5)
            pdf.drawString(margin, y, label)
            pdf.setFont("Helvetica-Bold", 11)
            pdf.drawString(margin + 39 * mm, y, str(value) if value is not None else "")

        field_y = height - 48 * mm
        for label, value in (
            ("Carbono %", etiqueta.carbono),
            ("Silicio %", etiqueta.silicio),
            ("Manganeso %", etiqueta.manganeso),
            ("Longitud (mts)", etiqueta.longitud),
            ("Peso (kg)", etiqueta.peso),
            ("Colada", etiqueta.colada),
            ("Lote", etiqueta.lote),
            ("Hornero", etiqueta.hornero),
            ("Supervisor", etiqueta.supervisor),
            ("Operador de CCM", etiqueta.operador_ccm),
        ):
            palanquilla_field(label, value, field_y)
            field_y -= 5 * mm

        pdf.setFont("Helvetica", 8.5)
        pdf.drawString(margin, field_y, "Cantidad")
        pdf.setFont("Helvetica-Bold", 31)
        pdf.drawString(margin + 39 * mm, field_y - 3 * mm, str(etiqueta.cantidad) if etiqueta.cantidad is not None else "")
        total_kg, total_ton = _palanquilla_total_weight(etiqueta)
        if total_kg is not None:
            pdf.setFillColor(colors.HexColor("#0b1f33"))
            pdf.setFont("Helvetica", 8.5)
            pdf.drawString(margin, field_y - 12 * mm, "Peso total colada")
            pdf.setFont("Helvetica-Bold", 10)
            pdf.drawString(margin + 39 * mm, field_y - 12 * mm, f"{total_kg:,.0f} kg".replace(",", "."))
            pdf.setFont("Helvetica-Bold", 9)
            pdf.drawString(margin + 39 * mm, field_y - 17 * mm, f"({total_ton:,.2f} t)".replace(",", "."))

        pdf.setFillColor(red)
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawRightString(width - margin, 8 * mm, numero)

        qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_L, box_size=5, border=2)
        qr.add_data(_build_palanquilla_label_qr(etiqueta, numero))
        qr.make(fit=True)
        qr_image = qr.make_image(fill_color="black", back_color="white")
        pdf.drawImage(ImageReader(qr_image.get_image()), width - 38 * mm, height - 69 * mm, width=28 * mm, height=28 * mm, preserveAspectRatio=True, mask="auto")
        if igp_logo.exists():
            watermark = PILImage.open(igp_logo).convert("RGBA")
            watermark.putdata([(red_value, green_value, blue_value, 42) for red_value, green_value, blue_value, _ in watermark.getdata()])
            pdf.drawImage(ImageReader(watermark), width / 2 - 20 * mm, 3 * mm, 40 * mm, 18 * mm, preserveAspectRatio=True, anchor="c", mask="auto")
        pdf.save()
        return

    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawCentredString(width / 2, height - 28 * mm, "PARA OBRAS EXCEPCIONALES")
    pdf.drawCentredString(width / 2, height - 33 * mm, "PRODUCTOS EXCEPCIONALES")
    pdf.setFont("Helvetica-Bold", 13)
    pdf.drawCentredString(width / 2, height - 39 * mm, "AP 500 S" if etiqueta.tipo_producto == "AP500S" else etiqueta.tipo_producto.upper())

    def field(label: str, value: str, x: float, y: float, value_x: float | None = None) -> None:
        pdf.setFillColor(navy)
        pdf.setFont("Helvetica", 8.5)
        pdf.drawString(x, y, label.upper())
        pdf.setFont("Helvetica-Bold", 11)
        pdf.drawString(value_x if value_x is not None else x, y - 5 * mm, str(value or "-"))

    top = height - 49 * mm
    field("Acero (calidad)", etiqueta.calidad, margin, top)
    if etiqueta.tipo_producto != "Palanquillas":
        field("Medida (mm)", etiqueta.medida, width / 2 + 3 * mm, top)
    field("Longitud (metros)", etiqueta.longitud, margin, top - 14 * mm)
    field("Peso (kg)", etiqueta.peso, width / 2 + 3 * mm, top - 14 * mm)

    pdf.setFillColor(navy)
    pdf.setFont("Helvetica", 8.5)
    pdf.drawString(margin, top - 27 * mm, "COLADA")
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(margin + 24 * mm, top - 27 * mm, str(etiqueta.colada or "-"))

    if etiqueta.tipo_producto == "Palanquillas":
        field("Fecha", etiqueta.fecha.strftime("%d/%m/%Y") if etiqueta.fecha else "-", width / 2 + 3 * mm, top - 28 * mm)
        field("Lote", etiqueta.lote, margin, top - 42 * mm)
        field("Carbono %", etiqueta.carbono, width / 2 + 3 * mm, top - 42 * mm)
        field("Silicio %", etiqueta.silicio, margin, top - 56 * mm)
        field("Manganeso %", etiqueta.manganeso, width / 2 + 3 * mm, top - 56 * mm)
        field("Cantidad", etiqueta.cantidad, margin, top - 70 * mm)
        field("Hornero", etiqueta.hornero, width / 2 + 3 * mm, top - 70 * mm)
        field("Supervisor", etiqueta.supervisor, margin, top - 84 * mm)
        field("Operador CCM", etiqueta.operador_ccm, width / 2 + 3 * mm, top - 84 * mm)
        if etiqueta.aprobado:
            pdf.setFillColor(colors.HexColor("#16803c"))
            pdf.setFont("Helvetica-Bold", 9)
            pdf.drawCentredString(width / 2, 25 * mm, "APROBADO")
    elif etiqueta.no_conforme or etiqueta.tipo_b:
        pdf.setFillColor(colors.HexColor("#c62828"))
        pdf.setFont("Helvetica-Bold", 10)
        pdf.drawCentredString(width / 2, 25 * mm, "NO CONFORME" if etiqueta.no_conforme else "TIPO B")

    igp_logo = QUALITY_LABELS_IMAGES_DIR / "igp_logo.png.jpg"
    if not igp_logo.exists():
        igp_logo = STATIC_LOGO_PATH
    if igp_logo.exists():
        watermark = PILImage.open(igp_logo).convert("RGBA")
        watermark_pixels = []
        for pixel in watermark.getdata():
            red_value, green_value, blue_value, _ = pixel
            luminance = (red_value + green_value + blue_value) / 3
            alpha = max(0, min(58, int((255 - luminance) * 0.38)))
            watermark_pixels.append((red_value, green_value, blue_value, alpha))
        watermark.putdata(watermark_pixels)
        pdf.drawImage(ImageReader(watermark), width / 2 - 20 * mm, 1 * mm, 40 * mm, 18 * mm, preserveAspectRatio=True, anchor="c", mask="auto")

    qr = qrcode.QRCode(error_correction=qrcode.constants.ERROR_CORRECT_L, box_size=5, border=2)
    qr.add_data(_build_quality_label_qr(etiqueta, numero))
    qr.make(fit=True)
    qr_image = qr.make_image(fill_color="black", back_color="white")
    pdf.drawImage(ImageReader(qr_image.get_image()), margin, 28 * mm, width=28 * mm, height=28 * mm, preserveAspectRatio=True, mask="auto")

    paraguay_logo = QUALITY_LABELS_IMAGES_DIR / "paraguay_logo.png"
    if paraguay_logo.exists():
        pdf.drawImage(str(paraguay_logo), width - 35 * mm, 29 * mm, 27 * mm, 27 * mm, preserveAspectRatio=True, anchor="c", mask="auto")
    onc = QUALITY_LABELS_IMAGES_DIR / "onc.png"
    if not onc.exists():
        onc = QUALITY_LABELS_IMAGES_DIR / "onc_logo.png.jpg"
    if etiqueta.tipo_producto == "AP500S" and onc.exists():
        pdf.drawImage(str(onc), width / 2 - 11 * mm, 30 * mm, 22 * mm, 22 * mm, preserveAspectRatio=True, anchor="c", mask="auto")

    pdf.setFillColor(red)
    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawRightString(width - margin, margin, numero)
    pdf.save()


def _regenerate_rod_labels_for_colada(colada: str) -> None:
    if not colada:
        return
    labels = EtiquetaCalidad.query.filter(
        EtiquetaCalidad.colada == colada,
        EtiquetaCalidad.tipo_producto.in_(["AP500S", "VarillasLisas"]),
    ).all()
    for label in labels:
        _draw_quality_label_pdf(label, QUALITY_LABELS_DIR / f"etiqueta_{label.id:07d}.pdf")


def _next_certificate_number() -> str:
    return f"CCA-{(CertificadoCalidad.query.count() + 1):06d}"


def _prefill_certificate(colada: str) -> dict:
    etiquetas = EtiquetaCalidad.query.filter_by(colada=colada).order_by(EtiquetaCalidad.created_at.desc()).all()
    varillas = [item for item in etiquetas if item.tipo_producto in {"AP500S", "VarillasLisas"}]
    palanquillas = [item for item in etiquetas if item.tipo_producto == "Palanquillas"]
    palanquilla = palanquillas[0] if palanquillas else None
    varilla_data = "; ".join(
        f"{'AP500 S' if item.tipo_producto == 'AP500S' else 'Varillas Lisas'}: {item.medida or '-'} mm x {item.longitud or '-'} m, {item.peso or '-'} kg"
        for item in varillas
    )
    product_options = []
    product_details = []
    seen_options = set()
    for item in varillas:
        option = f"{'AP500 S' if item.tipo_producto == 'AP500S' else 'Varillas Lisas'} - {item.medida or '-'} mm"
        if option not in seen_options:
            seen_options.add(option)
            product_options.append(option)
            product_details.append({
                "label": option,
                "longitudes": sorted({rod.longitud for rod in varillas if f"{'AP500 S' if rod.tipo_producto == 'AP500S' else 'Varillas Lisas'} - {rod.medida or '-'} mm" == option and rod.longitud}),
                "pesos": sorted({rod.peso for rod in varillas if f"{'AP500 S' if rod.tipo_producto == 'AP500S' else 'Varillas Lisas'} - {rod.medida or '-'} mm" == option and rod.peso}),
            })
    return {
        "tipo_producto": "AP 500 S" if any(item.tipo_producto == "AP500S" for item in varillas) else "Varillas Lisas",
        "has_varillas": bool(varillas),
        "has_palanquillas": bool(palanquillas),
        "carbono": palanquilla.carbono if palanquilla else "",
        "silicio": palanquilla.silicio if palanquilla else "",
        "manganeso": palanquilla.manganeso if palanquilla else "",
        "medidas": ", ".join(sorted({item.medida for item in varillas if item.medida})),
        "longitudes": ", ".join(sorted({item.longitud for item in varillas if item.longitud})),
        "cantidad_palanquillas": sum(item.cantidad or 0 for item in palanquillas),
        "longitudes_palanquillas": ", ".join(sorted({item.longitud for item in palanquillas if item.longitud})),
        "pesos_palanquillas": ", ".join(sorted({item.peso for item in palanquillas if item.peso})),
        "datos_varillas": varilla_data,
        "product_options": product_options,
        "product_details": product_details,
        "fecha_atado": (varillas[0].fecha or varillas[0].created_at.date()) if varillas else (palanquilla.fecha if palanquilla else datetime.utcnow().date()),
    }


def _draw_quality_certificate_pdf(certificado: CertificadoCalidad, output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    page_width, page_height = 210 * mm, 297 * mm
    pdf = canvas.Canvas(str(output_path), pagesize=(page_width, page_height))
    pdf.setTitle(f"Certificado {certificado.numero_certificado}")
    left, right = 9 * mm, page_width - 9 * mm
    watermark = IMAGENES_DIR / "marcadeagua.png"
    if watermark.exists():
        watermark_image = PILImage.open(watermark).convert("RGBA")
        watermark_pixels = []
        for red_value, green_value, blue_value, alpha_value in watermark_image.getdata():
            watermark_pixels.append((red_value, green_value, blue_value, min(alpha_value, 48)))
        watermark_image.putdata(watermark_pixels)
        pdf.drawImage(ImageReader(watermark_image), 0, 0, page_width, page_height, preserveAspectRatio=False, mask="auto")

    navy = colors.HexColor("#172b78")
    gray = colors.HexColor("#6b7280")
    logo = IMAGENES_DIR / "igp_logo.png.jpg"
    if logo.exists():
        pdf.drawImage(str(logo), left, page_height - 24 * mm, 42 * mm, 21 * mm, preserveAspectRatio=True, anchor="sw", mask="auto")
    pdf.setFillColor(navy)
    pdf.setFont("Helvetica-Bold", 16)
    pdf.drawCentredString(page_width / 2, page_height - 14 * mm, "CERTIFICADO DE CALIDAD")
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawRightString(right, page_height - 10 * mm, certificado.numero_certificado)
    pdf.setFont("Helvetica", 7)
    pdf.setFillColor(gray)
    pdf.drawRightString(right, page_height - 16 * mm, "CCA-FOR-012-REV.: 000  |  Vigencia: 13/07/2026")
    pdf.setStrokeColor(navy)
    pdf.setLineWidth(1.5)
    pdf.line(left, page_height - 29 * mm, right, page_height - 29 * mm)

    chart_x, chart_y = left, 153 * mm
    chart_width, chart_height = right - left, 91 * mm
    pdf.setFillColor(navy)
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawCentredString(page_width / 2, 258 * mm, "RESULTADO DE ANALISIS DE PRODUCTOS LAMINADOS")
    pdf.setFillColor(colors.white)
    pdf.setStrokeColor(colors.HexColor("#b8c0c9"))
    pdf.rect(chart_x, chart_y, chart_width, chart_height, fill=1, stroke=1)
    if certificado.grafico_path and Path(certificado.grafico_path).exists():
        pdf.drawImage(certificado.grafico_path, chart_x + 5 * mm, chart_y + 5 * mm, chart_width - 10 * mm, chart_height - 10 * mm, preserveAspectRatio=True, anchor="c", mask="auto")
    else:
        pdf.setFillColor(gray)
        pdf.setFont("Helvetica", 9)
        pdf.drawCentredString(page_width / 2, chart_y + chart_height / 2, "Gráfico no cargado")

    def box(label: str, value: str, x: float, y: float, box_width: float) -> None:
        pdf.setStrokeColor(colors.HexColor("#cbd5e1"))
        pdf.setFillColor(colors.white)
        pdf.rect(x, y - 7 * mm, box_width, 7 * mm, fill=1, stroke=1)
        pdf.setFillColor(navy)
        pdf.setFont("Helvetica-Bold", 7)
        pdf.drawString(x + 2 * mm, y - 2.8 * mm, label)
        pdf.setFillColor(colors.HexColor("#333333"))
        pdf.setFont("Helvetica", 8)
        pdf.drawRightString(x + box_width - 2 * mm, y - 2.8 * mm, str(value or "-"))

    pdf.setFillColor(navy)
    pdf.setFont("Helvetica-Bold", 11)
    pdf.drawString(left, 145 * mm, "INFORMACIÓN DEL PRODUCTO")
    selected_product_labels = [value.strip() for value in (certificado.productos_certificados or "").split("|") if value.strip()]
    selected_measures = sorted({match.group(1) for value in selected_product_labels if (match := re.search(r"-\s*([^ ]+)\s*mm", value))})
    display_measures = ", ".join(selected_measures) or certificado.medidas
    if not selected_measures:
        selected_measures = [value.strip() for value in re.split(r"[,;|]", certificado.medidas or "") if value.strip()]
    display_types = []
    for value in selected_product_labels:
        product_name = value.split(" - ", 1)[0]
        if product_name not in display_types:
            display_types.append(product_name)
    display_product = " / ".join(display_types) or certificado.tipo_producto

    def force_values(stress_value: str) -> str:
        try:
            stress = float(str(stress_value).replace(",", "."))
        except (TypeError, ValueError):
            return str(stress_value or "-")
        values = []
        for diameter in selected_measures:
            try:
                diameter_mm = float(diameter.replace(",", "."))
            except ValueError:
                continue
            force_newtons = stress * 3.141592653589793 * diameter_mm ** 2 / 4
            values.append(f"{force_newtons:,.0f} N".replace(",", "."))
        return f"{stress_value} MPa" + (f"  |  {'; '.join(values)}" if values else "")

    gap = 4 * mm
    box_width = (right - left - gap) / 2
    rows = [
        (("N° de certificado", certificado.numero_certificado), ("Carbono (%)", certificado.carbono or "0")),
        (("Colada N°", certificado.numero_colada), ("Silicio (%)", certificado.silicio)),
        (("Tipo de acero", display_product), ("Manganeso (%)", certificado.manganeso)),
        (("Diámetros (mm)", display_measures), ("LF (MPa / N)", force_values(certificado.limite_fluencia))),
        (("Fecha de producción / atado", certificado.fecha_atado.strftime("%d/%m/%Y") if certificado.fecha_atado else "-"), ("LR (MPa / N)", force_values(certificado.limite_resistencia))),
        (("Longitud(es) (m)", certificado.longitudes), ("Alargamiento A (%)", certificado.alargamiento)),
        (("", ""), ("Doblado 180°", "CUMPLE" if certificado.doblado_cumple else "NO CUMPLE")),
    ]
    y = 139 * mm
    for (left_label, left_value), (right_label, right_value) in rows:
        if left_label:
            box(left_label, left_value, left, y, box_width)
        box(right_label, right_value, left + box_width + gap, y, box_width)
        y -= 9 * mm

    signature = Path(get_config_value("calidad_firma_path", str(QUALITY_LABELS_IMAGES_DIR / "firma.png")))
    firma_nombre = get_config_value("calidad_firma_nombre", "Ing. Carlos Gonzalez")
    firma_cargo = get_config_value("calidad_firma_cargo", "Jefe de Control de Calidad")
    if signature.exists():
        pdf.drawImage(str(signature), right - 42 * mm, 22 * mm, 34 * mm, 14 * mm, preserveAspectRatio=True, anchor="c", mask="auto")
    pdf.setStrokeColor(gray)
    pdf.line(right - 48 * mm, 20 * mm, right, 20 * mm)
    pdf.setFillColor(gray)
    pdf.setFont("Helvetica", 7)
    pdf.drawCentredString(right - 24 * mm, 16 * mm, firma_nombre)
    pdf.drawCentredString(right - 24 * mm, 12 * mm, firma_cargo)
    pdf.drawCentredString(right - 24 * mm, 8 * mm, "IGP METALES")
    onc = QUALITY_LABELS_IMAGES_DIR / "onc.png"
    if not onc.exists():
        onc = QUALITY_LABELS_IMAGES_DIR / "onc_logo.png.jpg"
    if "AP500S" in certificado.tipo_producto.upper().replace(" ", "") and onc.exists():
        pdf.drawImage(str(onc), 88 * mm, 13 * mm, 28 * mm, 20 * mm, preserveAspectRatio=True, anchor="c", mask="auto")
    pdf.save()


@app.route("/admin")
@role_required("admin", "rrhh", "root")
def admin_dashboard():
    return render_template("admin_dashboard.html")


@app.get("/admin/diagrama-entidad-relacion.pdf")
@role_required("admin", "root")
def diagrama_entidad_relacion_pdf():
    return Response(
        "El diagrama entidad-relación aún no está disponible en esta instancia.",
        mimetype="text/plain",
        headers={"Content-Disposition": "attachment; filename=diagrama-entidad-relacion.txt"},
    )


@app.get("/admin/produccion")
@role_required("admin", "root")
def produccion_dashboard():
    return render_template("produccion_dashboard.html")


LATEST_SCALE_READING = {"value": "", "updated_at": ""}


@app.post("/bascula/lectura")
@admin_required
def publicar_lectura_bascula():
    payload = request.get_json(silent=True) or {}
    value = str(payload.get("value") or "").strip()
    if not value:
        return jsonify({"ok": False, "error": "Lectura vacía."}), 400
    LATEST_SCALE_READING.update(value=value, updated_at=datetime.utcnow().strftime("%H:%M:%S UTC"))
    return jsonify({"ok": True, **LATEST_SCALE_READING})


@app.get("/bascula/lectura")
@admin_required
def obtener_lectura_bascula():
    return jsonify(LATEST_SCALE_READING)


@app.route("/bascula", methods=["GET", "POST"])
@admin_required
def bascula():
    if request.method == "POST":
        accion = request.form.get("accion", "")
        if accion in {"producto", "producto_editar"}:
            producto_id = request.form.get("producto_id", type=int)
            codigo = (request.form.get("codigo") or "").strip().upper()
            nombre = (request.form.get("nombre") or "").strip()
            unidad = (request.form.get("unidad") or "kg").strip().lower()
            if not codigo or not nombre or unidad not in {"kg", "ton", "unidad"}:
                flash("Código, nombre y unidad son obligatorios.", "error")
            else:
                producto = ProductoBascula.query.get(producto_id) if producto_id else None
                if producto is None:
                    producto = ProductoBascula(codigo=codigo, nombre=nombre, unidad=unidad)
                    db.session.add(producto)
                else:
                    producto.codigo = codigo
                    producto.nombre = nombre
                    producto.unidad = unidad
                try:
                    db.session.commit()
                    flash("Producto de báscula guardado.", "success")
                except Exception:
                    db.session.rollback()
                    flash("El código o nombre del producto ya existe.", "error")
        elif accion == "producto_eliminar":
            producto = ProductoBascula.query.get_or_404(request.form.get("producto_id", type=int))
            producto.activo = False
            db.session.commit()
            flash("Producto desactivado; su historial se conserva.", "success")
        return redirect(url_for("bascula", sector="productos"))

    ensure_bascula_products()
    productos_todos = ProductoBascula.query.order_by(ProductoBascula.nombre.asc()).all()
    productos = [producto for producto in productos_todos if producto.activo]
    return render_template("bascula.html", productos=productos, productos_todos=productos_todos)


LAMINACION_PROBLEMAS = {
    "Horno de Recalentamiento": "Temperatura Baja|Temperatura alta|Material Frio|Palanquilla mal posicionada|Material no avanza|Falta de Velocidad de avance|Falla de refrigeracion|Falla de bomba de agua|Baja presion de agua|Refractario danado|Temperatura alta del tablero|Lote incorrecto o mal identificado|Perdida de agua|Reductora danada|Otros",
    "Mesa de rodillos/Entrada al Trio": "Rodillo detenido|Rodillo trabado|Cadena danada|Material desalineado|Material ingresa torcido|Palanquilla golpea canaleta|Atascamiento en mesa|Problema de arrastre|Sensor o tope fuera de posicion|Rodamiento roto|Otros",
    "Trio": "Material no toma el pase|Rotura de material|GAP fuera de medida|Calibracion incorrecta|Canal desgastado|Canal incorrecto|Rodillo danado|Rodamiento danado|Caja de rodamiento danada|Cardan danado|Brida de cardan danada|Cruceta danada|Barrones flojos|Barrones danados|Guia de entrada desalineada|Guia de entrada danada|Guia de salida desalineada|Guia de salida danada|Vibracion anormal|Falta de lubricacion|Burloneria floja|Base desalineada|Reductora danada",
    "Mesa Basculante": "Rodillo detenido|Rodillo trabado|Cadena danada|Material desalineado|Material ingresa torcido|Palanquilla golpea canaleta|Atascamiento en mesa|Problema de arrastre|Sensor o tope fuera de posicion|Rodamientos rotos|Manguera de aire rota|Cilindro neumatico roto|Otros",
    "Guias": "Guia floja|Rotura de guia|Roletes danados|Roletes trabados|Eje de rolete danado|Falta de refrigeracion en guia|Boquilla tapada|Sobrecalentamiento de guia|Material roza la guia|Otros",
    "Stand": "Atascamiento en stand|Material fuera de centro|Exceso de tension|Falta de traccion|Variacion de velocidad|RPM incorrecta|Frecuencia incorrecta|Gap cerrado|Gap abierto|Canaleta desalineada|Canaleta danada|Base de stand floja|Cambio de canal|Cambio de medida|Ajuste operativo|Cardan danado|Cruceta rota|Brida rota|Brida floja|Reductora danada",
    "Arrastradores": "Falla de arrastradores|Arrastradores desalineados|Rodillo arrastrador danado|Falta de presion en arrastrador|Material patina|Varilla ondulada|Varilla sale torcida|Mala sincronizacion de velocidad|Atascamiento en salida|Otros",
    "Tijera de Corte/Cizalla mesa enfriamiento": "Cizalla no corta|Corte irregular|Cuchilla gastada|Baja presion neumatica|Punta fria|Excedente de material|Despunte incorrecto|Tacho de despunte lleno|Acumulacion de chatarra|Retiro de despuntes para pesaje|Atascamiento en mesa|Mesa de rodillos danada|Correa rota|Cadena rota|Cuchilla rota|Cambio de cuchilla|Inversion de cuchilla",
    "Calidad": "Medida fuera de norma|Peso lineal fuera de rango|Friso alto|Friso discontinuo|Varilla ovalada|Varilla torcida|Varilla con marca superficial|Defecto de canal desgastado|Defecto por guia descentrada|Producto no conforme|Otros",
    "Mantenimiento": "Mantenimiento Electrico|Mantenimiento Mecanico|Falla mecanica|Falla electrica|Proteccion activada|Motor detenido|Variador en falla|Tablero con temperatura elevada|Sensor fuera de servicio|Reapriete de burloneria|Cambio de rodamiento|Cambio de componente critico|Otros",
}
LAMINACION_SECTORES = list(LAMINACION_PROBLEMAS)
LAMINACION_STANDS = [f"Stand {numero}" for numero in range(1, 13)]


def _laminacion_datetime(value):
    if not value:
        return datetime.utcnow()
    try:
        return datetime.fromisoformat(value)
    except ValueError:
        return datetime.utcnow()


def _laminacion_in_turno(fecha_hora, turno):
    if turno == "todos":
        return True
    if turno == "A":
        return 6 <= fecha_hora.hour < 14
    if turno == "B":
        return 14 <= fecha_hora.hour < 22
    return fecha_hora.hour >= 22 or fecha_hora.hour < 6


def _laminacion_turno_actual():
    ahora = datetime.utcnow()
    if 6 <= ahora.hour < 14:
        inicio = ahora.replace(hour=6, minute=0, second=0, microsecond=0)
    elif 14 <= ahora.hour < 22:
        inicio = ahora.replace(hour=14, minute=0, second=0, microsecond=0)
    else:
        inicio = ahora.replace(hour=22, minute=0, second=0, microsecond=0)
        if ahora.hour < 6:
            inicio -= timedelta(days=1)
    return inicio, ahora


def _laminacion_access(view):
    return role_required("admin", "root", "operador_horno", "operador_despunte")(view)


@app.route("/admin/laminacion", methods=["GET", "POST"])
@_laminacion_access
def laminacion_dashboard():
    if request.method == "POST":
        accion = request.form.get("accion", "")
        if accion == "iniciar_produccion":
            abierta = LaminacionProduccion.query.filter_by(estado="abierta").first()
            if abierta:
                flash(f"La colada {abierta.numero_colada} sigue abierta. Cierrela antes de iniciar otra.", "error")
            else:
                db.session.add(LaminacionProduccion(
                    numero_colada=(request.form.get("numero_colada") or "").strip(),
                    cantidad_palanquillas=request.form.get("cantidad_palanquillas", 0, type=int),
                    carbono=(request.form.get("carbono") or "").strip(),
                    manganeso=(request.form.get("manganeso") or "").strip(),
                    silicio=(request.form.get("silicio") or "").strip(),
                    temperatura=request.form.get("temperatura", 0, type=float),
                    velocidad=request.form.get("velocidad", 0, type=float),
                    creado_por_id=current_user.id,
                ))
                db.session.commit()
                flash("Produccion iniciada.", "success")
        elif accion == "cerrar_produccion":
            produccion = LaminacionProduccion.query.get_or_404(request.form.get("produccion_id", type=int))
            produccion.estado = "cerrada"
            produccion.fin = datetime.utcnow()
            db.session.commit()
            flash("Colada cerrada; puede iniciar la siguiente.", "success")
        elif accion == "despunte":
            produccion = LaminacionProduccion.query.get_or_404(request.form.get("produccion_id", type=int))
            db.session.add(LaminacionDespunte(produccion_id=produccion.id, pasaron=request.form.get("pasaron", 0, type=int), chatarra=request.form.get("chatarra", 0, type=int), observacion=(request.form.get("observacion") or "").strip(), creado_por_id=current_user.id))
            db.session.commit()
            flash("Despunte registrado.", "success")
        elif accion == "parada":
            sector = (request.form.get("sector") or "").strip()
            problema = (request.form.get("problema") or "").strip()
            if sector == "Stand":
                problema = f"{request.form.get('stand') or 'Stand sin seleccionar'} - {problema}"
            db.session.add(LaminacionParada(produccion_id=request.form.get("produccion_id", type=int) or None, accion=request.form.get("tipo_accion") or "parada", sector=sector, problema=problema, inicio=_laminacion_datetime(request.form.get("inicio")), chatarra=request.form.get("chatarra_parada", 0, type=int), observacion=(request.form.get("observacion") or "").strip(), creado_por_id=current_user.id))
            db.session.commit()
            flash("Parada registrada.", "success")
        elif accion == "cerrar_parada":
            parada = LaminacionParada.query.get_or_404(request.form.get("parada_id", type=int))
            parada.fin = _laminacion_datetime(request.form.get("fin"))
            if parada.fin < parada.inicio:
                flash("La finalizacion no puede ser anterior al inicio.", "error")
                return redirect(url_for("laminacion_dashboard"))
            db.session.commit()
            flash("Parada finalizada.", "success")
        return redirect(url_for("laminacion_dashboard"))
    produccion = LaminacionProduccion.query.filter_by(estado="abierta").order_by(LaminacionProduccion.inicio.desc()).first()
    producciones = LaminacionProduccion.query.order_by(LaminacionProduccion.inicio.desc()).limit(30).all()
    produccion_referencia = produccion or (producciones[0] if producciones else None)
    paradas_abiertas = LaminacionParada.query.filter_by(fin=None).order_by(LaminacionParada.inicio.desc()).all()
    turno_inicio, turno_fin = _laminacion_turno_actual()
    despuntes_turno = LaminacionDespunte.query.filter(LaminacionDespunte.fecha >= turno_inicio, LaminacionDespunte.fecha <= turno_fin).all()
    return render_template("laminacion_dashboard.html", produccion=produccion, produccion_referencia=produccion_referencia, producciones=producciones, paradas_abiertas=paradas_abiertas, sectores=LAMINACION_SECTORES, problemas=LAMINACION_PROBLEMAS, stands=LAMINACION_STANDS, turno_inicio=turno_inicio, turno_pasaron=sum(item.pasaron for item in despuntes_turno), turno_chatarra=sum(item.chatarra for item in despuntes_turno))


@app.get("/admin/laminacion/reporte")
@_laminacion_access
def laminacion_reporte():
    fecha = request.args.get("fecha") or datetime.utcnow().strftime("%Y-%m-%d")
    turno = request.args.get("turno") or "todos"
    inicio = datetime.strptime(fecha, "%Y-%m-%d")
    fin = inicio + timedelta(days=1)
    producciones = LaminacionProduccion.query.filter(LaminacionProduccion.inicio >= inicio, LaminacionProduccion.inicio < fin).order_by(LaminacionProduccion.inicio.asc()).all()
    paradas = LaminacionParada.query.filter(LaminacionParada.inicio >= inicio, LaminacionParada.inicio < fin).order_by(LaminacionParada.inicio.asc()).all()
    producciones = [item for item in producciones if _laminacion_in_turno(item.inicio, turno)]
    paradas = [item for item in paradas if _laminacion_in_turno(item.inicio, turno)]
    return render_template("laminacion_reporte.html", fecha=fecha, turno=turno, producciones=producciones, paradas=paradas)


@app.route("/horno-laminacion", methods=["GET", "POST"])
@admin_required
def horno_laminacion():
    if request.method == "POST":
        action = request.form.get("action", "")
        if action == "confirmar_colada":
            colada = (request.form.get("colada") or "").strip()
            if colada and EtiquetaCalidad.query.filter_by(tipo_producto="Palanquillas", colada=colada).first():
                set_config_value("horno_colada_activa", colada)
                flash(f"Colada {colada} cargada en el horno.", "success")
            else:
                flash("No se encontraron etiquetas de Palanquillas para esa colada.", "error")
        elif action == "cerrar_carga":
            set_config_value("horno_colada_activa", "")
            flash("Colada finalizada. Seleccione la nueva colada para continuar.", "success")
        return redirect(url_for("horno_laminacion"))
    etiquetas = EtiquetaCalidad.query.filter_by(tipo_producto="Palanquillas").order_by(EtiquetaCalidad.created_at.desc()).all()
    grupos = {}
    for etiqueta in etiquetas:
        grupos.setdefault(etiqueta.colada, []).append(etiqueta)
    cargas = []
    for colada, grupo in grupos.items():
        cantidad = sum(item.cantidad or 0 for item in grupo)
        peso_toneladas = 0.0
        for item in grupo:
            try:
                peso_toneladas += float(str(item.peso or "0").replace(",", ".")) * (item.cantidad or 0) / 1000
            except (TypeError, ValueError):
                pass
        cargas.append(SimpleNamespace(
            id=grupo[0].id,
            numero_colada=colada,
            cantidad_palanquillas=cantidad,
            tonelada=peso_toneladas,
            velocidad_horno=0.0,
            temperatura_horno=0.0,
            estado="activa" if len(cargas) == 0 else "cerrada",
            cantidad_enviada=0,
            fecha_inicio=grupo[-1].created_at,
            fecha_fin=None,
        ))
    config_colada_activa = Configuracion.query.filter_by(clave="horno_colada_activa").first()
    colada_activa = config_colada_activa.valor if config_colada_activa else None
    carga_actual = next((item for item in cargas if item.numero_colada == colada_activa), None) if colada_activa else (cargas[0] if colada_activa is None and cargas else None)
    return render_template(
        "horno_laminacion.html",
        carga_actual=carga_actual,
        ajustes_display={"velocidad": None, "temperatura": None},
        registros=cargas,
        laminadas=0,
        sobrantes=0,
        horno_display=False,
    )


@app.get("/admin/horno-laminacion/consultar-colada")
@admin_required
def consultar_colada_horno():
    metodo = request.args.get("metodo", "colada")
    valor = (request.args.get("etiqueta") if metodo == "etiqueta" else request.args.get("colada") or "").strip()
    etiqueta = None
    if metodo == "etiqueta":
        etiqueta_id = int(re.sub(r"\D", "", valor) or 0)
        etiqueta = EtiquetaCalidad.query.filter_by(id=etiqueta_id, tipo_producto="Palanquillas").first()
    else:
        etiqueta = EtiquetaCalidad.query.filter_by(colada=valor, tipo_producto="Palanquillas").order_by(EtiquetaCalidad.created_at.desc()).first()
    if not etiqueta:
        return jsonify({"ok": False, "error": "No se encontraron Palanquillas para ese dato."}), 404
    etiquetas = EtiquetaCalidad.query.filter_by(colada=etiqueta.colada, tipo_producto="Palanquillas").all()
    return jsonify({
        "ok": True,
        "colada": etiqueta.colada,
        "palanquillas": sum(item.cantidad or 0 for item in etiquetas),
        "grado_carbono": etiqueta.carbono or "",
        "longitud": etiqueta.longitud or "",
    })


@app.route("/syso")
@admin_required
def syso_dashboard():
    return render_template("syso_dashboard.html")


@app.route("/root/panel-control")
@root_required
def root_panel():
    return render_template("root_panel.html")


@app.route("/admin/qality", methods=["GET", "POST"])
@calidad_required
def quality_dashboard():
    if request.method == "POST":
        tipo_producto = (request.form.get("tipo_producto") or "").strip()
        calidad = (request.form.get("calidad") or "").strip()
        medida = (request.form.get("medida") or "").strip()
        longitud = (request.form.get("longitud") or "").strip()
        peso = (request.form.get("peso") or "").strip()
        colada = (request.form.get("colada") or "").strip()
        lote = (request.form.get("lote") or "").strip()
        cantidad = request.form.get("cantidad", "").strip()
        fecha_value = request.form.get("fecha", "").strip()
        carbono = (request.form.get("carbono") or "").strip()
        silicio = (request.form.get("silicio") or "").strip()
        manganeso = (request.form.get("manganeso") or "").strip()
        hornero = (request.form.get("hornero") or "").strip()
        supervisor = (request.form.get("supervisor") or "").strip()
        operador_ccm = (request.form.get("operador_ccm") or "").strip()
        aprobado = request.form.get("aprobado") == "on"
        tipo_b = request.form.get("tipo_b") == "on"
        no_conforme = request.form.get("no_conforme") == "on"
        observaciones = (request.form.get("observaciones") or "").strip()

        valid_types = {"AP500S", "VarillasLisas", "Palanquillas"}
        valid_palanquilla_qualities = {"AP500S", "Inoxidable", "A36"}
        if tipo_producto not in valid_types:
            flash("Seleccione un tipo de producto válido.", "error")
            return redirect(url_for("quality_dashboard"))
        faltantes_iniciales = [campo for campo, valor in (("Calidad", calidad), ("Número de colada", colada)) if not valor]
        if faltantes_iniciales:
            flash("Faltan estos datos: " + ", ".join(faltantes_iniciales) + ".", "error")
            return redirect(url_for("quality_dashboard", show="nueva"))

        if tipo_producto == "Palanquillas":
            try:
                cantidad_value = int(cantidad) if cantidad else 0
            except ValueError:
                flash("La cantidad de Palanquillas debe ser un número entero.", "error")
                return redirect(url_for("quality_dashboard"))
            if calidad not in valid_palanquilla_qualities:
                flash("Seleccione una calidad válida para Palanquillas.", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))
            faltantes = []
            for campo, valor in (("Longitud", longitud), ("Peso", peso), ("Lote", lote), ("Carbono", carbono), ("Silicio", silicio), ("Manganeso", manganeso), ("Operador de Horno", hornero), ("Supervisor", supervisor), ("Operador de CCM", operador_ccm)):
                if not valor:
                    faltantes.append(campo)
            if faltantes:
                flash("Faltan estos datos: " + ", ".join(faltantes) + ".", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))
            try:
                float(longitud)
                float(peso)
                float(carbono)
                float(silicio)
                float(manganeso)
            except ValueError:
                flash("Longitud, peso y composición química deben ser numéricos.", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))
            medida = ""
            try:
                fecha = datetime.strptime(fecha_value, "%Y-%m-%d").date() if fecha_value else datetime.utcnow().date()
            except ValueError:
                flash("La fecha no tiene un formato válido.", "error")
                return redirect(url_for("quality_dashboard"))
        else:
            if calidad not in {"AP500S", "A36"}:
                flash("Seleccione una calidad válida: AP500 S o A36 · Varillas Lisas.", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))
            cantidad_value = None
            lote = ""
            fecha = None
            carbono = ""
            silicio = ""
            manganeso = ""
            faltantes = []
            for campo, valor in (("Calidad", calidad), ("Medida", medida), ("Longitud", longitud), ("Peso", peso)):
                if not valor:
                    faltantes.append(campo)
            if faltantes:
                flash("Faltan estos datos: " + ", ".join(faltantes) + ".", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))
            aprobado = False
            if tipo_b and no_conforme:
                flash("Seleccione solo Tipo B o No conforme, no ambas opciones.", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))
            if medida not in QUALITY_ROD_DIAMETERS:
                flash("Seleccione una medida válida: 8, 10, 12, 16, 20 o 25 mm.", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))
            if not medida or not longitud or not peso:
                faltantes = [campo for campo, valor in (("Medida", medida), ("Longitud", longitud), ("Peso", peso)) if not valor]
                flash("Faltan estos datos: " + ", ".join(faltantes) + ".", "error")
                return redirect(url_for("quality_dashboard", show="nueva"))

        etiqueta = EtiquetaCalidad(
            tipo_producto=tipo_producto,
            calidad=calidad,
            medida=medida,
            longitud=longitud,
            peso=peso,
            colada=colada,
            lote=lote,
            cantidad=cantidad_value,
            fecha=fecha,
            carbono=carbono,
            silicio=silicio,
            manganeso=manganeso,
            hornero=hornero,
            supervisor=supervisor,
            operador_ccm=operador_ccm,
            aprobado=aprobado,
            tipo_b=tipo_b,
            no_conforme=no_conforme,
            observaciones=observaciones,
            created_by_id=current_user.id,
        )
        db.session.add(etiqueta)
        db.session.commit()
        if tipo_producto == "Palanquillas":
            _regenerate_rod_labels_for_colada(colada)
        pdf_path = QUALITY_LABELS_DIR / f"etiqueta_{etiqueta.id:07d}.pdf"
        try:
            _draw_quality_label_pdf(etiqueta, pdf_path)
        except Exception:
            app.logger.exception("No se pudo generar el PDF de la etiqueta %s", etiqueta.id)
            flash("Etiqueta guardada, pero no se pudo generar el PDF.", "warning")
            return redirect(url_for("quality_dashboard"))
        flash(f"Etiqueta ETQ-{etiqueta.id:07d} guardada y PDF generado.", "success")
        return redirect(url_for("vista_previa_etiqueta_qality", etiqueta_id=etiqueta.id))

    db.create_all()
    dirs, warning = ensure_quality_certificate_dirs()
    responsables = ResponsableCalidad.query.filter_by(activo=True).order_by(ResponsableCalidad.tipo.asc(), ResponsableCalidad.nombre.asc()).all()
    query = (request.args.get("q") or "").strip()
    certificados = list_quality_certificate_files(dirs)
    if query:
        query_lower = query.lower()
        certificados = [
            cert for cert in certificados
            if query_lower in cert["name"].lower()
        ]
    etiquetas = EtiquetaCalidad.query.order_by(EtiquetaCalidad.created_at.desc()).all()
    show_label = request.args.get("show") == "nueva"
    pdf_id = request.args.get("pdf_id", type=int)
    if not show_label and not pdf_id:
        return render_template("qality_dashboard.html")
    return render_template(
        "control_calidad.html",
        certificados=certificados,
        carpetas=dirs,
        warning=warning,
        filtro=query,
        etiquetas=etiquetas,
        pdf_id=pdf_id,
        show_label=show_label,
        responsables=responsables,
    )


@app.get("/admin/qality/existencias")
@calidad_required
def existencias_calidad():
    dispatched_ids = db.session.query(DetalleDespachoAcero.etiqueta_id).subquery()
    labels = (
        EtiquetaCalidad.query
        .filter(EtiquetaCalidad.no_conforme.is_(False), ~EtiquetaCalidad.id.in_(dispatched_ids))
        .order_by(EtiquetaCalidad.created_at.desc())
        .all()
    )
    varillas = {}
    palanquillas = {}
    for label in labels:
        if label.tipo_producto in {"AP500S", "VarillasLisas"}:
            key = (label.tipo_producto, label.medida or "-", label.longitud or "-")
            row = varillas.setdefault(key, {"producto": "AP 500 S" if label.tipo_producto == "AP500S" else "Varillas Lisas", "medida": label.medida or "-", "longitud": label.longitud or "-", "atados": 0, "peso": 0.0, "details": []})
            row["atados"] += 1
            row["details"].append({"id": label.id, "etiqueta": f"ETQ-{label.id:07d}", "colada": label.colada or "-", "calidad": label.calidad or "-", "medida": label.medida or "-", "longitud": label.longitud or "-", "peso": float(str(label.peso or "0").replace(",", ".")), "fecha": label.fecha.strftime("%d/%m/%Y") if label.fecha else "-", "created_at": label.created_at.strftime("%d/%m/%Y %H:%M") if label.created_at else "-", "lote": label.lote or "-"})
            try:
                row["peso"] += float(str(label.peso or "0").replace(",", "."))
            except ValueError:
                pass
        elif label.tipo_producto == "Palanquillas":
            key = (label.longitud or "-", label.calidad or "-")
            row = palanquillas.setdefault(key, {"longitud": label.longitud or "-", "calidad": label.calidad or "-", "etiquetas": 0, "piezas": 0, "peso": 0.0, "details": []})
            row["etiquetas"] += 1
            row["piezas"] += label.cantidad or 0
            row["details"].append({"id": label.id, "etiqueta": f"ETQ-{label.id:07d}", "colada": label.colada or "-", "calidad": label.calidad or "-", "longitud": label.longitud or "-", "cantidad": label.cantidad or 0, "peso": float(str(label.peso or "0").replace(",", ".")), "fecha": label.fecha.strftime("%d/%m/%Y") if label.fecha else "-", "created_at": label.created_at.strftime("%d/%m/%Y %H:%M") if label.created_at else "-", "lote": label.lote or "-"})
            try:
                row["peso"] += float(str(label.peso or "0").replace(",", ".")) * (label.cantidad or 0)
            except ValueError:
                pass
    return render_template("existencias_calidad.html", varillas=sorted(varillas.values(), key=lambda row: (row["producto"], row["medida"], row["longitud"])), palanquillas=sorted(palanquillas.values(), key=lambda row: (row["longitud"], row["calidad"])))


@app.route("/admin/qality/responsables", methods=["GET", "POST"])
@calidad_required
def responsables_calidad():
    tipos_validos = ("Operador de Horno", "Operador de CCM", "Supervisor")
    if request.method == "POST":
        nombre = (request.form.get("nombre") or "").strip()
        tipo = (request.form.get("tipo") or "").strip()
        if not nombre or tipo not in tipos_validos:
            flash("Ingrese un nombre y un tipo de responsable válido.", "error")
        elif ResponsableCalidad.query.filter_by(nombre=nombre, tipo=tipo).first():
            flash("Ese responsable ya está cargado para ese campo.", "warning")
        else:
            db.session.add(ResponsableCalidad(nombre=nombre, tipo=tipo))
            db.session.commit()
            flash("Responsable guardado correctamente.", "success")
        return redirect(url_for("responsables_calidad"))
    responsables = ResponsableCalidad.query.order_by(ResponsableCalidad.tipo.asc(), ResponsableCalidad.nombre.asc()).all()
    return render_template("responsables_calidad.html", responsables=responsables)


@app.post("/admin/qality/responsables/<int:responsable_id>/eliminar")
@calidad_required
def eliminar_responsable_calidad(responsable_id: int):
    responsable = ResponsableCalidad.query.get_or_404(responsable_id)
    responsable.activo = False
    db.session.commit()
    flash("Responsable retirado de las listas nuevas.", "success")
    return redirect(url_for("responsables_calidad"))


def _quality_shift_summary(etiquetas):
    resumen = []
    grupos = {}
    for etiqueta in etiquetas:
        usuario = etiqueta.created_by.full_name if etiqueta.created_by else "Sin usuario"
        grupos.setdefault(usuario, []).append(etiqueta)
    for usuario, filas in sorted(grupos.items(), key=lambda item: (-len(item[1]), item[0].lower())):
        resumen.append({
            "usuario": usuario,
            "etiquetas": len(filas),
            "piezas": sum(item.cantidad or 0 for item in filas),
            "conformes": sum(1 for item in filas if not item.no_conforme and not item.tipo_b),
            "tipo_b": sum(1 for item in filas if item.tipo_b),
            "no_conformes": sum(1 for item in filas if item.no_conforme),
        })
    return resumen


def _quality_label_status(etiqueta):
    if etiqueta.no_conforme:
        return "No conforme"
    if etiqueta.tipo_b:
        return "Tipo B"
    return "Conforme"


def _quality_label_operator(etiqueta):
    return etiqueta.operador_ccm or (etiqueta.created_by.full_name if etiqueta.created_by else "Sin operador")


def _quality_label_shift(fecha_hora):
    if 6 <= fecha_hora.hour < 14:
        return "A"
    if 14 <= fecha_hora.hour < 22:
        return "B"
    return "C"


def _build_quality_shift_pdf(etiquetas, resumen, fecha_desde, fecha_hasta, grupo):
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("quality_shift_title", parent=styles["Title"], fontSize=17, leading=21, alignment=1, spaceAfter=12)
    body = ParagraphStyle("quality_shift_body", parent=styles["BodyText"], fontSize=9, leading=12)
    header = ParagraphStyle("quality_shift_header", parent=body, fontName="Helvetica-Bold", textColor=colors.white)
    data = [[Paragraph("Usuario", header), Paragraph("Etiquetas", header), Paragraph("Piezas", header), Paragraph("Conformes", header), Paragraph("Tipo B", header), Paragraph("No conformes", header)]]
    for fila in resumen:
        data.append([fila["usuario"], str(fila["etiquetas"]), str(fila["piezas"]), str(fila["conformes"]), str(fila["tipo_b"]), str(fila["no_conformes"])])
    table = Table(data, repeatRows=1, colWidths=[170, 65, 65, 75, 55, 85])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155e75")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ALIGN", (1, 1), (-1, -1), "CENTER"),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("PADDING", (0, 0), (-1, -1), 5),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef6f8")]),
    ]))
    def detail_table(title, rows, include_diameter):
        columns = ["N° etiqueta", "Colada", "Tipo", "Peso (kg)", "Estado", "Operador", "Turno"]
        if include_diameter:
            columns.insert(3, "Diámetro")
        detail_data = [[Paragraph(label, header) for label in columns]]
        for etiqueta in rows:
            values = [f"ETQ-{etiqueta.id:07d}", etiqueta.colada or "-", "AP500 S" if etiqueta.tipo_producto == "AP500S" else "Varillas Lisas" if etiqueta.tipo_producto == "VarillasLisas" else "Palanquillas", etiqueta.peso or "-", _quality_label_status(etiqueta), _quality_label_operator(etiqueta), _quality_label_shift(etiqueta.created_at)]
            if include_diameter:
                values.insert(3, f"{etiqueta.medida or '-'} mm")
            detail_data.append([Paragraph(escape(str(value)), body) for value in values])
        widths = [78, 145, 82, 72, 88, 180, 48] if include_diameter else [78, 175, 105, 88, 190, 48]
        table = Table(detail_data, repeatRows=1, colWidths=widths)
        table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#155e75")),
            ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("FONTSIZE", (0, 0), (-1, -1), 8),
            ("PADDING", (0, 0), (-1, -1), 5),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef6f8")]),
        ]))
        return [Paragraph(title, ParagraphStyle(f"{title}_heading", parent=styles["Heading2"], fontSize=11, leading=14, textColor=colors.HexColor("#155e75"))), table]

    palanquillas = [item for item in etiquetas if item.tipo_producto == "Palanquillas"]
    varillas = [item for item in etiquetas if item.tipo_producto in {"AP500S", "VarillasLisas"}]
    total_piezas = sum(item.cantidad or 0 for item in etiquetas)
    elements = [
        Paragraph("REPORTE DE CIERRE DE TURNO - CONTROL DE CALIDAD", title),
        Paragraph(f"Grupo/turno: {escape(grupo or 'Sin especificar')}", body),
        Paragraph(f"Periodo: {fecha_desde.strftime('%d/%m/%Y %H:%M')} a {fecha_hasta.strftime('%d/%m/%Y %H:%M')}", body),
        Paragraph(f"Total de etiquetas: {len(etiquetas)} &nbsp;&nbsp; Total de piezas: {total_piezas} &nbsp;&nbsp; Usuarios: {len(resumen)}", body),
        Spacer(1, 12),
        table,
        Spacer(1, 14),
        *detail_table("DETALLE DE PALANQUILLAS", palanquillas, False),
        Spacer(1, 12),
        *detail_table("DETALLE DE VARILLAS / LAMINACIÓN", varillas, True),
        Spacer(1, 14),
        Paragraph(f"Generado: {datetime.utcnow().strftime('%d/%m/%Y %H:%M:%S')} UTC", body),
    ]
    doc.build(elements)
    return buffer.getvalue()


def _build_quality_shift_excel(etiquetas, resumen, fecha_desde, fecha_hasta, grupo) -> bytes:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary.append(["Reporte de cierre de turno", grupo])
    summary.append(["Fecha desde", fecha_desde.strftime("%d/%m/%Y %H:%M")])
    summary.append(["Fecha hasta", fecha_hasta.strftime("%d/%m/%Y %H:%M")])
    summary.append([])
    summary.append(["Usuario", "Etiquetas", "Piezas", "Conformes", "Tipo B", "No conformes"])
    for cell in summary[5]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="155e75")
    for row in resumen:
        summary.append([row["usuario"], row["etiquetas"], row["piezas"], row["conformes"], row["tipo_b"], row["no_conformes"]])
    for column, width in {"A": 28, "B": 14, "C": 12, "D": 14, "E": 12, "F": 16}.items():
        summary.column_dimensions[column].width = width
    columns = ["Área", "Producto", "Etiqueta", "Colada", "Calidad", "Medida", "Longitud", "Peso kg", "Cantidad", "Fecha y hora", "Turno", "Estado", "Operador CCM"]
    for sheet_name, sheet_rows in (("Aceria Palanquillas", [item for item in etiquetas if item.tipo_producto == "Palanquillas"]), ("Laminacion Varillas", [item for item in etiquetas if item.tipo_producto in {"AP500S", "VarillasLisas"}])):
        detail = workbook.create_sheet(sheet_name)
        detail.append(columns)
        for cell in detail[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="155e75")
        for etiqueta in sheet_rows:
            area = "Acería" if etiqueta.tipo_producto == "Palanquillas" else "Laminación"
            producto = "Palanquillas" if etiqueta.tipo_producto == "Palanquillas" else "AP500 S" if etiqueta.tipo_producto == "AP500S" else "Varillas Lisas"
            detail.append([area, producto, f"ETQ-{etiqueta.id:07d}", etiqueta.colada, etiqueta.calidad, etiqueta.medida or "", etiqueta.longitud or "", float(str(etiqueta.peso or "0").replace(",", ".")), etiqueta.cantidad or "", etiqueta.created_at, _quality_label_shift(etiqueta.created_at), _quality_label_status(etiqueta), etiqueta.operador_ccm or ""])
        for column, width in zip("ABCDEFGHIJKLM", [14, 18, 14, 16, 16, 12, 14, 12, 12, 20, 10, 16, 20]):
            detail.column_dimensions[column].width = width
        detail.freeze_panes = "A2"
        detail.auto_filter.ref = detail.dimensions
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


@app.route("/admin/qality/cierre-turno", methods=["GET", "POST"])
@calidad_required
def reporte_cierre_turno_qality():
    ahora = datetime.utcnow().replace(second=0, microsecond=0)
    fecha_desde = ahora.replace(hour=6, minute=0)
    fecha_hasta = ahora
    turno = "Global"
    area = "Global"
    if request.method == "POST":
        fecha_input = (request.form.get("fecha_desde") or "").strip()
        try:
            fecha_desde = datetime.strptime(fecha_input, "%Y-%m-%dT%H:%M") if "T" in fecha_input else datetime.strptime(fecha_input, "%Y-%m-%d")
            fecha_hasta = fecha_desde
        except ValueError:
            flash("Indique fechas válidas para el inicio y el fin del turno.", "error")
            return redirect(url_for("reporte_cierre_turno_qality"))
        turno = (request.form.get("turno") or "").strip()
        area = (request.form.get("area") or "").strip()
        areas = {"Aceria": {"Palanquillas"}, "Laminacion": {"AP500S", "VarillasLisas"}, "Global": {"Palanquillas", "AP500S", "VarillasLisas"}}
        if area not in areas:
            flash("Seleccione un área válida: ACERÍA, LAMINACIÓN o Global.", "error")
            return redirect(url_for("reporte_cierre_turno_qality"))
        turnos = {"A": (6, 14), "B": (14, 22), "C": (22, 6), "Global": (6, 6)}
        if turno not in turnos:
            flash("Seleccione un turno válido: A, B, C o Global.", "error")
            return redirect(url_for("reporte_cierre_turno_qality"))
        base_date = fecha_desde.date()
        start_hour, end_hour = turnos[turno]
        fecha_desde = datetime.combine(base_date, datetime.min.time()).replace(hour=start_hour)
        end_date = base_date + timedelta(days=1) if end_hour <= start_hour else base_date
        fecha_hasta = datetime.combine(end_date, datetime.min.time()).replace(hour=end_hour)
        if fecha_desde > fecha_hasta:
            flash("La fecha de inicio no puede ser posterior a la fecha de fin.", "error")
        else:
            etiquetas = EtiquetaCalidad.query.filter(
                EtiquetaCalidad.created_at >= fecha_desde,
                EtiquetaCalidad.created_at < fecha_hasta,
                EtiquetaCalidad.tipo_producto.in_(areas[area]),
            ).order_by(EtiquetaCalidad.created_at.asc()).all()
            resumen = _quality_shift_summary(etiquetas)
            grupo = f"{area} · {turno}"
            if request.form.get("formato") == "excel":
                excel = _build_quality_shift_excel(etiquetas, resumen, fecha_desde, fecha_hasta, grupo)
                return send_file(BytesIO(excel), mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"cierre_{area.lower()}_{turno.lower()}_{fecha_desde:%Y%m%d}.xlsx")
            pdf = _build_quality_shift_pdf(etiquetas, resumen, fecha_desde, fecha_hasta, grupo)
            nombre = f"cierre_{area.lower()}_{turno.lower()}_{fecha_desde:%Y%m%d_%H%M}_{fecha_hasta:%Y%m%d_%H%M}.pdf"
            return send_file(BytesIO(pdf), mimetype="application/pdf", as_attachment=True, download_name=nombre)

    etiquetas = EtiquetaCalidad.query.filter(
        EtiquetaCalidad.created_at >= fecha_desde,
        EtiquetaCalidad.created_at < fecha_hasta,
        EtiquetaCalidad.tipo_producto.in_({"Palanquillas", "AP500S", "VarillasLisas"}),
    ).order_by(EtiquetaCalidad.created_at.asc()).all()
    return render_template(
        "cierre_turno_qality.html",
        fecha_desde=fecha_desde.strftime("%Y-%m-%dT%H:%M"),
        fecha_hasta=fecha_hasta.strftime("%Y-%m-%dT%H:%M"),
        turno=turno,
        area=area,
        etiquetas=etiquetas,
        resumen=_quality_shift_summary(etiquetas),
        total_etiquetas=len(etiquetas),
    )


def _producto_terminado_turno(fecha_hora):
    hora = fecha_hora.hour
    if 6 <= hora < 14:
        return "Mañana"
    if 14 <= hora < 22:
        return "Tarde"
    return "Noche"


def _producto_terminado_estado(etiqueta):
    if etiqueta.no_conforme:
        return "Retenido"
    if etiqueta.tipo_b:
        return "Tipo B"
    if etiqueta.aprobado:
        return "Aprobado"
    return "Pendiente"


def _producto_terminado_query(source):
    tipo = (source.get("tipo") or "").strip()
    medida = (source.get("medida") or "").strip()
    turno = (source.get("turno") or "").strip()
    estado = (source.get("estado") or "").strip()
    fecha_desde = (source.get("fecha_desde") or "").strip()
    fecha_hasta = (source.get("fecha_hasta") or "").strip()
    query = EtiquetaCalidad.query.filter(EtiquetaCalidad.tipo_producto.in_(["AP500S", "VarillasLisas"]))
    if tipo in {"AP500S", "VarillasLisas"}:
        query = query.filter(EtiquetaCalidad.tipo_producto == tipo)
    else:
        tipo = ""
    if medida:
        query = query.filter(EtiquetaCalidad.medida == medida)
    if fecha_desde:
        try:
            query = query.filter(EtiquetaCalidad.created_at >= datetime.strptime(fecha_desde, "%Y-%m-%d"))
        except ValueError:
            fecha_desde = ""
    if fecha_hasta:
        try:
            limite = datetime.strptime(fecha_hasta, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(EtiquetaCalidad.created_at < limite)
        except ValueError:
            fecha_hasta = ""
    etiquetas = query.order_by(EtiquetaCalidad.created_at.desc()).all()
    if turno in {"Mañana", "Tarde", "Noche"}:
        etiquetas = [item for item in etiquetas if _producto_terminado_turno(item.created_at) == turno]
    else:
        turno = ""
    if estado in {"Aprobado", "Retenido", "Tipo B", "Pendiente"}:
        etiquetas = [item for item in etiquetas if _producto_terminado_estado(item) == estado]
    else:
        estado = ""
    return etiquetas, {"tipo": tipo, "medida": medida, "turno": turno, "estado": estado, "fecha_desde": fecha_desde, "fecha_hasta": fecha_hasta}


def _producto_terminado_rows(etiquetas):
    return [{
        "codigo": f"ETQ-{item.id:07d}",
        "id": item.id,
        "tipo": item.tipo_producto,
        "calidad": item.calidad,
        "medida": item.medida,
        "longitud": item.longitud,
        "peso": item.peso,
        "colada": item.colada,
        "fecha": item.created_at.strftime("%d/%m/%Y %H:%M"),
        "turno": _producto_terminado_turno(item.created_at),
        "estado": _producto_terminado_estado(item),
        "usuario": item.created_by.full_name if item.created_by else "Sin usuario",
    } for item in etiquetas]


def _build_producto_terminado_pdf(etiquetas, filtros):
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(letter), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("producto_terminado_title", parent=styles["Title"], fontSize=17, leading=20, alignment=1, spaceAfter=10)
    body = ParagraphStyle("producto_terminado_body", parent=styles["BodyText"], fontSize=8, leading=10)
    header = ParagraphStyle("producto_terminado_header", parent=body, fontName="Helvetica-Bold", textColor=colors.white, alignment=1)
    rows = _producto_terminado_rows(etiquetas)
    columns = ["Código", "Tipo", "Calidad", "mm", "Longitud", "Peso", "Colada", "Registro", "Turno", "Estado", "Usuario"]
    data = [[Paragraph(escape(column), header) for column in columns]]
    for row in rows:
        data.append([row[key] for key in ["codigo", "tipo", "calidad", "medida", "longitud", "peso", "colada", "fecha", "turno", "estado", "usuario"]])
    table = Table(data, repeatRows=1, colWidths=[58, 62, 62, 30, 55, 55, 65, 88, 48, 62, 105])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("ALIGN", (3, 1), (9, -1), "CENTER"),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("PADDING", (0, 0), (-1, -1), 4),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#eef8f7")]),
    ]))
    total = len(etiquetas)
    aprobadas = sum(1 for item in etiquetas if _producto_terminado_estado(item) == "Aprobado")
    retenidas = sum(1 for item in etiquetas if _producto_terminado_estado(item) == "Retenido")
    tipo_b = sum(1 for item in etiquetas if _producto_terminado_estado(item) == "Tipo B")
    logo = Image(str(STATIC_LOGO_PATH), width=105, height=45)
    filters_text = f"Tipo: {filtros['tipo'] or 'Todos'} | Medida: {filtros['medida'] or 'Todas'} | Turno: {filtros['turno'] or 'Todos'} | Estado: {filtros['estado'] or 'Todos'} | Desde: {filtros['fecha_desde'] or 'Todas'} | Hasta: {filtros['fecha_hasta'] or 'Todas'}"
    elements = [
        Table([[logo, Paragraph("EXISTENCIAS DE PRODUCTO TERMINADO", title)]], colWidths=[130, 630], style=[("VALIGN", (0, 0), (-1, -1), "MIDDLE")]),
        Paragraph(filters_text, body),
        Paragraph(f"Total: {total} | Aprobadas: {aprobadas} | Retenidas: {retenidas} | Tipo B: {tipo_b} | Generado: {datetime.utcnow():%d/%m/%Y %H:%M} UTC", body),
        Spacer(1, 10),
        table,
    ]
    doc.build(elements)
    return buffer.getvalue()


@app.route("/admin/producto-terminado")
@calidad_required
def producto_terminado():
    etiquetas, filtros = _producto_terminado_query(request.args)
    return render_template("producto_terminado.html", filas=_producto_terminado_rows(etiquetas), filtros=filtros, total=len(etiquetas))


@app.route("/admin/producto-terminado/pdf")
@calidad_required
def producto_terminado_pdf():
    etiquetas, filtros = _producto_terminado_query(request.args)
    pdf = _build_producto_terminado_pdf(etiquetas, filtros)
    return send_file(BytesIO(pdf), mimetype="application/pdf", as_attachment=True, download_name=f"existencias_producto_terminado_{datetime.utcnow():%Y%m%d_%H%M%S}.pdf")


@app.route("/admin/qality/etiqueta/<int:etiqueta_id>/pdf")
@calidad_required
def descargar_etiqueta_qality(etiqueta_id: int):
    etiqueta = EtiquetaCalidad.query.get_or_404(etiqueta_id)
    pdf_path = QUALITY_LABELS_DIR / f"etiqueta_{etiqueta.id:07d}.pdf"
    _draw_quality_label_pdf(etiqueta, pdf_path)
    response = send_file(pdf_path, mimetype="application/pdf", as_attachment=not request.args.get("preview"), download_name=pdf_path.name)
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    response.headers["Pragma"] = "no-cache"
    return response


@app.route("/admin/qality/etiqueta/<int:etiqueta_id>/vista-previa")
@calidad_required
def vista_previa_etiqueta_qality(etiqueta_id: int):
    etiqueta = EtiquetaCalidad.query.get_or_404(etiqueta_id)
    pdf_path = QUALITY_LABELS_DIR / f"etiqueta_{etiqueta.id:07d}.pdf"
    _draw_quality_label_pdf(etiqueta, pdf_path)
    return render_template("vista_previa_etiqueta.html", etiqueta=etiqueta, preview_version=int(datetime.utcnow().timestamp()))


@app.route("/admin/qality/etiqueta/<int:etiqueta_id>/editar", methods=["GET", "POST"])
@calidad_required
def editar_etiqueta_qality(etiqueta_id: int):
    etiqueta = EtiquetaCalidad.query.get_or_404(etiqueta_id)
    if request.method == "POST":
        tipo_producto = (request.form.get("tipo_producto") or "").strip()
        etiqueta.tipo_producto = tipo_producto
        etiqueta.calidad = (request.form.get("calidad") or "").strip()
        etiqueta.medida = (request.form.get("medida") or "").strip()
        etiqueta.longitud = (request.form.get("longitud") or "").strip()
        etiqueta.peso = (request.form.get("peso") or "").strip()
        etiqueta.colada = (request.form.get("colada") or "").strip()
        etiqueta.lote = (request.form.get("lote") or "").strip() if tipo_producto == "Palanquillas" else ""
        etiqueta.observaciones = (request.form.get("observaciones") or "").strip()
        etiqueta.tipo_b = request.form.get("tipo_b") == "on" if tipo_producto != "Palanquillas" else False
        etiqueta.no_conforme = request.form.get("no_conforme") == "on" if tipo_producto != "Palanquillas" else False
        if tipo_producto == "Palanquillas":
            etiqueta.medida = ""
            try:
                etiqueta.cantidad = int(request.form.get("cantidad") or 0)
                fecha_value = (request.form.get("fecha") or "").strip()
                etiqueta.fecha = datetime.strptime(fecha_value, "%Y-%m-%d").date() if fecha_value else datetime.utcnow().date()
            except ValueError:
                flash("Cantidad o fecha inválida para la Palanquilla.", "error")
                return render_template("editar_etiqueta.html", etiqueta=etiqueta, responsables=ResponsableCalidad.query.filter_by(activo=True).order_by(ResponsableCalidad.tipo.asc(), ResponsableCalidad.nombre.asc()).all())
            etiqueta.carbono = (request.form.get("carbono") or "").strip()
            etiqueta.silicio = (request.form.get("silicio") or "").strip()
            etiqueta.manganeso = (request.form.get("manganeso") or "").strip()
            etiqueta.hornero = (request.form.get("hornero") or "").strip()
            etiqueta.supervisor = (request.form.get("supervisor") or "").strip()
            etiqueta.operador_ccm = (request.form.get("operador_ccm") or "").strip()
            etiqueta.aprobado = request.form.get("aprobado") == "on"
        else:
            etiqueta.cantidad = None
            etiqueta.fecha = None
            etiqueta.carbono = etiqueta.silicio = etiqueta.manganeso = ""
            etiqueta.hornero = (request.form.get("hornero") or "").strip()
            etiqueta.supervisor = (request.form.get("supervisor") or "").strip()
            etiqueta.operador_ccm = (request.form.get("operador_ccm") or "").strip()
            etiqueta.aprobado = False
            if etiqueta.medida not in QUALITY_ROD_DIAMETERS:
                flash("Seleccione una medida válida: 8, 10, 12, 16, 20 o 25 mm.", "error")
                return render_template("editar_etiqueta.html", etiqueta=etiqueta, responsables=ResponsableCalidad.query.filter_by(activo=True).order_by(ResponsableCalidad.tipo.asc(), ResponsableCalidad.nombre.asc()).all())
            if not etiqueta.medida or not etiqueta.longitud or not etiqueta.peso:
                flash("Complete medida, longitud y peso para la etiqueta.", "error")
                return render_template("editar_etiqueta.html", etiqueta=etiqueta, responsables=ResponsableCalidad.query.filter_by(activo=True).order_by(ResponsableCalidad.tipo.asc(), ResponsableCalidad.nombre.asc()).all())
        if tipo_producto not in {"AP500S", "VarillasLisas", "Palanquillas"} or not etiqueta.calidad or not etiqueta.colada:
            flash("Complete el tipo, la calidad y la colada.", "error")
            return render_template("editar_etiqueta.html", etiqueta=etiqueta)
        db.session.commit()
        try:
            _draw_quality_label_pdf(etiqueta, QUALITY_LABELS_DIR / f"etiqueta_{etiqueta.id:07d}.pdf")
            flash("Etiqueta actualizada y PDF regenerado.", "success")
        except Exception:
            app.logger.exception("No se pudo regenerar el PDF de la etiqueta %s", etiqueta.id)
            flash("Etiqueta actualizada, pero no se pudo regenerar el PDF. Puede volver a descargarlo.", "warning")
        return redirect(url_for("vista_previa_etiqueta_qality", etiqueta_id=etiqueta.id))
    responsables = ResponsableCalidad.query.filter_by(activo=True).order_by(ResponsableCalidad.tipo.asc(), ResponsableCalidad.nombre.asc()).all()
    return render_template("editar_etiqueta.html", etiqueta=etiqueta, responsables=responsables)


@app.post("/admin/qality/etiqueta/<int:etiqueta_id>/eliminar")
@calidad_required
def eliminar_etiqueta_qality(etiqueta_id: int):
    etiqueta = EtiquetaCalidad.query.get_or_404(etiqueta_id)
    if DetalleDespachoAcero.query.filter_by(etiqueta_id=etiqueta.id).first():
        flash("No se puede eliminar una etiqueta que ya pertenece a un despacho.", "error")
        return redirect(url_for("consultar_etiquetas_qality"))
    pdf_path = QUALITY_LABELS_DIR / f"etiqueta_{etiqueta.id:07d}.pdf"
    db.session.delete(etiqueta)
    db.session.commit()
    pdf_path.unlink(missing_ok=True)
    flash("Etiqueta eliminada correctamente.", "success")
    return redirect(url_for("consultar_etiquetas_qality"))


@app.route("/admin/qality/consultar")
@calidad_required
def consultar_etiquetas_qality():
    numero_etiqueta = (request.args.get("numero_etiqueta") or "").strip()
    tipo_producto = (request.args.get("tipo_producto") or "").strip()
    colada = (request.args.get("colada") or "").strip()
    medida = (request.args.get("medida") or "").strip()
    longitud = (request.args.get("longitud") or "").strip()
    calidad = (request.args.get("calidad") or "").strip()
    estado = (request.args.get("estado") or "").strip()
    query = EtiquetaCalidad.query
    if numero_etiqueta:
        numero_normalizado = numero_etiqueta.upper().removeprefix("ETQ-").strip()
        if numero_normalizado.isdigit():
            query = query.filter(EtiquetaCalidad.id == int(numero_normalizado))
        else:
            query = query.filter(false())
    if tipo_producto in {"AP500S", "VarillasLisas", "Palanquillas"}:
        query = query.filter(EtiquetaCalidad.tipo_producto == tipo_producto)
    else:
        tipo_producto = ""
    if colada:
        query = query.filter(EtiquetaCalidad.colada.ilike(f"%{colada}%"))
    if medida:
        query = query.filter(EtiquetaCalidad.medida.ilike(f"%{medida}%"))
    if longitud:
        query = query.filter(EtiquetaCalidad.longitud.ilike(f"%{longitud}%"))
    if calidad:
        query = query.filter(EtiquetaCalidad.calidad.ilike(f"%{calidad}%"))
    if estado == "aprobada":
        query = query.filter(EtiquetaCalidad.aprobado.is_(True), EtiquetaCalidad.tipo_b.is_(False), EtiquetaCalidad.no_conforme.is_(False))
    elif estado == "tipo_b":
        query = query.filter(EtiquetaCalidad.tipo_b.is_(True))
    elif estado == "retenida":
        query = query.filter(EtiquetaCalidad.no_conforme.is_(True))
    elif estado == "registrada":
        query = query.filter(EtiquetaCalidad.aprobado.is_(False), EtiquetaCalidad.tipo_b.is_(False), EtiquetaCalidad.no_conforme.is_(False))
    else:
        estado = ""
    etiquetas = query.order_by(EtiquetaCalidad.created_at.desc()).all()
    palanquillas = [etiqueta for etiqueta in etiquetas if etiqueta.tipo_producto == "Palanquillas"]
    ap500s = [etiqueta for etiqueta in etiquetas if etiqueta.tipo_producto == "AP500S"]
    varillas_lisas = [etiqueta for etiqueta in etiquetas if etiqueta.tipo_producto == "VarillasLisas"]
    return render_template(
        "consultar_etiquetas.html",
        etiquetas=etiquetas,
        palanquillas=palanquillas,
        ap500s=ap500s,
        varillas_lisas=varillas_lisas,
        numero_etiqueta=numero_etiqueta,
        tipo_producto=tipo_producto,
        colada=colada,
        medida=medida,
        longitud=longitud,
        calidad=calidad,
        estado=estado,
    )


@app.get("/admin/qality/consultar/excel")
@calidad_required
def exportar_etiquetas_excel_qality():
    etiquetas = EtiquetaCalidad.query.order_by(EtiquetaCalidad.created_at.desc()).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Etiquetas"
    sheet.append(["Etiqueta", "Tipo", "Calidad", "Colada", "Medida", "Longitud", "Peso", "Estado"])
    for etiqueta in etiquetas:
        estado = "No conforme" if etiqueta.no_conforme else "Tipo B" if etiqueta.tipo_b else "Aprobada" if etiqueta.aprobado else "Conforme"
        sheet.append([f"ETQ-{etiqueta.id:07d}", etiqueta.tipo_producto, etiqueta.calidad, etiqueta.colada, etiqueta.medida, etiqueta.longitud, etiqueta.peso, estado])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name="etiquetas_calidad.xlsx")


@app.get("/admin/qality/consultar/pdf")
@calidad_required
def exportar_etiquetas_pdf_qality():
    etiquetas = EtiquetaCalidad.query.order_by(EtiquetaCalidad.created_at.desc()).all()
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(letter), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    header = ParagraphStyle("quality_export_header", parent=styles["BodyText"], fontName="Helvetica-Bold", textColor=colors.white, alignment=1)
    body = ParagraphStyle("quality_export_body", parent=styles["BodyText"], fontSize=8)
    rows = [[Paragraph(title, header) for title in ("Etiqueta", "Tipo", "Calidad", "Colada", "Medida", "Longitud", "Peso", "Estado")]]
    for etiqueta in etiquetas:
        estado = "No conforme" if etiqueta.no_conforme else "Tipo B" if etiqueta.tipo_b else "Aprobada" if etiqueta.aprobado else "Conforme"
        values = [f"ETQ-{etiqueta.id:07d}", etiqueta.tipo_producto, etiqueta.calidad, etiqueta.colada, etiqueta.medida or "-", etiqueta.longitud or "-", etiqueta.peso or "-", estado]
        rows.append([Paragraph(escape(str(value)), body) for value in values])
    table = Table(rows, repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")), ("GRID", (0, 0), (-1, -1), 0.35, colors.grey), ("FONTSIZE", (0, 0), (-1, -1), 8), ("PADDING", (0, 0), (-1, -1), 4)]))
    doc.build([Paragraph("ETIQUETAS DE CALIDAD", styles["Title"]), Spacer(1, 10), table])
    output.seek(0)
    return send_file(output, mimetype="application/pdf", as_attachment=True, download_name="etiquetas_calidad.pdf")


@app.post("/admin/qality/consultar/importar-excel")
@calidad_required
def importar_etiquetas_excel_qality():
    uploaded = request.files.get("archivo_excel")
    if not uploaded or not uploaded.filename:
        flash("Seleccione un archivo Excel (.xlsx).", "error")
        return redirect(url_for("consultar_etiquetas_qality"))
    try:
        workbook = load_workbook(uploaded, read_only=True, data_only=True)
        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            raise ValueError("El archivo Excel no contiene filas.")
        headers = [str(value or "").strip().lower() for value in rows[0]]
        aliases = {
            "id": {"id", "etq-id", "etiqueta", "numero"},
            "tipo_producto": {"tipo", "tipo_producto", "tipo de producto", "producto"},
            "calidad": {"calidad"}, "medida": {"medida", "tamaño", "tamano", "diametro", "diámetro"},
            "longitud": {"longitud", "dimension", "dimensión"}, "peso": {"peso"}, "colada": {"colada"},
            "lote": {"lote"}, "cantidad": {"cantidad"}, "fecha": {"fecha"}, "observaciones": {"observacion", "observación", "observaciones"},
        }
        indexes = {field: next((headers.index(alias) for alias in names if alias in headers), None) for field, names in aliases.items()}
        required = [field for field in ("tipo_producto", "calidad", "colada") if indexes[field] is None]
        if required:
            raise ValueError("Faltan columnas obligatorias: " + ", ".join(required))
        imported = 0
        updated = 0
        for values in rows[1:]:
            if not any(value not in (None, "") for value in values):
                continue
            def value(field):
                index = indexes[field]
                return values[index] if index is not None and index < len(values) else None
            raw_id = value("id")
            etiqueta = None
            if raw_id is not None:
                try:
                    etiqueta = db.session.get(EtiquetaCalidad, int(str(raw_id).replace("ETQ-", "")))
                except (TypeError, ValueError):
                    etiqueta = None
            etiqueta = etiqueta or EtiquetaCalidad(created_by_id=current_user.id)
            etiqueta.tipo_producto = str(value("tipo_producto") or "").strip()
            etiqueta.calidad = str(value("calidad") or "").strip()
            etiqueta.medida = str(value("medida") or "").strip()
            etiqueta.longitud = str(value("longitud") or "").strip()
            etiqueta.peso = str(value("peso") or "").strip()
            etiqueta.colada = str(value("colada") or "").strip()
            etiqueta.lote = str(value("lote") or "").strip()
            etiqueta.observaciones = str(value("observaciones") or "").strip()
            raw_quantity = value("cantidad")
            etiqueta.cantidad = int(raw_quantity) if raw_quantity not in (None, "") and str(raw_quantity).isdigit() else None
            raw_date = value("fecha")
            etiqueta.fecha = raw_date.date() if hasattr(raw_date, "date") else None
            if etiqueta.tipo_producto not in {"AP500S", "VarillasLisas", "Palanquillas"} or not etiqueta.calidad or not etiqueta.colada:
                continue
            if etiqueta.id:
                updated += 1
            else:
                db.session.add(etiqueta)
                imported += 1
        db.session.commit()
        flash(f"Excel procesado sin borrar datos: {imported} etiquetas nuevas y {updated} actualizadas.", "success")
    except (ValueError, TypeError, OSError, InvalidFileException, zipfile.BadZipFile) as error:
        db.session.rollback()
        flash(f"No se pudo importar el Excel: {error}", "error")
    return redirect(url_for("consultar_etiquetas_qality"))


@app.post("/admin/qality/consultar/importar-csv")
@calidad_required
def importar_etiquetas_csv_qality():
    uploaded = request.files.get("archivo_csv")
    if not uploaded or not uploaded.filename:
        flash("Seleccione un archivo CSV.", "error")
        return redirect(url_for("consultar_etiquetas_qality"))
    try:
        raw_content = uploaded.read()
        for encoding in ("utf-8-sig", "utf-16", "cp1252"):
            try:
                content = raw_content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ValueError("No se pudo reconocer la codificación del CSV.")
        sample = content[:4096]
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t").delimiter
        except csv.Error:
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        reader = csv.DictReader(StringIO(content), delimiter=delimiter)
        def normalize(value):
            text = unicodedata.normalize("NFKD", str(value or "")).encode("ascii", "ignore").decode().lower()
            return re.sub(r"[^a-z0-9]+", "", text)

        aliases = {
            "id": {"id", "etqid", "etiqueta", "numero"},
            "tipo_producto": {"tipo", "tipoproducto", "tipodeproducto", "producto"},
            "calidad": {"calidad"}, "medida": {"medida", "tamano", "diametro"},
            "longitud": {"longitud", "dimension"}, "peso": {"peso"}, "colada": {"colada"},
            "lote": {"lote"}, "cantidad": {"cantidad"}, "fecha": {"fecha"}, "carbono": {"carbono", "porcentajecarbono"},
            "silicio": {"silicio", "porcentajesilicio"}, "manganeso": {"manganeso", "porcentajemanganeso"},
            "hornero": {"hornero"}, "supervisor": {"supervisor"}, "operador_ccm": {"operadorccm", "operadordecm"},
            "aprobado": {"aprobado", "conforme"}, "tipo_b": {"tipob"}, "no_conforme": {"noconforme", "retenida"},
            "observaciones": {"observacion", "observaciones"},
        }
        fields = {}
        for field, names in aliases.items():
            fields[field] = next((header for header in reader.fieldnames or [] if normalize(header) in names), None)
        required = [field for field in ("tipo_producto", "calidad", "colada") if not fields[field]]
        if required:
            raise ValueError("Faltan columnas obligatorias: " + ", ".join(required))
        imported = updated = skipped = 0
        for row in reader:
            def value(field):
                return row.get(fields[field]) if fields[field] else None
            etiqueta = None
            raw_id = value("id")
            if raw_id:
                try:
                    etiqueta = db.session.get(EtiquetaCalidad, int(str(raw_id).replace("ETQ-", "")))
                except (TypeError, ValueError):
                    etiqueta = None
            etiqueta = etiqueta or EtiquetaCalidad(created_by_id=current_user.id)
            raw_type = normalize(value("tipo_producto"))
            etiqueta.tipo_producto = {"ap500s": "AP500S", "varillaslisas": "VarillasLisas", "palanquillas": "Palanquillas", "palanquilla": "Palanquillas"}.get(raw_type, str(value("tipo_producto") or "").strip())
            etiqueta.calidad = str(value("calidad") or "").strip()
            etiqueta.medida = str(value("medida") or "").strip()
            etiqueta.longitud = str(value("longitud") or "").strip()
            etiqueta.peso = str(value("peso") or "").strip()
            etiqueta.colada = str(value("colada") or "").strip()
            etiqueta.lote = str(value("lote") or "").strip()
            etiqueta.carbono = str(value("carbono") or "").strip()
            etiqueta.silicio = str(value("silicio") or "").strip()
            etiqueta.manganeso = str(value("manganeso") or "").strip()
            etiqueta.hornero = str(value("hornero") or "").strip()
            etiqueta.supervisor = str(value("supervisor") or "").strip()
            etiqueta.operador_ccm = str(value("operador_ccm") or "").strip()
            etiqueta.observaciones = str(value("observaciones") or "").strip()
            raw_quantity = str(value("cantidad") or "").strip()
            etiqueta.cantidad = int(raw_quantity) if raw_quantity.isdigit() else None
            def as_bool(raw_value):
                return normalize(raw_value) in {"1", "si", "yes", "true", "x", "aprobado", "conforme", "tipob", "noconforme", "retenida"}
            etiqueta.aprobado = as_bool(value("aprobado")) if fields["aprobado"] else False
            etiqueta.tipo_b = as_bool(value("tipo_b")) if fields["tipo_b"] else False
            etiqueta.no_conforme = as_bool(value("no_conforme")) if fields["no_conforme"] else False
            raw_date = str(value("fecha") or "").strip()
            try:
                etiqueta.fecha = datetime.strptime(raw_date, "%Y-%m-%d").date() if raw_date else None
            except ValueError:
                etiqueta.fecha = datetime.strptime(raw_date, "%d/%m/%Y").date() if raw_date else None
            if etiqueta.tipo_producto not in {"AP500S", "VarillasLisas", "Palanquillas"} or not etiqueta.calidad or not etiqueta.colada:
                skipped += 1
                continue
            if etiqueta.id:
                updated += 1
            else:
                db.session.add(etiqueta)
                imported += 1
        db.session.commit()
        flash(f"CSV procesado sin borrar datos: {imported} etiquetas nuevas, {updated} actualizadas y {skipped} filas omitidas.", "success")
    except (UnicodeDecodeError, ValueError, TypeError, OSError, csv.Error) as error:
        db.session.rollback()
        flash(f"No se pudo importar el CSV: {error}. Use columnas tipo_producto, calidad y colada.", "error")
    return redirect(url_for("consultar_etiquetas_qality"))


@app.post("/admin/qality/consultar/eliminar-todas")
@calidad_required
def eliminar_todas_etiquetas_qality():
    if (request.form.get("codigo_confirmacion") or "").strip() != QUALITY_LABEL_DELETE_CODE or (request.form.get("confirmacion_texto") or "").strip().upper() != "BORRAR TODAS LAS ETIQUETAS":
        flash("Debe ingresar el código y escribir BORRAR TODAS LAS ETIQUETAS. No se eliminó ninguna etiqueta.", "error")
        return redirect(url_for("consultar_etiquetas_qality"))
    if db.session.query(DetalleDespachoAcero.id).first():
        flash("No se pueden borrar todas las etiquetas porque existen etiquetas vinculadas a despachos de acero.", "error")
        return redirect(url_for("consultar_etiquetas_qality"))
    backup_path = QUALITY_LABELS_DIR / f"respaldo_etiquetas_antes_de_borrado_{datetime.utcnow():%Y%m%d_%H%M%S}.csv"
    backup_path.parent.mkdir(parents=True, exist_ok=True)
    with backup_path.open("w", newline="", encoding="utf-8-sig") as backup:
        writer = csv.writer(backup)
        writer.writerow(["id", "tipo_producto", "calidad", "medida", "longitud", "peso", "colada", "lote", "cantidad", "fecha", "observaciones"])
        for etiqueta in EtiquetaCalidad.query.order_by(EtiquetaCalidad.id.asc()).all():
            writer.writerow([etiqueta.id, etiqueta.tipo_producto, etiqueta.calidad, etiqueta.medida, etiqueta.longitud, etiqueta.peso, etiqueta.colada, etiqueta.lote, etiqueta.cantidad or "", etiqueta.fecha.isoformat() if etiqueta.fecha else "", etiqueta.observaciones])
    deleted = EtiquetaCalidad.query.count()
    EtiquetaCalidad.query.delete(synchronize_session=False)
    db.session.commit()
    for pdf_path in QUALITY_LABELS_DIR.glob("etiqueta_*.pdf"):
        pdf_path.unlink(missing_ok=True)
    flash(f"Se eliminaron {deleted} etiquetas. Se guardó un respaldo en {backup_path.name}.", "success")
    return redirect(url_for("consultar_etiquetas_qality"))


def _quality_backup_rows(model):
    return [
        {column.name: getattr(record, column.name) for column in model.__table__.columns}
        for record in model.query.order_by(model.id.asc()).all()
    ]


def _quality_backup_json_value(value):
    return value.isoformat() if isinstance(value, (date, datetime)) else value


def _quality_backup_archive() -> BytesIO:
    payload = {
        "version": 1,
        "etiquetas": [
            {name: _quality_backup_json_value(value) for name, value in row.items()}
            for row in _quality_backup_rows(EtiquetaCalidad)
        ],
        "certificados": [
            {name: _quality_backup_json_value(value) for name, value in row.items()}
            for row in _quality_backup_rows(CertificadoCalidad)
        ],
    }
    archive = BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as backup:
        backup.writestr("manifest.json", json.dumps(payload, ensure_ascii=False, indent=2))
        for directory, prefix in ((QUALITY_LABELS_DIR, "etiquetas"), (QUALITY_CERTIFICATES_PROJECT_DIR, "certificados")):
            if directory.exists():
                for pdf_path in directory.glob("*.pdf"):
                    backup.write(pdf_path, f"{prefix}/{pdf_path.name}")
    archive.seek(0)
    return archive


def _restore_quality_backup(uploaded_file) -> None:
    if not uploaded_file or not uploaded_file.filename:
        raise ValueError("Seleccione un archivo ZIP de respaldo.")
    with zipfile.ZipFile(uploaded_file) as backup:
        names = set(backup.namelist())
        if "manifest.json" not in names or any(Path(name).is_absolute() or ".." in Path(name).parts for name in names):
            raise ValueError("El archivo de respaldo no tiene un formato válido.")
        try:
            manifest = json.loads(backup.read("manifest.json"))
        except (json.JSONDecodeError, KeyError, UnicodeDecodeError) as error:
            raise ValueError("No se pudo leer el contenido del respaldo.") from error
        if manifest.get("version") != 1 or not isinstance(manifest.get("etiquetas"), list) or not isinstance(manifest.get("certificados"), list):
            raise ValueError("El respaldo no corresponde al módulo de calidad.")

        db.session.query(EtiquetaCalidad).delete()
        db.session.query(CertificadoCalidad).delete()
        for model, key in ((EtiquetaCalidad, "etiquetas"), (CertificadoCalidad, "certificados")):
            columns = {column.name: column for column in model.__table__.columns}
            rows = []
            for raw_row in manifest[key]:
                row = {}
                for name, value in raw_row.items():
                    column = columns.get(name)
                    if column is None:
                        continue
                    python_type = column.type.python_type
                    if value is not None and python_type is datetime:
                        value = datetime.fromisoformat(value)
                    elif value is not None and python_type is date:
                        value = date.fromisoformat(value)
                    row[name] = value
                rows.append(row)
            if rows:
                db.session.bulk_insert_mappings(model, rows)
        db.session.commit()
        QUALITY_LABELS_DIR.mkdir(parents=True, exist_ok=True)
        QUALITY_CERTIFICATES_PROJECT_DIR.mkdir(parents=True, exist_ok=True)
        for directory in (QUALITY_LABELS_DIR, QUALITY_CERTIFICATES_PROJECT_DIR):
            for pdf_path in directory.glob("*.pdf"):
                pdf_path.unlink()
        for name in names - {"manifest.json"}:
            path = Path(name)
            if path.parts[0] not in {"etiquetas", "certificados"} or path.suffix.lower() != ".pdf":
                continue
            destination_dir = QUALITY_LABELS_DIR if path.parts[0] == "etiquetas" else QUALITY_CERTIFICATES_PROJECT_DIR
            (destination_dir / path.name).write_bytes(backup.read(name))


@app.get("/admin/qality/backup")
@calidad_required
def descargar_respaldo_calidad():
    return send_file(_quality_backup_archive(), mimetype="application/zip", as_attachment=True, download_name=f"respaldo_calidad_{datetime.utcnow():%Y%m%d_%H%M%S}.zip")


@app.post("/admin/qality/restore")
@calidad_required
def restaurar_respaldo_calidad():
    try:
        _restore_quality_backup(request.files.get("respaldo"))
    except (ValueError, zipfile.BadZipFile, OSError, TypeError) as error:
        db.session.rollback()
        flash(f"No se pudo restaurar el respaldo: {error}", "error")
    else:
        flash("Respaldo de etiquetas y certificados restaurado correctamente.", "success")
    return redirect(url_for("consultar_etiquetas_qality"))


def _full_backup_archive() -> BytesIO:
    if not database_url.startswith("sqlite") or not DB_PATH.exists():
        raise ValueError("El respaldo completo está disponible para la base SQLite local.")
    with tempfile.TemporaryDirectory() as temporary_dir:
        snapshot_path = Path(temporary_dir) / "database.sqlite3"
        source = sqlite3.connect(DB_PATH)
        target = sqlite3.connect(snapshot_path)
        try:
            source.backup(target)
        finally:
            target.close()
            source.close()
        archive = BytesIO()
        with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as backup:
            backup.writestr("manifest.json", json.dumps({"version": 1, "type": "full", "database": "database.sqlite3"}, indent=2))
            backup.write(snapshot_path, "database.sqlite3")
            for directory_name in ("Reports", "instance"):
                directory = BASE_DIR / directory_name
                if not directory.exists():
                    continue
                for file_path in directory.rglob("*"):
                    if file_path.is_file():
                        backup.write(file_path, file_path.relative_to(BASE_DIR).as_posix())
        archive.seek(0)
        return archive


def _restore_full_backup(uploaded_file) -> None:
    if not uploaded_file or not uploaded_file.filename:
        raise ValueError("Seleccione un archivo ZIP de respaldo completo.")
    with tempfile.TemporaryDirectory() as temporary_dir:
        extract_dir = Path(temporary_dir)
        with zipfile.ZipFile(uploaded_file) as backup:
            names = backup.namelist()
            if "manifest.json" not in names or "database.sqlite3" not in names or any(Path(name).is_absolute() or ".." in Path(name).parts for name in names):
                raise ValueError("El archivo no es un respaldo completo válido.")
            manifest = json.loads(backup.read("manifest.json"))
            if manifest.get("version") != 1 or manifest.get("type") != "full":
                raise ValueError("El respaldo no corresponde a la base completa.")
            backup.extractall(extract_dir)
        validation = sqlite3.connect(extract_dir / "database.sqlite3")
        try:
            if validation.execute("PRAGMA integrity_check").fetchone()[0] != "ok":
                raise ValueError("La base de datos del respaldo está dañada.")
        finally:
            validation.close()
        db.session.remove()
        db.engine.dispose()
        if DB_PATH.exists():
            safety_copy = DB_PATH.with_name(f"{DB_PATH.stem}.before_full_restore_{datetime.utcnow():%Y%m%d_%H%M%S}{DB_PATH.suffix}")
            shutil.copy2(DB_PATH, safety_copy)
        shutil.copy2(extract_dir / "database.sqlite3", DB_PATH)
        for directory_name in ("Reports", "instance"):
            source_directory = extract_dir / directory_name
            if source_directory.exists():
                shutil.copytree(source_directory, BASE_DIR / directory_name, dirs_exist_ok=True)


@app.get("/admin/qality/full-backup")
@admin_required
def descargar_respaldo_completo():
    try:
        archive = _full_backup_archive()
    except ValueError as error:
        flash(str(error), "error")
        return redirect(url_for("consultar_etiquetas_qality"))
    return send_file(archive, mimetype="application/zip", as_attachment=True, download_name=f"respaldo_completo_{datetime.utcnow():%Y%m%d_%H%M%S}.zip")


@app.post("/admin/qality/full-restore")
@admin_required
def restaurar_respaldo_completo():
    try:
        _restore_full_backup(request.files.get("respaldo_completo"))
    except (ValueError, json.JSONDecodeError, zipfile.BadZipFile, OSError, sqlite3.Error) as error:
        db.session.rollback()
        flash(f"No se pudo restaurar el respaldo completo: {error}", "error")
    else:
        flash("Respaldo completo restaurado. Reinicie la aplicación para continuar usando la base restaurada.", "success")
    return redirect(url_for("consultar_etiquetas_qality"))


@app.route("/admin/qality/abrir/<path:filename>")
@calidad_required
def abrir_certificado_qality(filename: str):
    return abrir_certificado_calidad(filename)


@app.route("/admin/control-calidad", methods=["GET", "POST"])
@role_required("admin", "rrhh", "root", "calidad")
def control_calidad():
    if request.method == "POST":
        nombre = request.form.get("calidad_firma_nombre", "").strip()
        cargo = request.form.get("calidad_firma_cargo", "").strip()
        if not nombre or not cargo:
            flash("Indique el nombre y cargo del responsable de Control de Calidad.", "error")
        else:
            set_config_value("calidad_firma_nombre", nombre)
            set_config_value("calidad_firma_cargo", cargo)
            firma = request.files.get("calidad_firma")
            if firma and firma.filename:
                extension = Path(secure_filename(firma.filename)).suffix.lower()
                if extension not in {".png", ".jpg", ".jpeg"}:
                    flash("La firma debe estar en formato PNG, JPG o JPEG.", "error")
                    return redirect(url_for("control_calidad"))
                firma_path = QUALITY_LABELS_IMAGES_DIR / f"firma_calidad{extension}"
                firma.save(firma_path)
                set_config_value("calidad_firma_path", str(firma_path))
            flash("Responsable y firma de Control de Calidad actualizados.", "success")
        return redirect(url_for("control_calidad"))
    dirs, warning = ensure_quality_certificate_dirs()
    query = (request.args.get("q") or "").strip()
    certificados = list_quality_certificate_files(dirs)
    if query:
        query_lower = query.lower()
        certificados = [
            cert for cert in certificados
            if query_lower in cert["name"].lower()
        ]
    return render_template(
        "control_calidad.html",
        certificados=certificados,
        carpetas=dirs,
        warning=warning,
        filtro=query,
        calidad_firma_nombre=get_config_value("calidad_firma_nombre", "Ing. Carlos Gonzalez"),
        calidad_firma_cargo=get_config_value("calidad_firma_cargo", "Jefe de Control de Calidad"),
        calidad_firma_path=get_config_value("calidad_firma_path", str(QUALITY_LABELS_IMAGES_DIR / "firma.png")),
    )


@app.route("/admin/certificaciones", methods=["GET", "POST"])
@calidad_required
def certificaciones_calidad():
    colada = (request.values.get("numero_colada") or "").strip()
    palanquilla_exists = bool(colada and EtiquetaCalidad.query.filter_by(tipo_producto="Palanquillas", colada=colada).first())
    certificado = CertificadoCalidad.query.filter_by(numero_colada=colada).first() if colada else None
    prefill = _prefill_certificate(colada) if colada else {}
    if certificado:
        for field in ("carbono", "silicio", "manganeso", "longitudes", "medidas", "alargamiento", "limite_fluencia", "limite_resistencia"):
            if not prefill.get(field):
                prefill[field] = getattr(certificado, field, "") or ""
    if request.method == "POST":
        if not colada:
            flash("Ingrese el número de colada.", "error")
            return redirect(url_for("certificaciones_calidad"))
        if not palanquilla_exists:
            flash("No hay una colada de Palanquillas con el número indicado.", "error")
            return redirect(url_for("certificaciones_calidad"))
        prefill = _prefill_certificate(colada)
        if not prefill["has_palanquillas"] or not prefill["has_varillas"]:
            missing = []
            if not prefill["has_palanquillas"]:
                missing.append("Palanquillas")
            if not prefill["has_varillas"]:
                missing.append("Varillas Lisas o AP500 S")
            flash("No se encuentran datos de la colada en " + " ni en ".join(missing) + ".", "error")
            return redirect(url_for("certificaciones_calidad"))
        certificado = certificado or CertificadoCalidad(numero_colada=colada, numero_certificado=_next_certificate_number(), created_by_id=current_user.id)
        certificado.numero_colada = colada
        selected_products = list(dict.fromkeys(request.form.getlist("productos_certificados")))
        selected_details = [detail for detail in prefill["product_details"] if detail["label"] in selected_products]
        if not selected_details:
            flash("Seleccione al menos un producto o diámetro para certificar.", "error")
            return redirect(url_for("certificaciones_calidad", numero_colada=colada))
        selected_types = []
        for detail in selected_details:
            product_name = detail["label"].split(" - ", 1)[0]
            if product_name not in selected_types:
                selected_types.append(product_name)
        certificado.tipo_producto = " / ".join(selected_types)
        certificado.carbono = prefill["carbono"]
        certificado.silicio = prefill["silicio"]
        certificado.manganeso = prefill["manganeso"]
        selected_labels = {detail["label"] for detail in selected_details}
        selected_varillas = [item for item in EtiquetaCalidad.query.filter_by(colada=colada).all() if f"{'AP500 S' if item.tipo_producto == 'AP500S' else 'Varillas Lisas'} - {item.medida or '-'} mm" in selected_labels]
        certificado.medidas = ", ".join(sorted({item.medida for item in selected_varillas if item.medida}))
        certificado.longitudes = ", ".join(sorted({item.longitud for item in selected_varillas if item.longitud}))
        certificado.cantidad_palanquillas = prefill["cantidad_palanquillas"]
        certificado.longitudes_palanquillas = prefill["longitudes_palanquillas"]
        certificado.pesos_palanquillas = prefill["pesos_palanquillas"]
        certificado.datos_varillas = prefill["datos_varillas"]
        certificado.productos_certificados = " | ".join(selected_products)
        certificado.alargamiento = (request.form.get("alargamiento") or "").strip()
        certificado.limite_fluencia = (request.form.get("limite_fluencia") or "").strip()
        certificado.limite_resistencia = (request.form.get("limite_resistencia") or "").strip()
        certificado.doblado_cumple = request.form.get("doblado_cumple") == "on"
        fecha_value = (request.form.get("fecha_atado") or "").strip()
        try:
            certificado.fecha_atado = datetime.strptime(fecha_value, "%Y-%m-%d").date() if fecha_value else datetime.utcnow().date()
        except ValueError:
            flash("La fecha de atado no tiene un formato válido.", "error")
            return render_template("certificaciones_calidad.html", certificado=certificado, prefill=prefill, colada=colada, colada_encontrada=palanquilla_exists, palanquillas=[])
        fecha_certificacion = (request.form.get("fecha_certificacion") or "").strip()
        try:
            certificado.fecha_certificacion = datetime.strptime(fecha_certificacion, "%Y-%m-%d").date() if fecha_certificacion else datetime.utcnow().date()
        except ValueError:
            flash("La fecha de certificación no tiene un formato válido.", "error")
            return render_template("certificaciones_calidad.html", certificado=certificado, prefill=prefill, colada=colada, colada_encontrada=palanquilla_exists, palanquillas=[])
        graph = request.files.get("grafico")
        if graph and graph.filename:
            graph_name = secure_filename(graph.filename)
            graph_path = BASE_DIR / "Reports" / f"grafico_{certificado.numero_colada}_{graph_name}"
            graph.save(graph_path)
            certificado.grafico_path = str(graph_path)
        pasted_graph = (request.form.get("grafico_pegado") or "").strip()
        if pasted_graph.startswith("data:image/") and "," in pasted_graph:
            header, encoded = pasted_graph.split(",", 1)
            extension = "png" if "png" in header else "jpg"
            graph_path = BASE_DIR / "Reports" / f"grafico_{secure_filename(certificado.numero_colada)}_pegado.{extension}"
            try:
                graph_path.write_bytes(base64.b64decode(encoded))
                certificado.grafico_path = str(graph_path)
            except (OSError, ValueError):
                flash("No se pudo leer la captura pegada.", "error")
        if not certificado.firma_path:
            certificado.firma_path = str(QUALITY_LABELS_IMAGES_DIR / "firma.png")
        db.session.add(certificado)
        db.session.commit()
        pdf_path = QUALITY_CERTIFICATES_PROJECT_DIR / f"certificado_{secure_filename(colada)}.pdf"
        _draw_quality_certificate_pdf(certificado, pdf_path)
        return redirect(url_for("vista_previa_certificacion", certificado_id=certificado.id))
    palanquilla_rows = EtiquetaCalidad.query.filter_by(tipo_producto="Palanquillas").order_by(EtiquetaCalidad.colada.asc(), EtiquetaCalidad.created_at.desc()).all()
    palanquillas = []
    for item in palanquilla_rows:
        existing = next((row for row in palanquillas if row["colada"] == item.colada), None)
        if existing:
            existing["cantidad"] += item.cantidad or 0
            existing["longitudes"].add(item.longitud)
            existing["pesos"].add(item.peso)
        else:
            palanquillas.append({
                "colada": item.colada,
                "fecha": item.fecha or (item.created_at.date() if item.created_at else None),
                "cantidad": item.cantidad or 0,
                "longitudes": {item.longitud} if item.longitud else set(),
                "pesos": {item.peso} if item.peso else set(),
                "carbono": item.carbono,
                "silicio": item.silicio,
                "manganeso": item.manganeso,
            })
    for item in palanquillas:
        item["longitud"] = ", ".join(sorted(item.pop("longitudes")))
        item["peso"] = ", ".join(sorted(item.pop("pesos")))
    certificados = CertificadoCalidad.query.order_by(CertificadoCalidad.fecha_creacion.desc()).all()
    varilla_rows = EtiquetaCalidad.query.filter(EtiquetaCalidad.tipo_producto.in_(["AP500S", "VarillasLisas"])).order_by(EtiquetaCalidad.colada.asc(), EtiquetaCalidad.created_at.desc()).all()
    ap500s = [item for item in varilla_rows if item.tipo_producto == "AP500S"]
    varillas_lisas = [item for item in varilla_rows if item.tipo_producto == "VarillasLisas"]
    return render_template("certificaciones_calidad.html", certificado=certificado, prefill=prefill, colada=colada, colada_encontrada=palanquilla_exists, certificados=certificados, palanquillas=palanquillas, palanquilla_rows=palanquilla_rows, ap500s=ap500s, varillas_lisas=varillas_lisas)


def _steel_label_status(etiqueta):
    if etiqueta.no_conforme:
        return "Retenida"
    if etiqueta.tipo_b:
        return "Tipo B"
    if etiqueta.aprobado:
        return "Aprobada"
    return "Pendiente"


def _steel_weight(value):
    try:
        return float(str(value or "0").replace(",", "."))
    except (TypeError, ValueError):
        return 0.0


app.jinja_env.globals["steel_weight"] = _steel_weight


def _steel_label_json(etiqueta):
    not_informatized = (etiqueta.observaciones or "").startswith("NO INFORMATIZADO")
    return {
        "id": etiqueta.id,
        "codigo": f"NO-INF-{etiqueta.id:07d}" if not_informatized else f"ETQ-{etiqueta.id:07d}",
        "tipo_producto": etiqueta.tipo_producto,
        "tipo_nombre": "AP500 S" if etiqueta.tipo_producto == "AP500S" else "Varillas Lisas",
        "medida": etiqueta.medida or "-",
        "longitud": etiqueta.longitud or "-",
        "peso": etiqueta.peso or "-",
        "colada": etiqueta.colada or "-",
        "calidad": etiqueta.calidad or "-",
        "estado": _steel_label_status(etiqueta),
        "no_informatizado": not_informatized,
    }


def _parse_quality_qr(value):
    fields = {}
    for line in str(value or "").splitlines():
        if ":" not in line:
            continue
        label, field_value = line.split(":", 1)
        normalized_label = re.sub(
            r"[^a-z0-9]",
            "",
            unicodedata.normalize("NFKD", label.lower()).encode("ascii", "ignore").decode("ascii"),
        )
        fields[normalized_label] = field_value.strip()
    raw_tipo_producto = fields.get("tipodeproducto", "")
    normalized_tipo = re.sub(r"[^a-z0-9]", "", unicodedata.normalize("NFKD", raw_tipo_producto.lower()).encode("ascii", "ignore").decode("ascii"))
    tipo_producto = "AP500S" if normalized_tipo in {"ap500s", "ap500sacero"} else "VarillasLisas" if normalized_tipo in {"varillaslisas", "varillas", "varillal184"} else ""
    if not tipo_producto:
        return None
    return {
        "tipo_producto": tipo_producto,
        "calidad": fields.get("calidad", ""),
        "medida": fields.get("medidamm", ""),
        "longitud": fields.get("longitudm", ""),
        "peso": fields.get("pesokg", ""),
        "colada": fields.get("numerodecolada", fields.get("colada", "")),
        "observaciones": f"NO INFORMATIZADO - QR original: {str(value).strip()[:900]}",
    }


@app.route("/admin/despacho/carga", methods=["GET"])
@calidad_required
def despacho_acero():
    dispatched_ids = db.session.query(DetalleDespachoAcero.etiqueta_id).subquery()
    labels = EtiquetaCalidad.query.filter(
        EtiquetaCalidad.tipo_producto.in_(["AP500S", "VarillasLisas"]),
        EtiquetaCalidad.no_conforme.is_(False),
        ~EtiquetaCalidad.id.in_(dispatched_ids),
    ).order_by(EtiquetaCalidad.tipo_producto, EtiquetaCalidad.medida, EtiquetaCalidad.created_at).all()
    all_labels = EtiquetaCalidad.query.filter(EtiquetaCalidad.tipo_producto.in_(["AP500S", "VarillasLisas"])).all()
    stock = {}
    for label in all_labels:
        key = (label.tipo_producto, label.medida or "-")
        group = stock.setdefault(key, {"tipo": "AP500 S" if label.tipo_producto == "AP500S" else "Varillas Lisas", "medida": label.medida or "-", "aprobadas": 0, "tipo_b": 0, "retenidas": 0, "pendientes": 0, "total": 0})
        group["total"] += 1
        status_key = {"Aprobada": "aprobadas", "Tipo B": "tipo_b", "Retenida": "retenidas", "Pendiente": "pendientes"}[_steel_label_status(label)]
        group[status_key] += 1
    despachos = DespachoAcero.query.order_by(DespachoAcero.fecha.desc()).limit(12).all()
    clientes = ClienteDespacho.query.order_by(ClienteDespacho.nombre.asc()).all()
    return render_template("despacho_acero.html", labels=labels, stock=sorted(stock.values(), key=lambda row: (row["tipo"], row["medida"])), despachos=despachos, clientes=clientes)


@app.route("/admin/despacho")
@calidad_required
def panel_despacho_acero():
    return render_template("panel_despacho_acero.html")


@app.route("/admin/despacho/clientes", methods=["GET", "POST"])
@calidad_required
def clientes_despacho():
    if request.method == "POST":
        nombre = (request.form.get("nombre") or "").strip()
        destino = (request.form.get("destino") or "").strip()
        if not nombre:
            flash("Ingrese el nombre del cliente o empresa.", "error")
        elif ClienteDespacho.query.filter_by(nombre=nombre).first():
            flash("Ese cliente ya está registrado.", "warning")
        else:
            db.session.add(ClienteDespacho(nombre=nombre, ultimo_destino=destino))
            db.session.commit()
            flash("Cliente registrado correctamente.", "success")
        return redirect(url_for("clientes_despacho"))
    clientes = ClienteDespacho.query.order_by(ClienteDespacho.nombre.asc()).all()
    return render_template("clientes_despacho.html", clientes=clientes)


@app.route("/admin/despacho/etiqueta/<codigo>")
@calidad_required
def consultar_etiqueta_despacho(codigo):
    match = re.search(r"ETQ[-\s:]*(\d+)", codigo or "", re.IGNORECASE)
    if not match:
        match = re.search(r"Numero de etiqueta\s*:\s*ETQ[-\s:]*(\d+)", codigo or "", re.IGNORECASE)
    etiqueta = EtiquetaCalidad.query.get(int(match.group(1))) if match else None
    if not etiqueta or etiqueta.tipo_producto not in {"AP500S", "VarillasLisas"}:
        parsed = _parse_quality_qr(codigo)
        if not parsed:
            return jsonify({"ok": False, "legacy": True, "message": "El QR no contiene datos de una etiqueta de varillas.", "detail": "Verifique la etiqueta o ingrese el código manualmente."}), 404
        legacy_observation = parsed["observaciones"]
        etiqueta = EtiquetaCalidad.query.filter_by(
            tipo_producto=parsed["tipo_producto"],
            calidad=parsed["calidad"],
            medida=parsed["medida"],
            longitud=parsed["longitud"],
            peso=parsed["peso"],
            colada=parsed["colada"],
            observaciones=legacy_observation,
        ).first()
        if etiqueta is None:
            etiqueta = EtiquetaCalidad(**parsed)
            db.session.add(etiqueta)
            db.session.commit()
        return jsonify({"ok": True, "legacy": True, "message": "Atado anterior a la sistematización del despacho.", "detail": "Agregado como atado NO INFORMATIZADO.", "etiqueta": _steel_label_json(etiqueta)})
    if etiqueta.no_conforme:
        return jsonify({"ok": False, "message": "Esta etiqueta está retenida y no puede cargarse."}), 409
    if DetalleDespachoAcero.query.filter_by(etiqueta_id=etiqueta.id).first():
        return jsonify({"ok": False, "message": "Esta etiqueta ya fue despachada."}), 409
    return jsonify({"ok": True, "etiqueta": _steel_label_json(etiqueta)})


@app.route("/admin/despacho/guardar", methods=["POST"])
@calidad_required
def guardar_despacho_acero():
    payload = request.get_json(silent=True) or request.form
    try:
        label_ids = [int(value) for value in (payload.getlist("etiqueta_ids") if hasattr(payload, "getlist") else payload.get("etiqueta_ids", []))]
    except (TypeError, ValueError):
        label_ids = []
    if not label_ids and isinstance(payload.get("etiqueta_ids"), list):
        label_ids = [int(value) for value in payload["etiqueta_ids"] if str(value).isdigit()]
    destino = (payload.get("destino") or "").strip()
    cliente = (payload.get("cliente") or "").strip()
    chofer = (payload.get("chofer") or "").strip()
    cedula = (payload.get("cedula") or "").strip()
    chapa = (payload.get("chapa") or "").strip()
    if not label_ids or not cliente or not destino or not chofer or not cedula or not chapa:
        return jsonify({"ok": False, "message": "Complete cliente, destino, chofer, cédula, chapa y agregue al menos una etiqueta."}), 400
    if payload.get("cliente_confirmado") is not True:
        return jsonify({"ok": False, "message": "Confirme los datos del cliente antes de cerrar la carga."}), 400
    labels = EtiquetaCalidad.query.filter(EtiquetaCalidad.id.in_(set(label_ids))).all()
    already_dispatched = {row.etiqueta_id for row in DetalleDespachoAcero.query.filter(DetalleDespachoAcero.etiqueta_id.in_(set(label_ids))).all()}
    if len(labels) != len(set(label_ids)) or already_dispatched or any(label.no_conforme for label in labels):
        return jsonify({"ok": False, "message": "Una o más etiquetas no están disponibles para despacho."}), 409
    numero = f"D-{datetime.utcnow():%Y%m%d}-{DespachoAcero.query.count() + 1:04d}"
    cliente_codigo = (payload.get("cliente_codigo") or "").strip().upper()
    cliente_registro = None
    if cliente_codigo.startswith("CLI-") and cliente_codigo[4:].isdigit():
        cliente_registro = ClienteDespacho.query.get(int(cliente_codigo[4:]))
    cliente_registro = cliente_registro or ClienteDespacho.query.filter_by(nombre=cliente).first() or ClienteDespacho(nombre=cliente)
    cliente_registro.ultimo_destino = destino
    cliente_registro.ultimo_chofer = chofer
    cliente_registro.ultima_cedula = cedula
    cliente_registro.ultima_chapa = chapa
    db.session.add(cliente_registro)
    db.session.flush()
    despacho = DespachoAcero(numero=numero, cliente=cliente_registro.nombre, cliente_id=cliente_registro.id, destino=destino, chofer=chofer, cedula=cedula, chapa=chapa, responsable=current_user.full_name, observacion=(payload.get("observacion") or "").strip(), created_by_id=current_user.id)
    despacho.etiquetas = [DetalleDespachoAcero(etiqueta_id=label.id) for label in labels]
    db.session.add(despacho)
    db.session.commit()
    despacho_url = url_for("documento_despacho_acero", despacho_id=despacho.id, tipo="despacho")
    remision_url = url_for("documento_despacho_acero", despacho_id=despacho.id, tipo="remision")
    return jsonify({"ok": True, "id": despacho.id, "numero": despacho.numero, "preview_url": despacho_url, "despacho_url": despacho_url, "remision_url": remision_url})


@app.route("/admin/despacho/<int:despacho_id>/<tipo>")
@calidad_required
def documento_despacho_acero(despacho_id, tipo):
    if tipo not in {"despacho", "remision"}:
        abort(404)
    despacho = DespachoAcero.query.get_or_404(despacho_id)
    return render_template("documento_despacho_acero_v2.html", despacho=despacho, tipo=tipo)


@app.route("/admin/remision/<int:despacho_id>/<tipo>")
@calidad_required
def remision_despacho_acero_legacy(despacho_id, tipo):
    despacho = DespachoAcero.query.get_or_404(despacho_id)
    return render_template("documento_despacho_acero_v2.html", despacho=despacho, tipo="remision")


@app.post("/admin/despacho/<int:despacho_id>/etiqueta/<int:detalle_id>/devolver")
@calidad_required
def devolver_etiqueta_despacho(despacho_id, detalle_id):
    detalle = DetalleDespachoAcero.query.filter_by(id=detalle_id, despacho_id=despacho_id).first_or_404()
    codigo = f"ETQ-{detalle.etiqueta_id:07d}"
    db.session.delete(detalle)
    db.session.commit()
    flash(f"La etiqueta {codigo} fue devuelta a existencias.", "success")
    return redirect(url_for("historial_despacho_acero"))


@app.route("/admin/despacho/<int:despacho_id>/editar", methods=["GET", "POST"])
@calidad_required
def editar_despacho_acero(despacho_id):
    despacho = DespachoAcero.query.get_or_404(despacho_id)
    current_ids = {detalle.etiqueta_id for detalle in despacho.etiquetas}
    current_labels = [detalle.etiqueta for detalle in despacho.etiquetas if detalle.etiqueta]
    allowed_profiles = {(label.tipo_producto, label.medida or "-") for label in current_labels if label.tipo_producto in {"AP500S", "VarillasLisas"}}
    if request.method == "POST":
        requested_ids = set(current_ids)
        requested_ids.update(int(value) for value in request.form.getlist("etiqueta_ids") if value.isdigit())
        manual_values = re.findall(r"(?:ETQ-|NO-INF-)?\s*(\d+)", request.form.get("manual_etiquetas", ""), re.IGNORECASE)
        requested_ids.update(int(value) for value in manual_values)
        if not requested_ids:
            flash("El despacho debe conservar al menos un atado.", "error")
            return redirect(url_for("editar_despacho_acero", despacho_id=despacho.id))
        dispatched_elsewhere = {
            row.etiqueta_id
            for row in DetalleDespachoAcero.query.filter(DetalleDespachoAcero.etiqueta_id.in_(requested_ids)).all()
            if row.despacho_id != despacho.id
        }
        labels = EtiquetaCalidad.query.filter(EtiquetaCalidad.id.in_(requested_ids)).all()
        invalid_profiles = allowed_profiles and any((label.tipo_producto, label.medida or "-") not in allowed_profiles for label in labels)
        if len(labels) != len(requested_ids) or dispatched_elsewhere or any(label.no_conforme for label in labels) or invalid_profiles:
            flash("Uno o más atados ya no están disponibles para este despacho.", "error")
            return redirect(url_for("editar_despacho_acero", despacho_id=despacho.id))
        DetalleDespachoAcero.query.filter_by(despacho_id=despacho.id).delete(synchronize_session=False)
        despacho.etiquetas = [DetalleDespachoAcero(etiqueta_id=label_id) for label_id in sorted(requested_ids)]
        db.session.commit()
        flash(f"Despacho {despacho.numero} actualizado correctamente.", "success")
        return redirect(url_for("historial_despacho_acero"))

    dispatched_ids = {row.etiqueta_id for row in DetalleDespachoAcero.query.all()}
    available_labels = EtiquetaCalidad.query.filter(
        EtiquetaCalidad.tipo_producto.in_(["AP500S", "VarillasLisas"]),
        EtiquetaCalidad.no_conforme.is_(False),
    ).order_by(EtiquetaCalidad.tipo_producto, EtiquetaCalidad.medida, EtiquetaCalidad.id).all()
    available_labels = [label for label in available_labels if label.id in current_ids or label.id not in dispatched_ids]
    if allowed_profiles:
        available_labels = [label for label in available_labels if (label.tipo_producto, label.medida or "-") in allowed_profiles]
    allowed_measures = sorted({profile[1] for profile in allowed_profiles})
    return render_template("editar_despacho_acero.html", despacho=despacho, etiquetas=available_labels, current_ids=current_ids, allowed_profiles=allowed_profiles, allowed_measures=allowed_measures)


@app.post("/admin/despacho/<int:despacho_id>/eliminar")
@calidad_required
def eliminar_despacho_acero(despacho_id):
    despacho = DespachoAcero.query.get_or_404(despacho_id)
    numero = despacho.numero
    db.session.delete(despacho)
    db.session.commit()
    flash(f"Despacho {numero} eliminado. Sus atados volvieron a existencias.", "success")
    return redirect(url_for("historial_despacho_acero"))


@app.route("/admin/despacho/historial")
@calidad_required
def historial_despacho_acero():
    cliente_codigo = (request.args.get("cliente") or "").strip().upper()
    cliente = ClienteDespacho.query.get(int(cliente_codigo[4:])) if cliente_codigo.startswith("CLI-") and cliente_codigo[4:].isdigit() else None
    colada = (request.args.get("colada") or "").strip()
    fecha_desde = (request.args.get("fecha_desde") or "").strip()
    fecha_hasta = (request.args.get("fecha_hasta") or "").strip()
    query = DespachoAcero.query
    if cliente:
        query = query.filter((DespachoAcero.cliente_id == cliente.id) | ((DespachoAcero.cliente_id.is_(None)) & (DespachoAcero.cliente == cliente.nombre)))
    if colada:
        query = query.filter(DespachoAcero.etiquetas.any(DetalleDespachoAcero.etiqueta.has(EtiquetaCalidad.colada.ilike(f"%{colada}%"))))
    if fecha_desde:
        try:
            query = query.filter(DespachoAcero.fecha >= datetime.strptime(fecha_desde, "%Y-%m-%d"))
        except ValueError:
            fecha_desde = ""
    if fecha_hasta:
        try:
            limite = datetime.strptime(fecha_hasta, "%Y-%m-%d") + timedelta(days=1)
            query = query.filter(DespachoAcero.fecha < limite)
        except ValueError:
            fecha_hasta = ""
    despachos = query.order_by(DespachoAcero.fecha.desc()).all()
    clientes = ClienteDespacho.query.order_by(ClienteDespacho.nombre.asc()).all()
    return render_template("historial_despacho_acero.html", despachos=despachos, cliente=cliente, clientes=clientes, colada=colada, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)


@app.route("/admin/despacho/historico-atados")
@calidad_required
def historico_atados_despacho():
    busqueda = (request.args.get("busqueda") or "").strip()
    cliente = (request.args.get("cliente") or "").strip()
    medida = (request.args.get("medida") or "").strip()
    query = (
        DetalleDespachoAcero.query
        .join(DespachoAcero, DetalleDespachoAcero.despacho_id == DespachoAcero.id)
        .join(EtiquetaCalidad, DetalleDespachoAcero.etiqueta_id == EtiquetaCalidad.id)
    )
    if busqueda:
        numero = re.sub(r"\D", "", busqueda)
        conditions = [
            EtiquetaCalidad.colada.ilike(f"%{busqueda}%"),
            DespachoAcero.numero.ilike(f"%{busqueda}%"),
        ]
        if numero:
            conditions.append(EtiquetaCalidad.id == int(numero))
        query = query.filter(or_(*conditions))
    if cliente:
        query = query.filter(DespachoAcero.cliente.ilike(f"%{cliente}%"))
    if medida:
        query = query.filter(EtiquetaCalidad.medida == medida)
    detalles = query.order_by(DespachoAcero.fecha.desc(), EtiquetaCalidad.medida.asc(), EtiquetaCalidad.id.asc()).all()
    resumen_mm = {}
    for detalle in detalles:
        etiqueta = detalle.etiqueta
        key = etiqueta.medida or "-"
        row = resumen_mm.setdefault(key, {"medida": key, "atados": 0, "peso": 0.0})
        row["atados"] += 1
        row["peso"] += _steel_weight(etiqueta.peso)
    clientes = ClienteDespacho.query.order_by(ClienteDespacho.nombre.asc()).all()
    medidas = {label.medida for label in EtiquetaCalidad.query.filter(EtiquetaCalidad.medida.isnot(None), EtiquetaCalidad.medida != "").all()}
    medidas = sorted(medidas, key=lambda value: (0, float(value)) if str(value).replace(".", "", 1).isdigit() else (1, str(value)))
    return render_template("historico_atados_despacho.html", detalles=detalles, resumen_mm=sorted(resumen_mm.values(), key=lambda row: row["medida"]), clientes=clientes, medidas=medidas, busqueda=busqueda, cliente=cliente, medida=medida)


@app.route("/admin/certificaciones/<int:certificado_id>/vista-previa")
@calidad_required
def vista_previa_certificacion(certificado_id: int):
    certificado = CertificadoCalidad.query.get_or_404(certificado_id)
    pdf_path = QUALITY_CERTIFICATES_PROJECT_DIR / f"certificado_{secure_filename(certificado.numero_colada)}.pdf"
    _draw_quality_certificate_pdf(certificado, pdf_path)
    return render_template("vista_previa_certificacion.html", certificado=certificado)


@app.route("/admin/certificaciones/consultar")
@calidad_required
def consultar_certificados_calidad():
    query = (request.args.get("q") or "").strip()
    certificados_query = CertificadoCalidad.query
    if query:
        certificates_filter = or_(CertificadoCalidad.numero_colada.ilike(f"%{query}%"), CertificadoCalidad.numero_certificado.ilike(f"%{query}%"))
        certificados_query = certificados_query.filter(certificates_filter)
    certificados = certificados_query.order_by(CertificadoCalidad.fecha_creacion.desc()).all()
    return render_template("consultar_certificados.html", certificados=certificados, query=query)


@app.route("/admin/qality/centro-documental")
@calidad_required
def centro_documental_certificacion():
    query = (request.args.get("q") or "").strip()
    certificados_query = CertificadoCalidad.query
    if query:
        certificates_filter = or_(
            CertificadoCalidad.numero_colada.ilike(f"%{query}%"),
            CertificadoCalidad.numero_certificado.ilike(f"%{query}%"),
        )
        certificados_query = certificados_query.filter(certificates_filter)
    certificados = certificados_query.order_by(CertificadoCalidad.fecha_creacion.desc()).all()
    return render_template("centro_documental_certificacion.html", certificados=certificados, query=query)


def restore_certificates_from_external_json() -> tuple[int, int, int]:
    if not CERTIFICATE_JSON_RESTORE_PATH.exists():
        raise ValueError(f"No existe el JSON: {CERTIFICATE_JSON_RESTORE_PATH}")
    if not CERTIFICATE_ASSETS_RESTORE_DIR.exists():
        raise ValueError(f"No existe la carpeta de gráficos: {CERTIFICATE_ASSETS_RESTORE_DIR}")
    records = json.loads(CERTIFICATE_JSON_RESTORE_PATH.read_text(encoding="utf-8-sig"))
    if not isinstance(records, list):
        raise ValueError("El JSON de certificados debe contener una lista.")
    QUALITY_CERTIFICATES_PROJECT_DIR.mkdir(parents=True, exist_ok=True)
    imported = updated = missing_graphics = 0
    for record in records:
        colada = str(record.get("NumeroColada") or "").strip()
        certificado_numero = str(record.get("NumeroCertificado") or "").strip()
        if not colada or not certificado_numero:
            continue
        certificado = CertificadoCalidad.query.filter_by(numero_colada=colada).first()
        if certificado is None:
            certificado = CertificadoCalidad(numero_colada=colada, numero_certificado=certificado_numero)
            db.session.add(certificado)
            imported += 1
        else:
            updated += 1
        tipo = str(record.get("TipoProducto") or "").strip().replace(" ", "")
        certificado.numero_certificado = certificado_numero
        certificado.tipo_producto = "AP500S" if tipo.upper() == "AP500S" else tipo
        certificado.carbono = str(record.get("Carbono") or "")
        certificado.silicio = str(record.get("Silicio") or "")
        certificado.manganeso = str(record.get("Manganeso") or "")
        certificado.medidas = ", ".join(str(value) for value in (record.get("Medidas") or []))
        certificado.longitudes = ", ".join(str(value) for value in (record.get("Longitudes") or []))
        certificado.alargamiento = str(record.get("Alargamiento") or "")
        certificado.limite_fluencia = str(record.get("LimiteFluencia") or "")
        certificado.limite_resistencia = str(record.get("LimiteResistencia") or "")
        certificado.doblado_cumple = bool(record.get("DobladoCumple"))
        graph_name = Path(str(record.get("GraficoPath") or "")).name
        graph_source = CERTIFICATE_ASSETS_RESTORE_DIR / graph_name
        if graph_name and graph_source.exists():
            graph_target = BASE_DIR / "Reports" / f"grafico_{secure_filename(colada)}_{graph_name}"
            shutil.copy2(graph_source, graph_target)
            certificado.grafico_path = str(graph_target)
        else:
            missing_graphics += 1
        pdf_name = f"certificado_{secure_filename(colada)}.pdf"
        pdf_source = CERTIFICATE_ASSETS_RESTORE_DIR / pdf_name
        if pdf_source.exists():
            shutil.copy2(pdf_source, QUALITY_CERTIFICATES_PROJECT_DIR / pdf_name)
    db.session.commit()
    return imported, updated, missing_graphics


@app.post("/admin/certificaciones/consultar/restaurar-externos")
@calidad_required
def restaurar_certificados_externos():
    try:
        imported, updated, missing_graphics = restore_certificates_from_external_json()
        flash(f"Restauración completada: {imported} nuevos, {updated} actualizados y {missing_graphics} gráficos no encontrados.", "success")
    except (OSError, ValueError, json.JSONDecodeError, sqlite3.Error) as error:
        db.session.rollback()
        flash(f"No se pudo restaurar el respaldo externo: {error}", "error")
    return redirect(url_for("consultar_certificados_calidad"))


@app.route("/admin/certificaciones/<int:certificado_id>/pdf")
@calidad_required
def descargar_certificacion(certificado_id: int):
    certificado = CertificadoCalidad.query.get_or_404(certificado_id)
    pdf_path = QUALITY_CERTIFICATES_PROJECT_DIR / f"certificado_{secure_filename(certificado.numero_colada)}.pdf"
    _draw_quality_certificate_pdf(certificado, pdf_path)
    return send_file(pdf_path, mimetype="application/pdf", as_attachment=False)


@app.route("/admin/control-calidad/abrir/<path:filename>")
def abrir_certificado_calidad(filename: str):
    dirs, warning = ensure_quality_certificate_dirs()
    if warning:
        flash(warning, "warning")
    if not dirs:
        return redirect(url_for("control_calidad"))
    for directory in dirs:
        target = directory / filename
        if target.exists() and target.is_file() and target.suffix.lower() == ".pdf":
            return send_file(target, mimetype="application/pdf", as_attachment=False)
    flash("El archivo solicitado no existe o no está disponible en la carpeta de respaldo.", "error")
    return redirect(url_for("control_calidad"))


@app.route("/rrhh")
@rrhh_required
def rrhh():
    generar_recordatorios_rrhh()
    notifications = Notificacion.query.filter_by(usuario_id=current_user.id).order_by(Notificacion.fecha.desc()).limit(10).all()
    vacation_notifications = Notificacion.query.filter(
        Notificacion.usuario_id == current_user.id,
        Notificacion.titulo.like("%Vacaciones%"),
    ).order_by(Notificacion.fecha.desc()).limit(5).all()
    unread_notifications = sum(1 for notification in notifications if not notification.leida)
    return render_template("rrhh.html", notifications=notifications, unread_notifications=unread_notifications, vacation_notifications=vacation_notifications)


@app.get("/rrhh/notificaciones")
@rrhh_required
def rrhh_notificaciones():
    generar_recordatorios_rrhh()
    notifications = Notificacion.query.filter(
        Notificacion.usuario_id == current_user.id,
        Notificacion.url.startswith("/rrhh"),
    ).order_by(Notificacion.leida.asc(), Notificacion.fecha.desc()).limit(100).all()
    unread_notifications = sum(1 for notification in notifications if not notification.leida)
    notification_plan_urls = {}
    return render_template("rrhh_notificaciones.html", notifications=notifications, unread_notifications=unread_notifications, notification_plan_urls=notification_plan_urls)


class ManualModuleIllustration(Flowable):
    def __init__(self, number: str, title: str, color: str):
        super().__init__()
        self.number = number
        self.title = title
        self.color = colors.HexColor(color)
        self.width = 500
        self.height = 38

    def draw(self):
        self.canv.setFillColor(self.color)
        self.canv.roundRect(0, 0, self.width, self.height, 7, fill=1, stroke=0)
        self.canv.setFillColor(colors.white)
        self.canv.circle(20, 19, 12, fill=1, stroke=0)
        self.canv.setFillColor(self.color)
        self.canv.setFont("Helvetica-Bold", 10)
        self.canv.drawCentredString(20, 15.5, self.number)
        self.canv.setFillColor(colors.white)
        self.canv.setFont("Helvetica-Bold", 11)
        self.canv.drawString(42, 15, self.title)


def build_manual_usuario_rrhh_pdf() -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=36, leftMargin=36, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("manual_title", parent=styles["Title"], fontSize=20, leading=24, alignment=1, textColor=colors.HexColor("#173b63"), spaceAfter=8)
    subtitle = ParagraphStyle("manual_subtitle", parent=styles["Normal"], fontSize=10, leading=13, alignment=1, textColor=colors.HexColor("#52657d"), spaceAfter=18)
    heading = ParagraphStyle("manual_heading", parent=styles["Heading2"], fontSize=13, leading=16, textColor=colors.HexColor("#173b63"), spaceBefore=12, spaceAfter=6)
    body = ParagraphStyle("manual_body", parent=styles["BodyText"], fontSize=9.5, leading=13, spaceAfter=6)
    elements = [
        Image(str(STATIC_LOGO_PATH), width=140, height=60),
        Paragraph("MANUAL DE USUARIO", title),
        Paragraph("Módulo de Recursos Humanos · IGP Metales · Versión 1.0", subtitle),
        Paragraph("1. Acceso al módulo", heading),
        Paragraph("Abra la dirección del sistema, escriba su usuario y contraseña y pulse Ingresar. Desde el dashboard principal seleccione RRHH. Si no visualiza una opción, solicite al administrador la revisión de su rol.", body),
        Paragraph("Mapa rápido del módulo", heading),
        ManualModuleIllustration("01", "Padrón de colaboradores", "#2f7d68"),
        ManualModuleIllustration("02", "Vacaciones y notificaciones", "#3d8f91"),
        ManualModuleIllustration("03", "Evaluación de personal", "#477bb2"),
        ManualModuleIllustration("04", "Documentos RRHH", "#b27a35"),
        ManualModuleIllustration("05", "Amonestaciones y sanciones", "#a65353"),
        Paragraph("2. Padrón de colaboradores", heading),
        Paragraph("Para registrar una persona, abra Padrón de colaboradores y complete los campos obligatorios. Guarde y confirme que el registro aparezca en la tabla. Para actualizar datos, use Editar, revise la información y guarde. Utilice la búsqueda por legajo, documento o nombre para evitar registros duplicados.", body),
        Paragraph("3. Vacaciones", heading),
        Paragraph("Abra Vacaciones y busque al colaborador. Revise antigüedad, días correspondientes y saldo disponible. Seleccione las fechas de salida y retorno, agregue una observación si corresponde y guarde. Verifique que la agenda muestre el período correcto antes de cerrar la pantalla.", body),
        Paragraph("4. Notificación de Vacaciones", heading),
        Paragraph("Ingrese a Notificación de Vacaciones para revisar próximas salidas, vacaciones en curso y retornos. Use esta pantalla como control diario de la agenda y confirme cualquier cambio desde el módulo Vacaciones.", body),
        Paragraph("5. Agenda de RRHH", heading),
        Paragraph("Abra Agenda de RRHH para registrar reuniones, avisos, tareas, capacitaciones y recordatorios. Puede vincular cada evento con un colaborador, indicar lugar y responsable, y actualizar su estado o sus datos desde la grilla mensual.", body),
        Paragraph("6. Evaluación de personal", heading),
        Paragraph("Abra Evaluación de personal, indique legajo y formulario, y pulse Buscar. Complete cada criterio con una puntuación válida. Agregue fortalezas, aspectos a mejorar, plan de acción, recomendación y observaciones. Revise el resultado, guarde y use la opción de PDF para archivar o imprimir.", body),
        Paragraph("6. Documentos RRHH", heading),
        Paragraph("En Documentos RRHH cree o seleccione una carpeta antes de cargar un archivo. Compruebe el nombre, tipo y fecha del documento. Los archivos de texto y planillas pueden editarse desde el sistema; los demás deben descargarse para modificarlos externamente. Evite duplicar versiones.", body),
        Paragraph("7. Amonestaciones y sanciones", heading),
        Paragraph("Abra Amonestaciones y sanciones, seleccione al colaborador y registre tipo de medida, motivo, fecha y observaciones. Revise los antecedentes antes de guardar. La información disciplinaria es confidencial y debe ser utilizada solo por personal autorizado.", body),
        Paragraph("8. Notificaciones", heading),
        Paragraph("Las notificaciones aparecen debajo de los accesos principales y pueden desplegarse u ocultarse. Pulse el título para abrir el proceso relacionado. Use Marcar como leídas después de revisar los avisos pendientes.", body),
        Paragraph("Recomendaciones", heading),
        Paragraph("Mantenga actualizados los datos del personal, revise periódicamente vacaciones y evaluaciones pendientes, use fechas correctas y confirme los cambios antes de salir. No comparta sus credenciales. Ante un error de permisos, datos o archivos, contacte al administrador del sistema.", body),
    ]
    document.build(elements)
    return buffer.getvalue()


@app.get("/rrhh/manual.pdf")
@rrhh_required
def manual_usuario_rrhh():
    return send_file(BytesIO(build_manual_usuario_rrhh_pdf()), mimetype="application/pdf", as_attachment=True, download_name="Manual_Usuario_Modulo_RRHH.pdf")


def build_rrhh_report_pdf(filename: str) -> bytes:
    report_path = BASE_DIR / filename
    if not report_path.is_file():
        abort(404)
    buffer = BytesIO()
    document = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=38, bottomMargin=38)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("report_title", parent=styles["Title"], fontSize=18, leading=22, alignment=1, textColor=colors.HexColor("#173b63"), spaceAfter=6)
    metadata = ParagraphStyle("report_metadata", parent=styles["Normal"], fontSize=9, leading=12, alignment=1, textColor=colors.HexColor("#52657d"), spaceAfter=16)
    heading = ParagraphStyle("report_heading", parent=styles["Heading2"], fontSize=12, leading=15, textColor=colors.HexColor("#173b63"), spaceBefore=9, spaceAfter=5)
    body = ParagraphStyle("report_body", parent=styles["BodyText"], fontSize=8.8, leading=11.5, spaceAfter=4)
    lines = report_path.read_text(encoding="utf-8").splitlines()
    elements = [Image(str(STATIC_LOGO_PATH), width=120, height=51), Paragraph("INFORME OFICIAL DE ENTREGA", title), Paragraph("A nombre de: Ing. Isidro Vera · Fecha: 03/09/2026", metadata)]
    for line in lines:
        text = line.strip()
        if not text or text == "---":
            continue
        if text.startswith("# "):
            continue
        if text.startswith("## "):
            elements.append(Paragraph(escape(text[3:]), heading))
            continue
        if text.startswith("### "):
            elements.append(Paragraph(escape(text[4:]), heading))
            continue
        if text.startswith("**") and text.endswith("**"):
            elements.append(Paragraph(f"<b>{escape(text[2:-2])}</b>", body))
            continue
        text = text.replace("**", "").replace("`", "")
        elements.append(Paragraph(escape(text), body))
    document.build(elements)
    return buffer.getvalue()


@app.get("/rrhh/informes/tecnico.pdf")
@rrhh_required
def informe_tecnico_rrhh_pdf():
    return send_file(BytesIO(build_rrhh_report_pdf("INFORME_TECNICO_ENTREGA_MODULO_RRHH.md")), mimetype="application/pdf", as_attachment=True, download_name="Informe_Tecnico_RRHH_Ing_Isidro_Vera.pdf")


@app.get("/rrhh/informes/auditoria.pdf")
@rrhh_required
def informe_auditoria_rrhh_pdf():
    return send_file(BytesIO(build_rrhh_report_pdf("INFORME_AUDITORIA_ENTREGA_MODULO_RRHH.md")), mimetype="application/pdf", as_attachment=True, download_name="Informe_Auditoria_RRHH_Ing_Isidro_Vera.pdf")


def sumar_anios(fecha, anios):
    try:
        return fecha.replace(year=fecha.year + anios)
    except ValueError:
        return fecha.replace(year=fecha.year + anios, day=28)


def resumen_vacaciones(colaborador, hoy=None):
    hoy = hoy or datetime.utcnow().date()
    if not colaborador.fecha_ingreso:
        return {"anios_cumplidos": 0, "dias_correspondientes": 0, "dias_agendados": 0, "dias_disponibles": 0, "proximo_aniversario": None}
    anios = max(0, hoy.year - colaborador.fecha_ingreso.year)
    if sumar_anios(colaborador.fecha_ingreso, anios) > hoy:
        anios -= 1
    agendados = sum(item.dias for item in colaborador.vacaciones if item.estado == "Programada" and item.fecha_fin >= hoy)
    correspondientes = anios * 12
    return {"anios_cumplidos": anios, "dias_correspondientes": correspondientes, "dias_agendados": agendados, "dias_disponibles": max(0, correspondientes - agendados), "proximo_aniversario": sumar_anios(colaborador.fecha_ingreso, anios + 1)}


def notificar_evento_rrhh(evento):
    if not evento.notificar or not evento.responsable_usuario_id:
        return
    titulo = f"Agenda RRHH: {evento.titulo}"
    mensaje = f"{evento.tipo} programado para {evento.inicio.strftime('%d/%m/%Y %H:%M')}."
    if not Notificacion.query.filter_by(usuario_id=evento.responsable_usuario_id, titulo=titulo, mensaje=mensaje).first():
        db.session.add(Notificacion(usuario_id=evento.responsable_usuario_id, titulo=titulo, mensaje=mensaje, url=url_for("rrhh_agenda_editar", evento_id=evento.id)))


@app.route("/rrhh/agenda", methods=["GET", "POST"])
@rrhh_required
def rrhh_agenda():
    hoy = datetime.utcnow().date()
    try:
        year = int(request.args.get("year", hoy.year))
        month = int(request.args.get("month", hoy.month))
        mes_actual = date(year, month, 1)
    except (TypeError, ValueError):
        mes_actual = date(hoy.year, hoy.month, 1)

    if request.method == "POST":
        titulo = request.form.get("titulo", "").strip()
        tipo = request.form.get("tipo", "Tarea").strip()
        inicio_texto = request.form.get("inicio", "").strip()
        fin_texto = request.form.get("fin", "").strip()
        try:
            inicio = datetime.strptime(inicio_texto, "%Y-%m-%dT%H:%M")
            fin = datetime.strptime(fin_texto, "%Y-%m-%dT%H:%M") if fin_texto else None
        except ValueError:
            inicio = None
            fin = None
        colaborador_id = request.form.get("colaborador_id", type=int)
        colaborador = db.session.get(Colaborador, colaborador_id) if colaborador_id else None
        responsable_usuario_id = request.form.get("responsable_usuario_id", type=int)
        responsable_usuario = db.session.get(User, responsable_usuario_id) if responsable_usuario_id else None
        if not titulo or tipo not in {"Tarea", "Reunion", "Recordatorio", "Aviso", "Capacitacion", "Otro"} or not inicio or (fin and fin < inicio):
            flash("Complete un título, tipo y fecha válidos. La fecha final no puede ser anterior al inicio.", "error")
        else:
            evento = EventoRRHH(titulo=titulo, tipo=tipo, inicio=inicio, fin=fin, lugar=request.form.get("lugar", "").strip(), responsable=request.form.get("responsable", "").strip(), responsable_usuario_id=responsable_usuario.id if responsable_usuario else None, notificar=request.form.get("notificar") == "SI", colaborador_id=colaborador.id if colaborador else None, estado="Pendiente", notas=request.form.get("notas", "").strip(), creado_por_id=current_user.id)
            db.session.add(evento)
            db.session.commit()
            notificar_evento_rrhh(evento)
            db.session.commit()
            flash("Evento agendado correctamente.", "success")
            return redirect(url_for("rrhh_agenda", year=inicio.year, month=inicio.month))

    siguiente = (mes_actual.replace(day=28) + timedelta(days=4)).replace(day=1)
    anterior = (mes_actual - timedelta(days=1)).replace(day=1)
    eventos = EventoRRHH.query.filter(EventoRRHH.inicio < siguiente, or_(EventoRRHH.fin.is_(None), EventoRRHH.fin >= mes_actual)).order_by(EventoRRHH.inicio.asc()).all()
    dias = []
    primer_dia = mes_actual.weekday()
    for _ in range(primer_dia):
        dias.append(None)
    for numero in range(1, (siguiente - mes_actual).days + 1):
        dias.append(mes_actual.replace(day=numero))
    while len(dias) % 7:
        dias.append(None)
    colaboradores = Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all()
    supervisores = User.query.filter_by(role="supervisor").order_by(User.full_name.asc()).all()
    return render_template("rrhh_agenda.html", eventos=eventos, dias=dias, mes_actual=mes_actual, anterior=anterior, siguiente=siguiente, colaboradores=colaboradores, supervisores=supervisores, hoy=hoy)


@app.route("/rrhh/agenda/<int:evento_id>/editar", methods=["GET", "POST"])
@rrhh_required
def rrhh_agenda_editar(evento_id):
    evento = EventoRRHH.query.get_or_404(evento_id)
    colaboradores = Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all()
    supervisores = User.query.filter_by(role="supervisor").order_by(User.full_name.asc()).all()
    if request.method == "POST":
        titulo = request.form.get("titulo", "").strip()
        tipo = request.form.get("tipo", "Tarea").strip()
        inicio_texto = request.form.get("inicio", "").strip()
        fin_texto = request.form.get("fin", "").strip()
        try:
            inicio = datetime.strptime(inicio_texto, "%Y-%m-%dT%H:%M")
            fin = datetime.strptime(fin_texto, "%Y-%m-%dT%H:%M") if fin_texto else None
        except ValueError:
            inicio = None
            fin = None
        colaborador_id = request.form.get("colaborador_id", type=int)
        colaborador = db.session.get(Colaborador, colaborador_id) if colaborador_id else None
        responsable_usuario_id = request.form.get("responsable_usuario_id", type=int)
        responsable_usuario = db.session.get(User, responsable_usuario_id) if responsable_usuario_id else None
        estado = request.form.get("estado", evento.estado).strip()
        if not titulo or tipo not in {"Tarea", "Reunion", "Recordatorio", "Aviso", "Capacitacion", "Otro"} or not inicio or (fin and fin < inicio) or estado not in {"Pendiente", "En curso", "Completado", "Cancelado"}:
            flash("Revise título, fechas, tipo y estado del evento.", "error")
        else:
            evento.titulo = titulo
            evento.tipo = tipo
            evento.inicio = inicio
            evento.fin = fin
            evento.lugar = request.form.get("lugar", "").strip()
            evento.responsable = request.form.get("responsable", "").strip()
            evento.responsable_usuario_id = responsable_usuario.id if responsable_usuario else None
            evento.notificar = request.form.get("notificar") == "SI"
            evento.colaborador_id = colaborador.id if colaborador else None
            evento.estado = estado
            evento.notas = request.form.get("notas", "").strip()
            db.session.commit()
            notificar_evento_rrhh(evento)
            db.session.commit()
            flash("Evento actualizado correctamente.", "success")
            return redirect(url_for("rrhh_agenda", year=inicio.year, month=inicio.month))
    return render_template("rrhh_agenda_editar.html", evento=evento, colaboradores=colaboradores, supervisores=supervisores)


@app.post("/rrhh/agenda/<int:evento_id>/estado")
@rrhh_required
def rrhh_agenda_estado(evento_id):
    evento = EventoRRHH.query.get_or_404(evento_id)
    estado = request.form.get("estado", "").strip()
    if estado not in {"Pendiente", "En curso", "Completado", "Cancelado"}:
        flash("Estado de evento no válido.", "error")
    else:
        evento.estado = estado
        db.session.commit()
        flash("Estado del evento actualizado.", "success")
    return redirect(url_for("rrhh_agenda", year=evento.inicio.year, month=evento.inicio.month))


@app.post("/rrhh/agenda/<int:evento_id>/eliminar")
@rrhh_required
def rrhh_agenda_eliminar(evento_id):
    evento = EventoRRHH.query.get_or_404(evento_id)
    mes = evento.inicio
    db.session.delete(evento)
    db.session.commit()
    flash("Evento eliminado correctamente.", "success")
    return redirect(url_for("rrhh_agenda", year=mes.year, month=mes.month))


@app.route("/rrhh/vacaciones", methods=["GET", "POST"])
@rrhh_required
def rrhh_vacaciones():
    if request.method == "POST":
        colaborador = db.session.get(Colaborador, request.form.get("colaborador_id", type=int))
        inicio = parse_date(request.form.get("fecha_inicio", ""))
        try:
            dias = int(request.form.get("dias", "0"))
        except ValueError:
            dias = 0
        if not colaborador or not colaborador.activo or not colaborador.fecha_ingreso or dias != 12 or not inicio:
            flash("Seleccione un colaborador activo, una fecha de inicio y planifique exactamente 12 días.", "error")
        else:
            resumen = resumen_vacaciones(colaborador)
            if resumen["anios_cumplidos"] < 1:
                flash("El colaborador todavía no cumple un año de antigüedad.", "error")
            elif dias > resumen["dias_disponibles"]:
                flash(f"Los días solicitados superan el saldo disponible de {resumen['dias_disponibles']} días.", "error")
            else:
                fin = inicio + timedelta(days=dias - 1)
                agenda = Vacacion(colaborador_id=colaborador.id, fecha_inicio=inicio, fecha_fin=fin, fecha_retorno=fin + timedelta(days=1), dias=dias, observacion=request.form.get("observacion", "").strip(), creado_por_id=current_user.id)
                db.session.add(agenda)
                db.session.flush()
                for usuario in User.query.filter(User.role.in_(["admin", "rrhh"]), User.id != current_user.id).all():
                    db.session.add(Notificacion(
                        usuario_id=usuario.id,
                        titulo=f"Vacaciones agendadas: {colaborador.nombre}",
                        mensaje=f"Período del {inicio:%d/%m/%Y} al {fin:%d/%m/%Y}; retorna el {fin + timedelta(days=1):%d/%m/%Y}.",
                        url=url_for("rrhh_vacaciones_notificaciones"),
                    ))
                db.session.commit()
                flash("Vacaciones agendadas correctamente.", "success")
                return redirect(url_for("rrhh_vacaciones"))
    hoy = datetime.utcnow().date()
    colaboradores = Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all()
    for colaborador in colaboradores:
        colaborador.vacaciones_resumen = resumen_vacaciones(colaborador)
    elegibles = [colaborador for colaborador in colaboradores if colaborador.vacaciones_resumen["anios_cumplidos"] >= 1]
    planificables = [colaborador for colaborador in elegibles if colaborador.vacaciones_resumen["dias_disponibles"] >= 12]
    agendas = Vacacion.query.filter(Vacacion.estado != "Cancelada").order_by(Vacacion.fecha_inicio.asc()).all()
    return render_template("rrhh_vacaciones.html", colaboradores=colaboradores, elegibles=elegibles, planificables=planificables, agendas=agendas, hoy=hoy)


@app.post("/rrhh/vacaciones/<int:vacacion_id>/estado")
@rrhh_required
def actualizar_estado_vacaciones(vacacion_id):
    agenda = Vacacion.query.get_or_404(vacacion_id)
    estado = request.form.get("estado", "").strip()
    if estado not in {"Programada", "En curso", "Finalizada", "Cancelada"}:
        flash("Estado de vacaciones no válido.", "error")
    else:
        agenda.estado = estado
        db.session.commit()
        flash(f"Vacaciones marcadas como: {estado}.", "success")
    return redirect(url_for("rrhh_vacaciones"))


@app.get("/rrhh/vacaciones/notificaciones")
@rrhh_required
def rrhh_vacaciones_notificaciones():
    hoy = datetime.utcnow().date()
    agendas = Vacacion.query.filter(Vacacion.estado != "Cancelada").order_by(Vacacion.fecha_inicio.asc()).all()
    proximas = [item for item in agendas if item.fecha_inicio >= hoy]
    en_curso = [item for item in agendas if item.fecha_inicio <= hoy <= item.fecha_fin]
    notificaciones = Notificacion.query.filter(Notificacion.usuario_id == current_user.id, Notificacion.titulo.like("%Vacaciones%")).order_by(Notificacion.fecha.desc()).limit(100).all()
    return render_template("rrhh_vacaciones_notificaciones.html", proximas=proximas, en_curso=en_curso, notificaciones=notificaciones, hoy=hoy)


@app.post("/rrhh/vacaciones/<int:vacacion_id>/cancelar")
@rrhh_required
def cancelar_vacaciones(vacacion_id):
    agenda = Vacacion.query.get_or_404(vacacion_id)
    agenda.estado = "Cancelada"
    db.session.commit()
    flash("La agenda de vacaciones fue cancelada.", "success")
    return redirect(url_for("rrhh_vacaciones"))


@app.get("/supervisor")
@role_required("admin", "supervisor")
def supervisor_dashboard():
    return render_template("supervisor_dashboard.html")


@app.route("/supervisor/solicitudes", methods=["GET", "POST"])
@role_required("admin", "supervisor")
def supervisor_solicitudes():
    if request.method == "POST":
        productos = request.form.getlist("producto") or [request.form.get("producto", "")]
        cantidades = request.form.getlist("cantidad") or [request.form.get("cantidad", "0")]
        lineas = []
        if len(productos) != len(cantidades):
            flash("Cada producto o servicio debe tener una cantidad.", "error")
            return redirect(url_for("supervisor_solicitudes"))
        for producto, cantidad_texto in zip(productos, cantidades):
            producto = producto.strip()
            try:
                cantidad = float(cantidad_texto or "0")
            except (TypeError, ValueError):
                cantidad = 0
            if producto or cantidad:
                lineas.append((producto, cantidad))
        if not lineas or any(not producto or cantidad <= 0 for producto, cantidad in lineas):
            flash("Complete cada producto o servicio e indique una cantidad válida.", "error")
        else:
            pedido = PedidoInsumo(solicitado_por=current_user.full_name, solicitante_usuario_id=current_user.id, rango_monto="a_presupuestar", observacion=request.form.get("observacion", "").strip())
            for producto, cantidad in lineas:
                insumo = Insumo.query.filter(
                    Insumo.activo.is_(True),
                    or_(Insumo.nombre.ilike(producto), Insumo.codigo.ilike(producto)),
                ).first()
                if not insumo:
                    insumo = Insumo(codigo=f"SRV{secrets.token_hex(5).upper()}", nombre=f"Servicio: {producto} [{secrets.token_hex(4).upper()}]", categoria="Servicio", unidad="servicio", activo=False)
                    db.session.add(insumo)
                    db.session.flush()
                pedido.detalles.append(DetallePedidoInsumo(insumo_id=insumo.id, descripcion=producto, cantidad=cantidad, precio_unitario=0))
            db.session.add(pedido)
            db.session.flush()
            pedido.codigo = f"PED-{pedido.fecha:%Y%m%d}-{pedido.id:06d}"
            for usuario in User.query.filter_by(role="compras").all():
                db.session.add(Notificacion(
                    usuario_id=usuario.id,
                    titulo=f"Nuevo pedido {pedido.codigo}",
                    mensaje=f"{pedido.solicitado_por} envió un pedido para presupuestar.",
                    url=url_for("buzon_compras"),
                ))
            db.session.commit()
            flash("Solicitud enviada a Compras para presupuestar.", "success")
            return redirect(url_for("supervisor_solicitudes"))
    pedidos = PedidoInsumo.query.filter_by(solicitante_usuario_id=current_user.id).order_by(PedidoInsumo.fecha.desc()).all()
    insumos = Insumo.query.filter_by(activo=True).order_by(Insumo.nombre.asc()).all()
    return render_template("supervisor_solicitudes.html", pedidos=pedidos, insumos=insumos)


@app.route("/rrhh/disciplinario", methods=["GET", "POST"])
@rrhh_required
def rrhh_disciplinario():
    if request.method == "POST":
        colaborador = db.session.get(Colaborador, request.form.get("colaborador_id", type=int))
        tipo = request.form.get("tipo", "Amonestación").strip()
        motivo = request.form.get("motivo", "").strip()
        try:
            dias = int(request.form.get("dias_suspendidos", "0"))
        except ValueError:
            dias = -1
        if not colaborador or tipo not in {"Amonestación", "Sanción", "Advertencia", "Suspensión"} or not motivo or dias < 0:
            flash("Complete el colaborador, tipo, motivo y una cantidad válida de días.", "error")
        else:
            registro = RegistroDisciplinario(
                colaborador_id=colaborador.id,
                usuario_id=current_user.id,
                tipo=tipo,
                motivo=motivo,
                dias_suspendidos=dias,
                observacion=request.form.get("observacion", "").strip(),
            )
            db.session.add(registro)
            db.session.flush()
            registro.codigo = f"DISC-{registro.fecha:%Y%m%d}-{registro.id:06d}"
            db.session.commit()
            flash("Registro disciplinario guardado correctamente.", "success")
            return redirect(url_for("rrhh_disciplinario"))
    registros = RegistroDisciplinario.query.order_by(RegistroDisciplinario.fecha.desc()).all()
    colaboradores = Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all()
    return render_template("rrhh_disciplinario.html", registros=registros, colaboradores=colaboradores)


def build_disciplinary_pdf(registro: RegistroDisciplinario) -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("disciplinary_title", parent=styles["Title"], fontSize=16, leading=20, alignment=1, spaceAfter=16)
    body = ParagraphStyle("disciplinary_body", parent=styles["BodyText"], fontSize=10, leading=14, spaceAfter=8)
    small = ParagraphStyle("disciplinary_small", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#5f6b76"))
    logo = Image(str(STATIC_LOGO_PATH), width=125, height=54)
    header = Table([[logo, Paragraph("REGISTRO DISCIPLINARIO", title)]], colWidths=[160, 345])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, -1), 1, colors.HexColor("#123047"))]))
    data = [
        [Paragraph("Código", body), Paragraph(registro.codigo or f"DISC-{registro.id:06d}", body)],
        [Paragraph("Fecha", body), Paragraph(registro.fecha.strftime("%d/%m/%Y %H:%M"), body)],
        [Paragraph("Colaborador", body), Paragraph(registro.colaborador.nombre, body)],
        [Paragraph("Legajo / área", body), Paragraph(f"{registro.colaborador.legajo or '-'} / {registro.colaborador.area or '-'}", body)],
        [Paragraph("Tipo", body), Paragraph(registro.tipo, body)],
        [Paragraph("Días suspendidos", body), Paragraph(str(registro.dias_suspendidos), body)],
        [Paragraph("Motivo", body), Paragraph(registro.motivo, body)],
        [Paragraph("Observación", body), Paragraph(registro.observacion or "-", body)],
    ]
    detail = Table(data, colWidths=[145, 360])
    detail.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8c3cc")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf1f7")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 7)]))
    signatures = Table([["\n\n____________________________", "\n\n____________________________"], ["Firma y sello de RRHH", "Firma del colaborador"]], colWidths=[252, 252])
    signatures.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, 0), 20)]))
    doc.build([header, Spacer(1, 18), detail, Spacer(1, 28), Paragraph("El colaborador declara haber recibido comunicación de este registro.", small), Spacer(1, 18), signatures])
    return buffer.getvalue()


@app.get("/rrhh/disciplinario/<int:registro_id>/pdf")
@rrhh_required
def rrhh_disciplinario_pdf(registro_id: int):
    registro = RegistroDisciplinario.query.get_or_404(registro_id)
    return send_file(BytesIO(build_disciplinary_pdf(registro)), mimetype="application/pdf", as_attachment=True, download_name=f"{registro.codigo or 'registro-disciplinario'}.pdf")


@app.get("/rrhh/disciplinario/<int:registro_id>/pdf/preview")
@rrhh_required
def rrhh_disciplinario_pdf_preview(registro_id: int):
    registro = RegistroDisciplinario.query.get_or_404(registro_id)
    return send_file(BytesIO(build_disciplinary_pdf(registro)), mimetype="application/pdf", as_attachment=False)


def build_vacation_pdf(vacacion: Vacacion) -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("vacation_title", parent=styles["Title"], fontSize=16, leading=20, alignment=1, spaceAfter=16)
    body = ParagraphStyle("vacation_body", parent=styles["BodyText"], fontSize=10, leading=14)
    small = ParagraphStyle("vacation_small", parent=styles["BodyText"], fontSize=8, leading=10, textColor=colors.HexColor("#5f6b76"))
    logo = Image(str(STATIC_LOGO_PATH), width=125, height=54)
    header = Table([[logo, Paragraph("CONSTANCIA DE VACACIONES", title)]], colWidths=[160, 345])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, -1), 1, colors.HexColor("#123047"))]))
    colaborador = vacacion.colaborador
    data = [
        [Paragraph("Colaborador", body), Paragraph(escape(colaborador.nombre), body)],
        [Paragraph("Legajo / área", body), Paragraph(escape(f"{colaborador.legajo or '-'} / {colaborador.area or '-'}"), body)],
        [Paragraph("Fecha de ingreso", body), Paragraph(colaborador.fecha_ingreso.strftime("%d/%m/%Y"), body)],
        [Paragraph("Días correspondientes", body), Paragraph(str(resumen_vacaciones(colaborador)["dias_correspondientes"]), body)],
        [Paragraph("Inicio de vacaciones", body), Paragraph(vacacion.fecha_inicio.strftime("%d/%m/%Y"), body)],
        [Paragraph("Fin de vacaciones", body), Paragraph(vacacion.fecha_fin.strftime("%d/%m/%Y"), body)],
        [Paragraph("Retoma sus labores", body), Paragraph(vacacion.fecha_retorno.strftime("%d/%m/%Y"), body)],
        [Paragraph("Cantidad de días", body), Paragraph(str(vacacion.dias), body)],
        [Paragraph("Observación", body), Paragraph(escape(vacacion.observacion or "-"), body)],
    ]
    detail = Table(data, colWidths=[175, 330])
    detail.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#b8c3cc")), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf1f7")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("PADDING", (0, 0), (-1, -1), 7)]))
    signatures = Table([["\n\n____________________________", "\n\n____________________________"], ["Firma y sello de RRHH", "Firma del colaborador"]], colWidths=[252, 252])
    signatures.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, 0), 24)]))
    doc.build([header, Spacer(1, 18), detail, Spacer(1, 18), Paragraph("Por medio de la presente se comunica el período de descanso anual asignado y la fecha prevista de reintegro.", small), Spacer(1, 24), signatures])
    return buffer.getvalue()


@app.get("/rrhh/vacaciones/<int:vacacion_id>/pdf")
@rrhh_required
def rrhh_vacaciones_pdf(vacacion_id: int):
    vacacion = Vacacion.query.get_or_404(vacacion_id)
    return send_file(BytesIO(build_vacation_pdf(vacacion)), mimetype="application/pdf", as_attachment=True, download_name=f"vacaciones_{vacacion.colaborador.legajo or vacacion.colaborador.id}.pdf")


@app.post("/notificaciones/leer")
@login_required
def marcar_notificaciones_rrhh():
    Notificacion.query.filter_by(usuario_id=current_user.id, leida=False).update({"leida": True})
    db.session.commit()
    fallback = url_for("buzon_compras") if current_user.is_compras else url_for("supervisor_bandeja") if current_user.is_supervisor else url_for("rrhh")
    return redirect(request.form.get("next") or fallback)


@app.get("/supervisor/bandeja")
@role_required("admin", "supervisor", "rrhh")
def supervisor_bandeja():
    notifications = Notificacion.query.filter(
        Notificacion.usuario_id == current_user.id,
        Notificacion.url.like("%asignacion_id=%"),
    ).order_by(Notificacion.fecha.desc()).limit(100).all()
    assignments = {}
    for notification in notifications:
        match = re.search(r"asignacion_id=(\d+)", notification.url or "")
        if match:
            assignments[notification.id] = db.session.get(AsignacionEvaluacion, int(match.group(1)))
    return render_template("supervisor_bandeja.html", notifications=notifications, assignments=assignments)


@app.post("/supervisor/evaluaciones/<int:asignacion_id>/estado")
@role_required("supervisor")
def actualizar_estado_asignacion_evaluacion(asignacion_id):
    asignacion = AsignacionEvaluacion.query.get_or_404(asignacion_id)
    if asignacion.supervisor_id != current_user.id:
        abort(403)
    estado = request.form.get("estado", "").strip()
    estados_validos = {"Recibida", "En proceso"}
    if estado not in estados_validos:
        flash("Estado de evaluación no válido.", "error")
    else:
        asignacion.estado = estado
        mensaje = f"El supervisor {current_user.full_name} indicó que la evaluación de {asignacion.colaborador.nombre} fue {estado.lower()}."
        destinatarios = {asignacion.creado_por_id}
        destinatarios.update(user.id for user in User.query.filter_by(role="rrhh").all())
        for usuario_id in destinatarios:
            if usuario_id != current_user.id:
                db.session.add(Notificacion(
                    usuario_id=usuario_id,
                    titulo=f"Evaluación {estado.lower()}: {asignacion.colaborador.nombre}",
                    mensaje=mensaje,
                    url=url_for("rrhh_evaluaciones", legajo=asignacion.colaborador.legajo, formulario=asignacion.formulario, asignacion_id=asignacion.id, popup=1),
                ))
        db.session.commit()
        flash("RRHH fue notificado del estado de la evaluación.", "success")
    return redirect(url_for("supervisor_bandeja"))


@app.get("/modulo/compras/buzon")
@role_required("admin", "compras")
def buzon_compras():
    notificaciones = Notificacion.query.filter_by(usuario_id=current_user.id).order_by(Notificacion.fecha.desc()).limit(100).all()
    pendientes = PedidoInsumo.query.filter(
        PedidoInsumo.estado_aprobacion == "Pendiente",
        PedidoInsumo.solicitante_usuario_id.isnot(None),
        PedidoInsumo.rango_monto.in_(("mayor", "a_presupuestar")),
    ).order_by(PedidoInsumo.fecha.desc()).all()
    return render_template("buzon_compras.html", notificaciones=notificaciones, pendientes=pendientes)


def generar_recordatorios_rrhh():
    hoy = datetime.utcnow().date()
    usuarios = User.query.filter(User.role.in_(["admin", "rrhh", "supervisor"])).all()
    for colaborador in Colaborador.query.filter_by(activo=True).all():
        if not colaborador.fecha_ingreso:
            continue
        estado = calcular_estado_colaborador(colaborador)
        dias_aniversario = None
        if colaborador.fecha_ingreso.year < hoy.year:
            aniversario = colaborador.fecha_ingreso.replace(year=hoy.year)
            if aniversario < hoy:
                aniversario = aniversario.replace(year=hoy.year + 1)
            dias_aniversario = (aniversario - hoy).days
        fin_prueba = sumar_meses(colaborador.fecha_ingreso, 3)
        dias_prueba = (fin_prueba - hoy).days
        avisos = []
        if dias_aniversario in {0, 15}:
            momento = "hoy" if dias_aniversario == 0 else "en 15 días"
            formulario_anual = "administrativa" if es_personal_administrativo(colaborador) else "periodica"
            codigo_anual = EVALUATION_FORMS[formulario_anual]["codigo"]
            avisos.append((f"Evaluación anual: {colaborador.nombre}", f"{codigo_anual} para {colaborador.nombre}; cumple un año {momento}.", formulario_anual))
        if dias_prueba in {0, 15}:
            momento = "hoy" if dias_prueba == 0 else "en 15 días"
            avisos.append((f"Evaluación de prueba: {colaborador.nombre}", f"RRHH-FOR-007 para {colaborador.nombre}; finaliza el período de prueba {momento}.", "prueba"))
        for titulo, mensaje, formulario in avisos:
            url = f"/rrhh/evaluaciones?legajo={quote_plus(colaborador.legajo)}&formulario={formulario}&popup=1"
            for usuario in usuarios:
                if not Notificacion.query.filter_by(usuario_id=usuario.id, titulo=titulo, mensaje=mensaje).first():
                    db.session.add(Notificacion(usuario_id=usuario.id, titulo=titulo, mensaje=mensaje, url=url))
    for asignacion in AsignacionEvaluacion.query.filter(AsignacionEvaluacion.estado != "Realizada").all():
        if not asignacion.fecha_limite:
            continue
        dias_restantes = (asignacion.fecha_limite - hoy).days
        if dias_restantes not in {0, 15}:
            continue
        momento = "hoy" if dias_restantes == 0 else "en 15 días"
        codigo = EVALUATION_FORMS[asignacion.formulario]["codigo"]
        titulo = f"Evaluación planificada: {asignacion.colaborador.nombre}"
        mensaje = f"{codigo} para {asignacion.colaborador.nombre}; fecha planificada {momento}."
        url = url_for("rrhh_evaluaciones", legajo=asignacion.colaborador.legajo, formulario=asignacion.formulario, asignacion_id=asignacion.id, popup=1)
        destinatarios = {usuario.id for usuario in usuarios}
        destinatarios.add(asignacion.supervisor_id)
        for usuario_id in destinatarios:
            if not Notificacion.query.filter_by(usuario_id=usuario_id, titulo=titulo, mensaje=mensaje).first():
                db.session.add(Notificacion(usuario_id=usuario_id, titulo=titulo, mensaje=mensaje, url=url))
    db.session.commit()


def calcular_estado_colaborador(colaborador):
    if not colaborador.fecha_ingreso:
        return colaborador.estado or "Contratado"
    limite = sumar_meses(colaborador.fecha_ingreso, 3)
    return "Prueba" if datetime.utcnow().date() < limite else "Contratado"


def sumar_meses(fecha, meses):
    mes_total = fecha.month - 1 + meses
    anio = fecha.year + mes_total // 12
    mes = mes_total % 12 + 1
    dia = min(fecha.day, calendar.monthrange(anio, mes)[1])
    return fecha.replace(year=anio, month=mes, day=dia)


def parse_date(value):
    try:
        return datetime.strptime(value, "%Y-%m-%d").date() if value else None
    except ValueError:
        return None


def save_colaborador_photo(uploaded):
    if not uploaded or not uploaded.filename:
        return None
    extension = Path(secure_filename(uploaded.filename)).suffix.lower()
    if extension not in {".jpg", ".jpeg", ".png", ".webp"}:
        raise ValueError("La foto debe estar en formato JPG, PNG o WEBP.")
    try:
        uploaded.stream.seek(0)
        with PILImage.open(uploaded.stream) as image:
            image.verify()
        uploaded.stream.seek(0)
    except Exception as error:
        raise ValueError("El archivo seleccionado no es una imagen válida.") from error
    filename = f"{secrets.token_hex(16)}{extension}"
    uploaded.save(RRHH_PHOTOS_DIR / filename)
    return filename


@app.route("/rrhh/colaboradores", methods=["GET", "POST"])
@rrhh_required
def rrhh_colaboradores():
    if request.method == "POST":
        legajo = request.form.get("legajo", "").strip().upper()
        nombre = request.form.get("nombre", "").strip()
        fecha_ingreso = parse_date(request.form.get("fecha_ingreso", ""))
        fecha_nacimiento = parse_date(request.form.get("fecha_nacimiento", ""))
        vencimiento_cedula = parse_date(request.form.get("vencimiento_cedula", ""))
        try:
            hijos_cantidad = max(0, int(request.form.get("hijos_cantidad", "0") or 0))
        except ValueError:
            hijos_cantidad = 0
        if not legajo or not nombre or not fecha_ingreso:
            flash("Indique legajo, nombre y una fecha de ingreso válida.", "error")
        elif Colaborador.query.filter(func.lower(Colaborador.legajo) == legajo.lower()).first():
            flash("Ese legajo ya existe.", "error")
        else:
            try:
                foto_archivo = save_colaborador_photo(request.files.get("foto"))
                colaborador = Colaborador(codigo=f"COL-{legajo}", legajo=legajo, nombre=nombre, documento=request.form.get("documento", "").strip(), tipo_documento=request.form.get("tipo_documento", "CI").strip(), telefono=request.form.get("telefono", "").strip(), area=request.form.get("area", "").strip(), fecha_ingreso=fecha_ingreso, puesto=request.form.get("puesto", "").strip(), turno_linea=request.form.get("turno_linea", "").strip(), calle=request.form.get("calle", "").strip(), barrio=request.form.get("barrio", "").strip(), ciudad=request.form.get("ciudad", "").strip(), pais=request.form.get("pais", "").strip(), nacionalidad=request.form.get("nacionalidad", "").strip(), fecha_nacimiento=fecha_nacimiento, estado_civil=request.form.get("estado_civil", "").strip(), sexo=request.form.get("sexo", "").strip(), cargo=request.form.get("cargo", "").strip(), seccion=request.form.get("seccion", "").strip(), correo=request.form.get("correo", "").strip(), vencimiento_cedula=vencimiento_cedula, hijos_cantidad=hijos_cantidad, hijos=hijos_cantidad > 0 or request.form.get("hijos") == "SI", aporte_ips=request.form.get("aporte_ips", "SI").strip(), bonificacion_familiar=request.form.get("bonificacion_familiar") == "SI", foto_archivo=foto_archivo)
                colaborador.estado = calcular_estado_colaborador(colaborador)
                db.session.add(colaborador)
                db.session.commit()
                flash("Colaborador registrado correctamente.", "success")
                return redirect(url_for("rrhh_colaboradores"))
            except ValueError as error:
                flash(str(error), "error")
    colaboradores = Colaborador.query.order_by(Colaborador.nombre.asc()).all()
    for colaborador in colaboradores:
        colaborador.estado = calcular_estado_colaborador(colaborador)
    return render_template("rrhh_alta_colaborador.html")


@app.get("/rrhh/colaboradores/padron")
@rrhh_required
def rrhh_padron_colaboradores():
    colaboradores = Colaborador.query.all()
    def legajo_key(colaborador):
        legajo = (colaborador.legajo or "").strip()
        match = re.match(r"^\d+", legajo)
        return (0, int(match.group()), legajo.lower()) if match else (1, 0, legajo.lower())
    colaboradores.sort(key=lambda colaborador: (*legajo_key(colaborador), colaborador.nombre.lower()))
    for colaborador in colaboradores:
        colaborador.estado = calcular_estado_colaborador(colaborador)
    return render_template("rrhh_padron.html", colaboradores=colaboradores)


@app.get("/rrhh/colaboradores/exportar.xlsx")
@rrhh_required
def exportar_colaboradores_rrhh():
    columnas = [
        ("ID", "id"), ("Legajo", "legajo"), ("Nombre completo", "nombre"),
        ("Tipo de documento", "tipo_documento"), ("Documento", "documento"),
        ("Fecha de nacimiento", "fecha_nacimiento"), ("Estado civil", "estado_civil"),
        ("Sexo", "sexo"), ("Vencimiento de cédula", "vencimiento_cedula"),
        ("Teléfono", "telefono"), ("Correo", "correo"), ("Calle", "calle"),
        ("Barrio", "barrio"), ("Ciudad", "ciudad"), ("País", "pais"),
        ("Nacionalidad", "nacionalidad"), ("Cargo", "cargo"), ("Sección", "seccion"),
        ("Puesto", "puesto"), ("Área", "area"), ("Turno / línea", "turno_linea"),
        ("Fecha de ingreso", "fecha_ingreso"), ("Tiene hijos", "hijos"),
        ("Cantidad de hijos", "hijos_cantidad"), ("Aporte IPS", "aporte_ips"),
        ("Bonificación familiar", "bonificacion_familiar"), ("Estado", "estado"),
        ("Activo", "activo"), ("Foto", "foto_archivo"),
    ]
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Colaboradores"
    sheet.append([label for label, _ in columnas])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="087F5B")
    for colaborador in Colaborador.query.order_by(Colaborador.nombre.asc()).all():
        values = []
        for _, attribute in columnas:
            value = getattr(colaborador, attribute)
            if attribute in {"hijos", "bonificacion_familiar", "activo"}:
                value = "Sí" if value else "No"
            values.append(value)
        sheet.append(values)
    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 35)
        sheet.column_dimensions[column_cells[0].column_letter].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if hasattr(cell.value, "strftime"):
                cell.number_format = "dd/mm/yyyy"
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", as_attachment=True, download_name=f"colaboradores_rrhh_{datetime.utcnow():%Y%m%d_%H%M}.xlsx")


def construir_whatsapp_evaluacion(supervisor, colaborador, formulario, evaluation_url):
    telefono = re.sub(r"\D", "", supervisor.telefono or "")
    if telefono.startswith("0"):
        telefono = "595" + telefono[1:]
    elif telefono and not telefono.startswith("595"):
        telefono = "595" + telefono
    mensaje = f"Hola {supervisor.full_name}, RRHH te asignó la evaluación {EVALUATION_FORMS[formulario]['codigo']} de {colaborador.nombre}. Ingresa aquí: {evaluation_url}"
    return f"https://web.whatsapp.com/send?phone={telefono}&text={quote_plus(mensaje)}" if telefono else f"https://web.whatsapp.com/send?text={quote_plus(mensaje)}"


def construir_outlook_evaluacion(supervisor, colaborador, formulario, evaluation_url):
    subject = f"Evaluación asignada: {colaborador.nombre}"
    body = f"Hola {supervisor.full_name},\n\nRRHH te asignó la evaluación {EVALUATION_FORMS[formulario]['codigo']} de {colaborador.nombre}.\n\nIngresa aquí: {evaluation_url}\n\nSaludos."
    return "https://outlook.office.com/mail/deeplink/compose?" + urlencode({"to": supervisor.correo or "", "subject": subject, "body": body})


@app.post("/rrhh/evaluaciones/asignar")
@rrhh_required
def asignar_evaluacion_rrhh():
    colaborador = db.session.get(Colaborador, request.form.get("colaborador_id", type=int))
    supervisor = db.session.get(User, request.form.get("supervisor_id", type=int))
    formulario = request.form.get("formulario", "periodica").strip().lower()
    fecha_limite = parse_date(request.form.get("fecha_limite", ""))
    if not colaborador or not colaborador.activo or not supervisor or supervisor.role != "supervisor" or formulario not in EVALUATION_FORMS or (formulario == "administrativa" and not es_personal_administrativo(colaborador)):
        flash("Seleccione un colaborador, un supervisor y un formulario válidos.", "error")
    else:
        asignacion = AsignacionEvaluacion(colaborador_id=colaborador.id, supervisor_id=supervisor.id, formulario=formulario, fecha_limite=fecha_limite, creado_por_id=current_user.id)
        db.session.add(asignacion)
        db.session.flush()
        db.session.add(Notificacion(
            usuario_id=supervisor.id,
            titulo=f"Evaluación asignada: {colaborador.nombre}",
            mensaje=f"{EVALUATION_FORMS[formulario]['codigo']} pendiente para el legajo {colaborador.legajo}.",
            url=url_for("rrhh_evaluaciones", legajo=colaborador.legajo, formulario=formulario, asignacion_id=asignacion.id),
        ))
        db.session.commit()
        canal = request.form.get("canal_notificacion", "interno").strip().lower()
        if canal in {"correo", "ambos"}:
            return redirect(url_for("correo_evaluacion_rrhh", asignacion_id=asignacion.id))
        flash(f"Evaluación asignada a {supervisor.full_name}.", "success")
    return redirect(url_for("rrhh_evaluaciones"))


@app.route("/rrhh/evaluaciones/planificador", methods=["GET", "POST"])
@rrhh_required
def planificador_evaluaciones_rrhh():
    if request.method == "POST":
        colaborador = db.session.get(Colaborador, request.form.get("colaborador_id", type=int))
        supervisor = db.session.get(User, request.form.get("supervisor_id", type=int))
        formulario = request.form.get("formulario", "").strip().lower()
        fecha_limite = parse_date(request.form.get("fecha_limite", ""))
        formulario_esperado = formulario_evaluacion_colaborador(colaborador) if colaborador else None
        if not colaborador or not colaborador.activo or not supervisor or supervisor.role != "supervisor" or not fecha_limite or formulario != formulario_esperado:
            flash("Seleccione un colaborador, supervisor, formulario y fecha válidos.", "error")
        else:
            asignacion = AsignacionEvaluacion(colaborador_id=colaborador.id, supervisor_id=supervisor.id, formulario=formulario, fecha_limite=fecha_limite, creado_por_id=current_user.id)
            db.session.add(asignacion)
            db.session.flush()
            codigo = EVALUATION_FORMS[formulario]["codigo"]
            db.session.add(Notificacion(usuario_id=supervisor.id, titulo=f"Evaluación planificada: {colaborador.nombre}", mensaje=f"{codigo} programada para el {fecha_limite.strftime('%d/%m/%Y')}.", url=url_for("rrhh_evaluaciones", legajo=colaborador.legajo, formulario=formulario, asignacion_id=asignacion.id, popup=1)))
            db.session.commit()
            flash("Evaluación planificada correctamente.", "success")
        return redirect(url_for("planificador_evaluaciones_rrhh"))
    colaboradores = Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all()
    for colaborador in colaboradores:
        colaborador.estado = calcular_estado_colaborador(colaborador)
        colaborador.formulario_evaluacion = formulario_evaluacion_colaborador(colaborador)
    supervisores = User.query.filter_by(role="supervisor").order_by(User.full_name.asc()).all()
    asignaciones = AsignacionEvaluacion.query.order_by(AsignacionEvaluacion.fecha_limite.asc(), AsignacionEvaluacion.creado_en.desc()).all()
    return render_template("rrhh_planificador_evaluaciones.html", colaboradores=colaboradores, supervisores=supervisores, asignaciones=asignaciones, today=datetime.utcnow().date())


@app.get("/rrhh/evaluaciones/asignaciones/<int:asignacion_id>/correo")
@rrhh_required
def correo_evaluacion_rrhh(asignacion_id):
    asignacion = AsignacionEvaluacion.query.get_or_404(asignacion_id)
    evaluation_url = url_for("rrhh_evaluaciones", legajo=asignacion.colaborador.legajo, formulario=asignacion.formulario, asignacion_id=asignacion.id, _external=True)
    return redirect(construir_outlook_evaluacion(asignacion.supervisor, asignacion.colaborador, asignacion.formulario, evaluation_url))


@app.route("/rrhh/colaboradores/<int:colaborador_id>/editar", methods=["GET", "POST"])
@rrhh_required
def editar_colaborador_rrhh(colaborador_id):
    colaborador = Colaborador.query.get_or_404(colaborador_id)
    if request.method == "POST":
        legajo = request.form.get("legajo", "").strip().upper()
        fecha_ingreso = parse_date(request.form.get("fecha_ingreso", ""))
        fecha_nacimiento = parse_date(request.form.get("fecha_nacimiento", ""))
        vencimiento_cedula = parse_date(request.form.get("vencimiento_cedula", ""))
        try:
            hijos_cantidad = max(0, int(request.form.get("hijos_cantidad", "0") or 0))
        except ValueError:
            hijos_cantidad = 0
        duplicate = Colaborador.query.filter(Colaborador.id != colaborador.id, func.lower(Colaborador.legajo) == legajo.lower()).first()
        if not legajo or not request.form.get("nombre", "").strip() or not fecha_ingreso:
            flash("Indique legajo, nombre y una fecha de ingreso válida.", "error")
        elif duplicate:
            flash("Ese legajo ya existe.", "error")
        else:
            colaborador.legajo = legajo
            colaborador.codigo = f"COL-{legajo}"
            colaborador.nombre = request.form.get("nombre", "").strip()
            colaborador.documento = request.form.get("documento", "").strip()
            colaborador.telefono = request.form.get("telefono", "").strip()
            colaborador.area = request.form.get("area", "").strip()
            colaborador.puesto = request.form.get("puesto", "").strip()
            colaborador.turno_linea = request.form.get("turno_linea", "").strip()
            colaborador.tipo_documento = request.form.get("tipo_documento", "CI").strip()
            colaborador.calle = request.form.get("calle", "").strip()
            colaborador.barrio = request.form.get("barrio", "").strip()
            colaborador.ciudad = request.form.get("ciudad", "").strip()
            colaborador.pais = request.form.get("pais", "").strip()
            colaborador.nacionalidad = request.form.get("nacionalidad", "").strip()
            colaborador.fecha_nacimiento = fecha_nacimiento
            colaborador.estado_civil = request.form.get("estado_civil", "").strip()
            colaborador.sexo = request.form.get("sexo", "").strip()
            colaborador.cargo = request.form.get("cargo", "").strip()
            colaborador.seccion = request.form.get("seccion", "").strip()
            colaborador.correo = request.form.get("correo", "").strip()
            colaborador.vencimiento_cedula = vencimiento_cedula
            colaborador.hijos_cantidad = hijos_cantidad
            colaborador.hijos = hijos_cantidad > 0 or request.form.get("hijos") == "SI"
            colaborador.aporte_ips = request.form.get("aporte_ips", "SI").strip()
            colaborador.bonificacion_familiar = request.form.get("bonificacion_familiar") == "SI"
            colaborador.fecha_ingreso = fecha_ingreso
            colaborador.estado = calcular_estado_colaborador(colaborador)
            try:
                new_photo = save_colaborador_photo(request.files.get("foto"))
            except ValueError as error:
                flash(str(error), "error")
                return render_template("rrhh_colaborador_editar.html", colaborador=colaborador)
            if new_photo:
                old_photo = RRHH_PHOTOS_DIR / colaborador.foto_archivo if colaborador.foto_archivo else None
                colaborador.foto_archivo = new_photo
                if old_photo and old_photo.is_file():
                    old_photo.unlink()
            db.session.commit()
            flash("Colaborador actualizado correctamente.", "success")
            return redirect(url_for("rrhh_colaboradores"))
    return render_template("rrhh_colaborador_editar.html", colaborador=colaborador)


@app.get("/rrhh/colaboradores/<int:colaborador_id>/foto")
@rrhh_required
def foto_colaborador_rrhh(colaborador_id):
    colaborador = Colaborador.query.get_or_404(colaborador_id)
    if not colaborador.foto_archivo:
        abort(404)
    target = (RRHH_PHOTOS_DIR / colaborador.foto_archivo).resolve()
    if target.parent != RRHH_PHOTOS_DIR.resolve() or not target.is_file():
        abort(404)
    return send_file(target)


@app.get("/rrhh/colaboradores/<int:colaborador_id>/whatsapp")
@rrhh_required
def whatsapp_colaborador_rrhh(colaborador_id):
    colaborador = Colaborador.query.get_or_404(colaborador_id)
    telefono = re.sub(r"\D", "", colaborador.telefono or "")
    if telefono.startswith("0"):
        telefono = "595" + telefono[1:]
    elif telefono and not telefono.startswith("595"):
        telefono = "595" + telefono
    if not telefono:
        flash("Este colaborador no tiene un teléfono registrado.", "warning")
        return redirect(url_for("rrhh_padron_colaboradores"))
    return redirect(f"https://web.whatsapp.com/send?phone={telefono}")


@app.post("/rrhh/colaboradores/<int:colaborador_id>/eliminar")
@rrhh_required
def eliminar_colaborador_rrhh(colaborador_id):
    colaborador = Colaborador.query.get_or_404(colaborador_id)
    photo_path = RRHH_PHOTOS_DIR / colaborador.foto_archivo if colaborador.foto_archivo else None
    if colaborador.evaluaciones or colaborador.salidas or colaborador.retiros:
        colaborador.activo = False
        db.session.commit()
        flash("El colaborador tiene historial asociado y fue desactivado para conservar la trazabilidad.", "success")
    else:
        db.session.delete(colaborador)
        db.session.commit()
        if photo_path and photo_path.is_file():
            photo_path.unlink()
        flash("Colaborador eliminado correctamente.", "success")
    return redirect(url_for("rrhh_colaboradores"))


@app.route("/rrhh/evaluaciones", methods=["GET", "POST"])
@evaluacion_required
def rrhh_evaluaciones():
    legajo = request.args.get("legajo", "").strip()
    historial_search = request.args.get("historial", "").strip()
    historial_estado = request.args.get("historial_estado", "").strip()
    popup = request.args.get("popup") == "1"
    formulario = (request.form.get("formulario") if request.method == "POST" else request.args.get("formulario", "prueba")).strip().lower()
    form_config = EVALUATION_FORMS.get(formulario, EVALUATION_FORMS["prueba"])
    criterion_count = sum(len(criteria) for _, criteria in form_config["groups"])
    colaborador = None
    colaboradores_pendientes = []
    candidatos_query = Colaborador.query.filter_by(activo=True)
    candidatos = candidatos_query.order_by(Colaborador.nombre.asc()).all()
    evaluaciones_existentes = {
        (colaborador_id, formulario_nombre)
        for colaborador_id, formulario_nombre in EvaluacionPersonal.query.with_entities(EvaluacionPersonal.colaborador_id, EvaluacionPersonal.formulario).distinct().all()
    }
    for candidato in candidatos:
        candidato.estado = calcular_estado_colaborador(candidato)
        candidato_formulario = formulario_evaluacion_colaborador(candidato)
        candidato.formulario_evaluacion = candidato_formulario
        if (candidato.id, candidato_formulario) not in evaluaciones_existentes:
            ultima = None
            colaboradores_pendientes.append({"colaborador": candidato, "formulario": candidato_formulario, "ultima": ultima})
    colaboradores_pendientes_prueba = [item for item in colaboradores_pendientes if item["formulario"] == "prueba"]
    colaboradores_pendientes_periodica = [item for item in colaboradores_pendientes if item["formulario"] == "periodica"]
    if legajo:
        colaborador = Colaborador.query.filter(func.lower(Colaborador.legajo) == legajo.lower(), Colaborador.activo.is_(True)).first()
        if colaborador is None:
            if not current_user.is_supervisor:
                flash("No se encontró un colaborador activo con ese legajo.", "error")
        else:
            colaborador.estado = calcular_estado_colaborador(colaborador)
    if request.method == "POST":
        colaborador_id = request.form.get("colaborador_id", "")
        colaborador = db.session.get(Colaborador, int(colaborador_id)) if colaborador_id.isdigit() else None
        puntajes = {}
        criterion_index = 0
        for _, criteria in form_config["groups"]:
            for criterion in criteria:
                value = request.form.get(f"puntaje_{criterion_index}", "")
                if value not in {"1", "2", "3", "4", "5"}:
                    flash("Todos los criterios deben tener un puntaje entre 1 y 5.", "error")
                    return render_template("rrhh_evaluaciones.html", colaborador=colaborador, evaluaciones=EvaluacionPersonal.query.order_by(EvaluacionPersonal.fecha.desc()).all(), colaboradores_pendientes=colaboradores_pendientes, colaboradores_pendientes_prueba=colaboradores_pendientes_prueba, colaboradores_pendientes_periodica=colaboradores_pendientes_periodica, legajo=legajo, formulario=formulario, form_config=form_config, criterion_count=criterion_count, today=datetime.utcnow().date(), popup=popup)
                puntajes[criterion] = int(value)
                criterion_index += 1
        if colaborador:
            colaborador.estado = calcular_estado_colaborador(colaborador)
        formulario_esperado = formulario_evaluacion_colaborador(colaborador) if colaborador else None
        es_formulario_valido = formulario == formulario_esperado
        estado_invalido = form_config["estado"] != "Todos" and colaborador.estado.strip().lower() != form_config["estado"].lower() if colaborador else True
        if not colaborador or not colaborador.activo or estado_invalido or not es_formulario_valido:
            flash(f"Este formulario solo corresponde a colaboradores en estado {form_config['estado']}.", "error")
        else:
            total = sum(puntajes.values())
            promedio = round(total / criterion_count, 2)
            if formulario == "prueba":
                resultado = "Satisfactorio / Confirmación en el puesto" if promedio >= 3.8 else "Aceptable / En observación" if promedio >= 3 else "Insuficiente / No supera período de prueba"
            else:
                resultado = "Desempeño Destacado" if promedio >= 3.8 else "Desempeño Aceptable" if promedio >= 3 else "Desempeño Bajo"
            assignment_id = request.form.get("asignacion_id", type=int) or request.args.get("asignacion_id", type=int)
            asignacion = db.session.get(AsignacionEvaluacion, assignment_id) if assignment_id else None
            if asignacion and (asignacion.supervisor_id != current_user.id or asignacion.colaborador_id != colaborador.id):
                asignacion = None
            evaluacion = EvaluacionPersonal(
                colaborador_id=colaborador.id,
                usuario_id=current_user.id,
                formulario=formulario,
                codigo=form_config["codigo"],
                puntaje_total=total,
                puntajes=json.dumps(puntajes, ensure_ascii=False),
                promedio=promedio,
                resultado=resultado,
                fecha_evaluacion=parse_date(request.form.get("fecha_evaluacion", "")) or datetime.utcnow().date(),
                periodo_evaluado=request.form.get("periodo_evaluado", "").strip(),
                tipo_evaluacion=request.form.get("tipo_evaluacion", "A demanda / especial").strip(),
                jefe_directo=request.form.get("jefe_directo", "").strip(),
                fortalezas=request.form.get("fortalezas", "").strip(),
                aspectos_mejorar=request.form.get("aspectos_mejorar", "").strip(),
                plan_accion=request.form.get("plan_accion", "").strip(),
                fecha_revision=parse_date(request.form.get("fecha_revision", "")),
                recomendacion=request.form.get("recomendacion", "").strip(),
                estado=request.form.get("estado", "Borrador").strip() or "Borrador",
                observaciones=request.form.get("observaciones", "").strip(),
            )
            db.session.add(evaluacion)
            db.session.flush()
            if asignacion:
                asignacion.evaluacion_id = evaluacion.id
                asignacion.estado = "Realizada" if evaluacion.estado in {"Completada", "Revisada"} else "En proceso"
                estado_mensaje = "completó" if asignacion.estado == "Realizada" else "inició"
                mensaje = f"El supervisor {current_user.full_name} {estado_mensaje} la evaluación de {colaborador.nombre}. Resultado: {resultado}."
                destinatarios = {asignacion.creado_por_id}
                destinatarios.update(user.id for user in User.query.filter_by(role="rrhh").all())
                for usuario_id in destinatarios:
                    if usuario_id != current_user.id:
                        db.session.add(Notificacion(
                            usuario_id=usuario_id,
                            titulo=f"Evaluación {'realizada' if asignacion.estado == 'Realizada' else 'en proceso'}: {colaborador.nombre}",
                            mensaje=mensaje,
                            url=url_for("rrhh_evaluaciones", legajo=colaborador.legajo, formulario=formulario, asignacion_id=asignacion.id, popup=1),
                        ))
            for user in User.query.filter((User.role == "admin") | (User.role == "rrhh") | (User.role == "root"), User.id != current_user.id).all():
                if not asignacion:
                    db.session.add(Notificacion(usuario_id=user.id, titulo=f"Evaluar a {colaborador.nombre}", mensaje=f"{form_config['codigo']} pendiente para legajo {colaborador.legajo}.", url=url_for("rrhh_evaluaciones", legajo=colaborador.legajo, formulario=formulario, popup=1)))
            db.session.commit()
            flash("Evaluación guardada correctamente.", "success")
            return redirect(url_for("vista_previa_evaluacion_rrhh", evaluacion_id=evaluacion.id))
    historico_query = EvaluacionPersonal.query.join(Colaborador)
    if historial_search:
        pattern = f"%{historial_search.lower()}%"
        historico_query = historico_query.filter(or_(func.lower(Colaborador.nombre).like(pattern), func.lower(Colaborador.legajo).like(pattern), func.lower(Colaborador.area).like(pattern)))
    if historial_estado:
        historico_query = historico_query.filter(EvaluacionPersonal.estado == historial_estado)
    evaluaciones = historico_query.order_by(EvaluacionPersonal.fecha.desc()).all()
    supervisores = User.query.filter_by(role="supervisor").order_by(User.full_name.asc()).all()
    return render_template("rrhh_evaluaciones.html", colaborador=colaborador, evaluaciones=evaluaciones, colaboradores_pendientes=colaboradores_pendientes, colaboradores_pendientes_prueba=colaboradores_pendientes_prueba, colaboradores_pendientes_periodica=colaboradores_pendientes_periodica, colaboradores_asignables=candidatos, supervisores=supervisores, historial_search=historial_search, historial_estado=historial_estado, legajo=legajo, formulario=formulario, form_config=form_config, criterion_count=criterion_count, today=datetime.utcnow().date(), popup=popup)


@app.get("/rrhh/evaluaciones/<int:evaluacion_id>/vista-previa")
@evaluacion_required
def vista_previa_evaluacion_rrhh(evaluacion_id):
    evaluacion = EvaluacionPersonal.query.get_or_404(evaluacion_id)
    return send_file(BytesIO(build_evaluacion_rrhh_pdf(evaluacion)), mimetype="application/pdf", as_attachment=False, download_name=f"{evaluacion.codigo}_{evaluacion.colaborador.legajo}.pdf")


@app.get("/rrhh/evaluaciones/plantilla/<formulario>.pdf")
@evaluacion_required
def plantilla_evaluacion_rrhh(formulario):
    if formulario not in EVALUATION_FORMS or formulario == "documento_rrhh":
        abort(404)
    form_config = EVALUATION_FORMS[formulario]
    return send_file(BytesIO(build_evaluacion_rrhh_template_pdf(formulario)), mimetype="application/pdf", as_attachment=True, download_name=f"{form_config['codigo']}_plantilla.pdf")


@app.get("/rrhh/evaluaciones/registradas")
@rrhh_required
def rrhh_evaluaciones_registradas():
    historial_search = request.args.get("historial", "").strip()
    historial_estado = request.args.get("historial_estado", "").strip()
    historico_query = EvaluacionPersonal.query.join(Colaborador)
    if historial_search:
        pattern = f"%{historial_search.lower()}%"
        historico_query = historico_query.filter(or_(func.lower(Colaborador.nombre).like(pattern), func.lower(Colaborador.legajo).like(pattern), func.lower(Colaborador.area).like(pattern)))
    if historial_estado:
        historico_query = historico_query.filter(EvaluacionPersonal.estado == historial_estado)
    evaluaciones = historico_query.order_by(EvaluacionPersonal.fecha.desc()).all()
    estados = [estado for (estado,) in db.session.query(EvaluacionPersonal.estado).distinct().order_by(EvaluacionPersonal.estado).all() if estado]
    return render_template("rrhh_evaluaciones_registradas.html", evaluaciones=evaluaciones, estados=estados, historial_search=historial_search, historial_estado=historial_estado)


@app.route("/rrhh/documentos", methods=["GET", "POST"])
@rrhh_required
def rrhh_documentos():
    allowed_extensions = {"docx", "xlsx", "txt", "pdf", "doc", "xls", "csv"}
    if request.method == "POST":
        if request.form.get("accion") == "crear_carpeta":
            nombre_carpeta = secure_filename(request.form.get("nombre_carpeta", "").strip()).replace("_", " ")
            if not nombre_carpeta:
                flash("Indique un nombre válido para la carpeta.", "error")
            elif CarpetaRRHH.query.filter(db.func.lower(CarpetaRRHH.nombre) == nombre_carpeta.lower()).first():
                flash("Ya existe una carpeta con ese nombre.", "error")
            else:
                db.session.add(CarpetaRRHH(nombre=nombre_carpeta))
                db.session.commit()
                flash("Carpeta creada correctamente.", "success")
            return redirect(url_for("rrhh_documentos"))
        if request.form.get("accion") == "eliminar_carpeta":
            carpeta = db.session.get(CarpetaRRHH, request.form.get("carpeta_id", type=int))
            if not carpeta:
                flash("La carpeta no existe.", "error")
            elif carpeta.documentos:
                flash("No se puede eliminar una carpeta que contiene documentos. Mueva o elimine sus archivos primero.", "error")
            else:
                db.session.delete(carpeta)
                db.session.commit()
                flash("Carpeta eliminada correctamente.", "success")
            return redirect(url_for("rrhh_documentos"))
        uploaded = request.files.get("documento")
        if not uploaded or not uploaded.filename:
            flash("Seleccione un documento para subir.", "error")
        else:
            original_name = secure_filename(uploaded.filename)
            extension = Path(original_name).suffix.lower().removeprefix(".")
            if not original_name or extension not in allowed_extensions:
                flash("Formato no permitido. Use DOCX, XLSX, TXT, PDF, DOC, XLS o CSV.", "error")
            else:
                carpeta_id_value = next((value for value in request.form.getlist("carpeta_id") if value.strip()), "")
                carpeta_id = int(carpeta_id_value) if carpeta_id_value.isdigit() else None
                carpeta = db.session.get(CarpetaRRHH, carpeta_id) if carpeta_id else None
                if carpeta_id and carpeta is None:
                    flash("La carpeta seleccionada no existe.", "error")
                    return redirect(url_for("rrhh_documentos"))
                stored_name = f"{secrets.token_hex(12)}_{original_name}"
                target = RRHH_DOCUMENTS_DIR / stored_name
                uploaded.save(target)
                now = datetime.utcnow()
                db.session.add(DocumentoRRHH(
                    nombre=request.form.get("nombre", "").strip() or Path(original_name).stem,
                    nombre_archivo=stored_name,
                    tipo=uploaded.mimetype or extension,
                    categoria=request.form.get("categoria", "Otros").strip() or "Otros",
                                        carpeta_id=carpeta.id if carpeta else None,
                    tamano=target.stat().st_size,
                    fecha_creacion=now,
                    fecha_modificacion=now,
                    usuario_creacion_id=current_user.id,
                    usuario_modificacion_id=current_user.id,
                ))
                db.session.commit()
                flash("Documento guardado correctamente.", "success")
                return redirect(url_for("rrhh_documentos"))
    carpeta_id = request.args.get("carpeta_id", type=int)
    carpeta_actual = db.session.get(CarpetaRRHH, carpeta_id) if carpeta_id else None
    if carpeta_id and carpeta_actual is None:
        abort(404)
    documentos_query = DocumentoRRHH.query.filter_by(carpeta_id=carpeta_id) if carpeta_id else DocumentoRRHH.query.filter(DocumentoRRHH.carpeta_id.is_(None))
    return render_template("rrhh_documentos.html", documentos=documentos_query.order_by(DocumentoRRHH.fecha_modificacion.desc()).all(), carpetas=CarpetaRRHH.query.order_by(CarpetaRRHH.nombre.asc()).all(), carpeta_actual=carpeta_actual)


@app.route("/rrhh/documentos/<int:documento_id>/archivo")
@rrhh_required
def rrhh_documento_archivo(documento_id: int):
    documento = DocumentoRRHH.query.get_or_404(documento_id)
    target = (RRHH_DOCUMENTS_DIR / documento.nombre_archivo).resolve()
    if target.parent != RRHH_DOCUMENTS_DIR.resolve() or not target.is_file():
        abort(404)
    return send_file(target, mimetype=documento.tipo, as_attachment=False, download_name=documento.nombre)


@app.route("/rrhh/documentos/<int:documento_id>/editar", methods=["GET", "POST"])
@rrhh_required
def editar_documento_rrhh(documento_id: int):
    documento = DocumentoRRHH.query.get_or_404(documento_id)
    extension = Path(documento.nombre_archivo).suffix.lower()
    if extension not in {".txt", ".csv", ".docx", ".xlsx"}:
        flash("Este formato requiere un editor Office interno para editarse en línea.", "error")
        return redirect(url_for("rrhh_documentos"))
    target = (RRHH_DOCUMENTS_DIR / documento.nombre_archivo).resolve()
    if target.parent != RRHH_DOCUMENTS_DIR.resolve() or not target.is_file():
        abort(404)
    if request.method == "POST":
        nuevo_nombre = request.form.get("nombre", "").strip()
        nueva_carpeta_id = request.form.get("carpeta_id", type=int)
        nueva_carpeta = db.session.get(CarpetaRRHH, nueva_carpeta_id) if nueva_carpeta_id else None
        if nuevo_nombre:
            documento.nombre = nuevo_nombre
        if nueva_carpeta_id and not nueva_carpeta:
            flash("La carpeta seleccionada no existe.", "error")
            return redirect(url_for("editar_documento_rrhh", documento_id=documento.id))
        documento.carpeta_id = nueva_carpeta.id if nueva_carpeta else None
        if extension in {".txt", ".csv"}:
            target.write_text(request.form.get("contenido", ""), encoding="utf-8")
        elif extension == ".docx":
            word = WordDocument(str(target))
            paragraphs = request.form.get("contenido", "").splitlines()
            for index, paragraph in enumerate(word.paragraphs):
                paragraph.text = paragraphs[index] if index < len(paragraphs) else ""
            word.save(str(target))
        else:
            workbook = load_workbook(target)
            sheet = workbook.active
            for row in range(1, sheet.max_row + 1):
                for column in range(1, sheet.max_column + 1):
                    key = f"cell_{row}_{column}"
                    if key in request.form:
                        sheet.cell(row=row, column=column).value = request.form[key]
            workbook.save(target)
        documento.fecha_modificacion = datetime.utcnow()
        documento.tamano = target.stat().st_size
        documento.usuario_modificacion_id = current_user.id
        db.session.commit()
        flash("Documento actualizado correctamente.", "success")
        return redirect(url_for("rrhh_documentos"))
    if extension == ".xlsx":
        workbook = load_workbook(target, read_only=True, data_only=False)
        sheet = workbook.active
        max_rows = min(sheet.max_row or 1, 60)
        max_columns = min(sheet.max_column or 1, 15)
        rows = [list(row) for row in sheet.iter_rows(min_row=1, max_row=max_rows, min_col=1, max_col=max_columns, values_only=True)]
        rows = [["" if value is None else value for value in row] for row in rows]
        workbook.close()
        return render_template("rrhh_documento_editar.html", documento=documento, carpetas=CarpetaRRHH.query.order_by(CarpetaRRHH.nombre.asc()).all(), tipo_editor="xlsx", rows=rows, hoja=sheet.title, filas_mostradas=max_rows, filas_totales=sheet.max_row, columnas_mostradas=max_columns, columnas_totales=sheet.max_column)
    if extension == ".docx":
        word = WordDocument(str(target))
        contenido = "\n".join(paragraph.text for paragraph in word.paragraphs)
    else:
        contenido = target.read_text(encoding="utf-8")
    return render_template("rrhh_documento_editar.html", documento=documento, carpetas=CarpetaRRHH.query.order_by(CarpetaRRHH.nombre.asc()).all(), tipo_editor="text", contenido=contenido)


@app.route("/rrhh/documentos/<int:documento_id>/descargar")
@rrhh_required
def descargar_documento_rrhh(documento_id: int):
    documento = DocumentoRRHH.query.get_or_404(documento_id)
    target = (RRHH_DOCUMENTS_DIR / documento.nombre_archivo).resolve()
    if target.parent != RRHH_DOCUMENTS_DIR.resolve() or not target.is_file():
        abort(404)
    return send_file(target, as_attachment=True, download_name=documento.nombre)


@app.post("/rrhh/documentos/<int:documento_id>/eliminar")
@rrhh_required
def eliminar_documento_rrhh(documento_id: int):
    documento = DocumentoRRHH.query.get_or_404(documento_id)
    target = RRHH_DOCUMENTS_DIR / documento.nombre_archivo
    if target.is_file():
        target.unlink()
    db.session.delete(documento)
    db.session.commit()
    flash("Documento eliminado correctamente.", "success")
    return redirect(url_for("rrhh_documentos"))


@app.post("/rrhh/documentos/carpetas/<int:carpeta_id>/eliminar")
@rrhh_required
def eliminar_carpeta_rrhh(carpeta_id: int):
    carpeta = CarpetaRRHH.query.get_or_404(carpeta_id)
    if carpeta.documentos:
        flash("No se puede eliminar una carpeta que contiene documentos.", "error")
    else:
        db.session.delete(carpeta)
        db.session.commit()
        flash("Carpeta eliminada correctamente.", "success")
    return redirect(url_for("rrhh_documentos"))


@app.get("/admin/informe-auditado.pdf")
@admin_required
def informe_auditado_pdf():
    return send_file(
        BytesIO(build_system_audit_pdf()),
        mimetype="application/pdf",
        as_attachment=True,
        download_name="informe_auditado_sistema_IGP.pdf",
    )


@app.get("/admin/codigo-fuente.pdf")
@admin_required
def codigo_fuente_pdf():
    return send_file(
        BytesIO(build_source_code_pdf()),
        mimetype="application/pdf",
        as_attachment=True,
        download_name="codigo_fuente_sistema.pdf",
    )


@app.route("/inventario")
@inventory_required
def inventario():
    resumen = get_inventario_resumen()
    category_map = {}
    for item in resumen:
        category = item["insumo"].categoria or "General"
        group = category_map.setdefault(category, {"stock": 0.0, "insumos": 0, "bajos": 0})
        group["stock"] += item["stock"]
        group["insumos"] += 1
        group["bajos"] += int(item["bajo"])
    movimientos = []
    for entrada in EntradaInsumo.query.order_by(EntradaInsumo.fecha.desc()).limit(8).all():
        movimientos.append({"fecha": entrada.fecha, "tipo": "Entrada", "codigo": entrada.insumo.codigo, "nombre": entrada.insumo.nombre, "cantidad": entrada.cantidad, "unidad": entrada.insumo.unidad})
    for salida in SalidaInsumo.query.order_by(SalidaInsumo.fecha.desc()).limit(8).all():
        movimientos.append({"fecha": salida.fecha, "tipo": "Salida", "codigo": salida.insumo.codigo, "nombre": salida.insumo.nombre, "cantidad": salida.cantidad, "unidad": salida.insumo.unidad})
    movimientos.sort(key=lambda item: item["fecha"], reverse=True)
    return render_template(
        "inventario.html",
        resumen=resumen,
        category_data=sorted(({"nombre": nombre, **values} for nombre, values in category_map.items()), key=lambda item: item["nombre"]),
        movimientos=movimientos[:12],
        stock_total=sum(item["stock"] for item in resumen),
        low_count=sum(int(item["bajo"]) for item in resumen),
        pending_orders=PedidoInsumo.query.filter_by(estado="Pendiente").count(),
        total_entries=EntradaInsumo.query.count(),
        total_outputs=SalidaInsumo.query.count(),
    )


@app.route("/inventario/insumos")
@inventory_required
def inventario_insumos():
    return render_template("inventario_insumos.html", insumos=Insumo.query.filter_by(activo=True).order_by(Insumo.nombre.asc()).all(), categorias=CategoriaInsumo.query.filter_by(activo=True).order_by(CategoriaInsumo.nombre.asc()).all(), resumen_inventario=get_inventario_resumen())


@app.route("/inventario/insumos/<int:insumo_id>/editar", methods=["GET", "POST"])
@inventory_required
def editar_insumo(insumo_id: int):
    insumo = Insumo.query.get_or_404(insumo_id)
    categorias = CategoriaInsumo.query.filter_by(activo=True).order_by(CategoriaInsumo.nombre.asc()).all()
    if request.method == "POST":
        codigo = request.form.get("codigo", "").strip().upper()
        nombre = request.form.get("nombre", "").strip()
        categoria = request.form.get("categoria", "").strip()
        unidad = request.form.get("unidad", "").strip() or "unidad"
        try:
            stock_minimo = float(request.form.get("stock_minimo", "0"))
        except ValueError:
            stock_minimo = -1
        vencimiento = request.form.get("vencimiento", "").strip()
        try:
            vencimiento_date = datetime.strptime(vencimiento, "%Y-%m-%d").date() if vencimiento else None
        except ValueError:
            vencimiento_date = None
        categoria_obj = CategoriaInsumo.query.filter_by(nombre=categoria, activo=True).first()
        duplicate = Insumo.query.filter(func.lower(Insumo.nombre) == nombre.lower(), Insumo.id != insumo.id).first()
        duplicate_code = Insumo.query.filter(func.lower(Insumo.codigo) == codigo.lower(), Insumo.id != insumo.id).first() if codigo else None
        if not nombre or stock_minimo < 0:
            flash("Indique un nombre y mínimo válido para el insumo.", "error")
        elif not categoria_obj:
            flash("Seleccione una categoría registrada.", "error")
        elif duplicate:
            flash("Ya existe otro insumo con ese nombre.", "error")
        elif duplicate_code:
            flash("Ese código ya existe.", "error")
        elif vencimiento and vencimiento_date is None:
            flash("El vencimiento no tiene un formato válido.", "error")
        else:
            insumo.codigo = codigo or None
            insumo.nombre = nombre
            insumo.categoria = categoria
            insumo.unidad = unidad
            insumo.stock_minimo = stock_minimo
            insumo.vencimiento = vencimiento_date
            db.session.commit()
            flash("Insumo actualizado correctamente.", "success")
            return redirect(url_for("inventario_insumos"))
    return render_template("editar_insumo.html", insumo=insumo, categorias=categorias)


@app.post("/inventario/insumos/<int:insumo_id>/eliminar")
@inventory_required
def eliminar_insumo(insumo_id: int):
    insumo = Insumo.query.get_or_404(insumo_id)
    insumo.activo = False
    db.session.commit()
    flash(f"El insumo {insumo.nombre} fue retirado del catálogo activo.", "success")
    return redirect(url_for("inventario_insumos"))


@app.route("/inventario/categorias", methods=["GET", "POST"])
@inventory_required
def inventario_categorias():
    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        prefijo = normalize_prefix(request.form.get("prefijo", ""))
        if not nombre:
            flash("Indique el nombre de la categoría.", "error")
        elif not prefijo:
            flash("Indique un prefijo alfanumérico para la categoría.", "error")
        elif CategoriaInsumo.query.filter(func.lower(CategoriaInsumo.nombre) == nombre.lower()).first():
            flash("Esa categoría ya existe.", "error")
        elif CategoriaInsumo.query.filter(func.lower(CategoriaInsumo.prefijo) == prefijo.lower()).first():
            flash("Ese prefijo ya está asignado a otra categoría.", "error")
        else:
            db.session.add(CategoriaInsumo(nombre=nombre, prefijo=prefijo))
            db.session.commit()
            flash("Categoría registrada correctamente.", "success")
            return redirect(url_for("inventario_categorias"))
    return render_template("inventario_categorias.html", categorias=CategoriaInsumo.query.order_by(CategoriaInsumo.nombre.asc()).all())


@app.route("/inventario/categorias/<int:categoria_id>/editar", methods=["GET", "POST"])
@inventory_required
def editar_categoria_insumo(categoria_id: int):
    categoria = CategoriaInsumo.query.get_or_404(categoria_id)
    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        prefijo = normalize_prefix(request.form.get("prefijo", ""))
        duplicate_name = CategoriaInsumo.query.filter(
            func.lower(CategoriaInsumo.nombre) == nombre.lower(),
            CategoriaInsumo.id != categoria.id,
        ).first()
        duplicate_prefix = CategoriaInsumo.query.filter(
            func.lower(CategoriaInsumo.prefijo) == prefijo.lower(),
            CategoriaInsumo.id != categoria.id,
        ).first()
        if not nombre:
            flash("Indique el nombre de la categoría.", "error")
        elif not prefijo:
            flash("Indique un prefijo alfanumérico para la categoría.", "error")
        elif duplicate_name:
            flash("Esa categoría ya existe.", "error")
        elif duplicate_prefix:
            flash("Ese prefijo ya está asignado a otra categoría.", "error")
        else:
            old_name = categoria.nombre
            categoria.nombre = nombre
            categoria.prefijo = prefijo
            Insumo.query.filter_by(categoria=old_name).update({"categoria": nombre})
            db.session.commit()
            flash("Categoría actualizada correctamente.", "success")
            return redirect(url_for("inventario_categorias"))
    return render_template("editar_categoria.html", categoria=categoria)


@app.post("/inventario/categorias/<int:categoria_id>/eliminar")
@inventory_required
def eliminar_categoria_insumo(categoria_id: int):
    categoria = CategoriaInsumo.query.get_or_404(categoria_id)
    if Insumo.query.filter_by(categoria=categoria.nombre, activo=True).first():
        flash("No se puede eliminar una categoría que tiene insumos asociados.", "error")
    else:
        db.session.delete(categoria)
        db.session.commit()
        flash("Categoría eliminada correctamente.", "success")
    return redirect(url_for("inventario_categorias"))


@app.post("/inventario/insumos")
@inventory_required
def crear_insumo():
    codigo = request.form.get("codigo", "").strip().upper()
    nombre = request.form.get("nombre", "").strip()
    categoria = request.form.get("categoria", "General").strip() or "General"
    unidad = request.form.get("unidad", "unidad").strip() or "unidad"
    try:
        stock_minimo = float(request.form.get("stock_minimo", "0"))
    except ValueError:
        stock_minimo = -1
    vencimiento = request.form.get("vencimiento", "").strip()
    vencimiento_valido = True
    try:
        vencimiento_date = datetime.strptime(vencimiento, "%Y-%m-%d").date() if vencimiento else None
    except ValueError:
        vencimiento_date = None
        vencimiento_valido = False
        flash("El vencimiento no tiene un formato válido.", "error")
    categoria_obj = CategoriaInsumo.query.filter_by(nombre=categoria, activo=True).first()
    if not nombre or stock_minimo < 0:
        flash("Indique un nombre y mínimo válido para el insumo.", "error")
    elif not categoria_obj:
        flash("Seleccione una categoría registrada.", "error")
    elif not vencimiento_valido:
        pass
    elif Insumo.query.filter(func.lower(Insumo.nombre) == nombre.lower()).first():
        flash("Ese insumo ya existe.", "error")
    elif codigo and Insumo.query.filter(func.lower(Insumo.codigo) == codigo.lower()).first():
        flash("Ese código ya existe.", "error")
    else:
        db.session.add(Insumo(codigo=codigo or None, nombre=nombre, categoria=categoria, unidad=unidad, stock_minimo=stock_minimo, vencimiento=vencimiento_date))
        db.session.commit()
        flash("Insumo creado correctamente.", "success")
    return redirect(url_for("inventario_insumos"))


@app.post("/inventario/colaboradores")
@inventory_required
def crear_colaborador():
    legajo = request.form.get("legajo", "").strip().upper()
    nombre = request.form.get("nombre", "").strip()
    if not legajo or not nombre:
        flash("Indique el legajo y nombre del colaborador.", "error")
    elif Colaborador.query.filter(func.lower(Colaborador.legajo) == legajo.lower()).first():
        flash("Ese legajo ya existe.", "error")
    else:
        db.session.add(Colaborador(
            codigo=f"COL-{legajo}",
            legajo=legajo,
            nombre=nombre,
            documento=request.form.get("documento", "").strip(),
            area=request.form.get("area", "").strip(),
            estado=request.form.get("estado", "Contratado").strip() or "Contratado",
        ))
        db.session.commit()
        flash("Colaborador creado correctamente.", "success")
    return redirect(url_for("inventario_colaboradores"))


@app.post("/inventario/entradas")
@inventory_required
def registrar_entrada_insumo():
    if offline_operation_was_processed():
        return redirect(url_for("inventario_entradas"))
    try:
        insumo_id = int(request.form.get("insumo_id", "0"))
        cantidad = float(request.form.get("cantidad", "0"))
    except ValueError:
        insumo_id, cantidad = 0, 0
    insumo = db.session.get(Insumo, insumo_id)
    if not insumo or cantidad <= 0:
        flash("Seleccione un insumo e indique una cantidad válida.", "error")
    else:
        db.session.add(EntradaInsumo(
            insumo_id=insumo.id,
            cantidad=cantidad,
            observacion=request.form.get("observacion", "").strip(),
            usuario_id=current_user.id,
        ))
        remember_offline_operation()
        db.session.commit()
        flash("Entrada registrada correctamente.", "success")
    return redirect(url_for("inventario_entradas"))


@app.route("/inventario/colaboradores")
@inventory_required
def inventario_colaboradores():
    return render_template("inventario_colaboradores.html", colaboradores=Colaborador.query.order_by(Colaborador.nombre.asc()).all())


@app.route("/inventario/entradas")
@inventory_required
def inventario_entradas():
    product_query = request.args.get("producto", "").strip()
    selected_insumo_id = request.args.get("insumo_id", type=int)
    insumos_query = Insumo.query.filter_by(activo=True)
    if product_query:
        pattern = f"%{product_query}%"
        insumos_query = insumos_query.filter(or_(Insumo.codigo.ilike(pattern), Insumo.nombre.ilike(pattern)))
    return render_template(
        "inventario_entradas.html",
        insumos=insumos_query.order_by(Insumo.codigo.asc()).all(),
        entradas=EntradaInsumo.query.order_by(EntradaInsumo.fecha.desc()).all(),
        product_query=product_query,
        selected_insumo_id=selected_insumo_id,
    )


def get_retiro_cart():
    return {str(insumo_id): float(cantidad) for insumo_id, cantidad in session.get("retiro_cart", {}).items()}


@app.post("/inventario/salidas/carrito/agregar")
@inventory_required
def agregar_salida_carrito():
    if offline_operation_was_processed():
        return redirect(url_for("inventario_salidas"))
    try:
        insumo_id = int(request.form.get("insumo_id", "0"))
        colaborador_id = int(request.form.get("colaborador_id", "0"))
        cantidad = float(request.form.get("cantidad", "0"))
    except ValueError:
        insumo_id, colaborador_id, cantidad = 0, 0, 0
    insumo = db.session.get(Insumo, insumo_id)
    cart = get_retiro_cart()
    stock_actual = get_stock_insumo(insumo_id) if insumo else 0
    if not insumo or cantidad <= 0:
        flash("Seleccione un insumo e indique una cantidad válida.", "error")
    elif cantidad + cart.get(str(insumo_id), 0) > stock_actual:
        disponible = max(stock_actual - cart.get(str(insumo_id), 0), 0)
        flash(f"Stock insuficiente. Disponible para agregar: {disponible:.2f} {insumo.unidad}.", "error")
    else:
        cart[str(insumo_id)] = cart.get(str(insumo_id), 0) + cantidad
        session["retiro_cart"] = cart
        remember_offline_operation()
        db.session.commit()
        flash(f"{insumo.nombre} agregado al carrito.", "success")
    return redirect(url_for("inventario_salidas"))


@app.post("/inventario/salidas/carrito/quitar/<int:insumo_id>")
@inventory_required
def quitar_salida_carrito(insumo_id: int):
    cart = get_retiro_cart()
    cart.pop(str(insumo_id), None)
    session["retiro_cart"] = cart
    flash("Producto retirado del carrito.", "success")
    return redirect(url_for("inventario_salidas"))


@app.post("/inventario/salidas/carrito/vaciar")
@inventory_required
def vaciar_salida_carrito():
    session.pop("retiro_cart", None)
    flash("Carrito de retiro vaciado.", "success")
    return redirect(url_for("inventario_salidas"))


@app.post("/inventario/salidas/confirmar")
@inventory_required
def registrar_salida_insumo():
    if offline_operation_was_processed():
        return redirect(url_for("inventario_salidas"))
    colaborador_id = request.form.get("colaborador_id", "")
    observacion = request.form.get("observacion", "").strip()
    colaborador = db.session.get(Colaborador, int(colaborador_id)) if colaborador_id.isdigit() else None
    cart = get_retiro_cart()
    if not colaborador or not cart:
        flash("Seleccione un colaborador y agregue al menos un producto al carrito.", "error")
        return redirect(url_for("inventario_salidas"))
    items = []
    for insumo_id, cantidad in cart.items():
        insumo = db.session.get(Insumo, int(insumo_id))
        if not insumo or cantidad <= 0 or cantidad > get_stock_insumo(insumo.id):
            flash(f"El stock cambió para {insumo.nombre if insumo else 'un producto'}. Revise el carrito.", "error")
            return redirect(url_for("inventario_salidas"))
        items.append((insumo, cantidad))
    retiro = RetiroInsumo(
        codigo=f"TMP-{secrets.token_hex(8).upper()}",
        colaborador_id=colaborador.id,
        observacion=observacion,
        usuario_id=current_user.id,
    )
    db.session.add(retiro)
    db.session.flush()
    retiro.codigo = f"RET-{retiro.fecha:%Y%m%d}-{retiro.id:06d}"
    for insumo, cantidad in items:
        db.session.add(SalidaInsumo(
            retiro_id=retiro.id,
            insumo_id=insumo.id,
            colaborador_id=colaborador.id,
            cantidad=cantidad,
            observacion=observacion,
            usuario_id=current_user.id,
        ))
    remember_offline_operation()
    db.session.commit()
    session.pop("retiro_cart", None)
    flash(f"Retiro {retiro.codigo} registrado correctamente.", "success")
    return redirect(url_for("detalle_retiro_insumo", retiro_id=retiro.id))


@app.route("/inventario/salidas")
@inventory_required
def inventario_salidas():
    insumo_query = request.args.get("insumo", "").strip()
    colaborador_query = request.args.get("colaborador", "").strip()
    insumos_query = Insumo.query.filter_by(activo=True)
    colaboradores_query = Colaborador.query.filter_by(activo=True)
    cart = []
    for insumo_id, cantidad in get_retiro_cart().items():
        insumo = db.session.get(Insumo, int(insumo_id))
        if insumo:
            cart.append({"insumo": insumo, "cantidad": cantidad})
    return render_template(
        "inventario_salidas.html",
        insumos=insumos_query.order_by(Insumo.codigo.asc()).all(),
        colaboradores=colaboradores_query.order_by(Colaborador.nombre.asc()).all(),
        colaboradores_todos=Colaborador.query.filter_by(activo=True).order_by(Colaborador.nombre.asc()).all(),
        cart=cart,
        retiros=RetiroInsumo.query.order_by(RetiroInsumo.fecha.desc()).limit(50).all(),
        insumo_query=insumo_query,
        colaborador_query=colaborador_query,
    )


@app.route("/inventario/retiros/<int:retiro_id>")
@inventory_required
def detalle_retiro_insumo(retiro_id: int):
    retiro = RetiroInsumo.query.get_or_404(retiro_id)
    return render_template("retiro_detalle.html", retiro=retiro)


@app.route("/inventario/retiros/<int:retiro_id>/pdf")
@inventory_required
def pdf_retiro_insumo(retiro_id: int):
    retiro = RetiroInsumo.query.get_or_404(retiro_id)
    return send_file(
        BytesIO(build_retiro_pdf(retiro)),
        mimetype="application/pdf",
        as_attachment=False,
        download_name=f"retiro_{retiro.codigo}.pdf",
    )


@app.post("/inventario/pedidos")
@inventory_required
def crear_pedido_insumo():
    if offline_operation_was_processed():
        return redirect(url_for("inventario_pedidos"))
    try:
        insumo_id = int(request.form.get("insumo_id", "0"))
        cantidad = float(request.form.get("cantidad", "0"))
    except ValueError:
        insumo_id, cantidad = 0, 0
    insumo = db.session.get(Insumo, insumo_id)
    solicitado_por = request.form.get("solicitado_por", "").strip()
    if not insumo or cantidad <= 0 or not solicitado_por:
        flash("Complete el insumo, la cantidad y quién solicita el pedido.", "error")
    else:
        pedido = PedidoInsumo(solicitado_por=solicitado_por, estado_aprobacion="Aprobado", observacion=request.form.get("observacion", "").strip())
        try:
            precio_unitario = float(request.form.get("precio_unitario", "0"))
        except ValueError:
            precio_unitario = -1
        if precio_unitario < 0:
            flash("El precio unitario estimado no es válido.", "error")
            return redirect(url_for("inventario_pedidos"))
        pedido.detalles.append(DetallePedidoInsumo(insumo_id=insumo.id, cantidad=cantidad, precio_unitario=precio_unitario))
        pedido.fecha_evaluacion = datetime.utcnow()
        pedido.observacion_evaluacion = "Pedido de Inventario enviado directamente a Compras."
        db.session.add(pedido)
        db.session.flush()
        pedido.codigo = f"PED-{pedido.fecha:%Y%m%d}-{pedido.id:06d}"
        remember_offline_operation()
        db.session.commit()
        flash("Pedido registrado y enviado a Compras." if pedido.estado_aprobacion == "Aprobado" else "Pedido registrado y enviado a Gerencia para evaluación.", "success")
        return redirect(url_for("detalle_pedido_insumo", pedido_id=pedido.id))
    return redirect(url_for("inventario_pedidos"))


def pedido_orden_requerida(pedido):
    try:
        umbral = float(get_config_value("umbral_orden_compra_gs", "1000000"))
        cotizacion = float(get_config_value("cotizacion_usd", "7500"))
    except (TypeError, ValueError):
        umbral, cotizacion = 1000000.0, 7500.0
    presupuesto_mayor = max((presupuesto.monto for presupuesto in pedido.presupuestos), default=0.0)
    total_evaluable = max(pedido.total_estimado, presupuesto_mayor)
    return total_evaluable >= umbral, umbral, cotizacion


app.jinja_env.globals["pedido_orden_requerida"] = pedido_orden_requerida


@app.route("/gerencia/pedidos", methods=["GET", "POST"])
@gerencia_required
def gerencia_pedidos():
    if request.method == "POST":
        pedido = PedidoInsumo.query.get_or_404(request.form.get("pedido_id", type=int))
        decision = request.form.get("decision", "")
        presupuesto_id = request.form.get("presupuesto_id", type=int)
        presupuesto = PresupuestoPedido.query.filter_by(id=presupuesto_id, pedido_id=pedido.id).first() if presupuesto_id else None
        if decision not in {"Aprobado", "Rechazado"}:
            flash("Seleccione una decisión válida.", "error")
        elif decision == "Aprobado" and (pedido.rango_monto in {"mayor", "a_presupuestar"} or pedido_orden_requerida(pedido)[0]) and len(pedido.presupuestos) < 3:
            flash("Este pedido requiere como mínimo 3 presupuestos adjuntos antes de aprobarlo.", "error")
        elif decision == "Aprobado" and (pedido.rango_monto in {"mayor", "a_presupuestar"} or pedido_orden_requerida(pedido)[0]) and not presupuesto:
            flash("Debe seleccionar un presupuesto antes de aprobar este pedido.", "error")
        else:
            pedido.estado_aprobacion = decision
            pedido.evaluado_por_id = current_user.id
            pedido.fecha_evaluacion = datetime.utcnow()
            pedido.observacion_evaluacion = request.form.get("observacion_evaluacion", "").strip()
            pedido.presupuesto_elegido_id = presupuesto.id if presupuesto else None
            if decision == "Aprobado":
                for usuario in User.query.filter_by(role="compras").all():
                    db.session.add(Notificacion(
                        usuario_id=usuario.id,
                        titulo=f"Gerencia aprobó {pedido.codigo}",
                        mensaje=f"La orden de compra de {pedido.solicitado_por} está lista para gestionar.",
                        url=url_for("ordenes_compra"),
                    ))
            if pedido.solicitante_usuario_id:
                db.session.add(Notificacion(
                    usuario_id=pedido.solicitante_usuario_id,
                    titulo=f"Pedido {pedido.codigo}: {decision}",
                    mensaje=f"Gerencia {decision.lower()} su solicitud.",
                    url=url_for("supervisor_dashboard"),
                ))
            db.session.add(HistorialAprobacionPedido(
                pedido_id=pedido.id,
                decision=decision,
                fecha=pedido.fecha_evaluacion,
                usuario_id=current_user.id,
                observacion=pedido.observacion_evaluacion,
                presupuesto_id=presupuesto.id if presupuesto else None,
            ))
            if decision == "Aprobado" and (pedido.rango_monto in {"mayor", "a_presupuestar"} or pedido_orden_requerida(pedido)[0]) and not pedido.orden_compra:
                orden = OrdenCompra(pedido_id=pedido.id)
                db.session.add(orden)
                db.session.flush()
                orden.codigo = f"OC-{orden.fecha:%Y%m%d}-{orden.id:06d}"
            db.session.commit()
            flash(f"Pedido {decision.lower()} correctamente.", "success")
            return redirect(url_for("gerencia_pedidos"))
    pedidos = PedidoInsumo.query.filter_by(estado_aprobacion="Pendiente", presupuestos_enviados=True).order_by(PedidoInsumo.fecha.desc()).all()
    pendientes_compras = PedidoInsumo.query.filter_by(estado_aprobacion="Pendiente", presupuestos_enviados=False).filter(PedidoInsumo.solicitante_usuario_id.isnot(None)).order_by(PedidoInsumo.fecha.desc()).all()
    umbral, cotizacion = pedido_orden_requerida(pedidos[0]) [1:] if pedidos else (float(get_config_value("umbral_orden_compra_gs", "1000000")), float(get_config_value("cotizacion_usd", "7500")))
    return render_template("gerencia_pedidos.html", pedidos=pedidos, pendientes_compras=pendientes_compras, umbral=umbral, cotizacion=cotizacion)


@app.get("/modulo/compras/historial-aprobaciones")
@role_required("admin", "compras", "gerencia")
def historial_aprobaciones():
    # Conserva en el nuevo histórico las decisiones antiguas que solo tenían fecha en el pedido.
    pedidos_evaluados = PedidoInsumo.query.filter(
        PedidoInsumo.estado_aprobacion.in_(("Aprobado", "Rechazado")),
    ).all()
    cambios = 0
    for pedido in pedidos_evaluados:
        if not HistorialAprobacionPedido.query.filter_by(pedido_id=pedido.id).first():
            db.session.add(HistorialAprobacionPedido(
                pedido_id=pedido.id,
                decision=pedido.estado_aprobacion,
                fecha=pedido.fecha_evaluacion or pedido.fecha,
                usuario_id=pedido.evaluado_por_id,
                observacion=pedido.observacion_evaluacion or "",
                presupuesto_id=pedido.presupuesto_elegido_id,
            ))
            cambios += 1
    if cambios:
        db.session.commit()
    historial = HistorialAprobacionPedido.query.order_by(HistorialAprobacionPedido.fecha.desc()).all()
    return render_template("historial_aprobaciones.html", historial=historial)


@app.post("/modulo/compras/pedidos/<int:pedido_id>/presupuestos")
@role_required("admin", "compras")
def agregar_presupuesto_pedido(pedido_id):
    pedido = PedidoInsumo.query.get_or_404(pedido_id)
    if request.form.get("proveedor_1") or request.files.get("archivo_1"):
        cargados = 0
        for numero in range(1, 4):
            proveedor = request.form.get(f"proveedor_{numero}", "").strip()
            archivo = request.files.get(f"archivo_{numero}")
            try:
                monto = float(request.form.get(f"monto_{numero}", "0"))
            except (TypeError, ValueError):
                monto = 0
            if not proveedor and not archivo:
                continue
            extension = Path(secure_filename(archivo.filename)).suffix.lower() if archivo and archivo.filename else ""
            if not proveedor and monto <= 0 and not archivo:
                continue
            if not proveedor or monto <= 0 or (archivo and extension not in {".pdf", ".jpg", ".jpeg", ".png"}):
                flash(f"Complete correctamente el presupuesto {numero}.", "error")
                return redirect(url_for("solicitudes_compra"))
            nombre = ""
            if archivo:
                nombre = f"{secrets.token_hex(16)}{extension}"
                archivo.save(PURCHASE_QUOTES_DIR / nombre)
            db.session.add(PresupuestoPedido(pedido_id=pedido.id, proveedor=proveedor, monto=monto, archivo=nombre))
            cargados += 1
        if cargados:
            db.session.commit()
            flash(f"{cargados} presupuesto(s) adjuntado(s) correctamente.", "success")
        else:
            flash("Adjunte al menos un presupuesto.", "error")
        return redirect(url_for("solicitudes_compra"))
    proveedor = request.form.get("proveedor", "").strip()
    archivo = request.files.get("archivo")
    try:
        monto = float(request.form.get("monto", "0"))
    except ValueError:
        monto = 0
    extension = Path(secure_filename(archivo.filename)).suffix.lower() if archivo and archivo.filename else ""
    if not proveedor or monto <= 0 or (archivo and extension not in {".pdf", ".jpg", ".jpeg", ".png"}):
        flash("Indique proveedor y monto. El archivo adjunto es opcional.", "error")
    else:
        nombre = ""
        if archivo:
            nombre = f"{secrets.token_hex(16)}{extension}"
            archivo.save(PURCHASE_QUOTES_DIR / nombre)
        db.session.add(PresupuestoPedido(pedido_id=pedido.id, proveedor=proveedor, monto=monto, archivo=nombre))
        db.session.commit()
        flash("Presupuesto adjuntado correctamente.", "success")
    return redirect(url_for("modulo_compras"))


@app.post("/modulo/compras/pedidos/<int:pedido_id>/enviar-gerencia")
@role_required("admin", "compras")
def enviar_presupuestos_gerencia(pedido_id):
    pedido = PedidoInsumo.query.get_or_404(pedido_id)
    if len(pedido.presupuestos) < 3:
        flash("Debe adjuntar como mínimo 3 presupuestos antes de enviarlos a Gerencia.", "error")
    else:
        pedido.presupuestos_enviados = True
        for usuario in User.query.filter_by(role="gerencia").all():
            db.session.add(Notificacion(usuario_id=usuario.id, titulo=f"Solicitud {pedido.codigo} lista para evaluación", mensaje=f"Compras presentó 3 o más presupuestos para {pedido.codigo}.", url="/gerencia/pedidos"))
        db.session.commit()
        flash(f"Presupuestos del pedido {pedido.codigo} enviados a Gerencia.", "success")
    return redirect(url_for("modulo_compras"))


@app.route("/modulo/compras/solicitudes", methods=["GET", "POST"])
@role_required("admin", "compras")
def solicitudes_compra():
    if request.method == "POST":
        supervisor_id = request.form.get("supervisor_id") or request.form.get("supervisor")
        try:
            supervisor_id = int(supervisor_id or 0)
        except (TypeError, ValueError):
            supervisor_id = 0
        supervisor = User.query.filter(
            User.id == supervisor_id,
            func.lower(func.trim(User.role)) == "supervisor",
        ).first()

        producto = (request.form.get("producto") or request.form.get("product") or "").strip()
        insumo_id = request.form.get("insumo_id") or request.form.get("insumo")
        try:
            insumo_id = int(insumo_id or 0)
        except (TypeError, ValueError):
            insumo_id = 0
        insumo = db.session.get(Insumo, insumo_id)
        if not insumo and producto:
            insumo = Insumo.query.filter(
                Insumo.activo.is_(True),
                or_(Insumo.nombre.ilike(producto), Insumo.codigo.ilike(producto)),
            ).first()
        if not insumo and producto:
            insumo = Insumo(
                codigo=f"SRV{secrets.token_hex(5).upper()}",
                nombre=f"Servicio: {producto} [{secrets.token_hex(4).upper()}]",
                categoria="Servicio",
                unidad="servicio",
                activo=False,
            )
            db.session.add(insumo)
            db.session.flush()
        if not insumo and producto:
            insumo = Insumo(
                codigo=f"SRV{secrets.token_hex(5).upper()}",
                nombre=f"Servicio: {producto} [{secrets.token_hex(4).upper()}]",
                categoria="Servicio",
                unidad="servicio",
                activo=False,
            )
            db.session.add(insumo)
            db.session.flush()
        try:
            cantidad = float(request.form.get("cantidad", "0"))
        except (TypeError, ValueError):
            cantidad = 0
        if not supervisor:
            flash("Seleccione el Supervisor solicitante.", "error")
        elif not insumo or not insumo.activo:
            flash("Seleccione un producto válido del catálogo. Si no aparece, solicite a Depósito que lo cargue.", "error")
        elif cantidad <= 0:
            flash("Indique una cantidad mayor que cero.", "error")
        else:
            pedido = PedidoInsumo(solicitado_por=supervisor.full_name, solicitante_usuario_id=supervisor.id, rango_monto="mayor", observacion=request.form.get("observacion", "").strip())
            pedido.detalles.append(DetallePedidoInsumo(insumo_id=insumo.id, cantidad=cantidad, precio_unitario=0))
            db.session.add(pedido)
            db.session.flush()
            pedido.codigo = f"PED-{pedido.fecha:%Y%m%d}-{pedido.id:06d}"
            archivos_guardados = []
            for numero in range(1, 4):
                proveedor = request.form.get(f"proveedor_{numero}", "").strip()
                archivo = request.files.get(f"archivo_{numero}")
                if not proveedor and not archivo and not request.form.get(f"monto_{numero}"):
                    continue
                try:
                    monto = float(request.form.get(f"monto_{numero}", "0"))
                except (TypeError, ValueError):
                    monto = 0
                extension = Path(secure_filename(archivo.filename)).suffix.lower() if archivo and archivo.filename else ""
                if not proveedor or monto <= 0 or not archivo or extension not in {".pdf", ".jpg", ".jpeg", ".png"}:
                    flash(f"Complete correctamente el presupuesto {numero}.", "error")
                    db.session.rollback()
                    for ruta in archivos_guardados:
                        ruta.unlink(missing_ok=True)
                    return redirect(url_for("solicitudes_compra"))
                nombre = f"{secrets.token_hex(16)}{extension}"
                ruta = PURCHASE_QUOTES_DIR / nombre
                archivo.save(ruta)
                archivos_guardados.append(ruta)
                pedido.presupuestos.append(PresupuestoPedido(proveedor=proveedor, monto=monto, archivo=nombre))
            db.session.commit()
            flash(f"Solicitud {pedido.codigo} registrada para {supervisor.full_name}.", "success")
            return redirect(url_for("solicitudes_compra"))
    pedido_seleccionado_id = request.args.get("pedido_id", type=int)
    solicitudes_query = PedidoInsumo.query.filter(PedidoInsumo.estado_aprobacion == "Pendiente", PedidoInsumo.solicitante_usuario_id.isnot(None), PedidoInsumo.rango_monto.in_(("mayor", "a_presupuestar")))
    if pedido_seleccionado_id:
        solicitudes_query = solicitudes_query.filter(PedidoInsumo.id == pedido_seleccionado_id)
    solicitudes = solicitudes_query.order_by(PedidoInsumo.fecha.desc()).all()
    supervisores = User.query.filter_by(role="supervisor").order_by(User.full_name.asc()).all()
    insumos = Insumo.query.filter_by(activo=True).order_by(Insumo.nombre.asc()).all()
    return render_template("solicitudes_compra.html", solicitudes=solicitudes, supervisores=supervisores, insumos=insumos, pedido_seleccionado_id=pedido_seleccionado_id)


@app.get("/gerencia/pedidos/<int:pedido_id>/presupuestos/<int:presupuesto_id>/archivo")
@role_required("admin", "compras", "gerencia")
def ver_presupuesto_pedido(pedido_id, presupuesto_id):
    presupuesto = PresupuestoPedido.query.filter_by(id=presupuesto_id, pedido_id=pedido_id).first_or_404()
    target = (PURCHASE_QUOTES_DIR / presupuesto.archivo).resolve()
    if target.parent != PURCHASE_QUOTES_DIR.resolve() or not target.is_file():
        abort(404)
    return send_file(target, as_attachment=False, download_name=presupuesto.archivo)


@app.get("/modulo/compras/ordenes/<int:orden_id>/pdf")
@role_required("admin", "compras", "gerencia")
def orden_compra_pdf(orden_id):
    orden = OrdenCompra.query.get_or_404(orden_id)
    pedido = orden.pedido
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=36, bottomMargin=36)
    styles = getSampleStyleSheet()
    title = ParagraphStyle("oc_title", parent=styles["Title"], fontSize=17, alignment=1, spaceAfter=14)
    body = ParagraphStyle("oc_body", parent=styles["BodyText"], fontSize=9, leading=12)
    data = [["Código", orden.codigo or f"OC-{orden.id:06d}"], ["Pedido", pedido.codigo or f"PED-{pedido.id:06d}"], ["Fecha", orden.fecha.strftime("%d/%m/%Y %H:%M")], ["Solicitante", pedido.solicitado_por], ["Aprobado por", pedido.evaluado_por.full_name if getattr(pedido, "evaluado_por", None) else "Gerencia"]]
    data += [[detalle.insumo.nombre, f"{detalle.cantidad:g} {detalle.insumo.unidad} x {detalle.precio_unitario:,.0f} Gs"] for detalle in pedido.detalles]
    data.append(["TOTAL APROBADO", f"{pedido.total_aprobado:,.0f} Gs"])
    table = Table(data, colWidths=[170, 335])
    table.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), .5, colors.grey), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#eaf1f7")), ("PADDING", (0, 0), (-1, -1), 7)]))
    signatures = Table([["\n\n____________________________", "\n\n____________________________"], ["Firma de Gerencia", "Firma de Compras"]], colWidths=[252, 252])
    signatures.setStyle(TableStyle([("ALIGN", (0, 0), (-1, -1), "CENTER"), ("FONTSIZE", (0, 0), (-1, -1), 9), ("TOPPADDING", (0, 0), (-1, 0), 20)]))
    doc.build([Paragraph("ORDEN DE COMPRA", title), table, Spacer(1, 22), signatures])
    return send_file(BytesIO(buffer.getvalue()), mimetype="application/pdf", as_attachment=True, download_name=f"{orden.codigo or 'orden-compra'}.pdf")


def orden_compra_whatsapp_message(orden):
    pedido = orden.pedido
    proveedor = pedido.presupuesto_elegido.proveedor if pedido.presupuesto_elegido else "No informado"
    return "\n".join([
        f"📄 Orden de Compra {orden.codigo or f'OC-{orden.id:06d}'}",
        f"📋 Pedido: {pedido.codigo or f'PED-{pedido.id:06d}'}",
        f"👤 Solicitante: {pedido.solicitado_por}",
        f"🏢 Proveedor aprobado: {proveedor}",
        f"💰 Total aprobado: {pedido.total_aprobado:,.0f} Gs",
        f"✅ Estado: {orden.estado}",
        f"📅 Fecha: {orden.fecha:%d/%m/%Y %H:%M}",
    ])


app.jinja_env.globals["orden_compra_whatsapp_message"] = orden_compra_whatsapp_message


@app.get("/modulo/compras/ordenes")
@role_required("admin", "compras", "gerencia")
def ordenes_compra():
    pedidos_aprobados = PedidoInsumo.query.filter_by(estado_aprobacion="Aprobado").all()
    ordenes_faltantes = [pedido for pedido in pedidos_aprobados if not pedido.orden_compra]
    for pedido in ordenes_faltantes:
        orden = OrdenCompra(pedido_id=pedido.id)
        db.session.add(orden)
        db.session.flush()
        orden.codigo = f"OC-{orden.fecha:%Y%m%d}-{orden.id:06d}"
    if ordenes_faltantes:
        db.session.commit()
    ordenes = OrdenCompra.query.order_by(OrdenCompra.fecha.desc()).all()
    rechazados = PedidoInsumo.query.filter_by(estado_aprobacion="Rechazado").filter(PedidoInsumo.solicitante_usuario_id.isnot(None)).order_by(PedidoInsumo.fecha.desc()).all()
    return render_template("ordenes_compra.html", ordenes=ordenes, rechazados=rechazados)


@app.post("/modulo/compras/ordenes/<int:orden_id>/estado")
@role_required("admin", "compras")
def actualizar_estado_orden_compra(orden_id):
    orden = OrdenCompra.query.get_or_404(orden_id)
    estado = request.form.get("estado", "Abierta")
    estados_validos = {"Abierta", "Aún pendiente", "Presupuestado", "En proceso de compra", "Entregado", "Recibida", "Cerrada"}
    if estado not in estados_validos:
        flash("Seleccione un estado válido para la orden.", "error")
        return redirect(url_for("ordenes_compra"))
    orden.estado = estado
    pedido = orden.pedido
    pedido.estado = {"Abierta": "Pendiente", "Aún pendiente": "Pendiente", "Presupuestado": "Presupuestado", "En proceso de compra": "En proceso de compra", "Entregado": "Entregado", "Recibida": "Recibido", "Cerrada": "Entregado"}.get(estado, estado)
    if estado == "Cerrada":
        orden.fecha_cierre = datetime.utcnow()
        orden.observacion_cierre = request.form.get("observacion", "").strip()
    else:
        orden.fecha_cierre = None
        orden.observacion_cierre = ""
    if pedido.estado == "Entregado":
        Notificacion.query.filter(
            Notificacion.usuario_id == pedido.solicitante_usuario_id,
            Notificacion.titulo.like(f"%{pedido.codigo}%"),
        ).delete(synchronize_session=False)
    else:
        notificar_estado_pedido(pedido)
    db.session.commit()
    flash(f"{orden.codigo} actualizada a {estado}.", "success")
    return redirect(url_for("ordenes_compra"))


@app.post("/modulo/compras/ordenes/<int:orden_id>/cerrar")
@role_required("admin", "compras")
def cerrar_orden_compra(orden_id):
    orden = OrdenCompra.query.get_or_404(orden_id)
    orden.estado = "Cerrada"
    orden.pedido.estado = "Entregado"
    orden.fecha_cierre = datetime.utcnow()
    orden.observacion_cierre = request.form.get("observacion_cierre", "").strip()
    Notificacion.query.filter(
        Notificacion.usuario_id == orden.pedido.solicitante_usuario_id,
        Notificacion.titulo.like(f"%{orden.pedido.codigo}%"),
    ).delete(synchronize_session=False)
    db.session.commit()
    flash(f"{orden.codigo} cerrada correctamente.", "success")
    return redirect(url_for("ordenes_compra"))


@app.post("/inventario/pedidos/<int:pedido_id>/estado")
@inventory_required
def cambiar_estado_pedido(pedido_id: int):
    pedido = PedidoInsumo.query.get_or_404(pedido_id)
    estado = request.form.get("estado", "Pendiente")
    estados_permitidos = PEDIDO_ESTADOS if current_user.is_compras or current_user.is_admin else DEPOSITO_PEDIDO_ESTADOS
    if estado in estados_permitidos:
        pedido.estado = estado
        notificar_estado_pedido(pedido)
        db.session.commit()
        flash("Estado del pedido actualizado.", "success")
    next_url = request.form.get("next", "")
    if next_url.startswith("/modulo/compras"):
        return redirect(next_url)
    return redirect(url_for("inventario_pedidos"))


@app.route("/inventario/pedidos")
@inventory_required
def inventario_pedidos():
    pedidos_visibles = PedidoInsumo.query.order_by(PedidoInsumo.fecha.desc()).all()
    return render_template("inventario_pedidos.html", insumos=Insumo.query.filter_by(activo=True).order_by(Insumo.nombre.asc()).all(), pedidos=pedidos_visibles)


@app.route("/inventario/pedidos/<int:pedido_id>")
@inventory_required
def detalle_pedido_insumo(pedido_id: int):
    pedido = PedidoInsumo.query.get_or_404(pedido_id)
    return render_template("pedido_detalle.html", pedido=pedido, whatsapp_url=pedido_whatsapp_url(pedido), solo_compras=False)


@app.route("/modulo/compras/pedidos/<int:pedido_id>")
@role_required("admin", "compras", "gerencia")
def detalle_pedido_compras(pedido_id: int):
    pedido = PedidoInsumo.query.get_or_404(pedido_id)
    return render_template("pedido_detalle.html", pedido=pedido, whatsapp_url=pedido_whatsapp_url(pedido), solo_compras=True)


@app.route("/inventario/reportes")
@inventory_required
def reportes_inventario():
    return redirect(url_for("reporte_existencias"))


@app.route("/inventario/reportes/existencias")
@inventory_required
def reporte_existencias():
    filtros = reporte_existencias_query_params()
    resumen = filtrar_reporte_existencias(**filtros)
    categorias = sorted({item["insumo"].categoria or "General" for item in get_inventario_resumen()})
    return render_template("reporte_existencias.html", resumen_inventario=resumen, categorias=categorias, **filtros)


def build_existencias_pdf(rows, categoria: str = "", producto: str = "", estado: str = "") -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=24, leftMargin=24, topMargin=28, bottomMargin=28)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("existencias_title", parent=styles["Title"], fontSize=16, leading=20, alignment=1, spaceAfter=10)
    normal_style = ParagraphStyle("existencias_normal", parent=styles["BodyText"], fontSize=9, leading=12)
    data = [["Código", "Categoría", "Insumo", "Unidad", "Stock", "Mínimo", "Estado"]]
    for item in rows:
        insumo = item["insumo"]
        data.append([insumo.codigo or "-", insumo.categoria or "General", insumo.nombre, insumo.unidad, f"{item['stock']:.2f}", f"{insumo.stock_minimo:.2f}", "Reponer" if item["bajo"] else "Normal"])
    table = Table(data, repeatRows=1, colWidths=[58, 82, 150, 52, 58, 58, 58])
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8E6FF")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("PADDING", (0, 0), (-1, -1), 4),
        ("ALIGN", (4, 1), (5, -1), "RIGHT"),
    ]))
    logo = Image(str(STATIC_LOGO_PATH), width=125, height=54)
    header = Table([[logo, Paragraph("REPORTE DE EXISTENCIAS", title_style)]], colWidths=[160, 365])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    filtro = categoria or producto or estado or "Todos los productos"
    elements = [header, Spacer(1, 10), Paragraph(f"Filtro aplicado: {filtro}", normal_style), Paragraph(f"Generado: {datetime.utcnow():%d/%m/%Y %H:%M}", normal_style), Spacer(1, 10), table]
    doc.build(elements)
    return buffer.getvalue()


def build_system_audit_pdf() -> bytes:
    ensure_logo_exists()
    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=28, leftMargin=28, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("audit_title", parent=styles["Title"], fontSize=18, leading=22, alignment=1, spaceAfter=10)
    section_style = ParagraphStyle("audit_section", parent=styles["Heading2"], fontSize=13, leading=16, textColor=colors.HexColor("#123047"), spaceBefore=12, spaceAfter=7)
    normal_style = ParagraphStyle("audit_normal", parent=styles["BodyText"], fontSize=8.5, leading=11)
    small_style = ParagraphStyle("audit_small", parent=styles["BodyText"], fontSize=7, leading=9)

    def text(value):
        return Paragraph(str(value if value is not None else "-"), normal_style)

    def table(rows, widths, font_size=7):
        report_table = Table(rows, repeatRows=1, colWidths=widths)
        report_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D8E6FF")),
            ("GRID", (0, 0), (-1, -1), 0.45, colors.HexColor("#9AA8B5")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), font_size),
            ("PADDING", (0, 0), (-1, -1), 4),
        ]))
        return report_table

    resumen = get_inventario_resumen()
    entradas = EntradaInsumo.query.order_by(EntradaInsumo.fecha.asc()).all()
    salidas = SalidaInsumo.query.order_by(SalidaInsumo.fecha.asc()).all()
    retiros = RetiroInsumo.query.order_by(RetiroInsumo.fecha.asc()).all()
    pedidos = PedidoInsumo.query.order_by(PedidoInsumo.fecha.asc()).all()
    compras = Compra.query.order_by(Compra.fecha.asc()).all()
    despachos = Despacho.query.order_by(Despacho.fecha.asc()).all()
    usuarios = User.query.order_by(User.username.asc()).all()
    categorias = CategoriaInsumo.query.order_by(CategoriaInsumo.nombre.asc()).all()

    logo = Image(str(STATIC_LOGO_PATH), width=145, height=62)
    header = Table([[logo, Paragraph("INFORME AUDITADO DEL SISTEMA", title_style)]], colWidths=[175, 350])
    header.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LINEBELOW", (0, 0), (-1, -1), 1, colors.HexColor("#123047"))]))
    elements = [
        header,
        Spacer(1, 18),
        Paragraph("Informe de revisión y trazabilidad operativa", section_style),
        text("De: Isidro Vera"),
        text("Para: Ing. Nelson Valenzuela"),
        text(f"Fecha de emisión: {datetime.utcnow():%d/%m/%Y %H:%M}"),
        text("Alcance: revisión de los registros disponibles en los módulos de combustible, despacho e inventario/depósito. Este documento refleja los datos registrados en el sistema al momento de su generación."),
        Spacer(1, 8),
        Paragraph("Funcionamiento general del sistema", section_style),
        text("El sistema centraliza la gestión de combustible e inventario. Cada operación queda asociada a una fecha, usuario y registro relacionado. Las existencias se calculan con las entradas menos las salidas; los documentos de retiro conservan el detalle de productos y colaborador; y los módulos de consulta permiten filtrar y emitir reportes."),
        table([
            [text("Módulo"), text("Funcionamiento y alcance")],
            [text("Combustible"), text("Dashboard principal con stock, alertas, consumo y costos. Incluye Equipos, Compras, Despachos, Historial y Reportes." )],
            [text("Despacho"), text("Registra entregas de combustible por equipo, litros, responsable, chofer, hora, horómetro, kilometraje, tipo de carga y observaciones. El rol despacho puede cargar compras básicas; Admin completa los datos administrativos." )],
            [text("Inventario / Depósito"), text("Administra Insumos, Categorías, Entradas, Salidas, Pedidos, Reportes y Deducción. El stock se actualiza a partir de movimientos registrados." )],
            [text("Reportes"), text("Permite consultar existencias por categoría/producto/estado, retiros por colaborador y reportes operativos de combustible. Los resultados pueden imprimirse o descargarse en PDF cuando el módulo lo permite." )],
            [text("Administración"), text("Gestiona usuarios, roles, permisos y configuración visual. El acceso administrativo está restringido al rol Admin." )],
            [text("Modo offline"), text("Las operaciones compatibles se guardan temporalmente en el navegador cuando no hay conexión. Al recuperar internet se reenvían automáticamente al servidor y se evita duplicar una operación ya procesada." )],
        ], [145, 380], 7.5),
        Paragraph("Flujo de control y auditoría", section_style),
        text("1. Se registra el catálogo con código manual, nombre, categoría y unidad. 2. Las entradas incrementan la existencia. 3. Las salidas y retiros descuentan existencia y conservan el colaborador o la carga histórica general. 4. Los reportes consultan los movimientos sin modificar el origen. 5. Los documentos PDF sirven como respaldo imprimible para revisión y firma."),
        Spacer(1, 8),
        Paragraph("Resumen general", section_style),
        table([
            [text("Indicador"), text("Cantidad / valor")],
            [text("Usuarios registrados"), text(len(usuarios))],
            [text("Categorías"), text(len(categorias))],
            [text("Productos activos"), text(len(resumen))],
            [text("Entradas de inventario"), text(len(entradas))],
            [text("Salidas de inventario"), text(len(salidas))],
            [text("Documentos de retiro"), text(len(retiros))],
            [text("Pedidos de inventario"), text(len(pedidos))],
            [text("Compras de combustible"), text(len(compras))],
            [text("Despachos de combustible"), text(len(despachos))],
        ], [300, 225]),
        Paragraph("Usuarios y permisos", section_style),
        table([[text("Usuario"), text("Nombre"), text("Rol")]] + [[text(user.username), text(user.full_name), text(user.role)] for user in usuarios], [120, 245, 160]),
        Paragraph("Categorías registradas", section_style),
        table([[text("Categoría"), text("Prefijo"), text("Estado")]] + [[text(category.nombre), text(category.prefijo or "-"), text("Activa" if category.activo else "Inactiva")] for category in categorias], [270, 120, 135]),
        Paragraph("Catálogo y existencias actuales", section_style),
        table([[text("Código"), text("Producto"), text("Categoría"), text("Unidad"), text("Existencia"), text("Mínimo"), text("Estado")]] + [[text(item["insumo"].codigo or "-"), text(item["insumo"].nombre), text(item["insumo"].categoria), text(item["insumo"].unidad), text(f"{item['stock']:.2f}"), text(f"{item['insumo'].stock_minimo:.2f}"), text("Reponer" if item["bajo"] else "Normal")] for item in resumen], [52, 145, 85, 45, 62, 58, 58], 6.5),
        Paragraph("Movimientos de inventario: entradas", section_style),
        table([[text("Fecha"), text("Código"), text("Producto"), text("Cantidad"), text("Usuario"), text("Observación")]] + [[text(entry.fecha.strftime("%d/%m/%Y %H:%M")), text(entry.insumo.codigo), text(entry.insumo.nombre), text(f"{entry.cantidad:.2f} {entry.insumo.unidad}"), text(entry.usuario.username if entry.usuario else "-"), text(entry.observacion or "-")] for entry in entradas], [72, 55, 145, 75, 75, 103], 6.5),
        Paragraph("Movimientos de inventario: salidas", section_style),
        table([[text("Fecha"), text("Código"), text("Producto"), text("Cantidad"), text("Colaborador"), text("Documento")]] + [[text(output.fecha.strftime("%d/%m/%Y %H:%M")), text(output.insumo.codigo), text(output.insumo.nombre), text(f"{output.cantidad:.2f} {output.insumo.unidad}"), text(output.colaborador.nombre), text(output.retiro.codigo if output.retiro else "Legado")] for output in salidas], [72, 55, 145, 75, 110, 88], 6.5),
        Paragraph("Documentos de retiro", section_style),
        table([[text("Código"), text("Fecha"), text("Colaborador"), text("Productos"), text("Observación")]] + [[text(retiro.codigo), text(retiro.fecha.strftime("%d/%m/%Y %H:%M")), text(retiro.colaborador.nombre), text(len(retiro.salidas)), text(retiro.observacion or "-")] for retiro in retiros], [100, 82, 145, 65, 153], 6.5),
        Paragraph("Pedidos de inventario", section_style),
        table([[text("Código"), text("Fecha"), text("Solicitado por"), text("Estado"), text("Detalle")]] + [[text(pedido.codigo or f"PED-{pedido.id:06d}"), text(pedido.fecha.strftime("%d/%m/%Y %H:%M")), text(pedido.solicitado_por), text(pedido.estado), text("; ".join(f"{detail.insumo.codigo}: {detail.cantidad:g} {detail.insumo.unidad}" for detail in pedido.detalles))] for pedido in pedidos], [85, 82, 125, 75, 178], 6.5),
        Paragraph("Combustible: compras", section_style),
        table([[text("Fecha"), text("Litros"), text("Proveedor"), text("Precio/L"), text("Costo total"), text("Notas")]] + [[text(compra.fecha.strftime("%d/%m/%Y %H:%M")), text(f"{compra.litros:.2f}"), text(compra.proveedor), text(f"{compra.precio_litro:,.0f} Gs"), text(f"{compra.costo_total:,.0f} Gs"), text(compra.notas or "-")] for compra in compras], [82, 65, 125, 78, 85, 110], 6.5),
        Paragraph("Combustible: despachos", section_style),
        table([[text("Fecha"), text("Equipo"), text("Litros"), text("Responsable"), text("Chofer"), text("Carga"), text("Observación")]] + [[text(despacho.fecha.strftime("%d/%m/%Y %H:%M")), text(despacho.equipo.nombre), text(f"{despacho.litros:.2f}"), text(despacho.responsable), text(despacho.nombre_chofer or "-"), text(despacho.tipo_carga or "-"), text(despacho.observacion or "-")] for despacho in despachos], [72, 95, 55, 92, 80, 60, 90], 6.2),
        Paragraph("Configuración y controles", section_style),
        table([[text("Control"), text("Valor")]] + [[text(config.clave), text(config.valor)] for config in Configuracion.query.order_by(Configuracion.clave.asc()).all()], [220, 305]),
        Spacer(1, 18),
        Paragraph("Firma de revisión", section_style),
        Spacer(1, 28),
        text("__________________________________________        __________________________________________"),
        text("Isidro Vera                                                   Ing. Nelson Valenzuela"),
    ]
    doc.build(elements)
    return buffer.getvalue()


def build_source_code_pdf() -> bytes:
    code_extensions = {".py", ".html", ".css", ".js", ".json", ".ps1", ".bat", ".vbs", ".md", ".txt"}
    excluded_directories = {".git", ".venv", "__pycache__", "instance", "logs", "node_modules"}
    source_files = sorted(
        path for path in BASE_DIR.rglob("*")
        if path.is_file()
        and path.suffix.lower() in code_extensions
        and not excluded_directories.intersection(path.relative_to(BASE_DIR).parts)
    )

    buffer = BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=letter, rightMargin=30, leftMargin=30, topMargin=30, bottomMargin=30)
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle("source_title", parent=styles["Title"], fontSize=18, leading=22, alignment=1, spaceAfter=12)
    heading_style = ParagraphStyle("source_heading", parent=styles["Heading2"], fontSize=12, leading=15, textColor=colors.HexColor("#123047"), spaceBefore=8, spaceAfter=8)
    normal_style = ParagraphStyle("source_normal", parent=styles["BodyText"], fontSize=9, leading=12)
    code_style = ParagraphStyle("source_code", fontName="Courier", fontSize=6.5, leading=8, leftIndent=4, rightIndent=4)

    elements = [
        Paragraph("CÓDIGO FUENTE DEL SISTEMA", title_style),
        Paragraph(f"Generado: {datetime.now():%d/%m/%Y %H:%M}", normal_style),
        Paragraph(f"Archivos incluidos: {len(source_files)}", normal_style),
        Spacer(1, 12),
        Paragraph("Contenido", heading_style),
    ]
    elements.extend(Paragraph(str(path.relative_to(BASE_DIR)).replace("\\", "/"), normal_style) for path in source_files)

    for path in source_files:
        relative_path = str(path.relative_to(BASE_DIR)).replace("\\", "/")
        try:
            content = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            content = path.read_text(encoding="cp1252", errors="replace")
        content = content.replace("\x00", "")
        elements.extend([PageBreak(), Paragraph(relative_path, heading_style), Preformatted(content, code_style, maxLineLength=120)])

    doc.build(elements)
    return buffer.getvalue()


@app.route("/inventario/reportes/existencias/pdf")
@inventory_required
def pdf_reporte_existencias():
    filtros = reporte_existencias_query_params()
    rows = filtrar_reporte_existencias(**filtros)
    return send_file(BytesIO(build_existencias_pdf(rows, **filtros)), mimetype="application/pdf", as_attachment=True, download_name="reporte_existencias.pdf")


@app.route("/inventario/reportes/categorias")
@inventory_required
def reporte_categorias():
    categorias = {}
    for item in get_inventario_resumen():
        clave = (item["insumo"].categoria or "General", item["insumo"].unidad or "unidad")
        grupo = categorias.setdefault(clave, {"stock": 0.0, "insumos": 0, "bajos": 0})
        grupo["stock"] += item["stock"]
        grupo["insumos"] += 1
        grupo["bajos"] += int(item["bajo"])
    return render_template("reporte_categorias.html", categorias=sorted(categorias.items()))


@app.route("/inventario/reportes/retiros")
@inventory_required
def reporte_retiros():
    colaborador_query = request.args.get("colaborador", "").strip()
    historial = []
    retiros = []
    if colaborador_query:
        pattern = f"%{colaborador_query}%"
        colaborador_filter = or_(
            Colaborador.nombre.ilike(pattern),
            Colaborador.legajo.ilike(pattern),
            Colaborador.documento.ilike(pattern),
            Colaborador.codigo.ilike(pattern),
        )
        retiros = RetiroInsumo.query.join(Colaborador).filter(colaborador_filter).order_by(RetiroInsumo.fecha.desc()).all()
        historial = SalidaInsumo.query.join(Colaborador).filter(colaborador_filter).order_by(SalidaInsumo.fecha.desc()).all()
    return render_template("reporte_retiros.html", colaborador_query=colaborador_query, historial=historial, retiros=retiros)


@app.route("/inventario/reportes/vencimientos")
@inventory_required
def reporte_vencimientos():
    hoy = datetime.utcnow().date()
    insumos = Insumo.query.filter_by(activo=True).order_by(Insumo.vencimiento.asc(), Insumo.nombre.asc()).all()
    return render_template("reporte_vencimientos.html", insumos=insumos, hoy=hoy)


@app.route("/inventario/deduccion")
@inventory_required
def deduccion_inventario():
    return render_template("inventario_deduccion.html", deducciones=get_deduccion_inventario())


@app.route("/logout")
@login_required
def logout():
    logout_user()
    flash("Sesión cerrada correctamente.", "success")
    return redirect(url_for("login"))


@app.route("/")
@role_required("admin")
def index():
    period = request.args.get("period", "all")
    view = get_config_value("dashboard_theme_combustible", get_config_value("dashboard_theme", "classic"))
    stock = get_stock_actual()
    alert_threshold = float(get_config_value("stock_alert", "500"))
    compras = Compra.query.order_by(Compra.fecha.desc()).limit(5).all()
    despachos = Despacho.query.order_by(Despacho.fecha.desc()).limit(5).all()
    equipos = Equipo.query.order_by(Equipo.nombre.asc()).all()
    labels, values = build_stock_history(7)
    chart_data = build_equipment_consumption()
    total_consumo = get_total_consumo()
    chart_data_report = get_machine_chart_data(period)
    chart_data_report_detail = get_machine_report(period)
    if view == "pro":
        today = datetime.utcnow().date()
        month_start = today.replace(day=1)
        monthly_despachos = Despacho.query.filter(Despacho.fecha >= month_start).all()
        monthly_compras = Compra.query.filter(Compra.fecha >= month_start).all()
        return render_template(
            "index_pro.html",
            stock=stock,
            alert_threshold=alert_threshold,
            total_consumo=total_consumo,
            low_stock=stock <= alert_threshold,
            chart_data=chart_data,
            chart_data_report=chart_data_report,
            chart_data_report_detail=chart_data_report_detail,
            labels=labels,
            values=values,
            compras=compras,
            despachos=despachos,
            equipos=equipos,
            monthly_consumption=sum(item.litros or 0 for item in monthly_despachos),
            monthly_purchases=sum(item.litros or 0 for item in monthly_compras),
            active_equipment=Equipo.query.filter_by(estado="Activo").count(),
        )
    return render_template(
        "index.html",
        stock=stock,
        alert_threshold=alert_threshold,
        compras=compras,
        despachos=despachos,
        equipos=equipos,
        labels=labels,
        values=values,
        chart_data=chart_data,
        chart_data_report=chart_data_report,
        chart_data_report_detail=chart_data_report_detail,
        total_consumo=total_consumo,
        low_stock=stock <= alert_threshold,
    )


@app.route("/alerta", methods=["POST"])
@role_required("admin")
def set_alerta():
    valor = request.form.get("alerta", "500").strip()
    try:
        float(valor)
    except ValueError:
        flash("El valor de la alarma debe ser numérico.", "error")
        return redirect(url_for("index"))
    set_config_value("stock_alert", valor)
    flash("Alarma de stock actualizada.", "success")
    return redirect(url_for("index"))


@app.post("/combustible/configuracion")
@role_required("admin")
def configurar_combustible():
    try:
        stock_inicial = float(request.form.get("stock_inicial", "0"))
        stock_alert = float(request.form.get("stock_alert", "0"))
    except ValueError:
        flash("Los valores de stock deben ser numéricos.", "error")
        return redirect(url_for("reportes"))
    if stock_inicial < 0 or stock_alert < 0:
        flash("Los valores de stock no pueden ser negativos.", "error")
        return redirect(url_for("reportes"))
    set_config_value("stock_inicial", str(stock_inicial))
    set_config_value("stock_alert", str(stock_alert))
    flash("Stock inicial y stock saludable configurados correctamente.", "success")
    return redirect(url_for("reportes"))


@app.post("/tema-dashboard")
@role_required("admin")
def set_dashboard_theme():
    theme = request.form.get("theme", "classic").strip().lower()
    module = request.form.get("module", "combustible").strip().lower()
    valid_themes = {"classic", "pro", "sidebar"} | ({"inventory"} if module == "inventario" else set())
    if module not in {"combustible", "despacho", "inventario"} or theme not in valid_themes:
        flash("El tema seleccionado no es válido.", "error")
    else:
        set_config_value(f"dashboard_theme_{module}", theme)
        theme_label = {"classic": "Clásico", "pro": "Pro", "sidebar": "Navegación lateral", "inventory": "Panel de inventario"}[theme]
        flash(f"Tema {theme_label} configurado para {module}.", "success")
    return redirect(request.referrer or url_for("index"))


@app.route("/configuracion/tema", methods=["GET", "POST"])
@role_required("admin")
def configuracion_tema():
    if request.method == "POST":
        theme = request.form.get("theme", "").strip().lower()
        palette = request.form.get("palette", "").strip().lower()
        module = request.form.get("module", "combustible").strip().lower()
        valid_themes = {"classic", "pro", "sidebar"} | ({"inventory"} if module == "inventario" else set())
        valid_theme = module in {"combustible", "despacho", "inventario"} and theme in valid_themes
        valid_palette = palette in {"blue", "violet", "emerald", "amber"}
        if valid_theme:
            set_config_value(f"dashboard_theme_{module}", theme)
        if valid_palette:
            set_config_value("dashboard_palette", palette)
        if valid_theme or valid_palette:
            flash(f"Configuración actualizada para {module}.", "success")
        else:
            flash("El tema seleccionado no es válido.", "error")
        return redirect(url_for("configuracion_tema"))
    module = request.args.get("module", "combustible").strip().lower()
    if module not in {"combustible", "despacho", "inventario"}:
        module = "combustible"
    theme = get_config_value(f"dashboard_theme_{module}", get_config_value("dashboard_theme", "classic"))
    return render_template("configuracion_tema.html", theme=theme, palette=get_config_value("dashboard_palette", "blue"), module=module)


@app.route("/equipos", methods=["GET", "POST"])
@role_required("admin")
def equipos():
    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        tipo = request.form.get("tipo", "").strip()
        responsable = request.form.get("responsable", "").strip()
        estado = request.form.get("estado", "Activo").strip()
        if nombre and tipo and responsable:
            db.session.add(Equipo(nombre=nombre, tipo=tipo, responsable=responsable, estado=estado))
            db.session.commit()
            flash("Equipo registrado correctamente.", "success")
            return redirect(url_for("equipos"))
        flash("Complete todos los campos del equipo.", "error")
    listado = Equipo.query.order_by(Equipo.nombre.asc()).all()
    return render_template("equipos.html", equipos=listado)


@app.route("/equipos/<int:equipo_id>/editar", methods=["GET", "POST"])
@admin_required
def editar_equipo(equipo_id: int):
    equipo = Equipo.query.get_or_404(equipo_id)
    if request.method == "POST":
        nombre = request.form.get("nombre", "").strip()
        tipo = request.form.get("tipo", "").strip()
        responsable = request.form.get("responsable", "").strip()
        estado = request.form.get("estado", "Activo").strip()
        if not nombre or not tipo or not responsable:
            flash("Complete todos los campos del equipo.", "error")
        else:
            equipo.nombre = nombre
            equipo.tipo = tipo
            equipo.responsable = responsable
            equipo.estado = estado
            db.session.commit()
            flash("Equipo actualizado correctamente.", "success")
            return redirect(url_for("equipos"))
    return render_template("editar_equipo.html", equipo=equipo)


@app.post("/equipos/<int:equipo_id>/eliminar")
@admin_required
def eliminar_equipo(equipo_id: int):
    equipo = Equipo.query.get_or_404(equipo_id)
    if equipo.despachos:
        flash("No se puede eliminar un equipo que tiene despachos asociados.", "error")
        return redirect(url_for("equipos"))
    db.session.delete(equipo)
    db.session.commit()
    flash("Equipo eliminado correctamente.", "success")
    return redirect(url_for("equipos"))


@app.route("/compras", methods=["GET", "POST"])
@role_required("admin", "despacho", "compras")
def compras():
    if request.method == "POST":
        if offline_operation_was_processed():
            return redirect(url_for("compras"))
        litros = request.form.get("litros", "0").strip()
        proveedor = request.form.get("proveedor", "").strip() if current_user.is_admin else "Pendiente de completar"
        precio_litro = request.form.get("precio_litro", "0").strip() if current_user.is_admin else "0"
        fecha = request.form.get("fecha", "")
        notas = request.form.get("notas", "").strip()
        try:
            litros_float = float(litros)
            precio_float = float(precio_litro)
        except ValueError:
            flash("Los litros y el precio deben ser numéricos.", "error")
            return redirect(url_for("compras"))
        if litros_float <= 0:
            flash("Debe registrar más de 0 litros.", "error")
            return redirect(url_for("compras"))
        if (current_user.is_admin or current_user.is_compras) and not proveedor:
            flash("Debe indicar el proveedor.", "error")
            return redirect(url_for("compras"))
        compra = Compra(litros=litros_float, proveedor=proveedor, precio_litro=precio_float, costo_total=litros_float * precio_float, notas=notas)
        if fecha:
            compra.fecha = datetime.strptime(fecha, "%Y-%m-%d")
        db.session.add(compra)
        remember_offline_operation()
        db.session.commit()
        flash("Registro de compra agregado correctamente.", "success")
        return redirect(url_for("compras"))
    listado = Compra.query.order_by(Compra.fecha.desc()).all()
    return render_template("compras.html", compras=listado)


@app.get("/modulo/compras")
@role_required("admin", "compras")
def modulo_compras():
    query = request.args.get("q", "").strip()
    estado = request.args.get("estado", "").strip()
    pedidos_query = PedidoInsumo.query.filter_by(estado_aprobacion="Aprobado", estado="Pendiente")
    if query:
        pattern = f"%{query}%"
        pedidos_query = pedidos_query.filter(or_(PedidoInsumo.codigo.ilike(pattern), PedidoInsumo.solicitado_por.ilike(pattern)))
    if estado:
        pedidos_query = pedidos_query.filter_by(estado=estado)
    pedidos = pedidos_query.order_by(PedidoInsumo.fecha.desc()).all()
    pendientes_presupuesto = PedidoInsumo.query.filter(PedidoInsumo.estado_aprobacion == "Pendiente", PedidoInsumo.solicitante_usuario_id.isnot(None), PedidoInsumo.rango_monto.in_(("mayor", "a_presupuestar"))).order_by(PedidoInsumo.fecha.desc()).all()
    compras_recientes = Compra.query.order_by(Compra.fecha.desc()).limit(10).all()
    return render_template(
        "modulo_compras.html",
        pedidos=pedidos,
        pendientes_presupuesto=pendientes_presupuesto,
        compras_recientes=compras_recientes,
        query=query,
        estado=estado,
        total_pedidos=pedidos_query.count(),
        pendientes=pedidos_query.filter(PedidoInsumo.estado == "Pendiente").count(),
        en_proceso=pedidos_query.filter(PedidoInsumo.estado == "En proceso").count(),
        recibidos=pedidos_query.filter(PedidoInsumo.estado.in_(("Recibido", "Entregado", "Recibido en empresa"))).count(),
        no_gestionados=pedidos_query.filter(PedidoInsumo.estado == "No gestionado").count(),
        pendientes_aprobacion=PedidoInsumo.query.filter_by(estado_aprobacion="Pendiente").count(),
    )


@app.get("/modulo/compras/reportes")
@role_required("admin", "compras")
def reportes_compras():
    pedidos = PedidoInsumo.query.order_by(PedidoInsumo.fecha.desc()).all()
    directos = [pedido for pedido in pedidos if pedido.solicitante_usuario_id is None]
    menores = [pedido for pedido in pedidos if pedido.solicitante_usuario_id is not None and pedido.rango_monto == "menor"]
    gerencia = [pedido for pedido in pedidos if pedido.solicitante_usuario_id is not None and pedido.rango_monto in {"mayor", "a_presupuestar"}]
    return render_template("reportes_compras.html", directos=directos, menores=menores, gerencia=gerencia)


@app.get("/modulo/compras/historial.csv")
@role_required("admin", "compras")
def modulo_compras_csv():
    rows = [["Tipo", "Código", "Fecha", "Solicitante/Proveedor", "Detalle", "Estado"]]
    for pedido in PedidoInsumo.query.order_by(PedidoInsumo.fecha.desc()).all():
        detalle = "; ".join(f"{item.insumo.codigo} - {item.insumo.nombre}: {item.cantidad:g}" for item in pedido.detalles)
        rows.append(["Pedido de insumos", pedido.codigo or f"PED-{pedido.id:06d}", pedido.fecha.strftime("%d/%m/%Y %H:%M"), pedido.solicitado_por, detalle, pedido.estado])
    for compra in Compra.query.order_by(Compra.fecha.desc()).all():
        rows.append(["Compra de combustible", f"COMP-{compra.id:06d}", compra.fecha.strftime("%d/%m/%Y %H:%M"), compra.proveedor, f"{compra.litros:g} litros", "Registrada"])
    output = StringIO()
    csv.writer(output).writerows(rows)
    return Response("\ufeff" + output.getvalue(), mimetype="text/csv", headers={"Content-Disposition": "attachment; filename=historial_compras.csv"})


@app.route("/compras/<int:compra_id>/editar", methods=["GET", "POST"])
@admin_required
def editar_compra(compra_id: int):
    compra = Compra.query.get_or_404(compra_id)
    if request.method == "POST":
        try:
            litros = float(request.form.get("litros", "0"))
            precio_litro = float(request.form.get("precio_litro", "0"))
            fecha = request.form.get("fecha", "")
        except ValueError:
            flash("Los litros y el precio deben ser numéricos.", "error")
            return render_template("editar_compra.html", compra=compra)
        proveedor = request.form.get("proveedor", "").strip()
        if litros <= 0 or not proveedor:
            flash("Debe indicar litros válidos y el proveedor.", "error")
            return render_template("editar_compra.html", compra=compra)
        compra.litros = litros
        compra.precio_litro = precio_litro
        compra.costo_total = litros * precio_litro
        compra.proveedor = proveedor
        compra.notas = request.form.get("notas", "").strip()
        if fecha:
            compra.fecha = datetime.strptime(fecha, "%Y-%m-%d")
        db.session.commit()
        flash("Compra actualizada correctamente.", "success")
        return redirect(url_for("compras"))
    return render_template("editar_compra.html", compra=compra)


@app.post("/compras/<int:compra_id>/eliminar")
@admin_required
def eliminar_compra(compra_id: int):
    compra = Compra.query.get_or_404(compra_id)
    db.session.delete(compra)
    db.session.commit()
    flash("Compra eliminada correctamente. El stock fue recalculado.", "success")
    return redirect(url_for("compras"))


@app.route("/despacho")
@role_required("admin", "despacho")
def despacho_dashboard():
    return render_template(
        "despacho_dashboard.html",
        stock=get_stock_actual(),
        alert_threshold=float(get_config_value("stock_alert", "500")),
        total_despachos=Despacho.query.count(),
        total_compras=Compra.query.count(),
    )


@app.route("/despachos", methods=["GET", "POST"])
@role_required("admin", "despacho")
def despachos():
    if request.method == "POST":
        if offline_operation_was_processed():
            return redirect(url_for("despachos"))
        equipo_id = request.form.get("equipo_id", "")
        litros = request.form.get("litros", "0").strip()
        responsable = request.form.get("responsable", "").strip()
        nombre_chofer = request.form.get("nombre_chofer", "").strip()
        hora = request.form.get("hora", "").strip()
        horometro = request.form.get("horometro", "").strip()
        kilometraje = request.form.get("kilometraje", "").strip()
        tipo_carga = request.form.get("tipo_carga", "Media").strip()
        fecha = request.form.get("fecha", "")
        observacion = request.form.get("observacion", "").strip()
        try:
            litros_float = float(litros)
            hora_float = float(hora) if hora else 0.0
            horometro_float = float(horometro) if horometro else 0.0
            kilometraje_float = float(kilometraje) if kilometraje else 0.0
        except ValueError:
            flash("Los litros, hora, horómetro y kilometraje deben ser numéricos.", "error")
            return redirect(url_for("despachos"))
        if not equipo_id or litros_float <= 0 or not responsable:
            flash("Debe completar todos los datos del despacho.", "error")
            return redirect(url_for("despachos"))
        equipo = Equipo.query.get(int(equipo_id))
        if not equipo:
            flash("Equipo no encontrado.", "error")
            return redirect(url_for("despachos"))
        despacho = Despacho(
            equipo_id=equipo.id,
            litros=litros_float,
            responsable=responsable,
            nombre_chofer=nombre_chofer,
            hora=hora_float,
            horometro=horometro_float,
            kilometraje=kilometraje_float,
            tipo_carga=tipo_carga,
            observacion=observacion,
        )
        if fecha:
            despacho.fecha = datetime.strptime(fecha, "%Y-%m-%d")
        db.session.add(despacho)
        remember_offline_operation()
        db.session.commit()
        flash("Despacho registrado correctamente.", "success")
        return redirect(url_for("despachos"))
    listado = Despacho.query.order_by(Despacho.fecha.desc()).all()
    equipos = Equipo.query.order_by(Equipo.nombre.asc()).all()
    return render_template("despachos.html", despachos=listado, equipos=equipos)


@app.route("/despachos/<int:despacho_id>/editar", methods=["GET", "POST"])
@admin_required
def editar_despacho(despacho_id: int):
    despacho = Despacho.query.get_or_404(despacho_id)
    equipos = Equipo.query.order_by(Equipo.nombre.asc()).all()
    if request.method == "POST":
        try:
            equipo_id = int(request.form.get("equipo_id", "0"))
            litros = float(request.form.get("litros", "0"))
            hora = float(request.form.get("hora", "0") or 0)
            horometro = float(request.form.get("horometro", "0") or 0)
            kilometraje = float(request.form.get("kilometraje", "0") or 0)
        except ValueError:
            flash("Los datos numéricos del despacho no son válidos.", "error")
            return render_template("editar_despacho.html", despacho=despacho, equipos=equipos)
        responsable = request.form.get("responsable", "").strip()
        equipo = db.session.get(Equipo, equipo_id)
        if litros <= 0 or not responsable or equipo is None:
            flash("Complete correctamente el equipo, los litros y el responsable.", "error")
            return render_template("editar_despacho.html", despacho=despacho, equipos=equipos)
        despacho.equipo_id = equipo_id
        despacho.litros = litros
        despacho.responsable = responsable
        despacho.nombre_chofer = request.form.get("nombre_chofer", "").strip()
        despacho.hora = hora
        despacho.horometro = horometro
        despacho.kilometraje = kilometraje
        despacho.tipo_carga = request.form.get("tipo_carga", "Media").strip()
        despacho.observacion = request.form.get("observacion", "").strip()
        fecha = request.form.get("fecha", "")
        if fecha:
            despacho.fecha = datetime.strptime(fecha, "%Y-%m-%d")
        db.session.commit()
        flash("Despacho actualizado correctamente.", "success")
        return redirect(url_for("despachos"))
    return render_template("editar_despacho.html", despacho=despacho, equipos=equipos)


@app.post("/despachos/<int:despacho_id>/eliminar")
@admin_required
def eliminar_despacho(despacho_id: int):
    despacho = Despacho.query.get_or_404(despacho_id)
    db.session.delete(despacho)
    db.session.commit()
    flash("Despacho eliminado correctamente. El stock fue recalculado.", "success")
    return redirect(url_for("despachos"))


@app.route("/despachos/<int:despacho_id>")
@role_required("admin", "despacho")
def detalle_despacho(despacho_id: int):
    despacho = Despacho.query.get_or_404(despacho_id)
    return render_template("despacho_detalle.html", despacho=despacho)


@app.route("/despachos/<int:despacho_id>/pdf")
@role_required("admin", "despacho")
def pdf_despacho(despacho_id: int):
    despacho = Despacho.query.get_or_404(despacho_id)
    pdf_bytes = build_despacho_pdf(despacho)
    return send_file(
        BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"comprobante_despacho_{despacho.id}.pdf",
    )


@app.route("/historial")
@role_required("admin")
def historial():
    entries = build_history_entries()
    return render_template("historial.html", entries=entries)


@app.route("/reportes/maquinas")
@role_required("admin")
def reportes_maquinas():
    period = request.args.get("period", "all")
    fecha_desde = request.args.get("fecha_desde")
    fecha_hasta = request.args.get("fecha_hasta")
    start_dt, end_dt = get_date_range(period, fecha_desde, fecha_hasta)
    resumen = get_machine_report(period, None, fecha_desde, fecha_hasta)
    maquinas = []
    for item in resumen:
        query = Despacho.query.filter_by(equipo_id=item["equipo"].id)
        if start_dt is not None:
            query = query.filter(Despacho.fecha >= start_dt)
        if end_dt is not None:
            query = query.filter(Despacho.fecha <= end_dt)
        maquinas.append({"resumen": item, "cargas": query.order_by(Despacho.fecha.asc()).all()})
    return render_template("reportes_maquinas.html", maquinas=maquinas, period=period, fecha_desde=fecha_desde, fecha_hasta=fecha_hasta)


@app.route("/reportes")
@role_required("admin")
def reportes():
    period = request.args.get("period", "all")
    fecha_desde = request.args.get("fecha_desde")
    fecha_hasta = request.args.get("fecha_hasta")
    equipo_id = request.args.get("equipo_id", type=int)
    resumen = get_machine_report(period, equipo_id, fecha_desde, fecha_hasta) if equipo_id else get_machine_report(period, None, fecha_desde, fecha_hasta)
    selected_machine = Equipo.query.get(equipo_id) if equipo_id else None
    selected_data = next((item for item in resumen if item["equipo"].id == equipo_id), None) if equipo_id else None
    historial_equipo = []
    if selected_machine:
        historial_equipo = Despacho.query.filter_by(equipo_id=selected_machine.id).order_by(Despacho.fecha.desc()).all()
    chart_data = get_machine_chart_data(period, None, fecha_desde, fecha_hasta)
    stock_actual = get_stock_actual()
    stock_inicial = get_initial_stock_value()
    stock_alert = float(get_config_value("stock_alert", "500"))
    stock_fill_pct = min(100.0, max(0.0, (stock_actual / stock_alert) * 100.0)) if stock_alert > 0 else 0.0
    return render_template(
        "reportes.html",
        resumen=resumen,
        period=period,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        chart_data=chart_data,
        selected_machine=selected_machine,
        selected_data=selected_data,
        historial_equipo=historial_equipo,
        stock_actual=stock_actual,
        stock_inicial=stock_inicial,
        stock_alert=stock_alert,
        stock_fill_pct=stock_fill_pct,
        stock_saludable=stock_actual > stock_alert,
    )


@app.route("/reportes/export/<period>")
@role_required("admin")
def export_report_csv(period: str):
    fecha_desde = request.args.get("fecha_desde")
    fecha_hasta = request.args.get("fecha_hasta")
    csv_bytes = build_machine_report_csv(period, None, fecha_desde, fecha_hasta)
    return send_file(
        BytesIO(csv_bytes),
        mimetype="text/csv; charset=utf-8",
        as_attachment=True,
        download_name=f"reporte_combustible_{period}.csv",
    )


@app.route("/reportes/pdf/<period>")
@role_required("admin")
def export_report_pdf(period: str):
    fecha_desde = request.args.get("fecha_desde")
    fecha_hasta = request.args.get("fecha_hasta")
    pdf_bytes = build_machine_report_pdf(period, None, fecha_desde, fecha_hasta)
    return send_file(
        BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"reporte_combustible_{period}.pdf",
    )


@app.route("/reportes/pdf/historial/<period>")
@role_required("admin", "despacho")
def export_fuel_history_pdf(period: str):
    fecha_desde = request.args.get("fecha_desde")
    fecha_hasta = request.args.get("fecha_hasta")
    pdf_bytes = build_fuel_history_pdf(period, fecha_desde, fecha_hasta)
    return send_file(
        BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=False,
        download_name=f"historico_carga_combustible_{period}.pdf",
    )


@app.route("/reportes/excel/historial/<period>")
@role_required("admin", "despacho")
def export_fuel_history_excel(period: str):
    fecha_desde = request.args.get("fecha_desde")
    fecha_hasta = request.args.get("fecha_hasta")
    excel_bytes = build_fuel_history_excel(period, fecha_desde, fecha_hasta)
    return send_file(
        BytesIO(excel_bytes),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=f"historico_carga_combustible_{period}.xlsx",
    )


@app.route("/reportes/pdf/maquina/<int:equipo_id>")
@role_required("admin")
def export_machine_pdf(equipo_id: int):
    period = request.args.get("period", "all")
    fecha_desde = request.args.get("fecha_desde")
    fecha_hasta = request.args.get("fecha_hasta")
    pdf_bytes = build_machine_report_pdf(period, equipo_id, fecha_desde, fecha_hasta)
    equipo = Equipo.query.get_or_404(equipo_id)
    return send_file(
        BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=f"reporte_{equipo.nombre.lower().replace(' ', '_')}_{period}.pdf",
    )


@app.route("/historial/<string:tipo>/<int:item_id>/pdf")
@role_required("admin", "compras")
def export_history_pdf(tipo: str, item_id: int):
    if tipo == "compra":
        compra = Compra.query.get_or_404(item_id)
        pdf_bytes = build_compra_pdf(compra)
        download_name = f"compra_{compra.id}.pdf"
    elif tipo == "despacho":
        despacho = Despacho.query.get_or_404(item_id)
        pdf_bytes = build_despacho_pdf(despacho)
        download_name = f"despacho_{despacho.id}.pdf"
    else:
        return redirect(url_for("historial"))
    return send_file(
        BytesIO(pdf_bytes),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=download_name,
    )


with app.app_context():
    ensure_logo_exists()
    init_db()


if __name__ == "__main__":
    ssl_cert = os.environ.get("APP_SSL_CERT", str(BASE_DIR / "instance" / "https" / "control-combustible.crt"))
    ssl_key = os.environ.get("APP_SSL_KEY", str(BASE_DIR / "instance" / "https" / "control-combustible.key"))
    ssl_enabled = os.environ.get("APP_ENABLE_SSL", "0").strip().lower() in {"1", "true", "yes", "on"}
    ssl_context = (ssl_cert, ssl_key) if ssl_enabled and Path(ssl_cert).exists() and Path(ssl_key).exists() else None
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", "5000")), debug=False, threaded=True, ssl_context=ssl_context)
